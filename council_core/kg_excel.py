"""
council_core.kg_excel — Excel workbooks and Ledgers as knowledge-graph sources.

The knowledge graph (council_core/knowledge_graph.py) reads two kinds of
spreadsheet here, both with the Ledger's own code (sample_ledger.py,
source_extract.py, excel_open.py) rather than a second Excel reader:

LEDGERS — a parent CSV with its schema beside it (``<name>.schema.json``, made
by the Ledger tab). Every row holds values a person ACCEPTED in a review, so
its links are firm ('seeded'). The schema says what each column is; nothing is
guessed from a column's values:
  * the sample-ID column (the template's Part Number) is the PART, unless a
    confirmed field rule says its label means something else;
  * '<field> [JB]' columns (one inspector's own reading) and
    '<field> — Inspector Initials' columns name who INSPECTED the part;
  * a field whose label the Ledger reads as a person ('Operator',
    'Inspected by' ...) names one too;
  * a field whose label IS one of the graph's confirmed field rules
    ('Project', 'Program Lead') reads as that rule says.
  Each link's evidence is the CSV row (and the flaw), plus where the value
  came from when the Ledger kept it: the row's source workbooks
  ('source_file') and, from the Ledger's review workbooks beside the CSV, the
  workbook › sheet!cell of each accepted value. Values still waiting in a
  review workbook are not read: only the CSV holds the user's word.

WORKBOOKS — any .xlsx / .xlsm / .xls / .xlsb, read with the Ledger's reader
(source_extract.load_workbook_grids): every sheet, merged labels, text boxes,
charts, pictures read by Tesseract when it is installed, and .xls / .xlsb /
rights-protected files through the installed Excel (hidden, read-only, macros
off, closed unsaved: excel_open.py). A table's header may be on any row and a
sheet may hold several tables, side by side or stacked; labelled cells
('Program Lead: Dana Whitfield', or a label with its value beside it) count
too. Only the graph's CONFIRMED field rules make firm links. What only the
Ledger's own reading says — an 'Inspector' column, 'Operator: JB', a part
named by its sheet ('Part 101'), anything read from a picture — makes
'suggested' links for the user to accept. Long text (notes cells, text boxes,
pictures) becomes lines for the gazetteer and the free-text pass, cited by
sheet and row.

Initials ('JB') are never guessed. knowledge_graph resolves them through the
Ledger's people list (``<parent>.aliases.json``, "people": {"JB": "Jordan
Blake"}; source_extract.load_people), the user's answers and the names they
confirmed; any other initials are asked once, as a name question.

Nothing here writes a file, saves a picture or talks to a network. Excel is
used only for what openpyxl cannot open, and only when it is installed.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Iterator, List, Optional, Sequence, Tuple

from council_core import knowledge_graph as kgm

WORKBOOK_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".xlsb",
                     # slides: their tables laid out as sheets (pptx_open), as the
                     # Ledger reads them (synthetic test, 2026-10-27)
                     ".pptx", ".pptm")

#: Above this many bytes of uncompressed sheet XML a workbook is STREAMED
#: (openpyxl read-only, one row at a time) instead of loaded whole. MEASURED
#: (2026-10-16, this laptop): a 100,000 x 10 sheet (4.6 MB file, ~50 MB of
#: sheet XML) took +539 MB through the Ledger's loader, +5 MB streamed, in
#: about the same time (33-37 s). 12 MB of XML is roughly 250,000 cells, about
#: 130 MB loaded. Streaming reads no merged ranges and no pictures (openpyxl's
#: read-only mode has neither); text boxes and charts still come from the XML.
STREAM_XML_BYTES = 12 << 20
#: Free-text lines kept per workbook (a notes column on 100k rows would
#: otherwise put every one of them in the store's cache).
MAX_LINES = 20_000
MAX_LABELS = 5_000
#: ...and characters of it, as a text document's are capped (5 M): one
#: notes workbook kept 20,000 lines = 15.9 M characters (review, 2026-10-16).
MAX_TEXT_CHARS = 5_000_000
#: A cell is free text (for the gazetteer and the model) from this many words.
LONG_TEXT_WORDS = 4
#: Bigger workbooks are parsed from their file, smaller ones from a copy in
#: memory: openpyxl holds the file open while it parses (8.7 s for 20,000
#: rows), and Excel could not save the workbook meanwhile (review, 2026-10-16).
IN_MEMORY_BYTES = 64 << 20


class Unreadable(Exception):
    """A workbook or Ledger that cannot be read, in plain words.
    ``permanent``: reading the same bytes again cannot help (a password, a
    file that is not a workbook), so the graph remembers it by content hash
    instead of starting Excel for it on every rebuild."""

    def __init__(self, message: str, permanent: bool = False):
        super().__init__(message)
        self.permanent = permanent


@dataclass
class Reading:
    """What reading one workbook or Ledger gave. Inside a rebuild the records
    come as their cache lines (``record_lines``, kgm._record_json), written as
    each is found and ``records`` left empty: a 100k-row workbook's 100,000
    records held as objects, then turned into text, peaked at 274 MB
    (measured 2026-10-18); its cache lines are 46 MB."""
    records: List["kgm.Record"]
    lines: List[Tuple[Optional[str], int, str]] = field(default_factory=list)
    labels: List[str] = field(default_factory=list)
    summary: Dict[str, Any] = field(default_factory=dict)
    notes: List[str] = field(default_factory=list)
    record_lines: Optional[List[str]] = None


class _Sink:
    """Where a reader puts its records: objects, or (in a rebuild) their
    cache lines, with the people they name counted either way."""

    def __init__(self):
        self.as_lines = kgm.records_as_lines_now()
        self.records: List["kgm.Record"] = []
        self.lines: List[str] = []
        self.people: set = set()
        self.n = 0

    def add(self, rec: "kgm.Record") -> None:
        self.n += 1
        for v in rec.values:
            if v.rule.type == "PERSON":
                self.people.add(v.raw)
        if self.as_lines:
            self.lines.append(kgm._record_json(rec))
        else:
            self.records.append(rec)

    def reading(self, **kw) -> Reading:
        return Reading(self.records, record_lines=self.lines if self.as_lines else None, **kw)


# ── small helpers ────────────────────────────────────────────────────────
def is_workbook(p: Path) -> bool:
    return Path(p).suffix.lower() in WORKBOOK_SUFFIXES and not Path(p).name.startswith("~$")


def _se():
    import source_extract
    return source_extract


def _text(v) -> str:
    from sample_ledger import _text as t
    return t(v)


def _col_letter(c: int) -> str:
    s = ""
    while c > 0:
        c, rem = divmod(c - 1, 26)
        s = chr(65 + rem) + s
    return s


def _label_text(v) -> str:
    """'Program Lead:' -> 'Program Lead'. Not '[\\s:]+$': it rescanned a
    run of colons from each of them, quadratic on a cell of ':' (review,
    2026-10-20); the words have no other space to strip."""
    return " ".join(str(v).split()).rstrip(" :")


def _norm(v) -> str:
    import field_search as fs
    return fs._norm_key(_label_text(v))


#: Short words a results or initials column holds that are not anyone's
#: initials, however they are typed ('TBD', 'OK', 'NA'). Not 'EC': Eva Chen
#: signs as EC.
_NOT_INITIALS = {
    # outcome / status cells
    "seen", "ok", "okay", "yes", "no", "na", "n a", "tbd", "tba", "pass",
    "fail", "none", "nil", "good", "bad", "done", "open", "x", "y", "n", "nd",
    # method labels (deliberately NOT including 'VT', 'UT', 'ET', 'EC',
    # 'PT', 'MT', 'RT' — the comment above calls out EC as a real
    # inspector's initials, and the others collide with real names too)
    "nde", "fpi",
    # workbook-structure labels
    "date", "flaw", "part", "run", "item", "id", "all", "any",
    # stage / phase labels — the user hit "pre" and "post" being asked
    # as people (bug 2026-10-19); inspection sheets routinely use these
    # as column headers (short for Pre-Conceal / Post-Conceal), and all-
    # caps "PRE" / "POST" match the 1-4-letter initials regex. Short-
    # cut rejection here. No real inspector uses 'PRE' or 'POST' as
    # initials — if that ever happens the user can clear the stop-list.
    "pre", "post", "mid", "end", "start",
    # unit labels that are 1-4 letters and sometimes typed all-caps as
    # a column header on their own ("MM", "KG"). As with the method
    # labels, we skip 2-letter units that could be initials ("MM" vs
    # Matt M., "LB" vs Lee B.) — only the ones that are overwhelmingly
    # units. 'KW', 'MW', 'HZ' are real inspectors' initials (Mark Williams
    # became a person called 'MW' - review, 2026-10-19): not listed.
    "psi", "khz", "mhz", "rpm", "ipm", "fpm", "deg",
    "sec", "min", "max", "avg", "std", "ppm",
} | _se().NOT_INITIALS        # the Ledger's: result codes 'NRI', 'IND', 'REJ' ...
#: A cell that says nobody signed: 'N/A', 'n.a.', '-', '?', 'TBD'. The
#: Ledger's list: one there is no part number either (review, 2026-10-19).
_NOBODY = _se().PLACEHOLDER


def is_initials(s: Any, loose: bool = False) -> bool:
    """'JB', 'J.B.', 'JMB', 'j.b.' — what an inspector signs with: capitals,
    or letters with dots. ``loose`` (a cell under an 'Inspector' label) also
    takes 'jb', but never a word a results column holds ('seen', 'ok'). A
    short capitalised word ('Bob', 'Ann') is a name, not initials."""
    raw = " ".join(str(s or "").split())
    t = re.sub(r"[.\s]", "", raw)
    if not re.fullmatch(r"[A-Za-z]{1,4}", t) or re.fullmatch(r"[A-Z][a-z]{2,}", raw):
        return False
    if t.lower() in _NOT_INITIALS:          # 'TBD', 'OK', 'NA', 'seen'
        return False
    if re.fullmatch(r"(?:[A-Z]\.?\s?){1,4}", raw):          # JB, J.B., J B
        return True
    if re.fullmatch(r"(?:[A-Za-z]\.\s?){1,4}", raw):        # j.b.
        return True
    return loose and raw.islower()                         # jb


def initials_key(s: Any) -> str:
    return re.sub(r"[.\s]", "", str(s or "")).upper()


_NAME_WORD = re.compile(r"[A-Z][a-z'’\-]+\.?|[A-Z]\.?")


def person_ok(s: Any) -> bool:
    """Initials or a written name ('Jordan Blake', 'Lee, Carol', 'D. Whitfield'),
    never a reading or a column title ('EC Pre-Conceal', 'seen', '63')."""
    raw = " ".join(str(s or "").split())
    if not raw or len(raw) > 60 or any(ch.isdigit() for ch in raw):
        return False
    if is_initials(raw):
        return True
    words = raw.replace(",", " ").split()
    return (2 <= len(words) <= 4 and all(_NAME_WORD.fullmatch(w) for w in words)
            and any(len(w.rstrip(".")) > 1 for w in words))


def split_people(s: Any) -> List[str]:
    """Who signed, from a cell under an 'Inspector' / 'Initials' label: 'JB/MK'
    (two inspectors who agreed), 'JB, MK', 'JB & MK' -> each one, initials in
    capitals ('jb' -> 'JB'); a name stays whole ('Lee, Carol'). Anything else
    ('seen', '63', a date) is nobody."""
    raw = " ".join(str(s or "").split())
    # '(JB)', '[MK]': initials in brackets are someone's (the Ledger reads
    # them so - source_extract._initials; review, 2026-10-20)
    raw = re.sub(r"^[\(\[]\s*([^()\[\]]*?)\s*[\)\]]$", r"\1", raw)
    if not raw or _NOBODY.fullmatch(raw):
        return []
    parts = [p.strip() for p in re.split(r"\s*(?:/|&|;|\band\b)\s*", raw) if p.strip()]
    if len(parts) == 1 and "," in raw:
        bits = [b.strip() for b in raw.split(",") if b.strip()]
        if bits and all(is_initials(b, loose=True) for b in bits):
            parts = bits
    out = []
    for p in parts:
        if is_initials(p, loose=True):
            out.append(initials_key(p))
        elif person_ok(p):
            out.append(p)
    return out


def _who_head(v, taught: Any = None) -> bool:
    """A label the Ledger reads as 'who logged this' (Inspector, Initials,
    Operator, Inspected by ...): source_extract's own pattern - and, for a
    Ledger, the labels its Teach learned are an inspector's or are no one
    (``taught``: source_extract.Taught; council_core.ledger_teach)."""
    if not isinstance(v, str):
        return False
    if taught is not None and (taught.who or taught.not_who):
        from sample_ledger import normalize_label
        key = normalize_label(_label_text(v))
        if key in {normalize_label(x) for x in taught.not_who}:
            return False
        if key in {normalize_label(x) for x in taught.who}:
            return True
    return bool(_se()._WHO_HEAD.match(v))


def _who_inline(v) -> List[str]:
    """'Operator: JB', 'Inspected by MK', 'Inspector: JB/MK' -> the people,
    as the Ledger reads them."""
    if not isinstance(v, str):
        return []
    se = _se()
    m = se._WHO_INLINE.match(v) or se._WHO_BY.match(v)
    return split_people(m.group("who")) if m else []


# ── Ledgers ──────────────────────────────────────────────────────────────
def schema_path(csv_path: Path) -> Path:
    from sample_ledger import schema_path_for
    return schema_path_for(csv_path)


#: What the Ledger's schema file always holds (sample_ledger.TemplateSchema).
_LEDGER_SCHEMA_KEYS = {"template_path", "fields", "id_field"}


def is_ledger_csv(p: Path) -> bool:
    """A parent CSV: a .csv with the Ledger's own schema beside it. A
    '<name>.schema.json' is also how JSON Table Schema / Frictionless
    describe a plain CSV, so the file must BE a Ledger schema; one that
    cannot be parsed counts when the CSV ends with the Ledger's own columns
    (source_file, updated_at) - then it is a Ledger whose schema is damaged,
    reported as such. It used to be enough for the file to exist: a plain
    team.csv beside a Table Schema was refused as a broken Ledger (review,
    2026-10-16)."""
    p = Path(p)
    if p.suffix.lower() != ".csv":
        return False
    sp = schema_path(p)
    if not sp.is_file():
        return False
    try:
        data = json.loads(sp.read_text(encoding="utf-8"))
    except Exception:
        return _ends_with_ledger_columns(p)
    return (isinstance(data, dict) and _LEDGER_SCHEMA_KEYS <= set(data)
            and isinstance(data.get("fields"), list))


def _ends_with_ledger_columns(p: Path) -> bool:
    import sample_ledger as sl
    try:
        with open(p, "rb") as fh:
            head = _decode(fh.read(1 << 16)).splitlines()
    except OSError:
        return False
    for line in head:
        if line.strip():
            cols = [c.strip() for c in next(csv.reader([line]))]
            return tuple(cols[-len(sl.META_COLS):]) == tuple(sl.META_COLS)
    return False


def is_ledger_sidecar(p: Path) -> bool:
    """The Ledger's own files beside a parent CSV (its schema and its
    aliases): settings, not documents."""
    name = Path(p).name.lower()
    for suf in (".schema.json", ".aliases.json"):
        if name.endswith(suf):
            return is_ledger_csv(Path(p).with_name(Path(p).name[:-len(suf)] + ".csv"))
    return False


def is_review_workbook(p: Path) -> bool:
    """A review workbook the Ledger wrote (its hidden '_run' sheet): values
    WAITING for a person's decision, so never read as a source. Told by the
    sheet list in the zip, without opening the workbook."""
    p = Path(p)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        return False
    try:
        with zipfile.ZipFile(p) as z:
            xml = z.read("xl/workbook.xml")
    except Exception:
        return False
    return b'name="_run"' in xml and b'name="Values"' in xml


def _review_meta(p: Path) -> Optional[Dict[str, Any]]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if "_run" not in wb.sheetnames:
            return None
        for row in wb["_run"].iter_rows(min_row=1, max_row=1, max_col=1, values_only=True):
            return json.loads(row[0])
    except Exception:
        return None
    finally:
        wb.close()
    return None


def _reviews_of(csv_path: Path) -> List[Path]:
    """The review workbooks beside a parent CSV (where the Ledger writes
    them), newest first. Which CSV each belongs to is checked when read."""
    try:
        cands = [x for x in csv_path.parent.iterdir()
                 if x.is_file() and is_review_workbook(x)]
    except OSError:
        return []
    return sorted(cands, key=lambda x: (-x.stat().st_mtime, x.name))


def ledger_fingerprint(csv_path: Path) -> bytes:
    """Everything a Ledger's reading depends on besides the CSV's own bytes:
    its schema, and its review workbooks (name, size, time). The graph mixes
    this into the CSV's content hash, so a changed schema or a newly saved
    review is read again although the CSV did not change."""
    out = bytearray()
    try:
        out += schema_path(csv_path).read_bytes()
    except OSError:
        out += b"no schema"
    # the labels Teach learned are an inspector's: a column of this Ledger
    # reads as a person by them (read_ledger). Only when some are: a Ledger
    # with none keeps the reading it has. Labels taught as no one's are not
    # read here (read_ledger never lets them undo a template's own field).
    taught = _se().load_taught(csv_path)
    if taught.who:
        out += ("\0who\0" + json.dumps(sorted(taught.who))).encode("utf-8")
    for x in _reviews_of(csv_path):
        st = x.stat()
        out += f"\0{x.name}\0{st.st_size}\0{st.st_mtime_ns}".encode("utf-8")
    return bytes(out)


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", "replace")


#: (the review's sources as the CSV's source_file says them - or one
#: workbook of them, part, item, field, value) -> the cells found for it,
#: newest review first: (0 for the cell the review shows, 1 for one its
#: note says the value is 'also at'; when the review was made; the cell,
#: 'nde.xlsx › Part 101!C4')
Provenance = Dict[Tuple[str, str, str, str, str], List[Tuple[int, str, str]]]


def _applied_as(meta: Dict[str, Any]) -> str:
    """What apply_review wrote in the CSV's source_file for this review's
    rows until 2026-10-20: its sources' names, joined, cut at 200
    characters. Since then each row names its own workbooks."""
    return ", ".join(Path(str(s)).name for s in meta.get("sources", []))[:200]


def _provenance(csv_path: Path) -> Tuple[Provenance, int, List[str]]:
    """Where each value a person ACCEPTED (Accept = Y) in this Ledger's
    review workbooks came from, keyed by the review's sources as the CSV's
    source_file column names them - so a row cites only the review that was
    APPLIED to it, never a newer one saved and not applied (review,
    2026-10-16: a resent workbook was cited for a value it did not give).
    Keyed by one workbook too (a row now names its own), a newer review of
    a workbook of the same name is found as well: read_ledger.where ranks a
    review made after the row was written last.
    A review belongs to this CSV by its resolved path (by name only when that
    path is gone: the folder was moved). Only a place for a value already in
    the CSV is looked up here: nothing in a review is read as a fact."""
    index: Provenance = {}
    notes: List[str] = []
    reviews = _reviews_of(csv_path)
    if not reviews:
        return index, 0, notes
    try:
        import openpyxl
    except ImportError:
        return index, 0, ["review workbooks were not read (they need the openpyxl package, "
                          "which is not installed): no workbook cells for these values"]
    try:
        me = csv_path.resolve()
    except OSError:
        me = csv_path
    n = 0
    for x in reviews:
        meta = _review_meta(x)
        if not meta:
            continue
        theirs = Path(str(meta.get("csv", "")))
        if theirs.exists():
            try:
                if theirs.resolve() != me:
                    continue
            except OSError:
                continue
        elif theirs.name.lower() != csv_path.name.lower():
            continue
        applied = _applied_as(meta)
        try:
            wb = openpyxl.load_workbook(x, read_only=True, data_only=True)
        except Exception as exc:                      # held open, damaged
            notes.append(f"review {x.name} could not be read ({exc.__class__.__name__})")
            continue
        n += 1
        try:
            if "Values" not in wb.sheetnames:
                continue
            made = _made(wb)

            def put(key, rank: int, cell: str) -> None:
                got = index.setdefault(key, [])
                if (rank, made, cell) not in got:
                    got.append((rank, made, cell))
            for row in wb["Values"].iter_rows(min_row=2, values_only=True):
                if len(row) < 9 or str(row[0] or "").strip().upper() != "Y" or not row[8]:
                    continue
                rest = (_text(row[1]).lower(), _text(row[2]), str(row[3] or ""),
                        _text(row[4]))
                put((applied,) + rest, 0, str(row[8]))
                # Apply now names each row's own workbooks (source_extract.
                # rows_from), not the whole run's (review, 2026-10-20): the
                # value is also found by the workbook it was at - its cell,
                # or a cell its note says it is 'also at' (ranked after the
                # cell the review shows: read_ledger.where)
                note = str(row[10] or "") if len(row) > 10 else ""
                for rank, w in enumerate([str(row[8])] + [p[8:] for p in note.split("; ")
                                                          if p.startswith("also at ")]):
                    put((w.split(" › ", 1)[0],) + rest, min(rank, 1), w)
        finally:
            wb.close()
    return index, n, notes


def _made(wb) -> str:
    """When a review workbook was made ('Made: 2026-10-20T06:43:12' on its
    Read me sheet, as the CSV's updated_at is written), or ''."""
    if "Read me" not in wb.sheetnames:
        return ""
    for row in wb["Read me"].iter_rows(min_row=1, max_row=40, max_col=1, values_only=True):
        t = str(row[0] or "")
        if t.startswith("Made: "):
            return t[6:].strip()
    return ""


def load_schema(csv_path: Path):
    from sample_ledger import TemplateSchema
    sp = schema_path(csv_path)
    try:
        return TemplateSchema.load(sp)
    except Exception as exc:
        raise Unreadable(f"its Ledger schema {sp.name} could not be read "
                         f"({exc.__class__.__name__}), so the Ledger was not read",
                         permanent=True) from exc


def read_ledger(p: Path, rel: str, rules: Sequence["kgm.FieldRule"]) -> Reading:
    """One Record per parent-CSV row (kind 'ledger'), read by the schema's
    column roles. Streams the CSV: a ledger of 100k rows is never held as
    rows of dicts."""
    import sample_ledger as sl
    p = Path(p)
    schema = load_schema(p)
    by_norm = kgm._rule_index(rules)
    fields = {f.name: f for f in schema.fields}
    # the labels this Ledger's Teach learned are an inspector's ('Tech Sig'):
    # its person columns follow them (user, 2026-10-20). Never the labels
    # taught as no one's: they are a workbook's labels, and one answer made
    # the template's own 'Inspector Initials' fields no one's (review,
    # 2026-10-20)
    taught = _se().Taught(who=list(_se().load_taught(p).who))
    prov, n_reviews, notes = _provenance(p)
    sink = _Sink()
    try:
        text = _decode(p.read_bytes())
    except OSError as exc:
        raise Unreadable(f"could not be opened ({exc.__class__.__name__})") from exc
    reader = csv.reader(io.StringIO(text, newline=""))
    header: Optional[List[str]] = None
    roles: Dict[int, Tuple[str, Any]] = {}
    id_i = item_i = src_i = upd_i = None
    part_rule = kgm.LEDGER_PART
    parts, people = set(), set()
    n_rows = skipped = located = 0
    person_cols: List[str] = []
    rule_cols: List[str] = []

    def plain(name: str) -> str:
        f = fields.get(name)
        return name.split(": ", 1)[1] if f is not None and f.section and \
            name.startswith(f.section + ": ") else name

    for n, row in enumerate(reader, start=1):
        if header is None:
            if not any(c.strip() for c in row):
                continue
            header = [c.strip() for c in row]
            for i, col in enumerate(header):
                f = fields.get(col)
                if col == schema.id_field:
                    id_i = i
                    r = by_norm.get(_norm(f.label if f else col)) or by_norm.get(_norm(col))
                    if r is not None and r.type != "PART":
                        part_rule = r             # the user said what this ID is
                elif schema.item_field and col == schema.item_field:
                    item_i = i
                elif col == sl.SOURCE_COL:
                    src_i = i
                elif col == sl.UPDATED_COL:
                    upd_i = i
                elif col in sl.META_COLS:
                    continue
                elif schema.inspector_base(col):
                    who = sl.INSPECTOR_COL.match(col).group("who")
                    roles[i] = ("reading", who)
                    person_cols.append(col)
                elif col.endswith(sl.INITIALS_SUFFIX) and col in fields:
                    roles[i] = ("initials", col[:-len(sl.INITIALS_SUFFIX)])
                    person_cols.append(col)
                else:
                    r = (by_norm.get(_norm(f.label)) if f else None) or \
                        by_norm.get(_norm(plain(col)))
                    if r is not None:
                        roles[i] = ("rule", r)
                        rule_cols.append(col)
                    elif f is not None and _who_head(_label_text(f.label), taught):
                        roles[i] = ("person", None)
                        person_cols.append(col)
            if id_i is None:
                raise Unreadable(f"has no '{schema.id_field}' column although its "
                                 "Ledger schema keys on it", permanent=True)
            continue
        if not any(c.strip() for c in row):
            continue

        def cell(i):
            return row[i].strip() if i is not None and i < len(row) else ""
        part = cell(id_i)
        if not part:
            skipped += 1
            continue
        n_rows += 1
        parts.add(part)
        item = cell(item_i)
        sources = cell(src_i)
        updated = cell(upd_i)
        base: Dict[str, Any] = {"row": n}
        if item:
            base["item"] = f"{schema.item_field} {item}"
        where_part = f"{schema.id_field} {part}" + (f", {schema.item_field} {item}"
                                                    if item else "")
        rec = kgm.Record(rel, "ledger")
        rec.values.append(kgm.Value(part_rule, part,
                                    {**base, "column": schema.id_field}, where_part))
        who_vals: Dict[str, Dict[str, Any]] = {}

        def where(field: str, value: str) -> str:
            # a part-level value sits on every flaw row of its part, but the
            # review lists it once, under no flaw. The row's sources as one
            # review's whole run first (an older Apply wrote that), then each
            # workbook it names (_provenance). Of the cells found, one from a
            # review made before this row was written (the one applied to
            # it, not a newer one saved and not applied) first, then the
            # cell the review shows before one it is 'also at' - which was
            # cited when its workbook came first in source_file (review,
            # 2026-10-20)
            pl = part.lower()
            hits = [h for s in [sources] + [b for b in sources.split(", ") if b and b != sources]
                    for k in ((s, pl, item, field, value), (s, pl, "", field, value))
                    for h in prov.get(k, ())]
            if not hits:
                return ""
            return min(hits, key=lambda h: (bool(updated and h[1] and h[1] > updated),
                                            h[0]))[2]

        def add_person(who: str, col: str, shown: str, src_key: Tuple[str, str]):
            nonlocal located
            key = initials_key(who) if is_initials(who, loose=True) else who
            slot = who_vals.setdefault(key, {"raw": key if is_initials(who, loose=True) else who,
                                             "cols": [], "quotes": [], "from": []})
            if col not in slot["cols"]:
                slot["cols"].append(col)
                slot["quotes"].append(shown)
            w = where(*src_key)
            if w and w not in slot["from"]:
                slot["from"].append(w)
                located += 1

        for i, (role, x) in roles.items():
            v = cell(i)
            if not v:
                continue
            col = header[i]
            if role == "reading":
                add_person(x, col, f"{col} = {v}", (col, v))
            elif role == "initials":
                for who in split_people(v):
                    # each inspector's own reading, '<field> [MK]', says where
                    # MK's value came from; the template's one cell holds only
                    # the first agreeing inspector's (review, 2026-10-16)
                    own = sl.inspector_column(x, initials_key(who) if
                                              is_initials(who, loose=True) else who)
                    src_key = (own, cell(header.index(own))) if own in header and \
                        cell(header.index(own)) else (x, cell(header.index(x))
                                                      if x in header else "")
                    add_person(who, col, f"{col} = {v}", src_key)
            elif role == "person":
                for who in split_people(v):
                    add_person(who, col, f"{col} = {v}", (col, v))
            else:
                r = x
                import field_search as fs
                for one in (fs._split_values(_unpadded(v), r.label, r.type) or [v]):
                    loc = {**base, "column": col}
                    w = where(col, v)
                    if w:
                        loc["from"] = w
                    elif sources:
                        loc["sources"] = sources
                    rec.values.append(kgm.Value(r, one, loc, f"{col} = {v}",
                                                signed=r.type == "PERSON"))
        for slot in who_vals.values():
            loc = {**base, "columns": slot["cols"]}
            if slot["from"]:
                loc["from"] = "; ".join(slot["from"])
            elif sources:
                loc["sources"] = sources
            people.add(slot["raw"])
            rec.values.append(kgm.Value(kgm.INSPECTOR, slot["raw"], loc,
                                        "; ".join(slot["quotes"])))
        sink.add(rec)
    if header is None:
        raise Unreadable("is empty", permanent=True)
    if skipped:
        notes.append(f"{skipped} row(s) have no {schema.id_field} and were skipped")
    summary = {"kind": "ledger", "rows": n_rows, "parts": len(parts),
               "people": sorted(people), "person_columns": person_cols,
               "rule_columns": rule_cols, "reviews": n_reviews, "located": located,
               "id_field": schema.id_field, "item_field": schema.item_field,
               "skipped": skipped, "schema": schema_path(p).name}
    return sink.reading(labels=list(header), summary=summary, notes=notes)


# ── workbooks: the grid reader ──────────────────────────────────────────
@dataclass
class _Table:
    header_row: int
    c0: int
    c1: int
    cols: Dict[int, Tuple[str, "kgm.FieldRule"]]
    tentative: bool
    marker: Optional[Tuple[str, Dict[str, Any], str]] = None
    rows: int = 0


class _GridReader:
    """One sheet — or a text box, chart or picture read as a table — fed one
    row at a time, in row order, with the next row as a look-ahead. That is
    all it needs, so a streamed 100k-row sheet is read the same way as a
    loaded one, holding two rows."""

    def __init__(self, name: str, by_norm: Dict[str, "kgm.FieldRule"], rel: str, *,
                 tentative: bool = False, image: str = "", text_all: bool = False,
                 sheet: str = "", budget: Optional[Dict[str, Any]] = None,
                 sure: Optional[Dict[Tuple[int, int], int]] = None):
        #: ``name`` heads this grid's text lines ('Visual › picture F2');
        #: ``sheet`` is what a citation names ('Visual', with the picture).
        self.name, self.by_norm, self.rel = name, by_norm, rel
        self.sheet = sheet or name
        self.tentative, self.image, self.text_all = tentative, image, text_all
        #: A picture read by OCR: its values are marked from_image (the graph
        #: trusts them only when confirmed) and carry how sure the OCR was.
        self.from_image = bool(image)
        self.sure = sure
        self.kind = "picture" if image else "row"
        self.sink: Optional[_Sink] = None       # set by read_workbook
        self.records: List[kgm.Record] = []
        self.n_records = 0
        self.forms: List[kgm.Value] = []
        self.markers: List[Tuple[int, int, str, Dict[str, Any], str]] = []
        self.lines: List[Tuple[Optional[str], int, str]] = []
        self.labels: Dict[str, str] = {}
        self.tables: List[_Table] = []
        self.n_tables = 0
        self.n_forms = 0
        self.taken: set = set()          # cells a form above already used
        #: The free text the whole workbook may keep, shared by its sheets:
        #: lines and characters kept so far, and why any was dropped.
        self.budget = budget if budget is not None else {"lines": 0, "chars": 0,
                                                         "dropped": ""}
        self._last_r: Optional[int] = None
        self._pending: Optional[Tuple[int, Dict[int, Any]]] = None
        self._person_cols: Dict[int, Tuple[str, Dict[str, Any]]] = {}

    # -- cells --
    def _loc(self, r: int, c: int, label: str = "", header: int = 0,
             title: Optional[Tuple[int, int]] = None) -> Dict[str, Any]:
        """Where a value is. A picture's value carries how sure the OCR was
        of it AND of the title or label it hangs on (``title``): a column
        title read at 12% decided which person a link was about (review,
        2026-10-18)."""
        loc: Dict[str, Any] = {"sheet": self.sheet}
        if self.image:
            loc["image"] = self.image
            loc["row"] = r
            if self.sure is not None:
                loc["ocr"] = min(int(self.sure.get((r, c), 0)),
                                 int(self.sure.get(title, 100)) if title else 100)
        else:
            loc["cell"] = f"{_col_letter(c)}{r}"
        if label:
            loc["column"] = label
        if header:
            loc["row"] = r
            loc["header"] = header
        return loc

    def _rule(self, v) -> Optional["kgm.FieldRule"]:
        """The rule a label cell IS (exact label, as a table header must be),
        or the Ledger's 'who' label, or None."""
        if not isinstance(v, str) or self._merged(v):
            return None
        r = self.by_norm.get(_norm(v))
        # A ticked offered rule ('Inspector') does not take a 'who' label
        # from the Ledger's own reading: the initials still go through the
        # people list and the links stay suggested.
        if r is not None and not (r in kgm.PROPOSED_FIELD_RULES and _who_head(v)):
            return r
        return kgm.INSPECTOR if _who_head(v) else None

    @staticmethod
    def _merged(v) -> bool:
        return type(v).__name__ == "_Merged"

    def _labelish(self, v) -> bool:
        """A label, not a value: a rule's label, a 'who' label, 'Label: value'
        or 'Label:'."""
        if not isinstance(v, str):
            return False
        if self._rule(v) is not None:
            return True
        s = v.strip()
        if s.endswith(":") and len(s.split()) <= 5:
            return True
        lab, val = _se()._split_inline(s)
        return bool(val) and len(lab.split()) <= 5 and (self.by_norm.get(_norm(lab)) is not None
                                                        or _who_head(lab))

    # -- feeding --
    def feed(self, r: int, cells: Dict[int, Any]) -> None:
        if self._pending is not None:
            self._row(*self._pending, cells if r == self._pending[0] + 1 else {})
        self._pending = (r, cells)

    def close(self) -> None:
        if self._pending is not None:
            self._row(*self._pending, {})
            self._pending = None
        self.tables = []

    # -- one row --
    def _runs(self, cells: Dict[int, Any], skip: set) -> List[List[int]]:
        runs: List[List[int]] = []
        for c in sorted(cells):
            if c in skip:
                continue
            if runs and c == runs[-1][-1] + 1:
                runs[-1].append(c)
            else:
                runs.append([c])
        return runs

    def _header(self, r: int, run: List[int], cells, nxt) -> Optional[_Table]:
        """A run of cells that heads a table: all text, column titles that
        are rule (or 'who') labels, data under them in the next row and no
        label under a label.

        Firm: two titled columns side by side ('Project ID | Program Lead').
        Tentative (its links suggested): one titled column in a row of three
        or more titles whose next row fills most of them ('Part Number |
        Description | Qty'), or titles only every other cell ('Project |
        Helios | Program Lead | Dana' is a row of two labelled values as
        often as it is a header) when the next row fills every titled one.
        A label followed by a code ('P/N | PN-1234') is a labelled value."""
        real = [c for c in run if not self._merged(cells[c])]
        if len(real) < 2:
            return None
        cols = {c: (_label_text(cells[c]), self._rule(cells[c])) for c in real}
        cols = {c: (t, rule) for c, (t, rule) in cols.items() if rule is not None}
        if not cols:
            return None
        nexts = {real[i]: real[i + 1] for i in range(len(real) - 1)}
        for c in cols:
            # 'P/N | PN-1234' with nothing under the code is a labelled value;
            # a code-like TITLE ('Q1', 'Lot 3') has its column's data under it
            # (it used to lose the whole table - review, 2026-10-16)
            nc = nexts.get(c)
            nx = cells.get(nc) if nc is not None else None
            if isinstance(nx, str) and nc not in cols and kgm.looks_like_code(nx.strip()) \
                    and nxt.get(nc) is None:
                return None
        hits = 0
        for c in cols:
            x = nxt.get(c)
            if x is None or (isinstance(x, str) and not x.strip()):
                continue
            if self._labelish(x):
                return None
            hits += 1
        filled = sum(1 for c in real if nxt.get(c) is not None and not self._labelish(nxt[c]))
        adjacent = any(nexts.get(c) in cols for c in cols)
        if len(cols) >= 2 and adjacent and hits:
            tentative = self.tentative
        elif len(real) >= 3 and hits == len(cols) and 2 * filled >= len(real):
            tentative = True
        else:
            return None
        t = _Table(r, run[0], run[-1], cols, tentative)
        if self.markers:                   # 'Inspector: JB' above the table
            above = [m for m in self.markers if m[0] < r]
            if above:
                top = max(m[0] for m in above)
                m = min((m for m in above if m[0] == top), key=lambda m: abs(m[1] - run[0]))
                t.marker = (m[2], m[3], m[4])
        return t

    def _row(self, r: int, cells: Dict[int, Any], nxt: Dict[int, Any]) -> None:
        # A blank row ends every table and column of initials: a second table
        # stacked under the first was read as its rows (review, 2026-10-16).
        # Blank rows are never fed, so it shows as a jump in the row number.
        if self._last_r is not None and r > self._last_r + 1:
            self.tables = []
            self._person_cols = {}
        self._last_r = r
        used: set = {c for c in cells if (r, c) in self.taken}
        values_at: set = set()
        runs = self._runs(cells, used)
        for run in runs:
            self._harvest(run, cells, nxt)
        # new headers first: one overlapping a running table ends it
        cand = []
        i = 0
        while i < len(runs):
            run = runs[i]
            t = self._header(r, run, cells, nxt)
            # A title cell left empty ('Part Number | | Owner') splits a run;
            # the two halves are one header when neither is one alone (two
            # tables one empty column apart each are, and stay apart).
            if t is None and i + 1 < len(runs) and runs[i + 1][0] == run[-1] + 2 \
                    and self._header(r, runs[i + 1], cells, nxt) is None:
                t = self._header(r, run + runs[i + 1], cells, nxt)
                if t is not None:
                    i += 1
            if t is not None:
                cand.append(t)
            i += 1
        still = []
        for t in self.tables:
            if any(not (n.c1 < t.c0 or n.c0 > t.c1) for n in cand):
                continue
            if any(c in cells and c not in used for c in range(t.c0, t.c1 + 1)):
                values_at |= self._table_row(t, r, cells)
                used.update(c for c in range(t.c0, t.c1 + 1) if c in cells)
                still.append(t)
        self.tables = still
        for t in cand:
            self.tables.append(t)
            self.n_tables += 1
            used.update(range(t.c0, t.c1 + 1))
            for c in range(t.c0, t.c1 + 1):
                if isinstance(cells.get(c), str) and not self._merged(cells[c]):
                    self._label(cells[c])
        values_at |= self._forms(r, cells, used, nxt)
        self._free_text(r, cells, values_at)

    def _harvest(self, run: List[int], cells, nxt) -> None:
        """Every title of a header-shaped row (two or more text cells with
        values under most of them), so Fields… can suggest a drifted label
        ('POC') that no rule reads yet: the old reader harvested row 1's
        titles, this one only rule labels (review, 2026-10-16)."""
        texts = [c for c in run if isinstance(cells[c], str) and not self._merged(cells[c])]
        if len(texts) < 2:
            return
        under = sum(1 for c in texts if nxt.get(c) is not None)
        if 2 * under >= len(texts):
            for c in texts:
                if not any(ch.isdigit() for ch in cells[c]):
                    self._label(cells[c])

    def _table_row(self, t: _Table, r: int, cells) -> set:
        rec = kgm.Record(self.rel, self.kind)
        at = set()
        has_person = False
        for c, (label, rule) in t.cols.items():
            v = cells.get(c)
            if v is None:
                continue
            raw = _text(v)
            if not raw or self._labelish(v):
                continue
            loc = self._loc(r, c, label, header=t.header_row, title=(t.header_row, c))
            tent = t.tentative or rule is kgm.INSPECTOR
            # Under a firm header the column says what a cell is (as a CSV's
            # does); a tentative reading takes only values that read as one.
            ones = (split_people(raw) if rule is kgm.INSPECTOR else _split(raw, rule))
            if t.tentative:
                ones = [x for x in ones if plausible(rule, x)]
            if ones:
                at.add(c)
            has_person = has_person or (rule.type == "PERSON" and bool(ones))
            for one in ones:
                rec.values.append(kgm.Value(rule, one, loc, f"{label}: {raw}", tent,
                                            signed=rule.type == "PERSON",
                                            from_image=self.from_image))
        if t.marker and not has_person and any(v.rule.type == "PART" for v in rec.values):
            who, mloc, msnip = t.marker
            # a marker read from a picture is picture data like any other
            rec.values.append(kgm.Value(kgm.INSPECTOR, who, mloc, msnip, True,
                                        from_image=self.from_image))
        if rec.values and not any(v.rule.type != "PERSON" for v in rec.values):
            # People only (a chart's 'Inspector' column): they signed for
            # what the sheet is about, like 'Inspector: JB' above it.
            for v in rec.values:
                if v.rule is kgm.INSPECTOR:
                    self.markers.append((r, 0, v.raw, v.locator, v.snippet))
            return at
        if rec.values:
            t.rows += 1
            self.n_records += 1
            if self.sink is not None:
                self.sink.add(rec)
            else:
                self.records.append(rec)
        return at

    def _forms(self, r: int, cells, used: set, nxt) -> set:
        """Labelled cells outside the tables: 'Label: value' in one cell; a
        label with its value to the right (up to two cells over); a label
        with its value below. The Ledger's 'who' labels too ('Operator: JB',
        'Inspector Initials' with JB beside or under it), as markers."""
        at = set()
        for c in sorted(cells):
            if c in used:
                continue
            v = cells[c]
            if not isinstance(v, str) or self._merged(v):
                continue
            who = _who_inline(v)
            if who:
                for one in who:
                    self._marker(r, c, one, v.strip())
                at.add(c)
                continue
            lab, inline = _se()._split_inline(v)
            rule = self._rule(lab) if inline else self._rule(v)
            if inline and len(lab.split()) <= 5:
                self._label(lab)                   # for Fields…, rule or not
            elif not inline and v.strip().endswith(":"):
                self._label(v)
            if rule is None:
                continue
            label = _label_text(lab if inline else v)
            self._label(label)
            if inline:
                if plausible(rule, inline):
                    if rule is kgm.INSPECTOR:
                        for one in split_people(inline):
                            self._marker(r, c, one, v.strip())
                    else:
                        self._form(rule, inline, self._loc(r, c, label), v.strip())
                    at.add(c)
                continue
            # Its value: beside it (one or two cells over), else under it -
            # the first that reads as a value of its kind (a part number has
            # a digit, a person is a name or initials): 'Inspector Initials |
            # Date' over 'JB' is JB, not a person called Date.
            got = None
            for dc in (1, 2):
                x = cells.get(c + dc)
                if x is None:
                    continue
                if not self._merged(x) and not self._labelish(x) and c + dc not in used \
                        and plausible(rule, _text(x)):
                    got = (r, c + dc, x)
                break
            if got is None:
                x = nxt.get(c)
                if x is not None and not self._merged(x) and not self._labelish(x) \
                        and plausible(rule, _text(x)):
                    got = (r + 1, c, x)
                    self.taken.add((r + 1, c))
            if got is None:
                continue
            gr, gc, x = got
            raw = _text(x)
            if rule is kgm.INSPECTOR:
                for one in split_people(raw):
                    self._marker(gr, gc, one, f"{label}: {raw}")
                if gr == r + 1:             # a column of initials under its label
                    self._person_cols[c] = (label, self._loc(r, c, label))
            else:
                self._form(rule, raw, self._loc(gr, gc, label, title=(r, c)),
                           f"{label}: {raw}", tentative=gr != r)
            at.add(c)
            if gr == r:
                at.add(gc)
                used.add(gc)
        # the rest of a column of initials ('Inspector Initials' over JB, MK)
        for c, (label, _loc0) in list(self._person_cols.items()):
            if c in used or c in at:
                continue
            x = cells.get(c)
            if x is None:
                if (r, c) not in self.taken:
                    del self._person_cols[c]
                continue
            if (r, c) in self.taken:
                continue                       # the first one: read above
            if isinstance(x, str) and self._labelish(x):
                del self._person_cols[c]
                continue
            people = split_people(_text(x))
            if not people:
                del self._person_cols[c]
                continue
            for one in people:
                self._marker(r, c, one, f"{label}: {_text(x)}")
            at.add(c)
        return at

    def _form(self, rule, raw: str, loc, snippet: str, *, tentative: bool = False) -> None:
        ones = [x for x in _split(raw, rule) if plausible(rule, x)]
        for one in ones:
            self.forms.append(kgm.Value(rule, one, loc, snippet, tentative or self.tentative,
                                        signed=rule.type == "PERSON",
                                        from_image=self.from_image))
        if ones:
            self.n_forms += 1

    def _marker(self, r: int, c: int, who: str, snippet: str) -> None:
        loc = self._loc(r, c, "who")
        loc.pop("column", None)
        self.markers.append((r, c, who, loc, snippet))
        self.n_forms += 1

    def _label(self, text) -> None:
        t = _label_text(text)
        if t and len(t.split()) <= 5 and len(self.labels) < MAX_LABELS:
            self.labels.setdefault(_norm(t), t)

    def _free_text(self, r: int, cells, skip: set) -> None:
        bits = []
        for c in sorted(cells):
            v = cells[c]
            if c in skip or not isinstance(v, str) or self._merged(v):
                continue
            s = " ".join(v.split())
            if self.text_all or len(s.split()) >= LONG_TEXT_WORDS:
                bits.append(s)
        if not bits:
            return
        line = " | ".join(bits)
        b = self.budget
        if b["lines"] >= MAX_LINES:
            b["dropped"] = b["dropped"] or "lines"
        elif b["chars"] + len(line) > MAX_TEXT_CHARS:
            b["dropped"] = b["dropped"] or "chars"
        else:
            self.lines.append((self.name, r, line))
            b["lines"] += 1
            b["chars"] += len(line)

    # -- the sheet's own record --
    def sheet_record(self, sheet_part: str) -> Optional["kgm.Record"]:
        """The labelled cells outside tables, the 'who' markers and a part
        named by the sheet, as one record (like a text document's labelled
        header fields): firm when it names at most one project; what only the
        Ledger's reading says stays suggested."""
        vals = list(self.forms)
        seen = set()
        for _r, _c, who, loc, snip in self.markers:
            key = initials_key(who) if is_initials(who, loose=True) else who
            if key in seen:
                continue
            seen.add(key)
            vals.append(kgm.Value(kgm.INSPECTOR, key if is_initials(who, loose=True) else who,
                                  loc, snip, True, from_image=self.from_image))
        if sheet_part and not any(v.rule.type == "PART" for v in vals) and \
                any(v.rule.type in ("PERSON", "PROJECT") for v in vals):
            vals.append(kgm.Value(kgm.SHEET_PART, sheet_part, {"sheet": self.sheet},
                                  f"sheet '{self.sheet}'", True))
        if not vals:
            return None
        return kgm.Record(self.rel, "sheet", vals)


def _cell_confidence(cells: Dict[Tuple[int, int], Any], confs: Optional[List[List[float]]]
                     ) -> Optional[Dict[Tuple[int, int], int]]:
    """How sure the OCR was of each cell, by position (row r, column c of
    the reader's rows - source_extract.rows_with_confidence; a cell it gave
    no number for counts as 0). None when the reader says nothing. By word
    text, a glued 'PN-7'+'7' took 97 from a sure 'PN-77' elsewhere (review,
    2026-10-18)."""
    if confs is None:
        return None
    out = {}
    for (r, c) in cells:
        try:
            out[(r, c)] = int(confs[r - 1][c - 1])
        except (LookupError, TypeError, ValueError):
            out[(r, c)] = 0
    return out


def sheet_part(name: str) -> str:
    """The part a sheet is named for, the Ledger's way ('Part 101' -> 101,
    '101' -> 101) - but a year is not a part ('2025'), and 'PN-1001' is the
    part PN-1001, not '1001' (review, 2026-10-16)."""
    n = " ".join(str(name or "").split())
    if re.fullmatch(r"(?:19|20)\d\d", n):
        return ""
    if re.fullmatch(r"(?i)p/?n-?\d[\w\-/.]*", n):
        return n.upper()
    return _se()._part_from_name(n)


#: A run of spaces, tabs or line breaks this long is padding, not a gap
#: between values.
_PADDING = re.compile(r"\s{32}")


def _unpadded(raw: str) -> str:
    """``raw`` for field_search's split: as typed, unless it holds padding
    (_PADDING) - then its words, one space apart. The split's '\\s*;\\s*',
    '\\s*,\\s*' and '\\s+and\\s+' rescan a run of spaces from each of its
    characters: one cell of 'x', 32,000 spaces and 'y' held the GIL for 8 s,
    an 'Inspector:' padded cell for 22 s (review, 2026-10-20)."""
    return " ".join(raw.split()) if _PADDING.search(raw) else raw


def _split(raw: str, rule) -> List[str]:
    import field_search as fs
    return [x for x in (fs._split_values(_unpadded(raw), rule.label, rule.type) or [raw])
            if x]


def plausible(rule, raw: str) -> bool:
    """Could this cell be a value of the rule's kind? A workbook puts labels,
    titles and readings side by side, so a value is taken only when it reads
    as one: a person is a name or initials ('Jordan Blake', 'JB', 'Lee,
    Carol', 'D. Whitfield'); a part number has a digit in it; a project is
    any short text. A text document's 'Part: ...' lines are not checked
    here (their separator says which is the value)."""
    s = " ".join(str(raw or "").split())
    if not s or len(s) > 120:
        return False
    if rule.type == "PERSON":
        return person_ok(s) or kgm.is_initial_form(kgm.person_key(s))
    if rule.type == "PART":
        return any(ch.isdigit() for ch in s)
    return len(s.split()) <= 12


# ── workbooks: where the cells come from ────────────────────────────────
def _sheet_xml_bytes(p: Path) -> int:
    try:
        with zipfile.ZipFile(p) as z:
            return sum(i.file_size for i in z.infolist()
                       if i.filename.startswith("xl/worksheets/") and i.filename.endswith(".xml"))
    except Exception:
        return 0


def _grid_rows(cells: Dict[Tuple[int, int], Any]) -> Iterator[Tuple[int, Dict[int, Any]]]:
    by_row: Dict[int, Dict[int, Any]] = {}
    for (r, c), v in cells.items():
        by_row.setdefault(r, {})[c] = v
    for r in sorted(by_row):
        yield r, by_row[r]


def _stream_sheets(p: Path, notes: List[str], data: Optional[bytes] = None
                   ) -> Iterator[Tuple[str, Iterator]]:
    """A big .xlsx / .xlsm one sheet at a time, one row at a time (openpyxl
    read-only: no merged ranges, no pictures), from its bytes when given.
    The workbook is closed when the last sheet is done or the reader stops
    early."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data) if data is not None else p,
                                read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            def rows(ws=ws):
                for r, vals in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True),
                                         start=1):
                    cells = {}
                    for c, v in enumerate(vals, start=1):
                        if v is None or (isinstance(v, str) and not v.strip()):
                            continue
                        if isinstance(v, str) and v.strip().upper() in _se()._EXCEL_ERRORS:
                            continue
                        cells[c] = v
                    if cells:
                        yield r, cells
            yield ws.title, rows()
    finally:
        wb.close()


def missing_reader(p: Path) -> str:
    """Why a workbook cannot be read on this install, or ''."""
    import importlib.util as iu
    suf = Path(p).suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        return "" if iu.find_spec("openpyxl") else "needs the openpyxl package, which is not installed"
    if suf in (".xls", ".xlsb"):
        import excel_open
        if excel_open.excel_available():
            return ""
        if suf == ".xls" and iu.find_spec("xlrd") is not None:
            return ""
        return (f"a {suf} workbook opens only through Microsoft Excel (with the pywin32 "
                "package), which is not available here")
    return ""


def _legacy_grids(p: Path) -> List[Tuple[str, Dict[Tuple[int, int], Any]]]:
    """An .xls through xlrd, when Excel is not installed (the graph read
    .xls that way before the Ledger's reader)."""
    import pandas as pd
    out = []
    for name, fr in pd.read_excel(p, sheet_name=None, header=None, dtype=object).items():
        cells = {}
        for i, row in enumerate(fr.itertuples(index=False, name=None), start=1):
            for j, v in enumerate(row, start=1):
                if v is None or (isinstance(v, float) and v != v):
                    continue
                if isinstance(v, str) and not v.strip():
                    continue
                cells[(i, j)] = v
        out.append((str(name), cells))
    return out


def read_workbook(p: Path, rel: str, rules: Sequence["kgm.FieldRule"],
                  image_reader: Optional[Callable] = None) -> Reading:
    """A workbook through the Ledger's reader, as records (table rows and
    each sheet's labelled cells), free-text lines and the labels it uses."""
    import excel_open
    se = _se()
    p = Path(p)
    by_norm = kgm._rule_index(rules)
    try:
        kind = excel_open.kind_of(p)
    except OSError as exc:
        raise Unreadable(f"could not be opened ({exc.__class__.__name__}: it may be "
                         "open in another program)") from exc
    if kind == "password":
        raise Unreadable("has a password to open, so it was not read (open it in Excel "
                         "and save a copy without the password)", permanent=True)
    suf = p.suffix.lower()
    # Not a zip and not an Excel container: the Ledger hands such a file to
    # Excel (an old XML spreadsheet opens there); with no Excel, it is
    # unreadable for good.
    if kind == "unknown" and not excel_open.excel_available():
        raise Unreadable("is not a workbook openpyxl can read (and Excel, which might "
                         "open it, is not available here)", permanent=True)
    notes: List[str] = []
    via_excel = kind != "zip" or suf in (".xls", ".xlsb")
    summary: Dict[str, Any] = {"kind": "workbook", "sheets": 0, "text_boxes": 0,
                               "charts": 0, "pictures": 0, "pictures_read": 0,
                               "tables": 0, "rows": 0, "forms": 0, "people": [],
                               "image_reader": getattr(image_reader, "name", "") or ""}
    readers: List[Tuple[_GridReader, str]] = []
    budget: Dict[str, Any] = {"lines": 0, "chars": 0, "dropped": ""}
    sink = _Sink()

    def run(name: str, rows: Iterable, image: str = "", parent: str = "",
            sure: Optional[Dict[Tuple[int, int], int]] = None) -> None:
        is_extra = " › " in name and not image
        gr = _GridReader(name, by_norm, rel, image=image,
                         text_all=is_extra and "text box" in name, sheet=parent or name,
                         budget=budget, sure=sure)
        gr.sink = sink
        # the graph's Contents layer: every cell, whatever the rules say
        contents = kgm.contents_now()
        cg = contents.grid(name, loc_sheet=parent or name, image=image) \
            if contents is not None else None
        for r, cells in rows:
            gr.feed(r, cells)
            if cg is not None:
                cg.feed(r, cells)
        gr.close()
        if cg is not None:
            cg.close()
        readers.append((gr, parent or name.split(" › ")[0]))
        if image:
            return
        if " › text box" in name:
            summary["text_boxes"] += 1
        elif " › chart" in name:
            summary["charts"] += 1
        else:
            summary["sheets"] += 1

    images = []
    try:
        size = p.stat().st_size
    except OSError:
        size = 0
    data = None
    if not via_excel and size <= IN_MEMORY_BYTES:
        try:
            data = p.read_bytes()
        except OSError as exc:
            raise Unreadable(f"could not be opened ({exc.__class__.__name__}: it may be "
                             "open in another program)") from exc
    if not via_excel and _sheet_xml_bytes(p) > STREAM_XML_BYTES:
        summary["opened_with"] = "openpyxl, one row at a time (a big workbook)"
        notes.append("a big workbook: read one row at a time, without merged cells "
                     "or pictures")
        try:
            # read while the workbook is open: _stream_sheets closes it
            # after its last sheet
            for name, rows in _stream_sheets(p, notes, data):
                run(name, rows)
        except Exception as exc:
            raise Unreadable(f"could not be opened ({exc.__class__.__name__})",
                             permanent=True) from exc
        try:
            extras = excel_open.zip_extras(io.BytesIO(data) if data is not None else p)
        except Exception:
            extras = []
        for ex in extras:
            run(ex.name, _grid_rows(ex.cells))
    elif via_excel and suf == ".xls" and not excel_open.excel_available():
        summary["opened_with"] = "xlrd (Excel is not installed: no text boxes, charts "\
                                 "or pictures)"
        try:
            grids = _legacy_grids(p)
        except Exception as exc:
            raise Unreadable(f"could not be read ({exc.__class__.__name__})",
                             permanent=True) from exc
        for name, cells in grids:
            run(name, _grid_rows(cells))
    else:
        summary["opened_with"] = "Excel (hidden, read-only)" if via_excel else "openpyxl"
        # openpyxl lists an .xlsx's pictures for free. Excel gives a picture
        # only through the Windows clipboard: it is asked for them only when
        # something can read them and the clipboard is safe to use now (not
        # synced to other devices, holding no picture or files the copy would
        # replace - excel_open.clipboard_safe); otherwise they are listed and
        # counted. (The user, 2026-10-18: read them, trust them only once
        # confirmed - see knowledge_graph's picture rule.)
        copy_ok, copy_why = False, ""
        if via_excel and image_reader is not None:
            copy_ok, copy_why = excel_open.clipboard_safe()
        try:
            grids, images = se.load_workbook_grids(
                p, notes, want_pictures=not via_excel or copy_ok, data=data,
                careful=via_excel and copy_ok)
        except se.LedgerError as exc:
            # openpyxl refusing the bytes is for good, and so is a password;
            # Excel failing otherwise may not be (busy, asking to sign in),
            # so that is tried again next time. Not when Excel said these
            # bytes are no workbook ('the file format or file extension is
            # not valid') or the file is empty (ExcelOpenError.for_good): a
            # hidden Excel was started for it on every rebuild, unchanged or
            # not (review, 2026-10-20).
            raise Unreadable(str(exc), permanent=not via_excel or bool(
                re.search(r"(?i)password", str(exc))) or bool(
                getattr(exc.__cause__, "for_good", False))) from exc
        except Exception as exc:
            if not via_excel:
                raise
            # Excel itself stopped (RPC server unavailable, a COM error):
            # in words, and tried again next time.
            raise Unreadable(f"Excel stopped while it was being read ({exc.__class__.__name__}"
                             f": {str(exc)[:200]})") from exc
        for g in grids:
            run(g.sheet, _grid_rows(g.cells))
        del grids
        # Excel's notes name the file, and the graph names it too
        notes[:] = [n[len(p.name) + 3:] if n.startswith(f"{p.name} › ") else n for n in notes]
    summary["pictures"] = len(images)
    uncopied = [i for i in images if not i.data]
    images = [i for i in images if i.data]
    if image_reader is None and (images or uncopied):
        notes.append(f"{len(images) + len(uncopied)} picture(s) not read: install Tesseract "
                     "(local OCR) to read them")
    elif uncopied:
        # Read again at a later rebuild once the clipboard is safe (the
        # graph looks at this): an unchanged workbook was never read again,
        # so "rebuild again later" did not come true (review, 2026-10-18).
        summary["pictures_pending"] = len(uncopied)
        if copy_why:                  # else excel_open's own note says why
            notes.append(f"{len(uncopied)} picture(s) not read: {copy_why} (rebuild again "
                         "later to read them)")
    prog = kgm.progress_now()
    for i, img in enumerate(images, start=1):
        if prog is not None and image_reader is not None:
            prog.update(f"{rel}: reading picture {i} of {len(images)}", force=i == 1)
        confs: Optional[List[List[float]]] = None
        try:
            if image_reader is None:
                rows = None
            elif callable(getattr(image_reader, "with_confidence", None)):
                rows, confs = image_reader.with_confidence(img.data)
            else:
                rows = image_reader(img.data)
        except Exception as exc:                      # one bad picture is a note
            notes.append(f"the picture at {img.sheet}!{img.anchor} could not be read "
                         f"({exc.__class__.__name__})")
            continue
        if not rows:
            continue
        summary["pictures_read"] += 1
        g = se.rows_to_grid(rows, img)
        run(f"{img.sheet} › picture {img.anchor}", _grid_rows(g.cells), image=img.anchor,
            parent=img.sheet, sure=_cell_confidence(g.cells, confs))
    lines: List[Tuple[Optional[str], int, str]] = []
    labels: Dict[str, str] = {}
    for gr, parent in readers:
        rec = gr.sheet_record(sheet_part(parent))
        if rec is not None:
            sink.add(rec)
        lines += gr.lines
        for k, t in gr.labels.items():
            if len(labels) < MAX_LABELS:
                labels.setdefault(k, t)
        summary["tables"] += gr.n_tables
        summary["forms"] += gr.n_forms
        summary["rows"] += gr.n_records
    summary["people"] = sorted(sink.people)
    if budget["dropped"] == "chars":
        notes.append(f"only the first {MAX_TEXT_CHARS} characters of free text were kept")
    elif budget["dropped"] == "lines":
        notes.append(f"only the first {MAX_LINES} lines of free text were kept")
    return sink.reading(lines=lines, labels=list(labels.values()), summary=summary,
                        notes=notes)


# ── plain words for the Connections tab ─────────────────────────────────
def _n(k: int, one: str, many: str = "") -> str:
    return f"{k} {one if k == 1 else (many or one + 's')}"


def describe(summary: Dict[str, Any]) -> str:
    """One reading's summary as a sentence a person reads."""
    if not summary:
        return ""
    if summary.get("kind") == "ledger":
        s = (f"Ledger (parent CSV, values you accepted): {_n(summary.get('rows', 0), 'row')}, "
             f"{_n(summary.get('parts', 0), 'part')}")
        ppl = summary.get("people") or []
        if ppl:
            s += f", inspectors {', '.join(ppl[:8])}" + (" …" if len(ppl) > 8 else "")
        if summary.get("reviews"):
            s += (f"; {summary.get('located', 0)} value(s) traced to their workbook cell "
                  f"through {_n(summary['reviews'], 'review workbook')}")
        cols = summary.get("rule_columns") or []
        if cols:
            s += f"; also read: {', '.join(cols[:6])}"
        return s + "."
    bits = [_n(summary.get("sheets", 0), "sheet")]
    for key, one in (("text_boxes", "text box"), ("charts", "chart")):
        if summary.get(key):
            bits.append(_n(summary[key], one, "text boxes" if key == "text_boxes" else ""))
    if summary.get("pictures"):
        if summary.get("image_reader"):
            bits.append(f"{summary.get('pictures_read', 0)} of "
                        f"{_n(summary['pictures'], 'picture')} read by "
                        f"{summary['image_reader']}")
        else:
            bits.append(f"{_n(summary['pictures'], 'picture')} not read (no Tesseract)")
    s = f"Workbook ({summary.get('opened_with', 'openpyxl')}): " + ", ".join(bits)
    s += (f"; {_n(summary.get('tables', 0), 'table')}, {_n(summary.get('rows', 0), 'row')}"
          f" and {_n(summary.get('forms', 0), 'labelled cell')} with a person, part or project")
    ppl = summary.get("people") or []
    if ppl:
        s += f"; people: {', '.join(ppl[:8])}" + (" …" if len(ppl) > 8 else "")
    return s + "."


# ── the in-app preview of a cited spot ──────────────────────────────────
def preview(p: Path, locator: Dict[str, Any], quote: str = "") -> List[Tuple[str, str, bool]]:
    """The cited spot of a workbook: the table's header and the cited row
    ('B7  Part Number: PN-1001'), a labelled cell and its neighbours, or a
    text line. Only .xlsx / .xlsm are opened (openpyxl, read-only, streamed
    to the row): an .xls / .xlsb / rights-protected file would need Excel,
    which the preview never starts — it says so instead."""
    p = Path(p)
    sheet = str(locator.get("sheet") or locator.get("page") or "")
    out: List[Tuple[str, str, bool]] = []
    if not locator.get("image") and " › picture " in sheet:
        # a picture's text line: its "page" names the sheet and the picture
        locator = dict(locator, image=sheet.split(" › picture ", 1)[1])
    if locator.get("image"):
        out.append(("", f"a picture at {locator['image']} on sheet "
                        f"{sheet.split(' › ')[0]}, read by OCR", False))
        if quote:
            out.append(("read", quote, True))
        return out
    if p.suffix.lower() not in (".xlsx", ".xlsm") or not is_workbook(p):
        out.append(("", "This workbook opens only through Excel; use Open file to see "
                        "it. The cited text:", False))
        if quote:
            out.append(("", quote, True))
        return out
    import excel_open
    if " › " in sheet:                     # a text box or a chart
        try:
            extras = {ex.name: ex.cells for ex in excel_open.zip_extras(p)}
        except Exception:
            extras = {}
        cells = extras.get(sheet)
        if not cells:
            return []
        hit = int(locator.get("row") or locator.get("line") or 0)
        cell = str(locator.get("cell") or "")
        out.append(("", sheet, False))
        for (r, c), v in sorted(cells.items()):
            ref = f"{_col_letter(c)}{r}"
            out.append((ref, str(v), r == hit or ref == cell))
        return out
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        if sheet not in wb.sheetnames:
            return []
        ws = wb[sheet]
        row = int(locator.get("row") or locator.get("line") or 0)
        cell = str(locator.get("cell") or "")
        cells_hit = set(locator.get("cells") or ([cell] if cell else []))
        if not row and cell:
            m = re.match(r"([A-Z]+)(\d+)$", cell)
            row = int(m.group(2)) if m else 0
        if not row:
            return []
        header = int(locator.get("header") or 0)
        want = sorted({x for x in (header, row - 1, row, row + 1) if x > 0}) \
            if not header else sorted({header, row})
        heads: Dict[int, str] = {}
        got: Dict[int, Tuple[Any, ...]] = {}
        for r, vals in enumerate(ws.iter_rows(min_row=1, max_row=max(want), values_only=True),
                                 start=1):
            if r in want:
                got[r] = vals
        if header and header in got:
            heads = {c: _text(v) for c, v in enumerate(got[header], start=1) if v is not None}
        cols = set(locator.get("columns") or ([locator["column"]] if locator.get("column")
                                              else []))
        out.append(("", f"sheet {sheet}, row {row}", False))
        for r in want:
            if r == header:
                continue
            for c, v in enumerate(got.get(r, ()), start=1):
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                ref = f"{_col_letter(c)}{r}"
                head = heads.get(c, "")
                shown = f"{head}: {_text(v)}" if head else _text(v)
                out.append((ref, shown, ref in cells_hit or (r == row and (
                    head in cols or not cells_hit and not cols))))
        return out
    except Exception:
        return []
    finally:
        wb.close()
