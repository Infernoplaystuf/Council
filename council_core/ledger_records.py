"""
council_core.ledger_records — the Ledger for TABLE templates: each row of the
template's table is one RECORD, and the source workbooks hold many records
each (user, 2026-10-20, the BTB example: "the highlighted stuff should match
their highlights in the sample and template workbook").

The sample-ID Ledger (sample_ledger + source_extract) reads ONE part per
workbook into fixed template cells. This one reads tables:

    template      a header row of fields with empty (often coloured) rows
                  below -> TableTemplate (analyze_template)
    sources       every sheet of every workbook -> the tables in it
                  (find_tables: header rows found by shape, rows down to a
                  blank or a total row, tables written sideways too)
    mapping       each template field -> one column of each table, by
                  header words (exact, the user's taught words, similar
                  words), by the user's HIGHLIGHTS (the same fill colour in
                  the template and a workbook, learned automatically), or
                  by Council's model for headers nothing else placed (asked
                  once per header layout, never pre-accepted)
    other sources a field from the file or sheet name, a LOOKUP in another
                  table (join on the record's name, exact or "starts with"
                  on whole words), a label beside the table ("Inspector:
                  MK" above it), a constant, a template formula
    review        one row per record; amber cells are Council's guesses,
                  red ones clash between workbooks. A "Decisions" sheet
                  takes Y/N once per guessed column or rule, for every
                  record it touches (thousands of rows are not ticked one
                  by one)
    apply         the accepted records into the parent CSV (one row per
                  record), what the decisions taught into the recipe, and
                  optionally the template filled with every record (its
                  table grown as needed, its row styles, validations,
                  conditional formats and formulas carried down)

The recipe (<parent>.records.json) holds what was learned: header words per
field, rules per field, the model's answers per header layout. Reading
thousands of workbooks is cached per file (<parent>.records_cache/, gzip
JSON keyed by path, size and modification time), so a second Extract reads
only what changed.

Nothing a model or a fuzzy match proposes is accepted without the user:
those cells are amber and their records wait for a Y (on the record, or on
the Decisions row). Only exact header matches, the user's highlights and
answers, unique whole-word lookups and confirmed rules are pre-accepted.
No Qt here; the Ledger tab reaches this through council_core.ledger_ingest.
"""
from __future__ import annotations

import colorsys
import difflib
import fnmatch
import gzip
import hashlib
import json
import os
import re
from copy import copy
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from sample_ledger import (FILLED_MARK, SOURCE_COL, UPDATED_COL, LedgerError,
                           TemplateField, TemplateSchema, _atomic_write_text, _clean,
                           _split_unit, _typed, normalize_label, schema_path_for, xl_text)

#: (prompt, json_schema) -> (reply, model label): council_core.ledger_llm.Chat
Chat = Callable[[str, Dict[str, Any]], Tuple[str, str]]

RECIPE_VERSION = 1
#: bump when what a cached file holds changes: older caches are read again
READER_VERSION = 4
#: the parent CSV's columns after the fields (before source_file, updated_at)
SHEET_COL, ROW_COL, TABLE_COL = "source_sheet", "source_row", "source_table"

#: a sheet is read until this many empty rows follow its last value (a
#: formatted cell at row 1,048,576 must not cost a million rows)
EMPTY_ROWS_STOP = 2000
#: cells read from one sheet at most; the rest is said, not read
MAX_SHEET_CELLS = 2_000_000
MAX_COLS = 1000
#: header rows tried per sheet (a text-only table makes every row header-like)
MAX_HEADER_ROWS = 400
#: a header's text at most this long; longer is a note, not a column name
HEADER_CHARS = 120
#: header layouts asked of the model in one Extract at most
MODEL_MAX_LAYOUTS = 40
#: example values per column shown to the model
MODEL_EXAMPLES = 3
#: rows listed per table in the review's Tables sheet at most
TABLES_LISTED = 5000
#: Excel's rows per sheet, less the header
SHEET_ROWS = 1_048_575

#: a cell that holds "no value" in words: data, never a header
NULLISH = {"n a", "na", "none", "null", "nil", "tbd", "nan", "", "x"}
NULLISH_RAW = {"-", "--", "---", "?", "n/a", "#n/a"}
#: a row whose first text is one of these sums the table up: not a record
SUMMARY_WORDS = {"total", "totals", "sum", "average", "avg", "mean", "median", "min",
                 "minimum", "max", "maximum", "std", "stdev", "std dev", "count",
                 "range", "grand total", "subtotal", "sub total"}

#: how a value or a column was placed, and whether that is pre-accepted
SURE_HOWS = {"header", "words", "taught", "colour", "rule", "exact", "starts with",
             "your answer", "label", "formula", "constant", "file name", "sheet name",
             "near"}
AMBER = "FFF2CC"
RED = "F8CBAD"
GREY = "D9D9D9"


# ── small helpers ───────────────────────────────────────────────────────
def norm(text: Any) -> str:
    """'Laser Power (W):' -> 'laser power' (sample_ledger.normalize_label)."""
    if text is None:
        return ""
    return normalize_label(str(text)[:400])


def vnorm(text: Any) -> str:
    """A value for matching: lower case, runs of other characters one space
    ('Woop Woop #7' -> 'woop woop 7')."""
    return re.sub(r"[^a-z0-9]+", " ", str(text if text is not None else "").lower()).strip()


def cell_text(v: Any) -> str:
    """A value as the CSV keeps it: 12.0 -> '12', dates ISO, None -> ''."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v != v:                                       # NaN
            return ""
        return str(int(v)) if v.is_integer() and abs(v) < 1e16 else f"{v:.15g}"
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="seconds").replace(" 00:00:00", "")
    if isinstance(v, date):
        return v.isoformat()
    s = xl_text(str(v))
    return s.strip() if isinstance(s, str) else str(s)


_NUMBER = re.compile(r"[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?\s*%?")
_LETTER = re.compile(r"[A-Za-z]")
_NOT_ALNUM = re.compile(r"[^a-z0-9]+")


@lru_cache(maxsize=200_000)
def _str_kind(s: str) -> int:
    """0 nothing (blank, 'N/A'), 1 a number, 2 words that could name a
    column, 3 other text. Cached: the same values come back in every row."""
    t = s.strip()
    if not t or t.lower() in NULLISH_RAW or _NOT_ALNUM.sub(" ", t.lower()).strip() in NULLISH:
        return 0
    if _NUMBER.fullmatch(t.replace(",", "")):
        return 1
    return 2 if len(t) <= HEADER_CHARS and _LETTER.search(t) else 3


def is_number(v: Any) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    return isinstance(v, str) and _str_kind(v[:200]) == 1


def is_nullish(v: Any) -> bool:
    if v is None:
        return True
    return isinstance(v, str) and _str_kind(v[:200]) == 0


def is_text(v: Any) -> bool:
    """Words that could name a column: a letter in it, not 'N/A', not a
    number, not too long."""
    return isinstance(v, str) and len(v) <= 2 * HEADER_CHARS and _str_kind(v) == 2


def _col_letter(c: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(max(1, int(c)))


def _a1(r: int, c: int) -> str:
    return f"{_col_letter(c)}{r}"


def _sha(text: str, n: int = 12) -> str:
    return hashlib.sha1(text.encode("utf-8", "replace")).hexdigest()[:n]


class _Quiet:
    """A Progress stand-in when the caller gives none."""

    def stage(self, *_a, **_k) -> None:
        pass

    def update(self, *_a, **_k) -> None:
        pass

    def advance(self, *_a, **_k) -> None:
        pass

    def note(self, *_a, **_k) -> None:
        pass


# ── colours ─────────────────────────────────────────────────────────────
#: Excel's theme colour order as cells number them: lt1, dk1, lt2, dk2,
#: accent1..6, hlink, folHlink (the theme XML lists dk1 first)
_THEME_ORDER = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
                "accent5", "accent6", "hlink", "folHlink")


def theme_colours(wb) -> List[str]:
    """The workbook's theme colours as 'RRGGBB', in the order cells number
    them; [] when it has no theme."""
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return []
    try:
        text = raw.decode("utf-8", "replace") if isinstance(raw, bytes) else str(raw)
        scheme = re.search(r"<a:clrScheme\b.*?</a:clrScheme>", text, re.S)
        if not scheme:
            return []
        found: Dict[str, str] = {}
        for name, body in re.findall(r"<a:(\w+)>(.*?)</a:\1>", scheme.group(0), re.S):
            m = re.search(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', body)
            if m:
                found[name] = m.group(1).upper()
        return [found.get(k, "") for k in _THEME_ORDER]
    except Exception:                                       # noqa: BLE001
        return []


def _tinted(rgb: str, tint: float) -> str:
    if not tint:
        return rgb
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def colour_key(fill, theme: Sequence[str] = ()) -> str:
    """A cell's fill as 'RRGGBB' ('' when it has none, or white). Theme
    colours are resolved through the workbook's theme with their tint."""
    try:
        if fill is None or getattr(fill, "fill_type", None) in (None, "none"):
            return ""
        if fill.fill_type != "solid":
            c = fill.fgColor if getattr(fill.fgColor, "type", "") else fill.bgColor
        else:
            c = fill.fgColor
        kind = getattr(c, "type", "")
        rgb = ""
        if kind == "rgb" and isinstance(c.rgb, str):
            rgb = c.rgb[-6:].upper()
        elif kind == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            i = int(c.indexed)
            if 0 <= i < len(COLOR_INDEX) and i not in (64, 65):
                rgb = str(COLOR_INDEX[i])[-6:].upper()
        elif kind == "theme":
            i = int(c.theme)
            base = theme[i] if 0 <= i < len(theme) else ""
            rgb = _tinted(base, float(c.tint or 0)) if base else f"T{i}{float(c.tint or 0):+.2f}"
        if not rgb or rgb == "FFFFFF":
            return ""
        return rgb
    except Exception:                                       # noqa: BLE001
        return ""


def same_colour(a: str, b: str, tolerance: float = 48.0) -> bool:
    if not a or not b:
        return False
    if a == b:
        return True
    if not (re.fullmatch(r"[0-9A-F]{6}", a) and re.fullmatch(r"[0-9A-F]{6}", b)):
        return False
    da = [int(a[i:i + 2], 16) - int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return sum(x * x for x in da) ** 0.5 <= tolerance


# ── the template ────────────────────────────────────────────────────────
@dataclass
class RField:
    """One column of the template's table (or a labelled cell beside it)."""
    name: str                     # the parent CSV's column
    label: str                    # as written in the template
    col: int                      # template column (a form field: its value cell's)
    key: str                      # norm(label)
    keys: List[str] = field(default_factory=list)   # every way it may be written
    colour: str = ""
    unit: str = ""
    formula: str = ""             # the template computes it: never read
    form_row: int = 0             # a labelled cell outside the table: its row

    @property
    def is_form(self) -> bool:
        return self.form_row > 0


@dataclass
class TableTemplate:
    path: str
    sheet: str
    header_row: int
    first_row: int                # the first record row
    end_row: int                  # the last row the template laid out for records
    fields: List[RField]
    content_below: bool = False   # something sits under the table: rows are inserted
    forms: List[RField] = field(default_factory=list)

    @property
    def styled_rows(self) -> int:
        return max(0, self.end_row - self.first_row + 1)

    def field(self, name: str) -> RField:
        for f in self.fields + self.forms:
            if f.name == name:
                return f
        raise KeyError(name)

    @property
    def all_fields(self) -> List[RField]:
        return self.fields + self.forms

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "TableTemplate":
        d = dict(d)
        d["fields"] = [RField(**f) for f in d.get("fields", [])]
        d["forms"] = [RField(**f) for f in d.get("forms", [])]
        return cls(**d)


def _styled(cell) -> bool:
    try:
        if colour_key(cell.fill):
            return True
        b = cell.border
        return any(getattr(getattr(b, s, None), "style", None)
                   for s in ("left", "right", "top", "bottom"))
    except Exception:                                       # noqa: BLE001
        return False


def analyze_template(path) -> TableTemplate:
    """The record table a template lays out: the widest run of column
    headers with rows for records under it (empty, formatted, formulas or one
    example row), on any sheet. Labelled cells outside it ('Inspector:' with
    an empty cell beside it) are form fields: one value per source sheet.
    Raises LedgerError when the template has no such table."""
    import openpyxl
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:                                  # noqa: BLE001
        raise LedgerError(f"could not open the template {path.name}: {e}") from e
    theme = theme_colours(wb)
    best: Optional[Tuple[float, TableTemplate]] = None
    for ws in wb.worksheets:
        cells = {k: c for k, c in list(ws._cells.items())[:MAX_SHEET_CELLS]}
        values = {k: c.value for k, c in cells.items() if c.value not in (None, "")}
        last_row = max((r for r, _ in values), default=0)
        by_row: Dict[int, List[int]] = {}
        for (r, c), v in values.items():
            if is_text(v) and not str(v).startswith("="):
                by_row.setdefault(r, []).append(c)
        for r, cols in sorted(by_row.items())[:MAX_HEADER_ROWS]:
            for run in _runs(sorted(cols)):
                if len(run) < 2:
                    continue
                tt = _template_table(ws, path, r, run, cells, values, theme, last_row)
                if tt is None:
                    continue
                score = len(tt.fields) * (1 + min(tt.styled_rows, 60) / 10.0)
                if best is None or score > best[0]:
                    best = (score, tt)
    if best is None:
        raise LedgerError(
            f"{path.name} has no table to fill: a row of column headers (two or more) "
            "with empty rows for the records under it")
    tt = best[1]
    tt.forms = _form_fields(wb[tt.sheet], tt, theme)
    _unique_names(tt)
    return tt


def _runs(cols: List[int]) -> List[List[int]]:
    out: List[List[int]] = []
    for c in cols:
        if out and c == out[-1][-1] + 1:
            out[-1].append(c)
        else:
            out.append([c])
    return out


def _template_table(ws, path: Path, r: int, run: List[int], cells, values, theme,
                    last_row: int) -> Optional[TableTemplate]:
    first = r + 1
    rr = first
    example_ok = True
    while rr <= max(last_row, r) + 1 and rr - first < 100_000:
        vals = [values.get((rr, c)) for c in run]
        typed = [v for v in vals if v not in (None, "") and not str(v).startswith("=")]
        if typed:
            # one example row under the headers is written over; anything else
            # ends the table
            if example_ok and rr == first and len(typed) >= max(1, len(run) // 2):
                rr += 1
                example_ok = False
                continue
            break
        if rr > last_row and not any(_styled(cells[(rr, c)]) for c in run
                                     if (rr, c) in cells):
            break
        rr += 1
    end = rr - 1
    styled_end = first - 1
    for x in range(first, end + 1):
        if any((x, c) in cells and (_styled(cells[(x, c)]) or values.get((x, c)))
               for c in run):
            styled_end = x
    end = styled_end
    below = any(rw > end for (rw, _c) in values if rw > r)
    parent = {c: values.get((r - 1, c)) for c in run}
    fields: List[RField] = []
    for c in run:
        label = str(values[(r, c)]).strip()
        name, unit = _split_unit(_clean(label))
        hdr = cells.get((r, c))
        body = cells.get((first, c))
        colour = colour_key(hdr.fill, theme) if hdr is not None else ""
        if not colour and body is not None:
            colour = colour_key(body.fill, theme)
        f = RField(name=_clean(label), label=label, col=c, key=norm(label), unit=unit,
                   colour=colour)
        top = parent.get(c)
        f.keys = [k for k in {norm(label), norm(name), norm(_clean(label))} if k]
        if is_text(top):
            f.keys.append(norm(f"{top} {label}"))
        bv = values.get((first, c))
        if isinstance(bv, str) and bv.startswith("="):
            f.formula = bv
        fields.append(f)
    return TableTemplate(str(path), ws.title, r, first, max(end, first - 1), fields,
                         content_below=below)


def _form_fields(ws, tt: TableTemplate, theme) -> List[RField]:
    """'Inspector:' with an empty cell to its right (bordered, coloured, or
    the label ends with ':'), outside the table: one value per sheet."""
    out: List[RField] = []
    cols = {f.col for f in tt.fields}
    lo, hi = min(cols), max(cols)
    for (r, c), cell in list(ws._cells.items())[:MAX_SHEET_CELLS]:
        v = cell.value
        if not is_text(v) or str(v).startswith("="):
            continue
        if tt.header_row <= r <= max(tt.end_row, tt.header_row) and lo <= c <= hi:
            continue
        right = ws._cells.get((r, c + 1))
        if right is not None and right.value not in (None, ""):
            continue
        if not (str(v).rstrip().endswith(":") or (right is not None and _styled(right))):
            continue
        label = str(v).strip()
        out.append(RField(name=_clean(label), label=label, col=c + 1, key=norm(label),
                          keys=[norm(label)], unit=_split_unit(_clean(label))[1],
                          colour=colour_key(right.fill, theme) if right is not None else "",
                          form_row=r))
    return out


def _unique_names(tt: TableTemplate) -> None:
    seen: Dict[str, int] = {}
    for f in tt.fields + tt.forms:
        base = f.name or f"Column {_col_letter(f.col)}"
        n = seen.get(base.lower(), 0)
        seen[base.lower()] = n + 1
        f.name = base if n == 0 else f"{base} ({n + 1})"


def is_record_template(path) -> bool:
    try:
        analyze_template(path)
        return True
    except LedgerError:
        return False


# ── reading the source workbooks ────────────────────────────────────────
@dataclass
class SheetData:
    file: str                     # the workbook's name
    path: str
    name: str                     # the sheet's
    cells: Dict[Tuple[int, int], Any]
    colours: Dict[Tuple[int, int], str] = field(default_factory=dict)
    order: int = 0                # the sheet's place in its workbook
    #: a slide's text put over its table only by where it sits: cell -> how
    guessed: Dict[Tuple[int, int], str] = field(default_factory=dict)
    #: a text document's (docs_open): its records wait for the user
    document: bool = False


def read_book(path, notes: Optional[List[str]] = None) -> List[SheetData]:
    """Every sheet's values (and fill colours) of one workbook. .xlsx/.xlsm
    stream through openpyxl (read only: charts and pictures are not read -
    a chart is drawn from cells, which are); .csv as text; anything else
    (.xls, .xlsb, rights-protected) through source_extract's reader (the
    installed Excel on Windows), without colours."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in (".pptx", ".pptm"):
        return _read_pptx(path, notes)
    import docs_open
    if suffix in docs_open.SUFFIXES:
        return _read_doc(path, notes)
    if suffix in (".xlsx", ".xlsm"):
        try:
            return _read_xlsx(path, notes)
        except LedgerError:
            raise
        except Exception as e:                              # noqa: BLE001
            if notes is not None:
                notes.append(f"{path.name}: openpyxl could not read it ({e}); trying Excel")
    import source_extract as se
    grids, _ = se.load_workbook_grids(path, notes, want_pictures=False)
    out = []
    for i, g in enumerate(grids):
        if " › " in g.sheet:                               # a text box or a chart
            continue
        out.append(SheetData(path.name, str(path), g.sheet,
                             {k: (str(v) if type(v).__name__ == "_Merged" else v)
                              for k, v in g.cells.items()}, order=i))
    return out


def _read_pptx(path: Path, notes: Optional[List[str]]) -> List[SheetData]:
    """A deck: each slide's tables as sheets, the text that belongs to each
    above it (pptx_open), cell and text box fills as colours."""
    import pptx_open
    try:
        deck = pptx_open.read_pptx(path, want_pictures=False)
    except pptx_open.PptxError as e:
        raise LedgerError(str(e)) from e
    if notes is not None:
        notes.extend(deck.notes)
        notes.extend(deck.guesses(path.name))
    return [SheetData(path.name, str(path), s.name, {k: _num(v) for k, v in s.cells.items()},
                      dict(s.colours), order=i, guessed=dict(s.guessed))
            for i, s in enumerate(deck.sheets)]


def _read_doc(path: Path, notes: Optional[List[str]]) -> List[SheetData]:
    """A text document: its lines and each table as sheets (docs_open)."""
    import docs_open
    try:
        doc = docs_open.read_document(path, want_pictures=False)
    except docs_open.DocError as e:
        raise LedgerError(str(e)) from e
    if notes is not None:
        notes.extend(doc.notes)
    return [SheetData(path.name, str(path), s.name, {k: _num(v) for k, v in s.cells.items()},
                      order=i, document=True) for i, s in enumerate(doc.sheets)]


def _read_csv(path: Path) -> List[SheetData]:
    import csv
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff((text or "")[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    cells: Dict[Tuple[int, int], Any] = {}
    for r, row in enumerate(csv.reader((text or "").splitlines(), dialect), start=1):
        for c, v in enumerate(row[:MAX_COLS], start=1):
            v = v.strip()
            if v:
                cells[(r, c)] = _num(v)
    return [SheetData(path.name, str(path), path.stem, cells)]


def _num(v: str) -> Any:
    s = v.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", v) else v
    if re.fullmatch(r"-?\d+", s) and len(s) < 16:
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+([eE][-+]?\d+)?", s):
        try:
            return float(s)
        except ValueError:
            return v
    return v


def _read_xlsx(path: Path, notes: Optional[List[str]]) -> List[SheetData]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:                                  # noqa: BLE001
        if "encrypt" in str(e).lower() or "password" in str(e).lower():
            raise LedgerError(f"{path.name} has a password to open; save a copy "
                              "without it and add that copy") from e
        raise
    theme = theme_colours(wb)
    out: List[SheetData] = []
    try:
        for i, ws in enumerate(wb.worksheets):
            if getattr(ws, "sheet_state", "visible") == "veryHidden":
                continue
            cells: Dict[Tuple[int, int], Any] = {}
            colours: Dict[Tuple[int, int], str] = {}
            last, n, cut = 0, 0, False
            # from row 1, every row in turn (empty ones too): its number is
            # its place, as an empty cell carries none
            rows = ws.iter_rows(min_row=1, max_col=MAX_COLS)
            for rnum, row in enumerate(rows, start=1):
                any_here = False
                for ci, c in enumerate(row, start=1):
                    v = getattr(c, "value", None)
                    if v is None and not hasattr(c, "fill"):
                        continue
                    r, col = rnum, getattr(c, "column", None) or ci
                    k = colour_key(getattr(c, "fill", None), theme) \
                        if hasattr(c, "fill") else ""
                    if v is not None and not (isinstance(v, str) and not v.strip()):
                        if isinstance(v, str) and v.strip().upper() in (
                                "#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!",
                                "#NULL!"):
                            continue
                        if isinstance(v, (datetime, date)):
                            v = cell_text(v)
                        elif isinstance(v, str):
                            v = xl_text(v)
                        cells[(r, col)] = v
                        any_here = True
                        n += 1
                    if k:
                        colours[(r, col)] = k
                if any_here:
                    last = rnum
                elif last and rnum - last > EMPTY_ROWS_STOP:
                    break
                elif not last and rnum > EMPTY_ROWS_STOP * 5:
                    break
                if n >= MAX_SHEET_CELLS:
                    cut = True
                    break
            # closed here, while the file is open: left suspended, the read-only
            # generator's parser crashed the process when collected later
            rows.close()
            if cut and notes is not None:
                notes.append(f"{path.name} › {ws.title}: only its first "
                             f"{MAX_SHEET_CELLS:,} cells were read")
            # colours of empty cells far from any value are the sheet's style
            # (a whole row painted), kept only beside values
            if colours:
                colours = {k: v for k, v in colours.items()
                           if k in cells or (k[0] <= last + 1)}
            out.append(SheetData(path.name, str(path), ws.title, cells, colours, order=i))
    finally:
        wb.close()
    return out


class BookCache:
    """Sheets of a workbook already read, by its path, size and time
    (gzip JSON; no pickle). A file changed since is read again."""

    def __init__(self, folder: Optional[Path]):
        self.folder = Path(folder) if folder else None

    def _key(self, path: Path) -> Optional[str]:
        try:
            st = path.stat()
        except OSError:
            return None
        extra = ""
        if path.suffix.lower() in (".pptx", ".pptm"):
            # a deck is laid out by your text box <-> table links: read again
            # when they change
            import pptx_open
            mine = (getattr(pptx_open._LOCAL, "pairings", None) or {}).get(path.name) or {}
            extra = json.dumps(mine, sort_keys=True)
        return _sha(f"{path.resolve()}|{st.st_size}|{st.st_mtime_ns}|{READER_VERSION}|{extra}",
                    24)

    def get(self, path: Path) -> Optional[List[SheetData]]:
        if self.folder is None:
            return None
        k = self._key(path)
        f = self.folder / f"{k}.json.gz" if k else None
        if f is None or not f.is_file():
            return None
        try:
            data = json.loads(gzip.decompress(f.read_bytes()).decode("utf-8"))
            return [SheetData(path.name, str(path), s["name"],
                              {(r, c): v for r, c, v in s["cells"]},
                              {(r, c): k2 for r, c, k2 in s.get("colours", [])},
                              order=s.get("order", 0),
                              guessed={(r, c): h for r, c, h in s.get("guessed", [])},
                              document=bool(s.get("document")))
                    for s in data["sheets"]]
        except Exception:                                   # noqa: BLE001
            return None

    def put(self, path: Path, sheets: List[SheetData]) -> None:
        if self.folder is None:
            return
        k = self._key(path)
        if not k:
            return
        try:
            self.folder.mkdir(parents=True, exist_ok=True)
            data = {"file": str(path), "sheets": [
                {"name": s.name, "order": s.order,
                 "cells": [[r, c, v] for (r, c), v in s.cells.items()],
                 "colours": [[r, c, v] for (r, c), v in s.colours.items()],
                 "guessed": [[r, c, v] for (r, c), v in s.guessed.items()],
                 "document": s.document}
                for s in sheets]}
            tmp = self.folder / f"{k}.tmp"
            tmp.write_bytes(gzip.compress(json.dumps(data, default=str).encode("utf-8"), 5))
            os.replace(tmp, self.folder / f"{k}.json.gz")
        except OSError:
            pass


# ── tables in a sheet ───────────────────────────────────────────────────
@dataclass
class Table:
    file: str
    path: str
    sheet: str
    orient: str                   # 'rows': records are rows; 'cols': records are columns
    header_at: int                # the header's row (its column when sideways)
    cols: List[int]               # in the header's direction
    headers: List[str]
    keys: List[List[str]]         # per column, every way its header reads
    rows: List[int]               # record rows (columns when sideways)
    values: List[List[Any]]       # per record, per column
    colours: List[str] = field(default_factory=list)
    order: int = 0
    skipped_rows: int = 0         # total rows and repeated headers left out

    @property
    def id(self) -> str:
        return f"{self.file}|{self.sheet}|{self.orient}|{self.header_at}|{self.cols[0]}"

    @property
    def signature(self) -> str:
        return _sha("|".join(k[0] if k else "" for k in self.keys))

    def where(self, i: int, j: int) -> str:
        r, c = self.rows[i], self.cols[j]
        return _a1(r, c) if self.orient == "rows" else _a1(c, r)

    def row_label(self, i: int) -> int:
        return self.rows[i]

    def span(self) -> str:
        r1, r2 = self.header_at, self.rows[-1] if self.rows else self.header_at
        c1, c2 = self.cols[0], self.cols[-1]
        return f"{_a1(r1, c1)}:{_a1(r2, c2)}" if self.orient == "rows" else \
            f"{_a1(c1, r1)}:{_a1(c2, r2)}"

    def column(self, j: int) -> List[Any]:
        return [row[j] for row in self.values]

    def box(self) -> Tuple[int, int, int, int]:
        """(top, left, bottom, right) on the sheet, its header included."""
        r1, r2 = self.header_at, max(self.rows) if self.rows else self.header_at
        c1, c2 = self.cols[0], self.cols[-1]
        return (r1, c1, r2, c2) if self.orient == "rows" else (c1, r1, c2, r2)

    def cell_set(self) -> Set[Tuple[int, int]]:
        """The sheet cells of its header and records, the sheet's way up."""
        out = set()
        for r in [self.header_at] + list(self.rows):
            for c in self.cols:
                out.add((r, c) if self.orient == "rows" else (c, r))
        return out


def find_tables(sheet: SheetData, sideways: bool = True) -> List[Table]:
    """Every table on a sheet: a row of two or more headers (words, side by
    side; one unnamed column between them is kept), its records under it to
    a blank row or the next header. Rows that total the table up and a
    header repeated (a page break) are skipped. With ``sideways`` the sheet
    is also read turned (records across), for the mapping to choose."""
    out = _tables(sheet, sheet.cells, "rows")
    if sideways and sheet.cells:
        turned = {(c, r): v for (r, c), v in sheet.cells.items()}
        out += _tables(sheet, turned, "cols")
    return out


def _tables(sheet: SheetData, cells: Dict[Tuple[int, int], Any], orient: str) -> List[Table]:
    by_row: Dict[int, Dict[int, Any]] = {}
    for (r, c), v in cells.items():
        by_row.setdefault(r, {})[c] = v
    rows_sorted = sorted(by_row)
    header_like: List[int] = []
    for r in rows_sorted:
        row = by_row[r]
        texts = [c for c, v in row.items() if is_text(v)]
        if len(texts) >= 2 and len(texts) >= 0.6 * len(row):
            header_like.append(r)
        if len(header_like) >= MAX_HEADER_ROWS:
            break
    out: List[Table] = []
    # the cells a table already holds as records, by its columns: a row of
    # them is data of that table, not another header. Only the SAME columns
    # are passed over - a table beside it, or staggered against it (a grid of
    # tables), starts wherever its own header is
    taken: Dict[Tuple[int, ...], Set[int]] = {}
    colours = sheet.colours if orient == "rows" else {(c, r): k for (r, c), k
                                                      in sheet.colours.items()}
    for r in header_like:
        row = by_row[r]
        labels = _label_cols(row)
        texts = sorted(c for c, v in row.items() if is_text(v) and c not in labels)
        for run in _header_runs(texts, by_row, r):
            if r in taken.get(tuple(run), ()):
                continue
            t = _body(sheet, by_row, r, run, orient, colours)
            if t is None:
                continue
            out.append(t)
            taken.setdefault(tuple(run), set()).update(t.rows)
    return out


_LABEL_VALUE = re.compile(r"^\s*([^:]{1,60}?)\s*:\s*(\S.*)$")


def _is_label(v: Any) -> bool:
    """'Inspector:' or 'Inspector: MK' - a label and its value, not a
    column's header."""
    if not isinstance(v, str):
        return False
    t = v.strip()
    return (t.endswith(":") and len(t) > 1) or bool(
        _LABEL_VALUE.match(t) and re.search(r"[A-Za-z]", t.split(":", 1)[0])
        and not re.match(r"^\d{1,2}:\d{2}", t))


def _label_cols(row: Dict[int, Any]) -> Set[int]:
    """The columns of a row that hold labels ('Inspector:') and the value
    right after one: a form above a table, not its headers."""
    out: Set[int] = set()
    labels = [c for c, v in row.items() if _is_label(v)]
    texts = [c for c, v in row.items() if is_text(v)]
    # a header row may end its names with ':' ('Weight:'); a form line is
    # mostly labels ('Inspector: | JB | Inspector: | MK')
    if not labels or 2 * len(labels) < len(texts):
        return out
    for c in labels:
        out.add(c)
        if str(row[c]).strip().endswith(":") and (c + 1) in row:
            out.add(c + 1)
    return out


def _header_runs(texts: List[int], by_row, r: int) -> List[List[int]]:
    """Runs of header cells; one empty header between two is a column with
    no name when it has values under it."""
    runs: List[List[int]] = []
    for c in texts:
        if runs and c == runs[-1][-1] + 1:
            runs[-1].append(c)
        elif runs and c == runs[-1][-1] + 2 and _has_below(by_row, r, c - 1):
            runs[-1] += [c - 1, c]
        else:
            runs.append([c])
    out: List[List[int]] = []
    for run in runs:
        out += _split_repeats(run, by_row[r])
    return [x for x in out if sum(1 for c in x if is_text(by_row[r].get(c))) >= 2]


def _split_repeats(run: List[int], head: Dict[int, Any]) -> List[List[int]]:
    """Tables side by side with no empty column between: the first header
    coming back starts the next one ('Position Reading Position Reading')."""
    keys = [norm(head.get(c)) for c in run]
    first = keys[0]
    if not first:
        return [run]
    cuts = [i for i in range(2, len(run)) if keys[i] == first]
    if not cuts:
        return [run]
    out, start = [], 0
    for i in cuts + [len(run)]:
        out.append(run[start:i])
        start = i
    return out


def _has_below(by_row, r: int, c: int) -> bool:
    return any(c in by_row.get(r + i, {}) for i in range(1, 4))


def _body(sheet: SheetData, by_row, r: int, run: List[int], orient: str,
          colours) -> Optional[Table]:
    head = by_row[r]
    lo, hi = run[0], run[-1]
    header_keys = [norm(head.get(c)) for c in run]
    rows: List[int] = []
    values: List[List[Any]] = []
    numeric_cols: Set[int] = set()
    skipped = 0
    blank = 0
    x = r + 1
    last_row = max(by_row) if by_row else r
    while x <= last_row:
        row = by_row.get(x, {})
        vals = [row.get(c) for c in run]
        filled = [v for v in vals if v is not None]
        if not filled:
            blank += 1
            if blank >= 2:
                break
            x += 1
            continue
        if blank == 1 and rows and _new_header(vals, numeric_cols, run):
            break
        if rows and _form_row(row, lo, hi, len(run)):
            break                                         # 'Tank: 3' - the next block's labels
        if blank == 1 and rows and len(filled) < max(1, len(run) // 2) \
                and _typical_fill(values) > len(filled):
            break                                         # after a gap, not this table's shape
        texts = [v for v in vals if is_text(v)]
        if [norm(v) for v in vals] == header_keys:
            if blank:
                break                                     # another table like it, below it
            skipped += 1                                  # the header again (a page break)
            x += 1
            blank = 0
            continue
        if rows and len(texts) >= 2 and len(texts) == len(filled) and \
                _new_header(vals, numeric_cols, run):
            break                                         # the next table's header
        blank = 0
        rows.append(x)
        values.append(vals)
        for c, v in zip(run, vals):
            if is_number(v):
                numeric_cols.add(c)
        x += 1
    # rows that sum the table up ('Total', 'Average' ...) at its end are not
    # records; a record named 'Max' in the middle of it is
    while rows and _summary_row(values[-1]):
        rows.pop()
        values.pop()
        skipped += 1
    if not rows:
        return None
    if len(rows) == 1 and sum(v is not None for v in values[0]) < 2:
        return None
    # parents: a sparse row of words over the headers (a merged group heading)
    above = by_row.get(r - 1, {})
    not_parent = _label_cols(above)
    above = {c: v for c, v in above.items() if c not in not_parent}
    parent_texts = [(c, above.get(c)) for c in range(lo, hi + 1) if is_text(above.get(c))]
    parents: Dict[int, str] = {}
    if parent_texts and len(parent_texts) < len(run):
        cur = ""
        for c in range(lo, hi + 1):
            if is_text(above.get(c)):
                cur = str(above[c]).strip()
            parents[c] = cur
    headers, keys, cols_colour = [], [], []
    for c in run:
        h = head.get(c)
        text = str(h).strip() if is_text(h) else ""
        ks = [norm(text)] if text else [""]
        if parents.get(c):
            ks = [norm(f"{parents[c]} {text}")] + ks
        headers.append(text if not parents.get(c) else f"{parents[c]} › {text}")
        keys.append(ks)
        k = colours.get((r, c), "")
        if not k:
            body = [colours.get((x2, c), "") for x2 in rows[:200]]
            common = max(set(body), key=body.count) if body else ""
            if common and body.count(common) >= 0.5 * len(body):
                k = common
        cols_colour.append(k)
    return Table(sheet.file, sheet.path, sheet.name, orient, r, list(run), headers, keys,
                 rows, values, cols_colour, order=sheet.order, skipped_rows=skipped)


def _form_row(row: Dict[int, Any], lo: int, hi: int, width: int) -> bool:
    """A row of labels over the table's columns ('Tank: 3', 'Inspector: | AL'):
    the next block's form, not a record."""
    inside = {c: v for c, v in row.items() if lo <= c <= hi}
    if not inside:
        return False
    first = inside[min(inside)]
    return _is_label(first) and len(inside) <= max(2, width // 2)


def _typical_fill(values: List[List[Any]]) -> int:
    counts = sorted(sum(v is not None for v in row) for row in values[-20:])
    return counts[len(counts) // 2] if counts else 0


def _summary_row(vals) -> bool:
    first = next((v for v in vals if v is not None), None)
    return isinstance(first, str) and vnorm(first) in SUMMARY_WORDS


def _new_header(vals, numeric_cols: Set[int], run: List[int]) -> bool:
    """Words across most of the row, where the table so far had numbers:
    the next table's header (a remark or two in a row is not)."""
    words = [c for c, v in zip(run, vals) if is_text(v)]
    if len(words) < 2 or len(words) < 0.6 * len(run):
        return False
    return sum(1 for c in words if c in numeric_cols) >= max(1, len(words) // 2)


# ── the recipe: what this Ledger learned ────────────────────────────────
@dataclass
class Recipe:
    version: int = RECIPE_VERSION
    template: Dict[str, Any] = field(default_factory=dict)       # TableTemplate.to_dict
    #: header words -> field, the user's (taught, colour, a Decisions Y)
    aliases: Dict[str, str] = field(default_factory=dict)
    #: header words -> fields they are NOT (a Decisions N)
    not_aliases: Dict[str, List[str]] = field(default_factory=dict)
    #: field -> its rule when it is not a column: {"kind": "filename" |
    #: "sheetname" | "lookup" | "constant" | "label", ..., "confirmed": bool}
    sources: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    #: field -> {lookup value words: lookup key} the user confirmed
    key_aliases: Dict[str, Dict[str, str]] = field(default_factory=dict)
    #: the fields that tell one record from another (the CSV's identity)
    key_fields: List[str] = field(default_factory=list)
    #: header layout signature -> {header words: field or None} the model gave
    model_cache: Dict[str, Dict[str, Optional[str]]] = field(default_factory=dict)
    #: field -> {"numeric": share of numbers, "n": values seen, "cats": [...]}
    profiles: Dict[str, Dict[str, Any]] = field(default_factory=dict)
    #: (field|header words) whose odd-looking values the user accepted
    profile_ok: List[str] = field(default_factory=list)
    use_colours: bool = True
    extra: Dict[str, Any] = field(default_factory=dict)

    @property
    def table_template(self) -> TableTemplate:
        return TableTemplate.from_dict(self.template)


def recipe_path_for(csv_path) -> Path:
    return Path(csv_path).with_suffix(".records.json")


def cache_dir_for(csv_path) -> Path:
    p = Path(csv_path)
    return p.with_name(p.stem + ".records_cache")


def is_record_ledger(csv_path) -> bool:
    return bool(csv_path) and recipe_path_for(csv_path).is_file()


def load_recipe(csv_path) -> Recipe:
    p = recipe_path_for(csv_path)
    if not p.is_file():
        raise LedgerError(f"{Path(csv_path).name} is not a table Ledger (no "
                          f"{p.name} beside it); create it from a table template")
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception as e:                                  # noqa: BLE001
        raise LedgerError(f"{p.name} could not be read (fix or remove it): {e}") from e
    if not isinstance(data, dict):
        raise LedgerError(f"{p.name} is not a Ledger recipe")
    if int(data.get("version", 1)) > RECIPE_VERSION:
        raise LedgerError(f"{p.name} was made by a newer Council; update Council to use it")
    known = {k: data[k] for k in Recipe.__dataclass_fields__ if k in data and k != "extra"}
    extra = {k: v for k, v in data.items() if k not in Recipe.__dataclass_fields__}
    extra.update(data.get("extra") or {})
    return Recipe(**known, extra=extra)


def save_recipe(csv_path, recipe: Recipe) -> None:
    data = asdict(recipe)
    extra = data.pop("extra", {}) or {}
    for k, v in extra.items():
        data.setdefault(k, v)
    _atomic_write_text(recipe_path_for(csv_path), json.dumps(data, indent=1,
                                                             ensure_ascii=False))


def csv_columns(tt: TableTemplate) -> List[str]:
    return [f.name for f in tt.all_fields] + [SHEET_COL, TABLE_COL, ROW_COL, SOURCE_COL,
                                              UPDATED_COL]


def _schema_for(tt: TableTemplate, id_field: str) -> TemplateSchema:
    """A sample_ledger schema for the parent CSV, so the knowledge graph and
    the rest of Council read a table Ledger as any other (one 'table' field
    per column, its first record cell)."""
    fields = []
    for f in tt.all_fields:
        cell = _a1(f.form_row or tt.first_row, f.col)
        fields.append(TemplateField(name=f.name, label=f.label, sheet=tt.sheet,
                                    section="Records" if not f.is_form else "Form",
                                    layout="table", cells={"": cell}, unit=f.unit,
                                    key=f.key))
    return TemplateSchema(tt.path, fields, id_field)


@dataclass
class CreateResult:
    csv: str
    created: bool
    fields: List[str]
    forms: List[str]
    sheet: str
    first_row: int
    rows_laid_out: int
    key_fields: List[str]

    def describe(self) -> str:
        verb = "created" if self.created else "extended"
        s = (f"Parent CSV {verb}: {self.csv}\nTable template ({self.sheet}, records from "
             f"row {self.first_row}): one row per record, {len(self.fields)} fields: "
             + ", ".join(self.fields))
        if self.forms:
            s += "\nForm fields (one value per source sheet): " + ", ".join(self.forms)
        if self.key_fields:
            s += "\nRecords are told apart by: " + ", ".join(self.key_fields)
        return s


def create_ledger(template, csv_path, key_fields: Optional[Sequence[str]] = None
                  ) -> CreateResult:
    """A table Ledger: the parent CSV (one row per record), its schema, and
    the recipe. An existing one keeps what it learned; new template fields
    become new CSV columns."""
    tt = analyze_template(template)
    csv_path = Path(csv_path)
    names = [f.name for f in tt.all_fields]
    created = not csv_path.exists()
    if is_record_ledger(csv_path):
        recipe = load_recipe(csv_path)
        recipe.template = tt.to_dict()
    else:
        if csv_path.exists():
            raise LedgerError(f"{csv_path.name} exists and is not a table Ledger; pick "
                              "another parent CSV for this template")
        recipe = Recipe(template=tt.to_dict())
    if key_fields:
        bad = [k for k in key_fields if k not in names]
        if bad:
            raise LedgerError(f"key field(s) {bad} are not in the template; fields: {names}")
        recipe.key_fields = list(key_fields)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    cols = csv_columns(tt)
    if created:
        _write_rows(csv_path, cols, [])
    else:
        have_cols, rows = _read_rows(csv_path)
        if have_cols != cols:
            keep = [c for c in have_cols if c not in cols]
            _write_rows(csv_path, [c for c in cols if c not in (SOURCE_COL, UPDATED_COL)]
                        + keep + [SOURCE_COL, UPDATED_COL], rows)
    id_field = (recipe.key_fields or [names[0]])[0]
    _schema_for(tt, id_field).save(schema_path_for(csv_path))
    save_recipe(csv_path, recipe)
    return CreateResult(str(csv_path), created, [f.name for f in tt.fields],
                        [f.name for f in tt.forms], tt.sheet, tt.first_row, tt.styled_rows,
                        list(recipe.key_fields))


def _read_rows(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    import csv
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return list(reader.fieldnames or []), [dict(r) for r in reader]


def _write_rows(path: Path, columns: List[str], rows: List[Dict[str, str]]) -> None:
    import csv
    import io
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
    w.writeheader()
    for r in rows:
        w.writerow({c: r.get(c, "") for c in columns})
    _atomic_write_text(Path(path), buf.getvalue())


# ── one value, one record ───────────────────────────────────────────────
@dataclass
class Val:
    value: Any
    how: str                      # see SURE_HOWS; else a guess (amber)
    where: str = ""               # 'BTB_Woopentide.xlsx › Sheet1!B3'
    note: str = ""
    check: str = ""               # the Decisions id that accepts it, '' when sure

    @property
    def sure(self) -> bool:
        return not self.check


@dataclass
class Record:
    file: str
    path: str
    sheet: str
    row: int                      # in the source (its column number when sideways)
    table: str                    # Table.id
    values: Dict[str, Val] = field(default_factory=dict)
    accept: str = ""              # 'Y' pre-accepted, 'N' left out, '' to decide
    notes: List[str] = field(default_factory=list)
    flags: List[str] = field(default_factory=list)   # 'clash', 'duplicate'
    order: int = 0
    checks: List[str] = field(default_factory=list)  # Decisions the whole record waits on

    def key(self, key_fields: Sequence[str]) -> Tuple[str, ...]:
        return tuple(vnorm(cell_text(self.values[k].value)) if k in self.values else ""
                     for k in key_fields)


@dataclass
class Decision:
    id: str
    kind: str                     # 'column', 'values', 'rule', 'lookup'
    field: str
    detail: str
    examples: str = ""
    records: int = 0
    accept: str = ""


@dataclass
class ColMap:
    field: str
    j: int
    how: str
    score: float
    note: str = ""
    check: str = ""


@dataclass
class TableUse:
    table: Table
    use: str                      # 'records' | 'lookup' | 'skipped'
    why: str = ""
    maps: List[ColMap] = field(default_factory=list)
    score: float = 0.0
    check: str = ""               # a Decisions id every record of it waits on


@dataclass
class RecordRun:
    tt: TableTemplate
    records: List[Record] = field(default_factory=list)
    decisions: Dict[str, Decision] = field(default_factory=dict)
    tables: List[TableUse] = field(default_factory=list)
    problems: List[str] = field(default_factory=list)
    learned: List[str] = field(default_factory=list)
    files: int = 0
    files_used: int = 0
    from_cache: int = 0
    open_failures: List[Tuple[str, str]] = field(default_factory=list)
    timings: List[Tuple[str, float]] = field(default_factory=list)
    model: str = ""
    model_asked: int = 0
    #: what the text documents among the sources say labels mean (docs_open)
    explanations: List[Any] = field(default_factory=list)
    cancelled: bool = False
    key_fields: List[str] = field(default_factory=list)

    @property
    def cells(self) -> int:
        return sum(1 for r in self.records for v in r.values.values()
                   if cell_text(v.value) != "")

    @property
    def cells_pre_accepted(self) -> int:
        return sum(1 for r in self.records if r.accept == "Y" for v in r.values.values()
                   if cell_text(v.value) != "")

    def decide(self, d: Decision) -> str:
        have = self.decisions.get(d.id)
        if have is None:
            self.decisions[d.id] = d
            have = d
        have.records += 1
        return have.id


# ── mapping a table's columns to the fields ─────────────────────────────
def _words(key: str) -> Set[str]:
    return {w for w in key.split() if w}


@lru_cache(maxsize=200_000)
def _similar(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    wa, wb = _words(a), _words(b)
    jac = len(wa & wb) / max(1, len(wa | wb))
    if jac >= 0.84:
        return jac
    m = difflib.SequenceMatcher(None, a, b)
    if m.real_quick_ratio() < 0.84 or m.quick_ratio() < 0.84:
        return jac                                  # cannot reach the bar: skip the slow part
    return max(jac, m.ratio())


def map_table(t: Table, fields: Sequence[RField], recipe: Recipe,
              memo: Optional[Dict[tuple, List[ColMap]]] = None,
              explained: Optional[Dict[str, Tuple[str, Any]]] = None) -> List[ColMap]:
    """See _map_table; ``memo`` keeps the answer per header layout (the same
    report from a thousand workbooks is matched once). ``explained``: what
    the user's documents say header words mean (docs_open.explanations_for)."""
    if memo is None:
        return _map_table(t, fields, recipe, explained)
    key = (tuple(tuple(k) for k in t.keys), t.signature)
    if key not in memo:
        memo[key] = _map_table(t, fields, recipe, explained)
    return [ColMap(m.field, m.j, m.how, m.score, m.note, m.check) for m in memo[key]]


def _map_table(t: Table, fields: Sequence[RField], recipe: Recipe,
               explained: Optional[Dict[str, Tuple[str, Any]]] = None) -> List[ColMap]:
    """Each field's column in ``t`` (one column per field, one field per
    column), the surest first: the user's own words (taught, highlighted,
    accepted) and exact header words are sure; the same words in another
    order are sure; similar words, and the model's answers, are guesses."""
    cands: List[Tuple[float, int, ColMap]] = []
    model = recipe.model_cache.get(t.signature, {})
    for fi, f in enumerate(fields):
        fkeys = set(f.keys or [f.key]) | {norm(f.name)}
        fkeys.discard("")
        for j, ks in enumerate(t.keys):
            best: Optional[ColMap] = None
            for k in ks:
                if not k:
                    continue
                if f.name in recipe.not_aliases.get(k, []):
                    best = None
                    break
                if recipe.aliases.get(k) == f.name:
                    m = ColMap(f.name, j, "taught", 1.0, f"your answer: '{t.headers[j]}' "
                               f"is {f.name}")
                elif k in fkeys:
                    m = ColMap(f.name, j, "header", 0.99)
                elif any(_words(k) == _words(x) for x in fkeys):
                    m = ColMap(f.name, j, "words", 0.95)
                elif explained and explained.get(vnorm(k), ("",))[0] == f.name:
                    m = ColMap(f.name, j, "explained", 0.7,
                               f"explained in {explained[vnorm(k)][1].describe()}")
                elif model.get(k) == f.name:
                    m = ColMap(f.name, j, "model", 0.6,
                               f"Council's model read '{t.headers[j]}' as {f.name}")
                else:
                    s = max((_similar(k, x) for x in fkeys), default=0.0)
                    if s < 0.84 or k in recipe.aliases:
                        continue
                    m = ColMap(f.name, j, "similar", round(0.5 + s / 3, 3),
                               f"'{t.headers[j]}' is like {f.name}")
                if best is None or m.score > best.score:
                    best = m
            if best is not None:
                cands.append((best.score, -fi, best))
    out: List[ColMap] = []
    used_f, used_j = set(), set()
    for _s, _fi, m in sorted(cands, key=lambda x: (-x[0], -x[1], x[2].j)):
        if m.field in used_f or m.j in used_j:
            continue
        used_f.add(m.field)
        used_j.add(m.j)
        out.append(m)
    return out


def _qualifies(maps: List[ColMap], n_fields: int) -> bool:
    sure = sum(1 for m in maps if m.how in SURE_HOWS)
    return sure >= 2 or (sure >= 1 and len(maps) >= 2) or (n_fields == 1 and sure == 1)


# ── learning from the user's highlights ─────────────────────────────────
def learn_colours(tt: TableTemplate, tables: Sequence[Table], recipe: Recipe,
                  all_files: Sequence[str], sheets: Sequence[SheetData] = ()) -> List[str]:
    """What the user's highlights say (the same fill in the template and a
    workbook means 'this is that'): each painted column of a workbook's
    record table is the field of its colour (by name among fields of one
    colour, else by order when the counts agree); a painted column in a
    table with no other painted field is a LOOKUP for its field; a template
    colour found in no workbook comes from the file (or sheet) name - a
    guess until the user confirms it. Returns lines saying what was
    learned; the recipe is changed in place."""
    said: List[str] = []
    fields = [f for f in tt.fields if f.colour and not f.formula]
    if not fields or not recipe.use_colours:
        return said
    groups: List[Tuple[str, List[RField]]] = []
    for f in fields:
        for g in groups:
            if same_colour(g[0], f.colour):
                g[1].append(f)
                break
        else:
            groups.append((f.colour, [f]))
    # read the right way up only: a painted column read sideways is a row
    # of headers all one colour, which says nothing
    tables = [t for t in tables if t.orient == "rows"]
    hits: Dict[int, Dict[int, List[int]]] = {}       # table index -> group -> columns
    for ti, t in enumerate(tables):
        for j, k in enumerate(t.colours):
            for gi, (gc, _fs) in enumerate(groups):
                if same_colour(gc, k):
                    hits.setdefault(ti, {}).setdefault(gi, []).append(j)
    if not hits:
        return said
    # a record table: two highlighted fields of different colours, or two
    # columns of a colour several fields share; one field's colour alone
    # (however many columns) is a lookup when a record table is painted too
    record_like = {ti for ti, g in hits.items()
                   if len(g) >= 2 or any(len(groups[gi][1]) >= 2 and len(js) >= 2
                                         for gi, js in g.items())}
    if not record_like:
        record_like = set(hits)
    seen_group: Set[int] = set()
    for ti in sorted(record_like):
        t = tables[ti]
        taught: List[str] = []
        for gi, js in hits[ti].items():
            seen_group.add(gi)
            fs = list(groups[gi][1])
            cols = list(js)
            pairs: List[Tuple[RField, int]] = []
            for f in list(fs):                           # by name first
                fk = set(f.keys) | {norm(f.name)}
                j = next((j for j in cols if any(k in fk for k in t.keys[j] if k)), None)
                if j is None:
                    j = max(cols, key=lambda jj: max(_similar(k, x) for k in t.keys[jj]
                                                     for x in fk), default=None)
                    if j is not None and max(_similar(k, x) for k in t.keys[j]
                                             for x in fk) < 0.84:
                        j = None
                if j is not None:
                    pairs.append((f, j))
                    fs.remove(f)
                    cols.remove(j)
            if fs and len(fs) == len(cols):              # the rest by order
                pairs += list(zip(fs, sorted(cols)))
            for f, j in pairs:
                k = t.keys[j][0]
                if k and recipe.aliases.get(k) != f.name:
                    recipe.aliases[k] = f.name
                    if f.name in recipe.not_aliases.get(k, []):
                        recipe.not_aliases[k].remove(f.name)
                    taught.append(f"{f.name} ← '{t.headers[j]}'")
        if taught:
            said.append(f"Learned from your highlights in {t.file} › {t.sheet}: "
                        + ", ".join(taught))
    for ti, g in hits.items():
        if ti in record_like:
            continue
        t = tables[ti]
        for gi, js in g.items():
            if len(groups[gi][1]) != 1:
                continue
            f = groups[gi][1][0]
            seen_group.add(gi)
            j = js[-1]                                    # the value: the last painted column
            key_j = js[0] if _mostly_text(t.column(js[0])) else next(
                (jj for jj in range(len(t.cols)) if _mostly_text(t.column(jj))), j)
            rule = {"kind": "lookup", "file": t.file, "path": t.path, "sheet": t.sheet,
                    "key_header": t.keys[key_j][0], "value_header": t.keys[j][0],
                    "join": recipe.sources.get(f.name, {}).get("join", ""),
                    "confirmed": True, "how": "colour"}
            old = recipe.sources.get(f.name, {})
            if any(old.get(k) != rule[k] for k in ("kind", "file", "key_header",
                                                   "value_header")):
                recipe.sources[f.name] = rule
                said.append(f"Learned from your highlights: {f.name} is looked up in "
                            f"{t.file} › {t.sheet}, column '{t.headers[j]}'")
    seen_group |= _learn_painted_cells(groups, seen_group, tables, sheets, recipe, said)
    stems = {Path(x).stem for x in all_files}
    headed = {k for t in tables for ks in t.keys for k in ks if k}
    for gi, (gc, fs) in enumerate(groups):
        if gi in seen_group:
            continue
        for f in fs:
            if f.name in recipe.sources or any(v == f.name for v in recipe.aliases.values()):
                continue
            if headed & (set(f.keys) | {norm(f.name)}):
                continue                                  # a workbook has its column
            fkeys = (set(f.keys) | {norm(f.name)}) - {""}
            if any(_labelled_values(sh, fkeys, set()) for sh in sheets):
                continue                                  # a workbook labels it ('Tame: Spider')
            recipe.sources[f.name] = {
                "kind": "name", "pattern": "", "confirmed": False, "how": "colour",
                "why": f"{f.name} is highlighted in the template, but nothing in the "
                       "workbooks has its colour: taken from the file name (or the sheet "
                       "name when the file names are all alike)"}
            said.append(f"{f.name}: nothing in the workbooks has its highlight colour, so "
                        f"it is taken from the file name"
                        + ("" if len(stems) > 1 else " (one file only so far)")
                        + " — check it in the review (Decisions)")
    return said


def _learn_painted_cells(groups, seen_group: Set[int], tables: Sequence[Table],
                         sheets: Sequence[SheetData], recipe: Recipe,
                         said: List[str]) -> Set[int]:
    """A painted cell outside every table names its table (a grid's block
    title): 'Vessel: 3' teaches the label 'Vessel'; '3' beside the label
    'Vessel' teaches it too; a bare 'Tank 3' teaches its place over the
    table (the same place over every table, in every workbook)."""
    found: Set[int] = set()
    in_tables: Dict[Tuple[str, str], Set[Tuple[int, int]]] = {}
    for t in tables:
        in_tables.setdefault((t.path, t.sheet), set()).update(t.cell_set())
    for sh in sheets:
        if not sh.colours:
            continue
        mine = [t for t in tables if t.path == sh.path and t.sheet == sh.name]
        inside = in_tables.get((sh.path, sh.name), set())
        for (r, c), k in sorted(sh.colours.items()):
            v = sh.cells.get((r, c))
            if v is None or (r, c) in inside:
                continue
            for gi, (gc, fs) in enumerate(groups):
                if gi in seen_group or gi in found or len(fs) != 1 or not same_colour(gc, k):
                    continue
                f = fs[0]
                label = ""
                text = str(v).strip()
                if _is_label(text) and ":" in text.rstrip(":"):
                    label = text.split(":", 1)[0]
                elif isinstance(sh.cells.get((r, c - 1)), str) and \
                        _is_label(sh.cells[(r, c - 1)]):
                    label = str(sh.cells[(r, c - 1)]).rstrip(": ")
                if label:
                    if norm(label) not in set(f.keys) | {norm(f.name)}:
                        recipe.aliases[norm(label)] = f.name
                    said.append(f"Learned from your highlights in {sh.file} › {sh.name}: "
                                f"{f.name} is the value labelled '{label}' beside each table")
                    found.add(gi)
                    continue
                if not mine:
                    continue
                t = min(mine, key=lambda t: _distance(r, c, t.box()))
                top, left = t.box()[:2]
                recipe.sources[f.name] = {"kind": "near", "dr": r - top, "dc": c - left,
                                          "confirmed": True, "how": "colour"}
                said.append(f"Learned from your highlights in {sh.file} › {sh.name}: {f.name}"
                            f" is the cell {_a1(r, c)} ('{text[:40]}'), "
                            f"{_offset_words(r - top, c - left)} each table")
                found.add(gi)
    return found


def _offset_words(dr: int, dc: int) -> str:
    parts = []
    if dr:
        parts.append(f"{abs(dr)} row{'s' if abs(dr) > 1 else ''} "
                     + ("above" if dr < 0 else "below"))
    if dc:
        parts.append(f"{abs(dc)} column{'s' if abs(dc) > 1 else ''} "
                     + ("left of" if dc < 0 else "right of"))
    return " and ".join(parts) + (" the start of" if parts else "at the start of")


def _mostly_text(vals: Sequence[Any]) -> bool:
    vs = [v for v in vals if v is not None and not is_nullish(v)]
    return bool(vs) and sum(1 for v in vs if isinstance(v, str) and not is_number(v)) \
        >= 0.6 * len(vs)


# ── fields that are not columns ─────────────────────────────────────────
def infer_name_rule(names: Sequence[str]) -> str:
    """The part of the names that differs between them, as a pattern:
    ['BTB_Spider', 'BTB_Woopentide'] -> '^BTB_(.+?)$'. '' for one name."""
    names = sorted({n for n in names if n})
    if len(names) < 2:
        return ""
    pre = os.path.commonprefix(names)
    suf = os.path.commonprefix([n[::-1] for n in names])[::-1]
    seps = " _-.#()[]"
    while pre and pre[-1] not in seps:
        pre = pre[:-1]
    while suf and suf[0] not in seps:
        suf = suf[1:]
    if any(len(n) <= len(pre) + len(suf) for n in names):
        pre, suf = "", ""
    return "^" + re.escape(pre) + "(.+?)" + re.escape(suf) + "$"


def name_rule_from_example(name: str, value: str) -> str:
    """'Lot 101_Spider', 'Spider' -> '^Lot\\ \\d+_(.+?)$'."""
    i = name.lower().find(str(value).lower())
    if i < 0 or not value:
        raise LedgerError(f"'{value}' is not in '{name}'")
    gen = lambda s: re.sub(r"\\d+|\d+", r"\\d+", re.escape(s))     # noqa: E731
    return "^" + gen(name[:i]) + "(.+?)" + gen(name[i + len(value):]) + "$"


def name_value(rule: str, name: str) -> Optional[str]:
    if not rule:
        return None
    try:
        m = re.match(rule, name, re.I)
    except re.error:
        return None
    if not m:
        return None
    return (m.group(1) if m.groups() else m.group(0)).strip(" _-.")


@dataclass
class LookupIndex:
    field: str
    rule: Dict[str, Any]
    entries: List[Tuple[str, Any, str, str]] = field(default_factory=list)  # key, value, raw key, where
    tables: List[str] = field(default_factory=list)

    def match(self, v: Any, aliases: Dict[str, str]) -> Tuple[Optional[Tuple], str, str]:
        """(entry, how, note): exact, then the longest key the value starts
        with on whole words, then the user's confirmed answers, then a
        similar key (a guess)."""
        n = vnorm(v)
        if not n:
            return None, "", ""
        same = [e for e in self.entries if e[0] == n]
        if len(same) == 1 or (same and len({cell_text(e[1]) for e in same}) == 1):
            return same[0], "exact", ""
        if same:
            return same[0], "ambiguous", (f"'{same[0][2]}' is in the lookup "
                                          f"{len(same)} times with different values")
        if n in aliases:
            hit = [e for e in self.entries if e[0] == aliases[n]]
            if hit:
                return hit[0], "your answer", ""
        pre = [e for e in self.entries if e[0] and n.startswith(e[0] + " ")]
        if pre:
            longest = max(len(e[0]) for e in pre)
            top = [e for e in pre if len(e[0]) == longest]
            if len({e[0] for e in top}) == 1:
                return top[0], "starts with", ""
            return top[0], "ambiguous", ("starts with " + " and ".join(
                sorted({repr(e[2]) for e in top})))
        keys = [e[0] for e in self.entries if e[0]]
        close = difflib.get_close_matches(n, keys, n=2, cutoff=0.8)
        if not close:
            words = n.split()
            for cut in range(len(words) - 1, 0, -1):     # a mistyped tail ('#14' typed 'Fourt')
                close = difflib.get_close_matches(" ".join(words[:cut]), keys, n=2, cutoff=0.86)
                if close:
                    break
        if close:
            e = next(e for e in self.entries if e[0] == close[0])
            return e, "similar", f"'{v}' is like '{e[2]}'"
        return None, "", f"no entry for '{v}'"


def _labelled_values(sheet: SheetData, keys: Set[str], skip: Set[Tuple[int, int]]
                     ) -> List[Tuple[Any, int, int, str]]:
    """Every labelled value on a sheet for these label words: 'Inspector:'
    with the value to its right (or below), or 'Inspector: MK' in one cell,
    as (value, label row, label column, where the value is). ``skip``: the
    cells of the tables' headers and records (a header is not a label)."""
    out: List[Tuple[Any, int, int, str]] = []
    for (r, c), v in sheet.cells.items():
        if not isinstance(v, str) or (r, c) in skip:
            continue
        t = v.strip()
        if ":" in t:
            head, _, tail = t.partition(":")
            if norm(head) in keys and tail.strip():
                out.append((tail.strip(), r, c, _a1(r, c)))
                continue
        if norm(t) in keys:
            for dc in (1, 2, 3):
                x = sheet.cells.get((r, c + dc))
                if (r, c + dc) in skip:
                    break
                if x is not None and not (isinstance(x, str) and norm(x) in keys):
                    out.append((x, r, c, _a1(r, c + dc)))
                    break
            else:
                x = sheet.cells.get((r + 1, c))
                if x is not None and (r + 1, c) not in skip:
                    out.append((x, r, c, _a1(r + 1, c)))
    return out


def _distance(r: int, c: int, box: Tuple[int, int, int, int]) -> int:
    """How far a label is from a table (rows, columns); a label is looked for
    above and to the left of its table first, so below and right count
    double."""
    r1, c1, r2, c2 = box
    dr = 0 if r1 <= r <= r2 else (r1 - r if r < r1 else 2 * (r - r2))
    dc = 0 if c1 <= c <= c2 else (c1 - c if c < c1 else 2 * (c - c2))
    return dr + dc


def _label_for(hits: List[Tuple[Any, int, int, str]], box: Tuple[int, int, int, int],
               boxes: Sequence[Tuple[int, int, int, int]]
               ) -> Optional[Tuple[Any, str, Tuple[int, int]]]:
    """The labelled value that belongs to this table: one label on the sheet
    belongs to every table on it ('Inspector: MK' over the whole sheet); of
    several (a grid of tables, each with its own), the one nearest this
    table and nearer to it than to any other."""
    if not hits:
        return None
    if len(hits) == 1:
        return hits[0][0], hits[0][3], (hits[0][1], hits[0][2])
    mine = [h for h in hits
            if _distance(h[1], h[2], box) <= min(_distance(h[1], h[2], b) for b in boxes)]
    pick = min(mine or hits, key=lambda h: _distance(h[1], h[2], box))
    return pick[0], pick[3], (pick[1], pick[2])


# ── the model ───────────────────────────────────────────────────────────
def ask_model_headers(t: Table, unmapped_j: Sequence[int], free_fields: Sequence[str],
                      chat: Chat, recipe: Recipe, notes: Sequence[Any] = ()
                      ) -> Tuple[Dict[str, Optional[str]], str]:
    """Which field each of the table's unplaced headers is, with a few of
    its values: only header words, field names and up to MODEL_EXAMPLES
    values per column go to the (local) model. Its answers are guesses."""
    heads = []
    for j in unmapped_j:
        k = t.keys[j][0] or f"(no header, column {j + 1})"
        ex = [cell_text(v) for v in t.column(j) if v is not None][:MODEL_EXAMPLES]
        heads.append((k, t.headers[j] or k, ex))
    labels = [h[0] for h in heads]
    fields = list(free_fields)
    examples = [f"- {json.dumps(k)} -> {json.dumps(f)}"
                for k, f in list(recipe.aliases.items())[:20] if f in fields]
    schema = {"type": "object", "required": ["mappings"], "properties": {"mappings": {
        "type": "array", "items": {"type": "object", "required": ["header", "field"],
                                   "properties": {
                                       "header": {"type": "string", "enum": labels},
                                       "field": {"type": ["string", "null"],
                                                 "enum": fields + [None]}}}}}}
    prompt = ("These are column headers of a table in a spreadsheet, each with a few of its "
              "values. For EACH header give the ONE template field it holds, or null when "
              "none fits or you are unsure. Never give one field to two headers.\n"
              "Reply with JSON only: {\"mappings\": [{\"header\": ..., \"field\": ...}]}.\n\n"
              + ("The user answered these before; follow them:\n" + "\n".join(examples)
                 + "\n\n" if examples else "")
              + __import__("docs_open").notes_text(notes, [h[1] for h in heads])
              + f"Template fields:\n{json.dumps(fields)}\n\nHeaders (with example values):\n"
              + "\n".join(f"- {json.dumps(h[0])} (written {json.dumps(h[1])}): "
                          f"{json.dumps(h[2])}" for h in heads))
    text, model = chat(prompt, schema)
    from council_core.ledger_llm import _field_named, _json
    got = _json(text)
    out: Dict[str, Optional[str]] = {k: None for k in labels}
    pairs = []
    if isinstance(got, dict) and isinstance(got.get("mappings"), list):
        pairs = [(d.get("header"), d.get("field")) for d in got["mappings"]
                 if isinstance(d, dict)]
    used: Set[str] = set()
    for h, f in pairs:
        h = str(h or "").strip()
        name = _field_named(f, fields)
        if h in out and name and name not in used:
            out[h] = name
            used.add(name)
    return out, model


# ── Extract ─────────────────────────────────────────────────────────────
def extract(csv_path, files: Sequence, *, chat: Optional[Chat] = None, progress=None,
            cancel=None, use_cache: bool = True, report=None) -> RecordRun:
    """See _extract: decks are read with this Ledger's text box <-> table
    links (pptx_open.user_pairings)."""
    import pptx_open
    with pptx_open.user_pairings(csv_path):
        return _extract(csv_path, files, chat=chat, progress=progress, cancel=cancel,
                        use_cache=use_cache, report=report)


def _extract(csv_path, files: Sequence, *, chat: Optional[Chat] = None, progress=None,
             cancel=None, use_cache: bool = True, report=None) -> RecordRun:
    """Read every workbook's tables and turn them into records for the
    review: see the module docstring. ``chat``: Council's model, asked about
    headers nothing else placed (once per header layout; its answers are
    remembered in the recipe as guesses until the user accepts them)."""
    import time
    prog = progress or _Quiet()
    recipe = load_recipe(csv_path)
    tt = recipe.table_template
    run = RecordRun(tt)
    files = [Path(f) for f in files]
    run.files = len(files)
    cache = BookCache(cache_dir_for(csv_path) if use_cache else None)
    stopped = (lambda: bool(cancel is not None and cancel.is_set()))
    # 1. read (cached) and find the tables
    prog.stage("reading workbooks", total=len(files), unit="workbooks")
    tables: List[Table] = []
    sheets: Dict[Tuple[str, str], SheetData] = {}
    template_path = Path(tt.path)
    for p in files:
        if stopped():
            run.cancelled = True
            break
        t0 = time.monotonic()
        try:
            same = p.resolve() == template_path.resolve()
        except OSError:
            same = False
        if same:
            prog.advance(1)
            continue
        notes: List[str] = []
        got = cache.get(p)
        if got is not None:
            run.from_cache += 1
        else:
            try:
                # too slow, or stalled the app before: skipped, the run goes
                # on (council_core.file_guard)
                from council_core import file_guard
                got = file_guard.read(p, lambda p=p, notes=notes: read_book(p, notes),
                                      tab="Ledger", should_stop=stopped)
            except Exception as e:                          # noqa: BLE001
                if stopped():
                    run.cancelled = True
                    break
                kind = ("slow" if type(e).__name__ == "Skipped" else
                        getattr(getattr(e, "__cause__", None), "kind", "") or "other")
                run.open_failures.append((p.name, kind))
                run.problems.append(f"{p.name}: {e}")
                prog.advance(1)
                continue
            cache.put(p, got)
        run.problems += notes
        for s in got:
            sheets[(str(p), s.name)] = s
            tables += find_tables(s)
        run.timings.append((p.name, time.monotonic() - t0))
        prog.advance(1, p.name)
    # lookup workbooks named in the recipe but not among the sources
    for fname, rule in recipe.sources.items():
        lp = rule.get("path") if rule.get("kind") == "lookup" else None
        if lp and Path(lp).is_file() and not any(Path(x).resolve() == Path(lp).resolve()
                                                 for x in files):
            try:
                for s in cache.get(Path(lp)) or read_book(lp):
                    tables += [t for t in find_tables(s, sideways=False)]
            except Exception as e:                          # noqa: BLE001
                run.problems.append(f"lookup for {fname}: {Path(lp).name}: {e}")
    if run.cancelled:
        return run
    # what the text documents among the sources say header words mean
    import docs_open
    import source_extract as se
    explanations, said = se.document_explanations(files)
    run.problems += said
    run.explanations = explanations
    explained = docs_open.explanations_for(
        explanations, [(f.name, list(f.keys) + [f.label]) for f in tt.all_fields]) \
        if explanations else {}
    # 2. the user's highlights
    prog.stage("reading your highlights")
    before = json.dumps(asdict(recipe), sort_keys=True, default=str)
    run.learned += learn_colours(tt, tables, recipe, [str(f) for f in files],
                                 list(sheets.values()))
    # 3. lookups: their tables are not records
    lookups: Dict[str, LookupIndex] = {}
    lookup_tables: Set[str] = set()
    for fname, rule in recipe.sources.items():
        if rule.get("kind") != "lookup":
            continue
        idx = LookupIndex(fname, rule)
        for t in tables:
            if t.orient != "rows":
                continue
            if rule.get("file") and not fnmatch.fnmatch(t.file.lower(),
                                                        str(rule["file"]).lower()):
                continue
            heads = [k[0] for k in t.keys]
            if rule.get("key_header") in heads and rule.get("value_header") in heads:
                kj, vj = heads.index(rule["key_header"]), heads.index(rule["value_header"])
                for i, row in enumerate(t.values):
                    if row[kj] is None:
                        continue
                    idx.entries.append((vnorm(row[kj]), row[vj], cell_text(row[kj]),
                                        f"{t.file} › {t.sheet}!{t.where(i, vj)}"))
                idx.tables.append(t.id)
                lookup_tables.add(t.id)
        if not idx.entries:
            run.problems.append(f"{fname}: its lookup table ('{rule.get('key_header')}' in "
                                f"{rule.get('file') or 'any workbook'}) was not found")
        lookups[fname] = idx
    # 4. map each table's columns
    prog.stage("matching columns", total=len(tables), unit="tables")
    col_fields = [f for f in tt.all_fields if not f.formula and not f.is_form
                  and recipe.sources.get(f.name, {}).get("kind") not in
                  ("filename", "sheetname", "name", "lookup", "constant", "near")]
    uses: List[TableUse] = []
    memo: Dict[tuple, List[ColMap]] = {}
    upright: Set[Tuple[str, str]] = set()       # sheets with a table read the right way up
    for t in sorted(tables, key=lambda t: t.orient != "rows"):
        if stopped():
            run.cancelled = True
            break
        if t.id in lookup_tables:
            uses.append(TableUse(t, "lookup", "a lookup table"))
            continue
        if t.orient == "cols" and (t.path, t.sheet) in upright:
            uses.append(TableUse(t, "skipped", "the sheet's table reads the right way up"))
            prog.advance(1)
            continue
        maps = map_table(t, col_fields, recipe, memo, explained)
        if t.orient == "rows" and _qualifies(maps, len(col_fields)):
            upright.add((t.path, t.sheet))
        score = sum(m.score for m in maps) + 0.001 * len(t.rows)
        uses.append(TableUse(t, "records" if _qualifies(maps, len(col_fields)) else
                             "skipped", "" if maps else "no header is a template field",
                             maps, score))
        prog.advance(1)
    _one_per_place(uses)
    _weak_tables(uses, col_fields, run)
    # 4b. learn by example: columns no rule placed, whose values are one
    # field's in a filled copy of the template's table or the parent CSV
    _learn_by_example(uses, col_fields, recipe, csv_path, run)
    # 5. the model, for headers nothing placed (once per header layout)
    if chat is not None and not run.cancelled:
        _ask_model(uses, col_fields, recipe, chat, run, prog, stopped, report, explained)
    # 6. profiles: a column whose values look unlike the field's elsewhere
    _profiles(uses, recipe, run)
    # 7. records
    prog.stage("building records")
    rec_tables = [u for u in uses if u.use == "records"]
    run.files_used = len({u.table.path for u in rec_tables})
    stems = sorted({Path(u.table.file).stem for u in rec_tables})
    sheet_names = sorted({u.table.sheet for u in rec_tables})
    _name_rules(recipe, stems, sheet_names, run)
    join_rates = _joins(recipe, lookups, rec_tables, run)
    order = 0
    for u in sorted(rec_tables, key=lambda u: (_natural(u.table.path), u.table.order,
                                               u.table.header_at)):
        sheet = sheets.get((u.table.path, u.table.sheet))
        here = [x.table for x in uses if x.table.path == u.table.path
                and x.table.sheet == u.table.sheet and x.use != "skipped"] or [u.table]
        for rec in _records(u, tt, recipe, run, lookups, sheet, join_rates, here):
            rec.order = order
            order += 1
            run.records.append(rec)
    # 8. who is who: the same record in two workbooks
    run.key_fields = _key_fields(recipe, tt, rec_tables, run.records)
    _duplicates(run)
    for rec in run.records:
        if rec.accept != "N":
            rec.accept = "Y" if (all(v.sure for v in rec.values.values()) and not rec.checks
                                 and "clash" not in rec.flags and rec.values) else ""
    run.tables = uses
    for f in tt.all_fields:
        if f.formula:
            continue
        if not any(f.name in r.values for r in run.records) and run.records:
            src = recipe.sources.get(f.name, {}).get("kind")
            run.problems.append(
                f"{f.name}: found in no table" + (f" ({src} rule gave nothing)" if src else
                                                  "; if it comes from the file name, a "
                                                  "lookup or a label, set its rule "
                                                  "(ledger_records source …), or highlight "
                                                  "it in a sample with the template's "
                                                  "colour"))
    if json.dumps(asdict(recipe), sort_keys=True, default=str) != before:
        save_recipe(csv_path, recipe)
    return run


def _natural(s: str) -> tuple:
    from sample_ledger import sample_sort_key
    return sample_sort_key(s)


def _one_per_place(uses: List[TableUse]) -> None:
    """The same cells read two ways (sideways, or with its group heading):
    the reading that places the most fields stands."""
    def box(t: Table) -> Tuple[int, int, int, int]:
        r1, r2 = t.header_at, max(t.rows)
        c1, c2 = t.cols[0], t.cols[-1]
        return (r1, c1, r2, c2) if t.orient == "rows" else (c1, r1, c2, r2)
    by_sheet: Dict[Tuple[str, str], List[TableUse]] = {}
    for u in uses:
        if u.use == "records":
            by_sheet.setdefault((u.table.path, u.table.sheet), []).append(u)
    for group in by_sheet.values():
        group.sort(key=lambda u: -u.score)
        kept: List[TableUse] = []
        for u in group:
            a = box(u.table)
            clash = next((k for k in kept if not (a[2] < box(k.table)[0] or
                                                  box(k.table)[2] < a[0] or
                                                  a[3] < box(k.table)[1] or
                                                  box(k.table)[3] < a[1])), None)
            if clash is not None:
                u.use, u.why = "skipped", f"the same cells read better as {clash.table.span()}"
            else:
                kept.append(u)


#: a table read as a filled copy of the template's: this share of its fields
#: placed by their own headers
COPY_SHARE = 0.8


def _learn_by_example(uses: List[TableUse], col_fields, recipe: Recipe, csv_path,
                      run: RecordRun) -> None:
    """Worked answers - the records of tables that are filled copies of the
    template's (its own headers) and the parent CSV's - lined up with every
    other table by the record's key (its name): a column that holds one
    field's values for the same records is that field, as a suggestion
    (how 'example', amber; a Y on its Decisions row teaches it). A key
    column renamed too ('Creature' for Name) is found by the names it holds.
    (council_core.ledger_examples, user 2026-10-21.)"""
    from council_core import ledger_examples as lx
    names = [f.name for f in col_fields]
    if not names:
        return
    copies = [u for u in uses if u.use == "records" and u.table.orient == "rows"
              and sum(1 for m in u.maps if m.how in ("header", "words"))
              >= max(2, COPY_SHARE * len(names))]
    rows: List[Tuple[Dict[str, Any], str]] = []
    for u in copies:
        t = u.table
        for vals in t.values:
            rows.append(({m.field: vals[m.j] for m in u.maps if vals[m.j] is not None},
                         f"{t.file} › {t.sheet}"))
    try:
        _cols, csv_rows = _read_rows(Path(csv_path)) if Path(csv_path).is_file() else ([], [])
    except Exception:                                       # noqa: BLE001
        csv_rows = []
    rows += [({f: r.get(f) for f in names if r.get(f)}, "the parent CSV") for r in csv_rows]
    if not rows:
        return
    key = next((k for k in recipe.key_fields if k in names), None)
    if key is None:
        for f in names:                                     # a name each record has its own
            vals = [v for v in (cell_text(r.get(f)) for r, _ in rows) if v]
            if len(vals) >= 2 and _mostly_text(vals) and len(set(vals)) >= 0.9 * len(vals):
                key = f
                break
    if key is None:
        return
    ex = lx.Examples()
    for r, where in rows:
        if r.get(key) not in (None, ""):
            ex.add(r[key], "", r, where)
    if not ex:
        return
    copy_ids = {u.table.id for u in copies}
    for u in uses:
        t = u.table
        if t.id in copy_ids or u.use == "lookup" or t.orient != "rows" or \
                u.why.startswith("the same cells") or \
                u.why == "the sheet's table reads the right way up":
            continue
        mapped_f = {m.field for m in u.maps}
        mapped_j = {m.j for m in u.maps}
        added: List[ColMap] = []
        key_j = next((m.j for m in u.maps if m.field == key), None)
        if key_j is None:
            hit = lx.key_column([(j, t.column(j)) for j in range(len(t.cols))
                                 if j not in mapped_j], ex)
            if hit is None:
                continue
            key_j = hit[0]
            added.append(ColMap(key, key_j, "example", lx.EXAMPLE_SCORE,
                                f"{lx.EXAMPLE_NOTE}: '{t.headers[key_j]}' holds the "
                                f"{key}s of {hit[1]} of their records"))
            mapped_j.add(key_j)
            mapped_f.add(key)
        keys = t.column(key_j)
        free = [f for f in names if f not in mapped_f]
        for j in range(len(t.cols)):
            if j in mapped_j or not free:
                continue
            got = lx.learn_column(t.column(j), keys, ex, free)
            if got is None:
                continue
            added.append(ColMap(got.field, j, "example", lx.EXAMPLE_SCORE,
                                got.note(t.headers[j] or f"column {j + 1}")))
            free.remove(got.field)
        if not added:
            continue
        u.maps += added
        u.score = sum(m.score for m in u.maps) + 0.001 * len(t.rows)
        if u.use == "skipped" and len(u.maps) >= 2:
            u.use, u.why = "records", "its columns learned from your filled copies (check them)"
        run.learned.append(f"Learned by example in {t.file} › {t.sheet}: " + ", ".join(
            f"'{t.headers[m.j]}' is {m.field}" for m in added))


def _weak_tables(uses: List[TableUse], col_fields, run: RecordRun) -> None:
    """A table that names only a few of many fields (an unrelated workbook's
    'Name' and 'Date') is read, but its records wait for a Y."""
    n = len(col_fields)
    for u in uses:
        if u.use != "records":
            continue
        sure = sum(1 for m in u.maps if m.how in SURE_HOWS)
        if sure >= 4 or sure >= 0.5 * n:
            continue
        t = u.table
        u.check = f"T|{t.signature}"
        if u.check not in run.decisions:
            run.decisions[u.check] = Decision(
                u.check, "table", "", f"a table with only {len(u.maps)} of the {n} fields "
                f"read as records ({', '.join(m.field for m in u.maps)}); headers: "
                + ", ".join(h for h in t.headers if h)[:200], f"{t.file} › {t.sheet}")


def _ask_model(uses, col_fields, recipe, chat, run, prog, stopped, report,
               explained=None) -> None:
    asked = 0
    tried: Set[str] = set()
    names = [f.name for f in col_fields]
    files_with_records = {u.table.path for u in uses if u.use == "records"}
    for u in uses:
        if stopped():
            run.cancelled = True
            return
        t = u.table
        if u.use == "lookup" or t.orient != "rows":
            continue
        if u.use == "skipped" and (t.path in files_with_records or len(t.rows) < 2
                                   or len(t.cols) < 2):
            continue
        sig = t.signature
        if sig in tried:
            continue
        tried.add(sig)
        mapped_j = {m.j for m in u.maps}
        free = [n for n in names if n not in {m.field for m in u.maps}]
        todo = [j for j in range(len(t.cols)) if j not in mapped_j and t.keys[j][0]
                and t.keys[j][0] not in recipe.model_cache.get(sig, {})]
        if not free or not todo:
            continue
        if asked >= MODEL_MAX_LAYOUTS:
            run.problems.append(f"Council's model was asked about {MODEL_MAX_LAYOUTS} header "
                                "layouts; the rest wait for the next Extract")
            break
        prog.update(f"asking Council's model about {t.file} › {t.sheet}")
        try:
            got, model = ask_model_headers(t, todo, free, chat, recipe,
                                           notes=run.explanations)
        except Exception as e:                              # noqa: BLE001
            from council_core import ledger_llm
            if isinstance(e, ledger_llm.ModelUnavailable):
                break                                     # said once, by one_shot
            line = f"Council's model did not answer about {t.file} › {t.sheet}: {e}"
            if isinstance(e, ledger_llm.ModelFailedAgain):
                line = f"Council's model did not answer about {t.file} › {t.sheet} " \
                       "(the same error as above)"
            run.problems.append(line)
            if report is None:
                break
            continue
        asked += 1
        run.model = model
        recipe.model_cache.setdefault(sig, {}).update(got)
    run.model_asked = asked
    if asked:
        memo: Dict[tuple, List[ColMap]] = {}
        for u in uses:
            if u.use == "lookup" or u.why == "the sheet's table reads the right way up":
                continue
            u.maps = map_table(u.table, col_fields, recipe, memo, explained)
            u.score = sum(m.score for m in u.maps) + 0.001 * len(u.table.rows)
            if u.use == "skipped" and _qualifies(u.maps, len(col_fields)) and \
                    not u.why.startswith("the same cells"):
                u.use, u.why = "records", ""
            elif u.use == "skipped" and u.maps and any(m.how == "model" for m in u.maps) \
                    and not u.why.startswith("the same cells"):
                u.use, u.why = "records", "Council's model placed its columns (check them)"
        _one_per_place(uses)


def _profiles(uses: List[TableUse], recipe: Recipe, run: RecordRun) -> None:
    """Numbers where a field's other columns hold words (or the other way
    round), or words outside the few a field takes ('M', 'F'): that column
    is a guess, even with the right header."""
    stats: Dict[str, List[float]] = {}
    cats: Dict[str, Dict[str, int]] = {}
    for u in uses:
        if u.use != "records":
            continue
        for m in u.maps:
            if m.how not in SURE_HOWS:
                continue
            vals = [v for v in u.table.column(m.j) if v is not None and not is_nullish(v)]
            if not vals:
                continue
            stats.setdefault(m.field, []).append(sum(map(is_number, vals)) / len(vals))
            stats[m.field].append(len(vals))
            c = cats.setdefault(m.field, {})
            for v in vals[:500]:
                if not is_number(v):
                    c[vnorm(v)] = c.get(vnorm(v), 0) + 1
    prof = dict(recipe.profiles)
    for f, xs in stats.items():
        shares, ns = xs[0::2], xs[1::2]
        n = sum(ns)
        share = sum(s * k for s, k in zip(shares, ns)) / max(1, n)
        old = prof.get(f, {})
        if old.get("n", 0) > n:
            continue
        cat = cats.get(f, {})
        prof[f] = {"numeric": round(share, 3), "n": int(n),
                   "cats": sorted(cat) if 0 < len(cat) <= 12 and n >= 20 else []}
    recipe.profiles = prof
    for u in uses:
        if u.use != "records":
            continue
        for m in u.maps:
            p = prof.get(m.field)
            k = u.table.keys[m.j][0]
            if not p or p.get("n", 0) < 10 or f"{m.field}|{k}" in recipe.profile_ok:
                continue
            vals = [v for v in u.table.column(m.j) if v is not None and not is_nullish(v)]
            if len(vals) < 3:
                continue
            share = sum(map(is_number, vals)) / len(vals)
            odd = ""
            if abs(share - p["numeric"]) > 0.5:
                odd = ("numbers" if share > 0.5 else "words") + " where " + m.field + \
                      " elsewhere has " + ("numbers" if p["numeric"] > 0.5 else "words")
            elif p.get("cats") and share < 0.5:
                words = [vnorm(v) for v in vals if not is_number(v)]
                out = sum(1 for w in words if w not in set(p["cats"]))
                if words and out > 0.5 * len(words):
                    odd = f"values unlike {m.field}'s ({', '.join(p['cats'][:6])})"
            if odd:
                m.check = run.decide(Decision(
                    f"V|{m.field}|{k}", "values", m.field,
                    f"column '{u.table.headers[m.j]}' was read as {m.field}, but it holds "
                    f"{odd}",
                    ", ".join(cell_text(v) for v in vals[:4]))) or ""
                run.decisions[m.check].records -= 1      # counted per record below
                m.note = (m.note + "; " if m.note else "") + odd


def _name_rules(recipe: Recipe, stems: List[str], sheet_names: List[str],
                run: RecordRun) -> None:
    for fname, rule in recipe.sources.items():
        kind = rule.get("kind")
        if kind == "name":
            kind = "filename" if len(stems) > 1 or len(sheet_names) <= 1 else "sheetname"
            rule["kind"] = kind
        if kind in ("filename", "sheetname") and not rule.get("pattern"):
            pat = infer_name_rule(stems if kind == "filename" else sheet_names)
            if pat:
                rule["pattern"] = pat
                rule.setdefault("confirmed", False)
            else:
                run.problems.append(f"{fname}: from the {kind.replace('name', ' name')}, "
                                    "but more than one is needed to tell what part of it "
                                    "(set it with an example: ledger_records source … "
                                    "--example)")


def _joins(recipe: Recipe, lookups: Dict[str, LookupIndex], rec_tables: List[TableUse],
           run: RecordRun) -> Dict[str, float]:
    """Which record field each lookup joins on: the one whose values most
    often find an entry (exact or whole-word 'starts with')."""
    rates: Dict[str, float] = {}
    for fname, idx in lookups.items():
        rule = idx.rule
        if not idx.entries:
            continue
        cands: Dict[str, List[int]] = {}
        for u in rec_tables:
            for m in u.maps:
                vals = [v for v in u.table.column(m.j)[:400] if v is not None]
                if not vals or not _mostly_text(vals):
                    continue
                hit = sum(1 for v in vals
                          if idx.match(v, {})[1] in ("exact", "starts with"))
                c = cands.setdefault(m.field, [0, 0])
                c[0] += hit
                c[1] += len(vals)
        if rule.get("join") and rule["join"] in cands:
            hit, n = cands[rule["join"]]
            rates[fname] = hit / max(1, n)
            continue
        if not cands:
            run.problems.append(f"{fname}: no text column of the records to look it up by")
            continue
        best = max(cands, key=lambda k: cands[k][0] / max(1, cands[k][1]))
        rate = cands[best][0] / max(1, cands[best][1])
        if rate < 0.3:
            run.problems.append(f"{fname}: no record field matches its lookup's keys "
                                f"(best: {best}, {rate:.0%})")
            continue
        rule["join"] = best
        rule["join_auto"] = True
        rates[fname] = rate
        run.learned.append(f"{fname}: looked up by each record's {best} ({rate:.0%} found)")
    return rates


def _records(u: TableUse, tt: TableTemplate, recipe: Recipe, run: RecordRun,
             lookups: Dict[str, LookupIndex], sheet: Optional[SheetData],
             join_rates: Dict[str, float], tables_here: Sequence[Table] = ()
             ) -> Iterable[Record]:
    t = u.table
    mapped = {m.field: m for m in u.maps}
    # per sheet: the labelled values beside the table (form fields, and
    # fields this table has no column for)
    near: Dict[str, Tuple[Any, str, Tuple[int, int]]] = {}
    skip = set()
    if sheet is not None:
        for other in tables_here:
            skip |= other.cell_set()
        for f in tt.all_fields:
            if f.name in mapped or f.formula or recipe.sources.get(f.name, {}).get("kind") \
                    in ("filename", "sheetname", "lookup", "constant", "near"):
                continue
            keys = set(f.keys or [f.key]) | {norm(f.name)} | {
                k for k, v in recipe.aliases.items() if v == f.name}
            keys.discard("")
            got = _label_for(_labelled_values(sheet, keys, skip), t.box(),
                             [x.box() for x in tables_here])
            if got is not None:
                near[f.name] = got
    where0 = f"{t.file} › {t.sheet}!"
    for i, row in enumerate(t.values):
        rec = Record(t.file, t.path, t.sheet, t.rows[i], t.id)
        if u.check:
            rec.checks.append(u.check)
            run.decisions[u.check].records += 1
        if sheet is not None and sheet.document:
            # a text document is looser than a workbook: its records wait
            rec.checks.append(run.decide(Decision(
                f"D|{t.file}", "document", "",
                f"records read from the text document {t.file}", t.sheet)))
        for m in u.maps:
            v = row[m.j]
            if v is None:
                continue
            check = m.check
            if not check and m.how not in SURE_HOWS:
                hk = t.keys[m.j][0]
                check = run.decide(Decision(
                    f"C|{m.field}|{hk}", "column", m.field,
                    f"column '{t.headers[m.j]}' read as {m.field}"
                    + (f" ({m.note})" if m.note else f" ({m.how})"),
                    ", ".join(cell_text(x) for x in t.column(m.j)[:4] if x is not None)))
            elif check:
                run.decisions[check].records += 1
            rec.values[m.field] = Val(v, m.how, where0 + t.where(i, m.j), m.note, check)
        for fname, (v, at, cell) in near.items():
            check = ""
            if sheet is not None and cell in sheet.guessed:
                # a slide's text box put over this table by where it sits
                check = run.decide(Decision(
                    f"P|{t.file}|{t.sheet}|{fname}", "place", fname,
                    f"'{cell_text(sheet.cells.get(cell))}' was matched to {t.sheet}'s table "
                    "by its place on the slide; group the text box with its table in "
                    "PowerPoint (or give them the same colour) to say so",
                    cell_text(v)))
            rec.values[fname] = Val(v, "label", f"{where0}{at}",
                                    "the value beside its label on the sheet", check)
        for f in tt.all_fields:
            rule = recipe.sources.get(f.name)
            if not rule or f.formula:
                continue
            kind = rule.get("kind")
            if kind in ("filename", "sheetname"):
                name = Path(t.file).stem if kind == "filename" else t.sheet
                val = name_value(rule.get("pattern", ""), name)
                if val is None:
                    continue
                check = "" if rule.get("confirmed") else run.decide(Decision(
                    f"S|{f.name}", "rule", f.name,
                    rule.get("why") or f"{f.name} taken from the {kind.replace('name', ' name')}"
                    f" (pattern {rule.get('pattern')})",
                    f"{name} → {val}"))
                rec.values[f.name] = Val(val, "file name" if kind == "filename"
                                         else "sheet name", f"{t.file} › {t.sheet}",
                                         rule.get("pattern", ""), check)
            elif kind == "constant":
                rec.values[f.name] = Val(rule.get("value", ""), "constant")
            elif kind == "near" and sheet is not None:
                top, left = t.box()[:2]
                at = (top + int(rule.get("dr", 0)), left + int(rule.get("dc", 0)))
                v = sheet.cells.get(at)
                if v is not None and at not in skip:
                    rec.values[f.name] = Val(v, "near", f"{where0}{_a1(*at)}",
                                             "the cell in its place over the table")
            elif kind == "lookup" and f.name in lookups:
                idx = lookups[f.name]
                join = rule.get("join")
                if not join or join not in rec.values:
                    continue
                key_v = rec.values[join].value
                entry, how, note = idx.match(key_v, recipe.key_aliases.get(f.name, {}))
                if entry is None:
                    if note:
                        rec.notes.append(f"{f.name}: {note} in the lookup")
                    continue
                check = ""
                if how in ("similar", "ambiguous"):
                    check = run.decide(Decision(
                        f"L|{f.name}|{vnorm(key_v)}", "lookup", f.name,
                        f"'{cell_text(key_v)}' looked up as '{entry[2]}' ({note})",
                        cell_text(entry[1])))
                elif rule.get("join_auto") and not rule.get("join_confirmed") and \
                        join_rates.get(f.name, 1.0) < 0.95:
                    check = run.decide(Decision(
                        f"S|{f.name}", "rule", f.name,
                        f"{f.name} looked up by each record's {join} "
                        f"({join_rates.get(f.name, 0):.0%} of them found)",
                        f"{cell_text(key_v)} → {cell_text(entry[1])}"))
                rec.values[f.name] = Val(entry[1], how, entry[3],
                                         f"looked up by {join} '{cell_text(key_v)}'"
                                         + (f": {note}" if note else ""), check)
        if any(cell_text(v.value) for k, v in rec.values.items()
               if k in mapped):
            yield rec


#: how a value that names a whole table (not one record) was found
CONTEXT_HOWS = ("label", "near", "file name", "sheet name")


def _key_fields(recipe: Recipe, tt: TableTemplate, rec_tables: List[TableUse],
                records: Sequence[Record] = ()) -> List[str]:
    if recipe.key_fields:
        return list(recipe.key_fields)
    keys = [f for f, r in recipe.sources.items() if r.get("kind") in ("filename", "sheetname")]
    # what names each block of a grid ('Tank: 3' above it), when the blocks differ
    per_table: Dict[str, Dict[str, str]] = {}
    for rec in records:
        for f, v in rec.values.items():
            if v.how in CONTEXT_HOWS and f not in keys:
                per_table.setdefault(f, {})[rec.table] = cell_text(v.value)
    for f in tt.all_fields:
        if f.name in per_table and len(set(per_table[f.name].values())) > 1:
            keys.append(f.name)
    for f in tt.fields:
        cols = [(u, m) for u in rec_tables for m in u.maps if m.field == f.name]
        if not cols:
            continue
        # names every table keeps apart (the same name in two workbooks is
        # what the key is for)
        ok = True
        for u, m in cols:
            vals = [cell_text(v) for v in u.table.column(m.j)]
            vals = [v for v in vals if v]
            if not vals or not _mostly_text(vals) or len(set(vals)) < 0.9 * len(vals):
                ok = False
                break
        if ok:
            keys.append(f.name)
            break
    return keys


def _duplicates(run: RecordRun) -> None:
    """The same record (its key fields) from two workbooks: identical, the
    first is kept and the others left out (said); different, every one is
    red for the user to pick. Two tables of ONE workbook (a grid of blocks,
    a sheet per tank) are different things that may share names ('P1')."""
    keys = [k for k in run.key_fields]
    if not keys:
        return
    groups: Dict[Tuple[str, ...], List[Record]] = {}
    for rec in run.records:
        k = rec.key(keys)
        if not any(k):
            continue
        groups.setdefault(k, []).append(rec)
    for k, recs in groups.items():
        if len(recs) < 2:
            continue
        first = recs[0]
        for other in recs[1:]:
            if other.path == first.path:
                continue
            # a clash is two values for one field; a field one of them lacks
            # is no disagreement (a filled copy with more columns than a log)
            both = set(first.values) & set(other.values)
            differ = [f for f in both if cell_text(first.values[f].value)
                      != cell_text(other.values[f].value)]
            if not differ and set(other.values) <= set(first.values):
                other.accept = "N"
                other.flags.append("duplicate")
                other.notes.append(f"the same record as {first.file} › {first.sheet} row "
                                   f"{first.row}: left out")
            elif not differ:
                other.notes.append(f"also in {first.file} › {first.sheet} row {first.row}, "
                                   "which agrees where both have a value")
            else:
                for r in (first, other):
                    if "clash" not in r.flags:
                        r.flags.append("clash")
                    r.notes.append("differs from " + (other if r is first else first).file
                                   + f" ({', '.join(sorted(differ)[:5])}): accept one")


# ── the review workbook ─────────────────────────────────────────────────
README_LINES = [
    "Review of extracted records (table Ledger)",
    "",
    "1. Records: one row per record found. Accept = Y goes into the parent CSV; N does not.",
    "   Rows already Y were clear matches. Amber cells are Council's guesses (a similar",
    "   header, Council's model, a rule it inferred, a fuzzy lookup); red rows clash with",
    "   the same record in another workbook - accept one. You may correct any value here.",
    "2. Decisions: one row per guess, for every record it touches. Y there accepts that",
    "   guess for all of them (a record left blank on Records is then taken when all its",
    "   guesses have a Y) and teaches it for every later Extract; N rejects it.",
    "3. Tables: every table Council found, and whether it was read as records, used as a",
    "   lookup or skipped (and why).",
    "4. Save this workbook, then press Apply in the Ledger tab (or run:",
    "   python -m council_core.ledger_records apply \"<this file>\").",
]


def write_review(out, run: RecordRun, csv_path, sources: Sequence) -> Path:
    """The review workbook (streamed: thousands of records stay quick)."""
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    out = Path(out)
    tt = run.tt
    names = [f.name for f in tt.all_fields if not f.formula]
    wb = openpyxl.Workbook(write_only=True)
    amber, red, grey = (PatternFill("solid", fgColor=c) for c in (AMBER, RED, GREY))
    bold = Font(bold=True)

    def head(ws, cols, widths=None):
        row = []
        for c in cols:
            cell = WriteOnlyCell(ws, value=c)
            cell.font = bold
            row.append(cell)
        if widths:
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[_col_letter(i)].width = w
        ws.append(row)

    ws = wb.create_sheet("Read me")
    for line in README_LINES + ["", f"Parent CSV: {csv_path}",
                                f"Made: {datetime.now():%Y-%m-%dT%H:%M:%S}",
                                f"Workbooks: {run.files} ({run.files_used} with records, "
                                f"{run.from_cache} read from the cache)"]:
        ws.append([line])
    if run.learned:
        ws.append([])
        ws.append(["Learned in this Extract:"])
        for line in run.learned:
            ws.append(["  " + line])
    if run.key_fields:
        ws.append([])
        ws.append(["Records are told apart by: " + ", ".join(run.key_fields)])

    meta_cols = ["Accept", "Source", "Sheet", "Table", "Row"]
    tail = ["Notes", "_table", "_checks"]
    chunk = SHEET_ROWS
    parts = [run.records[i:i + chunk] for i in range(0, max(1, len(run.records)), chunk)] \
        or [[]]
    for pi, part in enumerate(parts):
        ws = wb.create_sheet("Records" if pi == 0 else f"Records {pi + 1}")
        ws.freeze_panes = "F2"
        ws.column_dimensions[_col_letter(len(meta_cols) + len(names) + 2)].hidden = True
        ws.column_dimensions[_col_letter(len(meta_cols) + len(names) + 3)].hidden = True
        head(ws, meta_cols + names + tail, [8, 28, 14, 8, 6] + [14] * len(names) + [50])
        for rec in part:
            checks = sorted({v.check for v in rec.values.values() if v.check} | set(rec.checks))
            row = [rec.accept, rec.file, rec.sheet, table_anchor(rec.table), rec.row]
            for n in names:
                v = rec.values.get(n)
                cell = WriteOnlyCell(ws, value=_cell_value(v.value) if v else None)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"                 # a text, never a formula
                if v is not None and (v.check or "clash" in rec.flags):
                    cell.fill = red if "clash" in rec.flags else amber
                    text = f"{v.how}: {v.where}" + (f"\n{v.note}" if v.note else "")
                    cell.comment = Comment(text[:1000], "Council")
                elif "duplicate" in rec.flags:
                    cell.fill = grey
                row.append(cell)
            flags = [f"!{f}" for f in rec.flags]
            row += ["; ".join(rec.notes)[:1000], rec.table, ";".join(checks + flags)]
            ws.append(row)

    ws = wb.create_sheet("Decisions")
    head(ws, ["Id", "Accept", "Kind", "Field", "What Council did", "Example values",
              "Records"], [24, 8, 9, 18, 80, 40, 9])
    for d in sorted(run.decisions.values(), key=lambda d: (d.kind, d.field, d.id)):
        ws.append([d.id, d.accept, d.kind, d.field, d.detail, d.examples, d.records])

    ws = wb.create_sheet("Tables")
    head(ws, ["Source", "Sheet", "Cells", "Read as", "Records", "Columns → fields", "Why"],
          [28, 14, 14, 9, 8, 90, 50])
    for u in run.tables[:TABLES_LISTED]:
        t = u.table
        maps = ", ".join(f"'{t.headers[m.j]}' → {m.field}" + ("" if m.how in SURE_HOWS
                                                              else f" ({m.how})")
                         for m in sorted(u.maps, key=lambda m: m.j))
        if u.use == "skipped" and not maps:
            maps = "headers: " + ", ".join(h for h in t.headers if h)[:300]
        ws.append([t.file, t.sheet, t.span() + (" (sideways)" if t.orient == "cols" else ""),
                   u.use, len(t.rows) if u.use == "records" else "", maps[:2000], u.why])
    if len(run.tables) > TABLES_LISTED:
        ws.append([f"… and {len(run.tables) - TABLES_LISTED} more tables"])

    ws = wb.create_sheet("Problems")
    head(ws, ["Problem"], [120])
    for p in run.problems:
        ws.append([p])

    # what the knowledge graph and the Ledger's source scan look for in a
    # review: it must never be read back as a source workbook
    ws = wb.create_sheet("Values")
    head(ws, ["Accept", "Part", "Item", "Field", "Value", "Confidence", "Matched by",
              "Source label", "Source"])
    ws.append(["(a table Ledger's review: see the Records sheet)"])

    ws = wb.create_sheet("_run")
    ws.sheet_state = "hidden"
    meta = {"mode": "records", "csv": str(csv_path), "sources": [str(s) for s in sources],
            "fields": names, "key_fields": run.key_fields,
            "made": f"{datetime.now():%Y-%m-%dT%H:%M:%S}"}
    ws.append([json.dumps(meta)])
    wb.properties.keywords = REVIEW_MARK
    tmp = out.with_name(out.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, out)
    return out


def table_anchor(table_id: str) -> str:
    """Where a table starts on its sheet ('E12'), from Table.id."""
    try:
        _f, _s, orient, at, col = table_id.rsplit("|", 4)
        r, c = int(at), int(col)
        return _a1(r, c) if orient == "rows" else _a1(c, r)
    except ValueError:
        return ""


def _cell_value(v: Any) -> Any:
    return xl_text(v) if isinstance(v, str) else v


#: the keyword a table Ledger's review carries in its document properties
REVIEW_MARK = "council-ledger-records-review"


def is_record_review(path) -> bool:
    """A table Ledger's review, told from its document properties in the
    zip (no workbook is opened: Apply asks on a worker thread)."""
    import zipfile
    try:
        with zipfile.ZipFile(Path(path)) as z:
            return REVIEW_MARK.encode() in z.read("docProps/core.xml")
    except Exception:                                       # noqa: BLE001
        return False


# ── Apply ───────────────────────────────────────────────────────────────
@dataclass
class ApplyResult:
    accepted: int                 # values
    records: int
    rows: int                     # in the parent CSV after
    csv_result: str
    learned: Dict[str, str] = field(default_factory=dict)
    ignored: List[str] = field(default_factory=list)
    filled: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


def _yes(v: Any) -> Optional[bool]:
    s = str(v or "").strip().upper()
    return True if s in ("Y", "YES", "TRUE", "1") else False if s in ("N", "NO", "FALSE", "0") \
        else None


def apply_review(review, fill_dir=None, progress=None) -> ApplyResult:
    """The accepted records into the parent CSV, what the Decisions taught
    into the recipe, and (with ``fill_dir``) the template filled with every
    record of the Ledger."""
    import openpyxl
    prog = progress or _Quiet()
    review = Path(review)
    try:
        wb = openpyxl.load_workbook(review, read_only=True, data_only=True)
    except Exception as e:                                  # noqa: BLE001
        raise LedgerError(f"could not open the review {review.name}: {e}") from e
    try:
        # read whole (list): a suspended read-only generator crashes lxml
        # when it is collected after the workbook closed
        meta = json.loads(list(wb["_run"].iter_rows(max_row=1, max_col=1,
                                                    values_only=True))[0][0])
        if meta.get("mode") != "records":
            raise LedgerError(f"{review.name} is not a table Ledger's review")
        csv_path = Path(meta["csv"])
        recipe = load_recipe(csv_path)
        tt = recipe.table_template
        decisions: Dict[str, Optional[bool]] = {}
        details: Dict[str, Tuple[str, str]] = {}
        if "Decisions" in wb.sheetnames:
            for row in wb["Decisions"].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    decisions[str(row[0])] = _yes(row[1])
                    details[str(row[0])] = (str(row[4] or ""), str(row[5] or ""))
        taken: List[Dict[str, str]] = []
        cells = 0
        result = ApplyResult(0, 0, 0, "")
        prog.stage("reading the review")
        for sheet in [s for s in wb.sheetnames if s == "Records" or s.startswith("Records ")]:
            rows = iter(list(wb[sheet].iter_rows(values_only=True)))
            header = [str(h) if h is not None else "" for h in next(rows, [])]
            ix = {h: i for i, h in enumerate(header)}
            for row in rows:
                if not row or all(v is None for v in row):
                    continue
                get = (lambda name: row[ix[name]] if name in ix and ix[name] < len(row)
                       else None)
                ok = _yes(get("Accept"))
                checks = [c for c in str(get("_checks") or "").split(";") if c]
                if ok is None:
                    ok = (not any(c.startswith("!") for c in checks)
                          and all(decisions.get(c) is True for c in checks))
                if not ok:
                    continue
                rec = {f.name: cell_text(get(f.name)) for f in tt.all_fields
                       if f.name in ix}
                rec[SOURCE_COL] = str(get("Source") or "")
                rec[SHEET_COL] = str(get("Sheet") or "")
                rec[ROW_COL] = cell_text(get("Row"))
                rec[TABLE_COL] = str(get("Table") or "")
                cells += sum(1 for f in tt.all_fields if rec.get(f.name))
                taken.append(rec)
    finally:
        wb.close()
    # what the decisions teach
    learned = _learn_decisions(recipe, decisions, details)
    result.learned = learned
    _profiles_from_rows(recipe, tt, taken)
    save_recipe(csv_path, recipe)
    # the parent CSV: one row per record, the same record replaced
    prog.stage("writing the parent CSV")
    cols, rows = _read_rows(csv_path) if csv_path.exists() else (csv_columns(tt), [])
    want = csv_columns(tt)
    cols = [c for c in want if c not in (SOURCE_COL, UPDATED_COL)] + \
        [c for c in cols if c not in want] + [SOURCE_COL, UPDATED_COL]
    keys = list(meta.get("key_fields") or recipe.key_fields)

    # a record is the same one again by its workbook, sheet and key fields; when
    # that is not enough (the blocks of a grid share names), by its table's
    # place too - never by a guess that would write one record over another
    def base(r: Dict[str, str]) -> Tuple[str, ...]:
        k = tuple(vnorm(r.get(x, "")) for x in keys)
        where = (vnorm(r.get(SOURCE_COL, "")), vnorm(r.get(SHEET_COL, "")))
        if keys and all(k):
            return ("k",) + where + k
        return ("r",) + where + (str(r.get(TABLE_COL, "")), str(r.get(ROW_COL, "")))

    def full(r: Dict[str, str]) -> Tuple[str, ...]:
        b = base(r)
        return b if b[0] == "r" else b + ("@", str(r.get(TABLE_COL, "")))
    from collections import Counter
    many = {k for k, n in Counter(base(r) for r in taken).items() if n > 1}
    many |= {k for k, n in Counter(base(r) for r in rows).items() if n > 1}

    def ident(r: Dict[str, str]) -> Tuple[str, ...]:
        b = base(r)
        return full(r) if b in many else b
    at = {ident(r): i for i, r in enumerate(rows)}
    now = f"{datetime.now():%Y-%m-%d %H:%M:%S}"
    added = updated = 0
    for rec in taken:
        rec[UPDATED_COL] = now
        k = ident(rec)
        if k in at:
            rows[at[k]] = {**rows[at[k]], **rec}
            updated += 1
        else:
            at[k] = len(rows)
            rows.append(rec)
            added += 1
    _write_rows(csv_path, cols, rows)
    result.accepted, result.records, result.rows = cells, len(taken), len(rows)
    result.csv_result = (f"{csv_path.name} ({added} records added, {updated} updated, "
                         f"{len(rows)} in all)")
    if fill_dir:
        prog.stage("filling the template")
        out = Path(fill_dir) / f"{Path(tt.path).stem}_filled.xlsx"
        result.warnings += fill_records(tt, rows, out)
        result.filled.append(str(out))
    return result


def _learn_decisions(recipe: Recipe, decisions: Dict[str, Optional[bool]],
                     details: Dict[str, Tuple[str, str]]) -> Dict[str, str]:
    learned: Dict[str, str] = {}
    for did, ok in decisions.items():
        if ok is None:
            continue
        kind, _, rest = did.partition("|")
        fname, _, key = rest.partition("|")
        if kind == "C":
            if ok:
                recipe.aliases[key] = fname
                if fname in recipe.not_aliases.get(key, []):
                    recipe.not_aliases[key].remove(fname)
                learned[key] = fname
            else:
                lst = recipe.not_aliases.setdefault(key, [])
                if fname not in lst:
                    lst.append(fname)
                if recipe.aliases.get(key) == fname:
                    del recipe.aliases[key]
                for sig in recipe.model_cache.values():
                    if sig.get(key) == fname:
                        sig[key] = None
        elif kind == "V" and ok:
            if f"{fname}|{key}" not in recipe.profile_ok:
                recipe.profile_ok.append(f"{fname}|{key}")
        elif kind == "S":
            rule = recipe.sources.get(fname)
            if rule is None:
                continue
            if rule.get("kind") == "lookup":
                rule["join_confirmed"] = bool(ok)
                if not ok:
                    rule["join"] = ""
            elif ok:
                rule["confirmed"] = True
                learned[f"{fname} rule"] = rule.get("pattern", rule.get("kind", ""))
            else:
                recipe.sources.pop(fname, None)
        elif kind == "L" and ok:
            value = details.get(did, ("", ""))
            # the entry it was matched to is in the detail: "'x' looked up as 'Y' (...)"
            m = re.search(r"looked up as '(.*?)'", value[0]) if value[0] else None
            if m:
                recipe.key_aliases.setdefault(fname, {})[key] = vnorm(m.group(1))
    return learned


def _profiles_from_rows(recipe: Recipe, tt: TableTemplate, rows: List[Dict[str, str]]) -> None:
    for f in tt.fields:
        vals = [r.get(f.name, "") for r in rows if r.get(f.name, "")]
        if len(vals) < 10:
            continue
        old = recipe.profiles.get(f.name, {})
        if old.get("n", 0) > len(vals):
            continue
        words = sorted({vnorm(v) for v in vals if not is_number(v)})
        recipe.profiles[f.name] = {
            "numeric": round(sum(map(is_number, vals)) / len(vals), 3), "n": len(vals),
            "cats": words if 0 < len(words) <= 12 and len(vals) >= 20 else []}


# ── filling the template ────────────────────────────────────────────────
def fill_records(tt: TableTemplate, rows: Sequence[Dict[str, str]], out_path) -> List[str]:
    """A copy of the template with one table row per record: rows past the
    ones the template laid out take their formatting (banded rows keep
    their bands), formulas are carried down, data validation and
    conditional formatting are stretched over the new rows, and when
    something sits under the table rows are inserted for it. Form fields are
    filled when every record agrees. Returns warnings."""
    import openpyxl
    from openpyxl.formula.translate import Translator
    src, out_path = Path(tt.path), Path(out_path)
    if out_path.resolve() == src.resolve():
        raise LedgerError("refusing to write over the template itself")
    warnings: List[str] = []
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    ws = wb[tt.sheet] if tt.sheet in wb.sheetnames else wb.worksheets[0]
    n = len(rows)
    first = tt.first_row
    laid = tt.styled_rows
    if tt.content_below and n > laid:
        extra = n - laid
        at = first + laid
        _insert_rows(ws, at, extra, warnings)
    last = first + max(n, 1) - 1
    period = _band(ws, tt)
    for i, rec in enumerate(rows):
        r = first + i
        if i >= laid:
            sr = first + (i % period if period else 0)
            if laid and sr > tt.end_row:
                sr = first
            for f in tt.fields:
                s = ws.cell(sr, f.col)
                d = ws.cell(r, f.col)
                if s.has_style:
                    d._style = copy(s._style)
            if ws.row_dimensions[sr].height:
                ws.row_dimensions[r].height = ws.row_dimensions[sr].height
        for f in tt.fields:
            cell = ws.cell(r, f.col)
            if f.formula:
                if i > 0:
                    try:
                        cell.value = Translator(f.formula, origin=_a1(first, f.col)) \
                            .translate_formula(_a1(r, f.col))
                    except Exception:                       # noqa: BLE001
                        cell.value = f.formula
                continue
            v = rec.get(f.name, "")
            cell.value = _typed(xl_text(v)) if v not in ("", None) else None
    for f in tt.forms:
        vals = {rec.get(f.name, "") for rec in rows} - {""}
        if len(vals) == 1:
            ws.cell(f.form_row, f.col).value = _typed(vals.pop())
        elif len(vals) > 1:
            warnings.append(f"{f.name}: the records differ ({len(vals)} values), so the "
                            "template's cell is left empty")
    if n > laid:
        _stretch(ws, tt, last, warnings)
    try:
        wb.properties.keywords = FILLED_MARK
    except Exception:                                       # noqa: BLE001
        pass
    tmp = out_path.with_name(out_path.stem + ".tmp" + out_path.suffix)
    wb.save(tmp)
    os.replace(tmp, out_path)
    return warnings


def _band(ws, tt: TableTemplate) -> int:
    """How many laid-out rows repeat (2 for banded rows), 0 when none."""
    rows = tt.styled_rows
    if rows < 1:
        return 0

    def sig(r):
        return tuple(colour_key(ws.cell(r, f.col).fill) for f in tt.fields)
    for p in (1, 2, 3):
        if rows >= 2 * p and all(sig(tt.first_row + i) == sig(tt.first_row + i % p)
                                 for i in range(rows)):
            return p
    return rows


def _insert_rows(ws, at: int, amount: int, warnings: List[str]) -> None:
    """Rows inserted for the records; merged cells under them move down
    with them (openpyxl moves values and styles, not merges)."""
    from openpyxl.worksheet.cell_range import CellRange
    merged = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in merged:
        if m.min_row >= at:
            ws.merged_cells.remove(m)
    ws.insert_rows(at, amount)
    for m in merged:
        if m.min_row >= at:
            m.shift(row_shift=amount)
            ws.merge_cells(m.coord)
    if ws.conditional_formatting or ws.data_validations.dataValidation:
        warnings.append(f"{amount} rows were inserted under the table: check the "
                        "validation and conditional formats below it")


def _stretch(ws, tt: TableTemplate, last: int, warnings: List[str]) -> None:
    """Data validation and conditional formats that cover the first record
    row of a column, stretched down to the last record."""
    cols = {f.col for f in tt.fields}
    for dv in list(ws.data_validations.dataValidation):
        try:
            for rng in list(dv.sqref.ranges):
                if rng.min_row <= tt.first_row <= rng.max_row and \
                        any(rng.min_col <= c <= rng.max_col for c in cols) and rng.max_row < last:
                    dv.sqref.add(f"{_a1(tt.first_row, rng.min_col)}:{_a1(last, rng.max_col)}")
        except Exception:                                   # noqa: BLE001
            continue
    try:
        cf = ws.conditional_formatting
        adds = []
        for entry in list(cf):
            for rng in list(entry.sqref.ranges):
                if rng.min_row <= tt.first_row <= rng.max_row and rng.max_row < last:
                    adds.append((f"{_a1(rng.max_row + 1, rng.min_col)}:"
                                 f"{_a1(last, rng.max_col)}", list(entry.rules)))
        for rng, rules in adds:
            for rule in rules:
                cf.add(rng, rule)
    except Exception:                                       # noqa: BLE001
        warnings.append("conditional formatting could not be stretched over the new rows")


# ── Teach: the highlights of a few samples, and what Extract will read ──
def teach(csv_path, workbooks: Sequence, progress=None, cancel=None) -> str:
    """Learn the user's highlights in these sample workbooks (the same fill
    colour in the template and a sample links them; learned into the
    recipe) and say, in plain lines, what the next Extract reads from them.
    ``progress``: a Ledger tab listener (line, fraction)."""
    from council_core.ledger_ingest import collect_sources
    from progress_report import Progress
    recipe = load_recipe(csv_path)
    files, _ = collect_sources(workbooks, template_path=recipe.template.get("path"))
    if not files:
        raise LedgerError("no sample workbook to teach from")
    with Progress("ledger teach", listener=progress) as prog:
        run = extract(csv_path, files, progress=prog, cancel=cancel)
    return describe_run(run) + _pairings_text(files, csv_path)


def _pairings_text(files: Sequence[Path], csv_path=None) -> str:
    """Which text box went with which table on each slide of the decks, and
    how that was told, for the user to check (and fix by grouping)."""
    import pptx_open
    out: List[str] = []
    for f in files:
        if Path(f).suffix.lower() not in pptx_open.SUFFIXES:
            continue
        try:
            with pptx_open.user_pairings(csv_path):
                deck = pptx_open.read_pptx(f, want_pictures=False)
        except pptx_open.PptxError:
            continue
        if deck.pairings:
            out.append(f"Text boxes and tables in {Path(f).name}:")
            out += [f"  {p.describe()}" for p in deck.pairings]
    return ("\n" + "\n".join(out)) if out else ""


def describe_run(run: RecordRun, examples: int = 3) -> str:
    lines: List[str] = []
    if run.learned:
        lines += run.learned
    else:
        lines.append("No highlight in these workbooks has a colour of the template's "
                     "fields: columns are matched by their header words.")
    recipe_fields = {f.name for f in run.tt.all_fields}
    for u in run.tables:
        t = u.table
        if u.use == "records":
            got = ", ".join(f"{m.field} ← '{t.headers[m.j]}'" + (
                "" if m.how in SURE_HOWS and not m.check else f" ({m.how}, to check)")
                for m in sorted(u.maps, key=lambda m: m.j))
            lines.append(f"{t.file} › {t.sheet} {t.span()}: {len(t.rows)} records; {got}")
        elif u.use == "lookup":
            lines.append(f"{t.file} › {t.sheet} {t.span()}: a lookup table, not records")
    found = {k for r in run.records for k in r.values}
    missing = sorted(recipe_fields - found - {f.name for f in run.tt.fields if f.formula})
    if missing:
        lines.append("Not found in these workbooks: " + ", ".join(missing))
    for rec in run.records[:max(0, examples)]:
        lines.append(f"  e.g. {rec.file} row {rec.row}: " + ", ".join(
            f"{k} {cell_text(v.value)}" + ("?" if v.check else "")
            for k, v in rec.values.items())[:400])
    for p in run.problems[:20]:
        lines.append(f"  ! {p}")
    return "\n".join(lines)


# ── setting a field's rule by hand (CLI, tests, a future window) ─────────
def set_source(csv_path, field_name: str, kind: str, *, pattern: str = "",
               example: Optional[Tuple[str, str]] = None, lookup_file: str = "",
               key_header: str = "", value_header: str = "", join: str = "",
               value: str = "", confirmed: bool = True) -> Dict[str, Any]:
    """Say where a field comes from when it is not a column: 'filename' /
    'sheetname' (``pattern``, or ``example`` = (name, value)), 'lookup' (a
    workbook and its key and value headers; ``join`` the record field it is
    looked up by, '' = found by Council), 'constant' (``value``), or
    'column' (back to reading it from a column)."""
    recipe = load_recipe(csv_path)
    tt = recipe.table_template
    names = [f.name for f in tt.all_fields]
    if field_name not in names:
        raise LedgerError(f"{field_name!r} is not a field; fields: {names}")
    if kind == "column":
        recipe.sources.pop(field_name, None)
    elif kind in ("filename", "sheetname"):
        if example:
            pattern = name_rule_from_example(Path(example[0]).stem
                                             if kind == "filename" else example[0], example[1])
        recipe.sources[field_name] = {"kind": kind, "pattern": pattern,
                                      "confirmed": confirmed, "how": "your answer"}
    elif kind == "lookup":
        p = Path(lookup_file)
        recipe.sources[field_name] = {
            "kind": "lookup", "file": p.name, "path": str(p) if p.is_file() else "",
            "key_header": norm(key_header), "value_header": norm(value_header or key_header),
            "join": join, "join_confirmed": bool(join), "confirmed": True,
            "how": "your answer"}
    elif kind == "constant":
        recipe.sources[field_name] = {"kind": "constant", "value": value, "confirmed": True}
    else:
        raise LedgerError(f"unknown kind {kind!r}")
    save_recipe(csv_path, recipe)
    return recipe.sources.get(field_name, {})


# ── command line ────────────────────────────────────────────────────────
def main(argv=None) -> int:
    import argparse
    ap = argparse.ArgumentParser(prog="python -m council_core.ledger_records",
                                 description="The Ledger for table templates (one row "
                                             "per record).")
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("create", help="template -> parent CSV")
    c.add_argument("template")
    c.add_argument("--csv")
    c.add_argument("--key", action="append", help="a field that tells records apart")
    e = sub.add_parser("extract", help="workbooks/folders -> review workbook")
    e.add_argument("csv")
    e.add_argument("sources", nargs="+")
    e.add_argument("--review")
    e.add_argument("--model", help="Council's model: 'ollama:<name>' or a GGUF path")
    t = sub.add_parser("teach", help="learn the highlights in sample workbooks")
    t.add_argument("csv")
    t.add_argument("samples", nargs="+")
    a = sub.add_parser("apply", help="review -> parent CSV (+ filled template)")
    a.add_argument("review")
    a.add_argument("--fill-dir")
    s = sub.add_parser("source", help="where a field comes from when it is not a column")
    s.add_argument("csv")
    s.add_argument("field")
    s.add_argument("kind", choices=["filename", "sheetname", "lookup", "constant", "column"])
    s.add_argument("--pattern", default="")
    s.add_argument("--example", nargs=2, metavar=("NAME", "VALUE"))
    s.add_argument("--file", default="")
    s.add_argument("--key-header", default="")
    s.add_argument("--value-header", default="")
    s.add_argument("--join", default="")
    s.add_argument("--value", default="")
    args = ap.parse_args(argv)
    try:
        if args.cmd == "create":
            from council_core.ledger_ingest import default_csv_for
            csv = args.csv or default_csv_for(args.template)
            print(create_ledger(args.template, csv, args.key).describe())
        elif args.cmd == "extract":
            from council_core import ledger_ingest
            s = ledger_ingest.run_extract(args.csv, args.sources, args.review,
                                          use_images=False,
                                          assist=args.model or False)
            print(s.describe())
        elif args.cmd == "teach":
            print(teach(args.csv, args.samples))
        elif args.cmd == "apply":
            r = apply_review(args.review, args.fill_dir)
            print(f"{r.records} records ({r.accepted} values) -> {r.csv_result}")
            for f in r.filled:
                print(f"  filled {f}")
            for w in r.warnings:
                print(f"  ! {w}")
        elif args.cmd == "source":
            print(set_source(args.csv, args.field, args.kind, pattern=args.pattern,
                             example=tuple(args.example) if args.example else None,
                             lookup_file=args.file, key_header=args.key_header,
                             value_header=args.value_header, join=args.join,
                             value=args.value))
    except LedgerError as e:
        print(f"error: {e}")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
