"""
council_core.ledger_linking — the click-to-link window's data side (no Qt):
the user clicks where a value goes in the template, then where it comes
from in a sample workbook or deck, and the Ledger learns it (user,
2026-10-21: "instead of highlighting we can add some easier linking
system?").

    template_view(csv)        the template's sheets, and which field each cell is
    sample_view(path, csv)    a sample's sheets, the tables found on them, and
                              (for a deck) each text box's table
    link_cell(...)            what a click on a sample cell means for a field,
                              as a Link with a plain sentence
    link_file_name(...)       the field is (part of) the file's name
    link_lookup(...)          the field is looked up in this column
    link_pairing(...)         this text box goes with these tables
    save(csv, links)          the links into what Extract reads

A sample-ID Ledger (sample_ledger) learns a label: the same as an answer on
the review's Labels to map (<parent>.aliases.json). A table Ledger
(ledger_records) learns a column's header words, a label beside a table, a
cell's place over each table, the file name or a lookup (its recipe). A
deck's text box <-> table links go to <parent>.pairings.json for both.
Every link is the user's own answer: what it places is pre-accepted.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import pptx_open
from council_core import ledger_records as lr
from sample_ledger import LedgerError, TemplateSchema, normalize_label, schema_path_for

#: what the window shows of a sheet at most (it reads all of it to link)
VIEW_ROWS, VIEW_COLS = 3000, 120


@dataclass
class ViewSheet:
    name: str
    cells: Dict[Tuple[int, int], Any]
    #: (row, col) -> the field that cell is (the template), or the table and
    #: column it is in (a sample: 'T3:2')
    fields: Dict[Tuple[int, int], str] = field(default_factory=dict)
    rows: int = 0
    cols: int = 0


@dataclass
class TemplateView:
    kind: str                     # 'records' (a table Ledger) | 'samples'
    sheets: List[ViewSheet]
    fields: List[str]             # every field, in template order


@dataclass
class SampleView:
    path: str
    sheets: List[ViewSheet]
    tables: Dict[str, List[lr.Table]]              # sheet name -> its tables
    data: Dict[str, lr.SheetData]                  # sheet name -> its cells
    pairings: List[pptx_open.Pairing] = field(default_factory=list)


@dataclass
class Link:
    kind: str                     # label | column | near | filename | lookup | pairing
    field: str
    detail: Dict[str, Any]
    sentence: str

    @property
    def key(self) -> Tuple:
        """Two links with the same key are the same answer (the newer wins)."""
        if self.kind == "pairing":
            return ("pairing", self.detail["deck"], self.detail["slide"], self.detail["key"])
        return (self.kind, self.field)


def ledger_kind(csv_path) -> str:
    if lr.is_record_ledger(csv_path):
        return "records"
    if schema_path_for(csv_path).is_file():
        return "samples"
    raise LedgerError(f"{Path(csv_path).name} has no schema beside it: run Create first")


# ── the template ────────────────────────────────────────────────────────
def _cell(a1: str) -> Optional[Tuple[int, int]]:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    try:
        col, row = coordinate_from_string(a1)
        return row, column_index_from_string(col)
    except Exception:                                       # noqa: BLE001
        return None


def template_view(csv_path) -> TemplateView:
    """The template's sheets as the window shows them, each cell that is a
    field's marked with it: a table Ledger's header and record cells, a
    sample Ledger's value cells (one per flaw for a grid)."""
    import openpyxl
    kind = ledger_kind(csv_path)
    marks: Dict[Tuple[str, int, int], str] = {}
    if kind == "records":
        tt = lr.load_recipe(csv_path).table_template
        path, names = tt.path, [f.name for f in tt.all_fields]
        last = max(tt.end_row, tt.first_row + 3)
        for f in tt.fields:
            for r in range(tt.header_row, last + 1):
                marks.setdefault((tt.sheet, r, f.col), f.name)
        for f in tt.forms:
            marks[(tt.sheet, f.form_row, f.col)] = f.name
            marks.setdefault((tt.sheet, f.form_row, f.col - 1), f.name)
    else:
        schema = TemplateSchema.load(schema_path_for(csv_path))
        path, names = schema.template_path, [f.name for f in schema.fields]
        for f in schema.fields:
            for a1 in f.cells.values():
                rc = _cell(a1)
                if rc:
                    marks.setdefault((f.sheet, rc[0], rc[1]), f.name)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:                                  # noqa: BLE001
        raise LedgerError(f"could not open the template {Path(path).name}: {e}") from e
    sheets: List[ViewSheet] = []
    try:
        for ws in wb.worksheets:
            cells: Dict[Tuple[int, int], Any] = {}
            for r, row in enumerate(list(ws.iter_rows(min_row=1, max_row=VIEW_ROWS,
                                                      max_col=VIEW_COLS, values_only=True)),
                                    start=1):
                for c, v in enumerate(row, start=1):
                    if v is not None and str(v).strip():
                        cells[(r, c)] = v
            mine = {(r, c): f for (s, r, c), f in marks.items() if s == ws.title}
            keys = list(cells) + list(mine)
            sheets.append(ViewSheet(ws.title, cells, mine,
                                    max((r for r, _ in keys), default=1) + 3,
                                    max((c for _, c in keys), default=1) + 2))
    finally:
        wb.close()
    return TemplateView(kind, sheets, names)


# ── a sample ────────────────────────────────────────────────────────────
def sample_view(path, csv_path=None) -> SampleView:
    """A sample workbook or deck as the window shows it: its sheets (a deck:
    each slide's tables, laid out as Extract reads them), the tables found
    on each, and a deck's text box <-> table pairings."""
    path = Path(path)
    with pptx_open.user_pairings(csv_path):
        data = lr.read_book(path)
        pairings: List[pptx_open.Pairing] = []
        if path.suffix.lower() in pptx_open.SUFFIXES:
            pairings = pptx_open.read_pptx(path, want_pictures=False).pairings
    sheets, tables, by_name = [], {}, {}
    for s in data:
        found = [t for t in lr.find_tables(s, sideways=False)]
        tables[s.name] = found
        by_name[s.name] = s
        marks = {}
        for ti, t in enumerate(found):
            for j, c in enumerate(t.cols):
                for r in [t.header_at] + t.rows:
                    marks[(r, c)] = f"T{ti}:{j}"
        keys = list(s.cells)
        sheets.append(ViewSheet(s.name, dict(s.cells), marks,
                                min(VIEW_ROWS, max((r for r, _ in keys), default=1) + 3),
                                min(VIEW_COLS, max((c for _, c in keys), default=1) + 2)))
    return SampleView(str(path), sheets, tables, by_name, pairings)


def _table_at(view: SampleView, sheet: str, r: int, c: int
              ) -> Optional[Tuple[lr.Table, int]]:
    for t in view.tables.get(sheet, []):
        if c in t.cols and (r == t.header_at or r in t.rows):
            return t, t.cols.index(c)
    return None


def _label_of(cells: Dict[Tuple[int, int], Any], r: int, c: int) -> Optional[str]:
    """The words that name the clicked cell: itself when it is words (a
    label, 'Inspector: JB' -> 'Inspector'), else the nearest words to its
    left on its row, else above it in its column."""
    v = cells.get((r, c))
    if lr.is_text(v):
        t = str(v).strip()
        if lr._is_label(t) and ":" in t.rstrip(":"):
            return t.split(":", 1)[0].strip()
        return t.rstrip(": ")
    for cc in range(c - 1, max(0, c - 8), -1):
        w = cells.get((r, cc))
        if lr.is_text(w):
            return str(w).strip().rstrip(": ")
    for rr in range(r - 1, max(0, r - 30), -1):
        w = cells.get((rr, c))
        if lr.is_text(w):
            return str(w).strip().rstrip(": ")
    return None


def link_cell(kind: str, field_name: str, view: SampleView, sheet: str, r: int, c: int
              ) -> Link:
    """What clicking this sample cell for this field means. A sample Ledger
    learns the words that name the cell (they are found wherever they are
    written). A table Ledger learns the column's header when the cell is in
    a table, the label when it is labelled, else the cell's place over the
    nearest table (a grid's block title). Raises LedgerError when nothing
    there can be learned."""
    cells = view.data[sheet].cells if sheet in view.data else {}
    name = Path(view.path).name
    if kind == "records":
        hit = _table_at(view, sheet, r, c)
        if hit is not None:
            t, j = hit
            key = t.keys[j][0]
            if not key:
                raise LedgerError("that column has no header to learn; click a column "
                                  "with a header")
            return Link("column", field_name, {"header_key": key, "header": t.headers[j]},
                        f"Column '{t.headers[j]}' is {field_name}, in every workbook "
                        f"(learned from {name} › {sheet})")
        v = cells.get((r, c))
        if v is None:
            raise LedgerError("that cell is empty")
        label = None
        if lr.is_text(v) and lr._is_label(str(v)):
            label = _label_of(cells, r, c)
        elif lr.is_text(cells.get((r, c - 1))) and lr._is_label(str(cells.get((r, c - 1)))):
            label = str(cells[(r, c - 1)]).strip().rstrip(": ")
        if label:
            return Link("label", field_name, {"label": label},
                        f"The value labelled '{label}' beside each table is {field_name}")
        mine = view.tables.get(sheet, [])
        if not mine:
            raise LedgerError("no table on this sheet to place it by: click a table's "
                              "column, or a labelled value")
        t = min(mine, key=lambda t: lr._distance(r, c, t.box()))
        top, left = t.box()[:2]
        return Link("near", field_name, {"dr": r - top, "dc": c - left},
                    f"{field_name} is the cell {lr._offset_words(r - top, c - left)} each "
                    f"table (here '{lr.cell_text(v)[:40]}')")
    label = _label_of(cells, r, c)
    if not label:
        raise LedgerError("nothing names that cell: click its label or its column's header")
    return Link("label", field_name, {"label": label},
                f"A value labelled '{label}' is {field_name}, in every workbook "
                f"(learned from {name} › {sheet})")


def template_fields(sheet: ViewSheet, cells: Sequence[Tuple[int, int]]) -> List[str]:
    """The fields of the template cells drawn over, in reading order (left to
    right, top to bottom); cells that are no field's are passed over."""
    return [sheet.fields[rc] for rc in sorted(cells) if rc in sheet.fields]


def link_cells(kind: str, fields: Sequence[str], view: SampleView, sheet: str,
               cells: Sequence[Tuple[int, int]]) -> List[Link]:
    """Template cells drawn over, linked to the same number of sample cells
    drawn over, one to one in reading order (user, 2026-10-21: "a portion of a
    row that lines up to a portion of a row in the template"): a row of the
    template's headers to a row of a sample's, say. When the template cells
    are one field repeated (a grid row across its flaws) they count once
    each. Raises LedgerError when the numbers differ."""
    order = sorted(cells)
    fields = list(fields)
    if len(fields) != len(order):
        distinct = list(dict.fromkeys(fields))
        if len(distinct) == len(order):
            fields = distinct
        elif len(distinct) == 1:
            fields = distinct * len(order)                 # one field, many places of it
        else:
            raise LedgerError(
                f"{len(fields)} template cell(s) ({len(distinct)} field(s)) and "
                f"{len(order)} sample cell(s): draw as many on each side")
    out: List[Link] = []
    for f, (r, c) in zip(fields, order):
        link = link_cell(kind, f, view, sheet, r, c)
        link.detail["_at"] = (view.path, sheet, r, c)
        out.append(link)
    # one answer per field: the same field drawn twice keeps its last
    seen: Dict[Tuple, Link] = {}
    for link in out:
        seen[link.key] = link
    return list(seen.values())


def link_file_name(field_name: str, file_name: str, value: str) -> Link:
    stem = Path(file_name).stem
    rule = lr.name_rule_from_example(stem, value)
    return Link("filename", field_name, {"name": file_name, "value": value, "pattern": rule},
                f"{field_name} is taken from the file name: '{stem}' gives '{value}'")


def link_lookup(field_name: str, view: SampleView, sheet: str, r: int, c: int) -> Link:
    hit = _table_at(view, sheet, r, c)
    if hit is None:
        raise LedgerError("click a column of a table to look the value up in")
    t, j = hit
    key_j = j if lr._mostly_text(t.column(j)) else next(
        (jj for jj in range(len(t.cols)) if lr._mostly_text(t.column(jj))), j)
    return Link("lookup", field_name,
                {"file": t.file, "path": t.path, "key_header": t.keys[key_j][0],
                 "value_header": t.keys[j][0]},
                f"{field_name} is looked up in {t.file} › {t.sheet}: the row whose "
                f"'{t.headers[key_j]}' each record's name starts with, its "
                f"'{t.headers[j]}'")


def link_pairing(view: SampleView, pairing: pptx_open.Pairing, numbers: Sequence[int]) -> Link:
    to = ", ".join(pairing.choices[i - 1] for i in numbers) if numbers \
        else "every table on the slide"
    return Link("pairing", "", {"deck": Path(view.path).name, "slide": pairing.slide,
                                "key": pairing.key, "tables": list(numbers)},
                f"Slide {pairing.slide}: '{pairing.text}' goes with {to}")


# ── saving ──────────────────────────────────────────────────────────────
def save(csv_path, links: Sequence[Link]) -> List[str]:
    """The links into what the next Extract reads; returns what was saved,
    one line each. Links are the user's own answers."""
    kind = ledger_kind(csv_path)
    said: List[str] = []
    labels = [k for k in links if k.kind == "label"]
    pairs = [k for k in links if k.kind == "pairing"]
    if kind == "samples":
        bad = [k for k in links if k.kind not in ("label", "pairing")]
        if bad:
            raise LedgerError("this Ledger reads one sample per workbook: link a label or a "
                              "column header (a file name, lookup or place over a table is "
                              "for a table Ledger)")
        if labels:
            import source_extract as se
            mapping, ignore = se.load_aliases(csv_path)
            for k in labels:
                mapping[k.detail["label"]] = k.field
                ignore = [x for x in ignore
                          if normalize_label(x) != normalize_label(k.detail["label"])]
                said.append(k.sentence)
            se.save_aliases(csv_path, mapping, ignore)
    else:
        recipe = lr.load_recipe(csv_path)
        for k in links:
            if k.kind in ("column", "label"):
                words = lr.norm(k.detail.get("header_key") or k.detail.get("label"))
                recipe.aliases[words] = k.field
                if k.field in recipe.not_aliases.get(words, []):
                    recipe.not_aliases[words].remove(k.field)
                if k.kind == "column" and recipe.sources.get(k.field, {}).get("kind") in (
                        "name", "filename", "sheetname", "near"):
                    recipe.sources.pop(k.field, None)     # it is a column after all
            elif k.kind == "near":
                recipe.sources[k.field] = {"kind": "near", "dr": k.detail["dr"],
                                           "dc": k.detail["dc"], "confirmed": True,
                                           "how": "your link"}
            elif k.kind == "filename":
                recipe.sources[k.field] = {"kind": "filename", "pattern": k.detail["pattern"],
                                           "confirmed": True, "how": "your link"}
            elif k.kind == "lookup":
                recipe.sources[k.field] = {
                    "kind": "lookup", "file": k.detail["file"], "path": k.detail["path"],
                    "key_header": k.detail["key_header"],
                    "value_header": k.detail["value_header"], "join": "",
                    "confirmed": True, "how": "your link"}
            if k.kind != "pairing":
                said.append(k.sentence)
        lr.save_recipe(csv_path, recipe)
    if pairs:
        stored = pptx_open.load_pairings(csv_path)
        for k in pairs:
            d = k.detail
            stored.setdefault(d["deck"], {}).setdefault(str(d["slide"]), {})[d["key"]] = \
                list(d["tables"])
            said.append(k.sentence)
        pptx_open.save_pairings(csv_path, stored)
    return said
