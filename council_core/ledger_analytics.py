"""
council_core.ledger_analytics — analysis of what a Ledger holds (user,
2026-10-30: "add a data analytics function for example the ledger output
file we can do some analysis on it"; a window inside the Ledger, no split of
fields into inputs and results, no spec limits).

Read only. Everything comes from files the Ledger already writes: the parent
CSV (a row per part, or per part per flaw), its schema (each field's unit,
section and whether it is per flaw) and the saved Ledger's groups.

    load(csv)                  -> Data: the rows as a DataFrame, each field's kind
    coverage(data, group)      -> what is filled and what is missing, per field
    missing_by_part(...)       -> the parts with the most fields left blank
    summary(data, group, per)  -> count / mean / spread / percentiles, or top values
    trend(data, field, ...)    -> the field over a date (or part order), with
                                  individuals-chart control limits and the run
                                  rules each point breaks
    correlation(data, group)   -> every pair of number fields: r, how sure (p,
                                  Benjamini-Hochberg across the pairs), how many
                                  values; and the pairs strongest first
    outliers(data, group)      -> values far from the rest of their field
                                  (median and MAD, so one wild value does not
                                  hide another)
    export(data, out)          -> <csv stem>_analysis.xlsx, a sheet per view
                                  and a chart per numeric field

A group's fields that are not tracked are left out of it; a Yes / No field is
counted as Yes / No. "Per part" takes a per-flaw field's average over the
part's flaws, as the template's own AVERAGE columns do.

Nothing here imports Qt.
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import numpy as np
import pandas as pd

from council_core import ledger_groups as lg
from sample_ledger import (FILLED_MARK, META_COLS, SOURCE_COL, LedgerError, ParentCSV,
                           mark_filled, sample_sort_key)

NUMBER, DATE, TEXT, EMPTY = "number", "date", "text", "empty"
PER_ROW, PER_PART = "row", "part"
#: a field is a number (or a date) when this share of its filled cells reads as one
KIND_SHARE = 0.8
#: robust z (0.6745 * distance from the median / MAD) above which a value is an outlier
OUTLIER_Z = 3.5
#: values a pair of fields needs in common before it is correlated
CORR_MIN_N = 5
#: charts written into the analysis workbook, at most
CHART_LIMIT = 40

_NUM = re.compile(r"^\s*([-+]?(?:\d+(?:,\d{3})*(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?)\s*([^\W\d_]\S{0,5}|%|°\S{0,4})?\s*$")


@dataclass
class FieldInfo:
    name: str
    kind: str                     # NUMBER | DATE | TEXT
    unit: str = ""
    section: str = ""
    per_item: bool = False


@dataclass
class Data:
    csv: str
    frame: pd.DataFrame           # the parent CSV's rows, text as written
    id_field: str
    item_field: str
    fields: Dict[str, FieldInfo]  # analysed fields, in template order
    groups: List[lg.Group] = field(default_factory=list)
    _facts: Optional[Dict[str, lg.PartFacts]] = None

    @property
    def parts(self) -> List[str]:
        ids = [p for p in self.frame[self.id_field].unique() if str(p).strip()]
        return sorted(ids, key=sample_sort_key)

    @property
    def date_fields(self) -> List[str]:
        return [f for f, i in self.fields.items() if i.kind == DATE]

    @property
    def number_fields(self) -> List[str]:
        return [f for f, i in self.fields.items() if i.kind == NUMBER]

    def group(self, name: str = "") -> Optional[lg.Group]:
        key = (name or "").strip().lower()
        return next((g for g in self.groups if g.name.strip().lower() == key), None)

    def group_names(self) -> List[str]:
        return [g.name for g in self.groups]


def to_number(v: Any) -> Optional[float]:
    """'12.5', '1,200', '12.5 um', '-3' -> a float; None when not a number."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return None if isinstance(v, float) and math.isnan(v) else float(v)
    m = _NUM.match(str(v))
    if not m:
        return None
    try:
        return float(m.group(1).replace(",", ""))
    except ValueError:
        return None


def _kind(name: str, values: Sequence[str]) -> str:
    filled = [v for v in values if str(v).strip()]
    if not filled:
        return EMPTY
    nums = sum(to_number(v) is not None for v in filled)
    dates = sum(lg.parse_date(v) is not None for v in filled)
    if "date" in name.lower() and dates >= KIND_SHARE * len(filled):
        return DATE
    if nums >= KIND_SHARE * len(filled):
        return NUMBER
    if dates >= KIND_SHARE * len(filled):
        return DATE
    return TEXT


def _varies_within_parts(frame: pd.DataFrame, id_field: str, col: str) -> bool:
    """A column the schema does not know: per flaw when a part's rows differ
    (a blank on some rows included)."""
    for _part, sub in frame.groupby(id_field, sort=False):
        if sub[col].nunique() > 1:
            return True
    return False


def load(csv) -> Data:
    """The parent CSV, ready to analyse."""
    ledger = ParentCSV(csv)
    frame = pd.DataFrame(ledger.rows, columns=ledger.columns).fillna("").astype(str)
    schema = ledger.schema
    item_field = schema.item_field if schema else ""
    skip = {ledger.id_field, item_field, *META_COLS}
    info: Dict[str, FieldInfo] = {}
    by_name = {f.name: f for f in schema.fields} if schema else {}
    for col in ledger.columns:
        if col in skip or not col:
            continue
        # an inspector's own reading ('… [JB]') is its field's twin
        tf = by_name.get(col) or by_name.get(base_field(col))
        per_item = (tf.per_item if tf else
                    bool(item_field) and _varies_within_parts(frame, ledger.id_field, col))
        info[col] = FieldInfo(col, _kind(col, frame[col].tolist()),
                              unit=tf.unit if tf else "", section=tf.section if tf else "",
                              per_item=per_item)
    return Data(csv=str(csv), frame=frame, id_field=ledger.id_field,
                item_field=item_field, fields=info, groups=lg.load_groups(csv))


# ── slicing ─────────────────────────────────────────────────────────────
def _facts(data: Data) -> Dict[str, lg.PartFacts]:
    if data._facts is None:
        data._facts = lg.facts_from_ledger(ParentCSV(data.csv))
    return data._facts


def rows(data: Data, group: str = "") -> pd.DataFrame:
    """The rows of ``group`` ('' or 'all parts': every row); its fields not
    tracked blanked, its Yes / No fields written Yes / No."""
    g = data.group(group) if group and group.strip().lower() != lg.ALL_PARTS else None
    df = data.frame
    if g is None:
        return df
    keep = {p.lower() for p in lg.members(g, _facts(data), data.csv)}
    df = df[df[data.id_field].str.lower().isin(keep)].copy()
    for f in data.fields:
        mode = g.mode(f)
        if mode == lg.OFF:
            df[f] = ""
        elif mode == lg.YES_NO:
            df[f] = [lg.as_yes_no(v, g.yes_no.get(f, "")) for v in df[f]]
    return df


def fields_of(data: Data, group: str = "") -> List[str]:
    """The fields analysed for ``group``: not those it does not track."""
    g = data.group(group)
    return [f for f in data.fields if not g or g.mode(f) != lg.OFF]


def _kind_in(data: Data, f: str, group: str) -> str:
    g = data.group(group)
    if g and g.mode(f) == lg.YES_NO:
        return TEXT
    return data.fields[f].kind


def per_part(data: Data, df: pd.DataFrame) -> pd.DataFrame:
    """One row per part: a number field's average over its flaws, any other
    field's first filled value."""
    out = []
    for part, sub in df.groupby(data.id_field, sort=False):
        row: Dict[str, Any] = {data.id_field: part}
        for f, info in data.fields.items():
            vals = [v for v in sub[f] if str(v).strip()]
            if info.kind == NUMBER:
                nums = [n for n in map(to_number, vals) if n is not None]
                row[f] = (sum(nums) / len(nums)) if nums else ""
            else:
                row[f] = vals[0] if vals else ""
        row[SOURCE_COL] = ", ".join(sorted({s for s in sub.get(SOURCE_COL, []) if s}))
        out.append(row)
    res = pd.DataFrame(out, columns=[data.id_field, *data.fields, SOURCE_COL])
    if len(res):
        res = res.iloc[sorted(range(len(res)),
                              key=lambda i: sample_sort_key(res.iloc[i][data.id_field]))]
    return res.reset_index(drop=True)


def _part_level(data: Data, df: pd.DataFrame, f: str) -> List[str]:
    """A part-level field's values, once per part (each flaw row repeats them)."""
    out, seen = [], set()
    for part, v in zip(df[data.id_field], df[f]):
        key = str(part).lower()
        if key not in seen and str(v).strip():
            seen.add(key)
            out.append(v)
    return out


def _numbers(series) -> np.ndarray:
    return np.array([n for n in map(to_number, series) if n is not None], dtype=float)


# ── views ───────────────────────────────────────────────────────────────
def coverage(data: Data, group: str = "") -> pd.DataFrame:
    """Per field: how many it could hold (parts for a part-level field, the
    flaws with any per-flaw value for a per-flaw one), how many are filled,
    and the share."""
    df = rows(data, group)
    parts = df[data.id_field].str.lower().nunique()
    # a flaw row is real when any per-flaw field has a value on it: the
    # template's 30 flaw slots are not all used
    flaw_fields = [f for f, i in data.fields.items() if i.per_item]
    real = (df[flaw_fields].apply(lambda c: c.str.strip() != "").any(axis=1)
            if flaw_fields else pd.Series(False, index=df.index))
    out = []
    for f in fields_of(data, group):
        info = data.fields[f]
        if info.per_item:
            expected = int(real.sum())
            filled = int((df[f].str.strip() != "").sum())
        else:
            expected = parts
            filled = int(df[df[f].str.strip() != ""][data.id_field].str.lower().nunique())
        out.append({"Field": f, "Section": info.section, "Per": "flaw" if info.per_item
                    else "part", "Kind": info.kind, "Filled": filled, "Could hold": expected,
                    "Filled %": round(100.0 * filled / expected, 1) if expected else 0.0})
    return pd.DataFrame(out, columns=["Field", "Section", "Per", "Kind", "Filled",
                                      "Could hold", "Filled %"])


def missing_by_part(data: Data, group: str = "") -> pd.DataFrame:
    """Per part: fields left blank on every one of its rows, most first.
    Fields no part has a value in are not counted."""
    df = rows(data, group)
    fs = [f for f in fields_of(data, group) if (df[f].str.strip() != "").any()]
    out = []
    for part, sub in df.groupby(data.id_field, sort=False):
        blank = [f for f in fs if not (sub[f].str.strip() != "").any()]
        out.append({"Part": part, "Fields blank": len(blank), "Of": len(fs),
                    "Which": ", ".join(blank[:12]) + (" …" if len(blank) > 12 else "")})
    res = pd.DataFrame(out, columns=["Part", "Fields blank", "Of", "Which"])
    if len(res):
        order = sorted(range(len(res)), key=lambda i: (-res.iloc[i]["Fields blank"],
                                                       sample_sort_key(res.iloc[i]["Part"])))
        res = res.iloc[order].reset_index(drop=True)
    return res


def summary(data: Data, group: str = "", per: str = PER_ROW) -> pd.DataFrame:
    """Per field: a number's count, mean, spread and percentiles (with its
    unit); a text's count, distinct values and the most common ones."""
    df = rows(data, group)
    if per == PER_PART:
        df = per_part(data, df)
    out = []
    for f in fields_of(data, group):
        info = data.fields[f]
        kind = _kind_in(data, f, group)
        filled = ([v for v in df[f] if str(v).strip()] if info.per_item or per == PER_PART
                  else _part_level(data, df, f))
        row: Dict[str, Any] = {"Field": f, "Unit": info.unit, "Kind": kind, "Count": len(filled)}
        if kind == NUMBER:
            a = _numbers(filled)
            if len(a):
                q = np.percentile(a, [25, 50, 75])
                row.update({"Mean": a.mean(), "Std dev": a.std(ddof=1) if len(a) > 1 else 0.0,
                            "Min": a.min(), "25%": q[0], "Median": q[1], "75%": q[2],
                            "Max": a.max()})
        elif kind == DATE:
            ds = sorted(d for d in map(lg.parse_date, filled) if d)
            if ds:
                row.update({"Min": ds[0].isoformat(), "Max": ds[-1].isoformat()})
        elif kind == TEXT:
            counts = pd.Series(filled, dtype=str).value_counts()
            row["Distinct"] = int(len(counts))
            row["Most common"] = ", ".join(f"{k} ({v})" for k, v in counts.head(5).items())
        out.append(row)
    cols = ["Field", "Unit", "Kind", "Count", "Mean", "Std dev", "Min", "25%", "Median",
            "75%", "Max", "Distinct", "Most common"]
    res = pd.DataFrame(out, columns=cols, dtype=object)
    for c in ("Mean", "Std dev", "25%", "Median", "75%"):
        res[c] = [round(v, 4) if isinstance(v, float) else v for v in res[c]]
    for c in ("Min", "Max"):
        res[c] = [round(v, 4) if isinstance(v, float) else v for v in res[c]]
    return res.fillna("")


@dataclass
class Trend:
    field: str
    unit: str
    x_label: str                       # the date field, or 'Part order'
    points: pd.DataFrame               # Part, (Item), X, Value, Source, Rules
    center: Optional[float] = None
    ucl: Optional[float] = None
    lcl: Optional[float] = None
    note: str = ""

    @property
    def flagged(self) -> pd.DataFrame:
        return self.points[self.points["Rules"] != ""]


def trend(data: Data, fname: str, group: str = "", by: str = "",
          per: str = PER_PART) -> Trend:
    """``fname`` in date order of ``by`` (a date field; '' = part order),
    with individuals-chart limits and the run rules each point breaks."""
    from analyst_helpers.spc import control_chart_limits, western_electric_rules
    if fname not in data.fields:
        raise LedgerError(f"no field {fname!r} in this Ledger")
    by_part = per == PER_PART or not data.fields[fname].per_item
    df = rows(data, group)
    df = per_part(data, df) if by_part else df
    pts = []
    for i, r in enumerate(df.to_dict("records")):
        v = to_number(r.get(fname))
        if v is None:
            continue
        x = lg.parse_date(r.get(by)) if by else None
        if by and x is None:
            continue
        pts.append({"Part": r[data.id_field],
                    "Item": "" if by_part else r.get(data.item_field, ""),
                    "X": x, "Value": v, "Source": r.get(SOURCE_COL, ""), "_i": i})
    if by:
        pts.sort(key=lambda p: (p["X"], sample_sort_key(p["Part"]), p["_i"]))
    else:
        pts.sort(key=lambda p: (sample_sort_key(p["Part"]), p["_i"]))
    for p in pts:
        p.pop("_i")
    frame = pd.DataFrame(pts, columns=["Part", "Item", "X", "Value", "Source"])
    frame["Rules"] = ""
    t = Trend(fname, data.fields[fname].unit, by or "Part order", frame)
    if len(frame) < 3:
        t.note = "Too few values to draw control limits (3 or more are needed)."
        return t
    try:
        lim = control_chart_limits(frame["Value"].to_numpy(float), chart_type="i")
        t.center, t.ucl, t.lcl = lim["center"], lim["ucl"], lim["lcl"]
        hits = western_electric_rules(frame["Value"].to_numpy(float), t.ucl, t.lcl, t.center)
    except ValueError:
        t.note = "Every value is the same: no spread to draw limits from."
        return t
    rules: Dict[int, List[str]] = {}
    for h in hits.to_dict("records"):
        rules.setdefault(int(h["index"]), []).append(str(h["rule_number"]))
    frame["Rules"] = [", ".join(sorted(set(rules.get(i, [])))) for i in range(len(frame))]
    n = int((frame["Rules"] != "").sum())
    t.note = (f"{n} point(s) break a run rule (1: beyond the limits; 2: 2 of 3 past 2σ; "
              "3: 4 of 5 past 1σ; 4: 8 in a row on one side)." if n else
              "Every point is inside the limits and no run rule is broken.")
    return t


def outliers(data: Data, group: str = "", per: str = PER_ROW,
             z: float = OUTLIER_Z) -> pd.DataFrame:
    """Values far from the rest of their field: robust z = 0.6745 x distance
    from the median / MAD. Sorted farthest first."""
    flat = rows(data, group)
    whole = per_part(data, flat)
    out = []
    for f in fields_of(data, group):
        if _kind_in(data, f, group) != NUMBER:
            continue
        # a part-level field once per part, whatever ``per`` says
        by_part = per == PER_PART or not data.fields[f].per_item
        df = whole if by_part else flat
        vals = [(i, to_number(v)) for i, v in enumerate(df[f])]
        vals = [(i, v) for i, v in vals if v is not None]
        if len(vals) < 5:
            continue
        a = np.array([v for _, v in vals])
        med = float(np.median(a))
        mad = float(np.median(np.abs(a - med)))
        if mad == 0:
            continue
        for i, v in vals:
            rz = 0.6745 * (v - med) / mad
            if abs(rz) >= z:
                r = df.iloc[i]
                out.append({"Part": r[data.id_field],
                            "Item": "" if by_part else r.get(data.item_field, ""),
                            "Field": f, "Value": v, "Unit": data.fields[f].unit,
                            "Median": round(med, 4), "Robust z": round(rz, 2),
                            "Source": r.get(SOURCE_COL, "")})
    res = pd.DataFrame(out, columns=["Part", "Item", "Field", "Value", "Unit", "Median",
                                     "Robust z", "Source"])
    if len(res):
        res = res.reindex(res["Robust z"].abs().sort_values(ascending=False).index)
    return res.reset_index(drop=True)


@dataclass
class Correlation:
    method: str
    fields: List[str]
    r: pd.DataFrame                   # field x field
    p: pd.DataFrame
    n: pd.DataFrame
    sure: pd.DataFrame                # bool: significant after the FDR correction
    note: str = ""

    def pairs(self, most: int = 0, twins: bool = False) -> pd.DataFrame:
        """Each pair once, strongest first. ``twins``: also a field against
        an inspector's own reading of it ('… [JB]'), which only says the
        readings agree; left out by default so they do not crowd the top."""
        out = []
        for i, a in enumerate(self.fields):
            for b in self.fields[i + 1:]:
                r = self.r.loc[a, b]
                if pd.isna(r) or (not twins and base_field(a) == base_field(b)):
                    continue
                out.append({"Field A": a, "Field B": b, "r": round(float(r), 3),
                            "Strength": strength(r), "p": float(self.p.loc[a, b]),
                            "Sure": "yes" if self.sure.loc[a, b] else "no",
                            "Values": int(self.n.loc[a, b])})
        res = pd.DataFrame(out, columns=["r", "Strength", "Sure", "Field A", "Field B",
                                         "Values", "p"])
        if len(res):
            res = res.reindex(res["r"].abs().sort_values(ascending=False).index)
            res["p"] = [float(f"{v:.3g}") for v in res["p"]]
        res = res.reset_index(drop=True)
        return res.head(most) if most else res


def base_field(fname: str) -> str:
    """A field without its inspector's initials ('… [JB]' -> '…')."""
    return re.sub(r"\s*\[[^\]]*\]$", "", fname)


def strength(r: float) -> str:
    a = abs(float(r))
    word = ("very strong" if a >= 0.8 else "strong" if a >= 0.6 else "moderate"
            if a >= 0.4 else "weak" if a >= 0.2 else "none")
    return word if word == "none" else f"{word} {'+' if r > 0 else '-'}"


def correlation(data: Data, group: str = "", per: str = PER_ROW,
                method: str = "pearson") -> Correlation:
    """Every pair of number fields (each with 2+ different values). Per
    flaw, a pair of part-level fields is still taken once per part, so its
    count is not inflated by each part's flaws; per part, a per-flaw field is
    its part's average. Pairs with fewer than CORR_MIN_N values in common are
    left blank."""
    from analyst_helpers.stats import correlation_with_significance
    flat = rows(data, group)
    whole = per_part(data, flat)
    fs = []
    for f in fields_of(data, group):
        if _kind_in(data, f, group) != NUMBER:
            continue
        if len(set(_numbers(whole[f]).tolist())) >= 2:
            fs.append(f)
    empty = pd.DataFrame(index=fs, columns=fs, dtype=float)
    if len(fs) < 2:
        return Correlation(method, fs, empty, empty, empty, empty.astype(bool),
                           note="Fewer than 2 number fields with different values: "
                                "nothing to correlate.")

    def run(df):
        import warnings
        num = pd.DataFrame({f: [to_number(v) for v in df[f]] for f in fs}, dtype=float)
        with warnings.catch_warnings():   # a pair constant where both are filled: left blank
            warnings.simplefilter("ignore")
            return correlation_with_significance(num, fs, method=method)

    res = run(whole)
    if per == PER_ROW and any(data.fields[f].per_item for f in fs):
        flaws = run(flat)
        for a in fs:
            for b in fs:
                if data.fields[a].per_item or data.fields[b].per_item:
                    for k in ("corr", "p_values", "n", "significant_fdr"):
                        res[k].loc[a, b] = flaws[k].loc[a, b]
    r, p = res["corr"].copy(), res["p_values"].copy()
    n, sure = res["n"].astype(int), res["significant_fdr"].astype(bool)
    few = n < CORR_MIN_N
    r, p = r.mask(few), p.mask(few)
    for f in fs:
        r.loc[f, f] = 1.0
    sure = sure & ~few
    c = Correlation(method, fs, r, p, n, sure)
    strong = c.pairs()
    strong = strong[(strong["r"].abs() >= 0.6) & (strong["Sure"] == "yes")]
    c.note = (f"{len(strong)} pair(s) strongly correlated and sure (|r| 0.6 or more, "
              "significant after correcting for the number of pairs; a field against its own "
              "inspectors' readings not counted). A correlation is "
              "not a cause: check the pair's chart." if len(strong) else
              "No pair is both strongly correlated (|r| 0.6 or more) and sure.")
    return c


def correlation_figure(c: Correlation):
    """A heatmap of r, -1 blue to +1 red; a dot marks a sure pair."""
    from matplotlib.figure import Figure
    k = len(c.fields)
    side = min(11.0, max(4.5, 0.42 * k + 2.5))
    fig = Figure(figsize=(side, side * 0.85), dpi=100)
    ax = fig.add_subplot(111)
    if k < 2:
        ax.text(0.5, 0.5, c.note, ha="center", va="center", wrap=True, fontsize=9)
        ax.axis("off")
        return fig
    im = ax.imshow(c.r.to_numpy(float), cmap="RdBu_r", vmin=-1, vmax=1)
    short = [label(f, "", 30) for f in c.fields]
    for i, f in enumerate(c.fields):          # 'Media Run Time' of Concealment 1, 2 …
        if short.count(short[i]) > 1:
            short[i] = label(f.replace(": ", " "), "", 30) if ": " in f else short[i]
    ax.set_xticks(range(k))
    ax.set_yticks(range(k))
    ax.set_xticklabels(short, rotation=60, ha="right", fontsize=7)
    ax.set_yticklabels(short, fontsize=7)
    for i in range(k):
        for j in range(k):
            v = c.r.iat[i, j]
            if i != j and not pd.isna(v):
                if k <= 14:
                    ax.text(j, i, f"{v:.2f}" + ("*" if c.sure.iat[i, j] else ""),
                            ha="center", va="center", fontsize=6.5,
                            color="white" if abs(v) > 0.6 else "black")
                elif c.sure.iat[i, j]:
                    ax.plot(j, i, marker=".", color="black", markersize=3)
    fig.colorbar(im, ax=ax, fraction=0.04, pad=0.02, label=f"r ({c.method})")
    ax.set_title("Correlation (* sure)", fontsize=10)
    fig.tight_layout()
    return fig


def scatter_figure(data: Data, a: str, b: str, group: str = "", per: str = PER_ROW):
    """Field ``a`` against ``b``, with a least-squares line."""
    from matplotlib.figure import Figure
    flat = rows(data, group)
    both_part = not (data.fields[a].per_item or data.fields[b].per_item)
    df = per_part(data, flat) if per == PER_PART or both_part else flat
    pts = [(to_number(x), to_number(y)) for x, y in zip(df[a], df[b])]
    pts = [(x, y) for x, y in pts if x is not None and y is not None]
    fig = Figure(figsize=(5.5, 4), dpi=100)
    ax = fig.add_subplot(111)
    if pts:
        xs, ys = np.array(pts).T
        ax.scatter(xs, ys, s=14, color="#3b6ea5", alpha=0.8)
        if len(pts) >= 3 and np.std(xs) > 0:
            m, k = np.polyfit(xs, ys, 1)
            line = np.linspace(xs.min(), xs.max(), 2)
            ax.plot(line, m * line + k, color="#c0392b", linewidth=1)
    for setter, f in ((ax.set_xlabel, a), (ax.set_ylabel, b)):
        setter(label(f, data.fields[f].unit, 48), fontsize=8)
    ax.tick_params(labelsize=7)
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    return fig


# ── charts ──────────────────────────────────────────────────────────────
def label(fname: str, unit: str = "", most: int = 30) -> str:
    """A chart's name for a field: its section dropped, its unit once, the
    middle cut when long (so '… Pre-Conceal [JB]' and '[MK]' still differ)."""
    text = fname.split(": ", 1)[-1] if ": " in fname else fname
    if unit and f"({unit})" not in text:
        text += f" ({unit})"
    if len(text) > most:
        keep = most - 1
        text = text[:keep // 2] + "…" + text[-(keep - keep // 2):]
    return text


def trend_figure(t: Trend):
    """A matplotlib Figure (no pyplot: safe off the main thread)."""
    from matplotlib.figure import Figure
    fig = Figure(figsize=(8, 3.6), dpi=100)
    ax = fig.add_subplot(111)
    pts = t.points
    xs = list(pts["X"]) if t.x_label != "Part order" else list(range(1, len(pts) + 1))
    if len(pts):
        ax.plot(xs, pts["Value"], marker="o", markersize=3.5, linewidth=1, color="#3b6ea5")
        bad = pts["Rules"] != ""
        if bad.any():
            ax.scatter([x for x, b in zip(xs, bad) if b], pts["Value"][bad], s=36,
                       color="#c0392b", zorder=3, label="breaks a run rule")
            ax.legend(loc="best", fontsize=8, frameon=False)
    for y, style, name in ((t.center, "-", "center"), (t.ucl, "--", "UCL"),
                           (t.lcl, "--", "LCL")):
        if y is not None:
            ax.axhline(y, linestyle=style, linewidth=0.9, color="#7f8c8d")
            ax.annotate(f"{name} {y:.4g}", xy=(1, y), xycoords=("axes fraction", "data"),
                        fontsize=7, color="#555", ha="right", va="bottom")
    ax.set_title(label(t.field, t.unit, 80), fontsize=10)
    ax.set_xlabel(t.x_label, fontsize=8)
    ax.tick_params(labelsize=7)
    if t.x_label == "Part order" and len(pts) <= 30:
        ax.set_xticks(xs)
        ax.set_xticklabels(list(pts["Part"]), rotation=60, ha="right", fontsize=6)
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    return fig


def figure_png(fig) -> bytes:
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    FigureCanvasAgg(fig)
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    return buf.getvalue()


# ── export ──────────────────────────────────────────────────────────────
def export_path(csv) -> Path:
    p = Path(csv)
    return p.with_name(f"{p.stem}_analysis.xlsx")


def _ours(path: Path) -> bool:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return FILLED_MARK in str(wb.properties.keywords or "")
        finally:
            wb.close()
    except Exception:                                       # noqa: BLE001
        return False


def export(data: Data, out=None, group: str = "", per: str = PER_ROW,
           by: Optional[str] = None) -> Path:
    """The analysis workbook: Coverage, Missing, Summary, Trends, Outliers and
    a Charts sheet. Marked as the Ledger's own, so no scan reads it back, and
    never written over a workbook the Ledger did not write."""
    import os
    import tempfile

    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    out = Path(out) if out else export_path(data.csv)
    if out.exists() and not _ours(out):
        raise LedgerError(f"{out.name} is there already and the Ledger did not write it; "
                          "move it or pick another name.")
    if by is None:                    # the first date field; '' = part order
        by = data.date_fields[0] if data.date_fields else ""
    elif by not in data.fields:
        by = ""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)

    def sheet(title: str, df: pd.DataFrame, intro: str = ""):
        ws = wb.create_sheet(title)
        r = 1
        if intro:
            ws.cell(1, 1, intro).font = Font(italic=True)
            r = 3
        for c, name in enumerate(df.columns, start=1):
            ws.cell(r, c, str(name)).font = bold
        for row in df.itertuples(index=False):
            r += 1
            for c, v in enumerate(row, start=1):
                ws.cell(r, c, v if not hasattr(v, "isoformat") else v.isoformat())
        for c, name in enumerate(df.columns, start=1):
            width = max([len(str(name))] + [len(str(v)) for v in df[name].head(200)])
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, width + 2))
        ws.freeze_panes = ws.cell(r - len(df) if len(df) else r, 1)
        return ws

    label = group if group and data.group(group) else "all parts"
    sheet("Coverage", coverage(data, group), f"{Path(data.csv).name} · {label}")
    sheet("Missing", missing_by_part(data, group))
    sheet("Summary", summary(data, group, per),
          "Per flaw" if per == PER_ROW else "Per part (a per-flaw field's average)")
    trends, charts = [], []
    for f in [f for f in fields_of(data, group) if _kind_in(data, f, group) == NUMBER]:
        t = trend(data, f, group, by=by)
        charts.append(t)
        for p in t.flagged.to_dict("records"):
            trends.append({"Field": f, "Part": p["Part"], "X": p["X"], "Value": p["Value"],
                           "Rules": p["Rules"], "Source": p["Source"]})
    sheet("Trends", pd.DataFrame(trends, columns=["Field", "Part", "X", "Value", "Rules",
                                                  "Source"]),
          f"Points breaking a run rule, per part, in order of {by or 'part'}")
    corr = correlation(data, group, per)
    cs = sheet("Correlation", corr.pairs(), corr.note)
    if len(corr.fields) >= 2:
        r0 = len(corr.pairs()) + 6
        cs.cell(r0, 1, f"r matrix ({corr.method})").font = bold
        for j, f in enumerate(corr.fields, start=2):
            cs.cell(r0, j, f).font = bold
        for i, a in enumerate(corr.fields, start=1):
            cs.cell(r0 + i, 1, a).font = bold
            for j, b in enumerate(corr.fields, start=2):
                v = corr.r.loc[a, b]
                cs.cell(r0 + i, j, None if pd.isna(v) else round(float(v), 3))
    sheet("Outliers", outliers(data, group, per),
          f"Robust z of {OUTLIER_Z} or more (distance from the median over the MAD)")
    ws = wb.create_sheet("Charts")
    row = 1
    if len(corr.fields) >= 2:
        fig = correlation_figure(corr)
        ws.add_image(XLImage(io.BytesIO(figure_png(fig))), f"A{row}")
        row += int(fig.get_figheight() * 5.2) + 2
    for t in charts[:CHART_LIMIT]:
        if not len(t.points):
            continue
        img = XLImage(io.BytesIO(figure_png(trend_figure(t))))
        ws.add_image(img, f"A{row}")
        row += 19
    if len(charts) > CHART_LIMIT:
        ws.cell(row, 1, f"{len(charts) - CHART_LIMIT} more field(s) not charted.")
    mark_filled(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(out.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return out
