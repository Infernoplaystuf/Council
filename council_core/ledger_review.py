"""
council_core.ledger_review — the review workbook read as a list of things to
check, each with where its value came from, and the user's answers written
back into it (user, 2026-10-22: "the review document spit out to review
connections is way too clunky to actually review the connections and see if
they are right or wrong" ... "can we do something similar to the existing
link by clicking? ... Highlight the cells in question in that").

The Ledger tab's "Review by clicking…" window (council_qt/tabs/
ledger_review_dialog.py) shows one item at a time: the template with the
field's cell highlighted, the source workbook with the cell the value came
from highlighted. Accept / Reject / a corrected value are written back into
the SAME review workbook (its Accept and value cells), so Apply, and the
Excel review, work exactly as before.

Both kinds of review:

  * a sample Ledger's (source_extract.write_review): one item per row of
    Values; its Source says 'file › sheet!B7' (or a picture);
  * a table Ledger's (ledger_records.write_review): one item per record on
    Records (and Records 2, …); Source, Sheet, Table and Row say where.

Nothing here imports Qt.
"""
from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from sample_ledger import LedgerError

Cell = Tuple[int, int]

#: The sample review's Values columns (source_extract.VAL_HEAD), 1-based.
V_ACCEPT, V_PART, V_ITEM, V_FIELD, V_VALUE, V_SCORE, V_VIA, V_LABEL, V_SOURCE, \
    V_OTHERS, V_NOTE, V_PART_FROM, V_INSPECTOR = range(1, 14)

#: The table review's leading columns (ledger_records.write_review).
R_META = ["Accept", "Source", "Sheet", "Table", "Row"]
R_TAIL = ["Notes", "_table", "_checks"]

_WHERE_RE = re.compile(r"^(?P<file>.+?) › (?P<sheet>.+)!(?P<cell>\$?[A-Z]{1,3}\$?\d+)$")


@dataclass
class Item:
    sheet: str                    # the review's sheet the item is on
    row: int                      # its row there (1-based)
    accept: str                   # 'Y', 'N' or ''
    needs: bool                   # amber/red: a decision to make
    clash: bool
    title: str                    # one line for the list
    detail: str                   # what to check, in words
    source_file: str = ""         # a full path when the review knows it
    source_sheet: str = ""
    source_cells: List[Cell] = field(default_factory=list)
    #: the template fields this item fills
    fields: List[str] = field(default_factory=list)
    value: str = ""               # a sample item's value (editable)
    value_col: int = 0            # the review column of ``value``; 0 = not editable
    where: str = ""               # as the review wrote it
    item: str = ""                # a sample item's flaw / item ("1A"), "" part-level
    #: a table record's: field -> its value
    values: Dict[str, Any] = field(default_factory=dict)
    part: str = ""                # a sample item's part number
    score: float = 0.0            # how sure the match is (the review's Score)


@dataclass
class Review:
    path: str
    kind: str                     # 'samples' | 'records'
    csv: str
    sources: List[str]
    items: List[Item]

    def source_path(self, name: str) -> str:
        """The full path of a source given by name in the review."""
        if not name:
            return ""
        for s in self.sources:
            if Path(s).name == name or s == name:
                return s
        return ""


def _a1(text: str) -> Optional[Cell]:
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    try:
        col, row = coordinate_from_string(text.replace("$", ""))
        return row, column_index_from_string(col)
    except Exception:                                       # noqa: BLE001
        return None


def parse_where(where: str) -> Tuple[str, str, Optional[Cell]]:
    """'laser log.xlsx › Sheet1!D7' -> ('laser log.xlsx', 'Sheet1', (7, 4));
    a picture or anything else -> (file, '', None)."""
    w = str(where or "").strip()
    m = _WHERE_RE.match(w)
    if m:
        return m.group("file"), m.group("sheet"), _a1(m.group("cell"))
    return w.split(" › ", 1)[0], "", None


def _text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load(path) -> Review:
    """The review workbook at ``path`` as items to check."""
    import openpyxl
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:                                  # noqa: BLE001
        raise LedgerError(f"could not open the review {path.name}: {e}") from e
    try:
        try:
            meta = json.loads(wb["_run"]["A1"].value or "{}")
        except Exception:                                   # noqa: BLE001
            meta = {}
        sources = [str(s) for s in meta.get("sources") or []]
        csv = str(meta.get("csv") or "")
        if meta.get("mode") == "records" or "Records" in wb.sheetnames:
            items = _records_items(wb, sources)
            kind = "records"
        elif "Values" in wb.sheetnames:
            items = _values_items(wb, sources)
            kind = "samples"
        else:
            raise LedgerError(f"{path.name} is not a Ledger review (no Values or Records sheet)")
    finally:
        wb.close()
    review = Review(str(path), kind, csv, sources, items)
    if not csv:
        review.csv = _csv_from_readme(path)
    return review


def _csv_from_readme(path: Path) -> str:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for (v,) in wb["Read me"].iter_rows(max_col=1, values_only=True):
                if isinstance(v, str) and v.startswith("Parent CSV: "):
                    return v[len("Parent CSV: "):].strip()
        finally:
            wb.close()
    except Exception:                                       # noqa: BLE001
        pass
    return ""


def _find(sources: Sequence[str], name: str) -> str:
    for s in sources:
        if Path(s).name == name:
            return s
    return ""


def _values_items(wb, sources: Sequence[str]) -> List[Item]:
    out: List[Item] = []
    rows = list(wb["Values"].iter_rows(min_row=2, values_only=True))
    for i, row in enumerate(rows, start=2):
        row = list(row) + [None] * (13 - len(row))
        where = _text(row[V_SOURCE - 1])
        if not where:
            continue
        acc = _text(row[V_ACCEPT - 1]).upper()
        part, item, fname, value = (_text(row[k - 1]) for k in (V_PART, V_ITEM, V_FIELD, V_VALUE))
        clash = bool(_text(row[V_OTHERS - 1]))
        file, sheet, cell = parse_where(where)
        label = _text(row[V_LABEL - 1])
        score = _text(row[V_SCORE - 1])
        via = _text(row[V_VIA - 1])
        note = _text(row[V_NOTE - 1])
        lines = [f"{fname or '(no field)'} = {value!r}" + (f" for part {part}" if part else "")
                 + (f", {item}" if item else ""),
                 f"read from the label {label!r} at {where}" if label else f"read at {where}",
                 f"confidence {score}, matched by {via}" if score or via else ""]
        if clash:
            lines.append(f"other values offered: {_text(row[V_OTHERS - 1])}")
        if note:
            lines.append(f"note: {note}")
        part_from = _text(row[V_PART_FROM - 1])
        if part_from:
            lines.append(f"part from: {part_from}")
        if cell is None and "image" in where:
            lines.append("from a picture: the cell cannot be shown")
        out.append(Item(
            # 'N': set aside for another reading of the same workbook (decide)
            sheet="Values", row=i, accept=acc, needs=acc not in ("Y", "N") or clash,
            clash=clash,
            title=" · ".join(x for x in (value, fname, " ".join(y for y in (part, item) if y))
                             if x) or where,
            detail="\n".join(x for x in lines if x),
            source_file=_find(sources, file), source_sheet=sheet,
            source_cells=[cell] if cell else [], fields=[fname] if fname else [],
            value=value, value_col=V_VALUE, where=where, item=item, part=part,
            score=_score(row[V_SCORE - 1])))
    return out


def _score(v: Any) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _records_items(wb, sources: Sequence[str]) -> List[Item]:
    out: List[Item] = []
    for name in wb.sheetnames:
        if not (name == "Records" or name.startswith("Records ")):
            continue
        rows = wb[name].iter_rows(values_only=True)
        head = [_text(h) for h in next(rows, [])]
        if head[:len(R_META)] != R_META:
            continue
        tail_at = head.index("Notes") if "Notes" in head else len(head)
        names = head[len(R_META):tail_at]
        checks_at = head.index("_checks") if "_checks" in head else None
        for i, row in enumerate(rows, start=2):
            row = list(row) + [None] * (len(head) - len(row))
            acc, file, sheet, table, srow = (row[k] for k in range(len(R_META)))
            values = {n: row[len(R_META) + j] for j, n in enumerate(names)
                      if row[len(R_META) + j] not in (None, "")}
            checks = _text(row[checks_at]) if checks_at is not None else ""
            notes = _text(row[tail_at]) if tail_at < len(head) else ""
            clash = "!clash" in checks
            acc = _text(acc).upper()
            at = _a1(_text(table)) if table else None
            r = int(srow) if isinstance(srow, (int, float)) else None
            cells: List[Cell] = []
            if r is not None and at is not None:
                cells = [(r, at[1])]      # the table's first column: widened to its row on show
            shown = ", ".join(f"{k}={_text(v)}" for k, v in list(values.items())[:4])
            lines = [f"record from {file} › {sheet}, row {srow} (table at {table})",
                     "values: " + "; ".join(f"{k} = {_text(v)}" for k, v in values.items())]
            if checks:
                lines.append("to check: " + checks.replace(";", ", "))
            if notes:
                lines.append(f"notes: {notes}")
            out.append(Item(
                sheet=name, row=i, accept=acc, needs=acc != "Y" or bool(checks), clash=clash,
                title=shown or f"{file} row {srow}", detail="\n".join(lines),
                source_file=_find(sources, _text(file)), source_sheet=_text(sheet),
                source_cells=cells, fields=list(values), values=values,
                where=f"{file} › {sheet}!{table}"))
    return out


def template_cells(csv_path) -> Dict[Tuple[str, str], Tuple[str, Cell]]:
    """A sample Ledger's template: (field, item) -> (sheet, cell) of the one
    cell that value goes in (item "" for a part-level field); {} for a table
    Ledger or a schema that cannot be read."""
    from sample_ledger import TemplateSchema, schema_path_for
    try:
        schema = TemplateSchema.load(schema_path_for(csv_path))
    except Exception:                                       # noqa: BLE001
        return {}
    out: Dict[Tuple[str, str], Tuple[str, Cell]] = {}
    for f in schema.fields:
        for item, a1 in (f.cells or {}).items():
            rc = _a1(a1) if a1 else None
            if rc:
                out[(f.name, str(item))] = (f.sheet, rc)
    return out


def record_row_cells(view, item: Item) -> List[Cell]:
    """A table record's whole row on its source sheet (its table's columns),
    from a ledger_linking.SampleView; the anchor cell alone when no table is
    found there."""
    if not item.source_cells:
        return []
    r, c = item.source_cells[0]
    tables = [t for t in view.tables.get(item.source_sheet, []) if r in t.rows]
    for t in sorted(tables, key=lambda t: c not in t.cols):
        return [(r, cc) for cc in t.cols]
    return item.source_cells


def save(path, answers: Dict[Tuple[str, int], Dict[str, Any]]) -> int:
    """Write the answers into the review workbook: {(sheet, row):
    {'accept': 'Y'|'N'|'', 'value': str, 'value_col': int}}. Every other
    cell, fill, list and hidden sheet stays. Atomic: a temporary file, then
    a replace. Returns how many rows changed."""
    import openpyxl
    path = Path(path)
    if not answers:
        return 0
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:                                  # noqa: BLE001
        raise LedgerError(f"could not open the review {path.name}: {e}") from e
    changed = 0
    try:
        for (sheet, row), a in answers.items():
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            if "accept" in a:
                ws.cell(row=row, column=1).value = a["accept"] or None
            if a.get("value_col") and "value" in a:
                ws.cell(row=row, column=int(a["value_col"])).value = a["value"]
            changed += 1
        tmp = path.with_name(path.stem + ".saving.xlsx")
        try:
            wb.save(tmp)
        except PermissionError as e:
            raise LedgerError(f"could not write next to {path.name}: {e}") from e
    finally:
        wb.close()
    try:
        os.replace(tmp, path)
    except PermissionError as e:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise LedgerError(f"{path.name} is open in another program (Excel?): close it "
                          "and save again") from e
    return changed


# ── the filled template (user, 2026-10-23: "have the ledger output the
# template filled out with the values entered into it highlighted by their
# confidence values ... click on cells and see their source") ─────────────
#: a match at least this sure is high confidence (source_extract.ACCEPT_AT)
CONF_HIGH = 0.85
#: under this, low
CONF_LOW = 0.6
TINT_HIGH, TINT_MID, TINT_LOW, TINT_CLASH = "#C6EFCE", "#FFEB9C", "#F8CBAD", "#FF7C80"
#: what each tint means, for a legend
LEGEND = [(TINT_HIGH, f"sure (score {CONF_HIGH:g} or more)"),
          (TINT_MID, f"fairly sure ({CONF_LOW:g}–{CONF_HIGH:g})"),
          (TINT_LOW, f"unsure (under {CONF_LOW:g})"),
          (TINT_CLASH, "sources disagree")]


def confidence_tint(item: Item) -> str:
    if item.clash:
        return TINT_CLASH
    if item.score >= CONF_HIGH:
        return TINT_HIGH
    return TINT_MID if item.score >= CONF_LOW else TINT_LOW


def parts(review: Review) -> List[str]:
    """The parts a sample review has values for, in their review order."""
    return list(dict.fromkeys(i.part for i in review.items if i.part))


def cell_candidates(items: Sequence[Item], exact: Dict[Tuple[str, str], Tuple[str, Cell]],
                    part: str, accept_of=None) -> Dict[Tuple[str, Cell], List[Item]]:
    """One part's values by the template cell each goes in: (sheet, cell) ->
    the items offered for it, the one shown first - an accepted one, then
    the surest. A value said to be wrong ('N') is left out. ``accept_of``:
    an item's Accept as the window has it (answers not saved yet)."""
    acc = accept_of or (lambda it: it.accept)
    out: Dict[Tuple[str, Cell], List[Item]] = {}
    for it in items:
        if it.part != part or not it.fields or acc(it) == "N":
            continue
        at = exact.get((it.fields[0], it.item)) or exact.get((it.fields[0], ""))
        if at is not None:
            out.setdefault(at, []).append(it)
    for cands in out.values():
        cands.sort(key=lambda it: (acc(it) != "Y", -it.score))
    return out


def _cell_value(text: str) -> Any:
    """A number written as one, so the filled template computes."""
    try:
        f = float(text)
    except (TypeError, ValueError):
        return text
    return int(f) if f.is_integer() and "." not in text else f


def write_filled(review: Review, part: str, out_path,
                 exact: Optional[Dict[Tuple[str, str], Tuple[str, Cell]]] = None) -> str:
    """A copy of the template with ``part``'s values in their cells, each
    tinted by its confidence (confidence_tint) with a note of where it was
    read and how sure the match is. The template itself is never written.
    Atomic. Returns the path."""
    import shutil

    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill
    from sample_ledger import TemplateSchema, schema_path_for
    schema = TemplateSchema.load(schema_path_for(review.csv))
    template = Path(schema.template_path)
    out = Path(out_path)
    if out.resolve() == template.resolve():
        raise LedgerError("refusing to write over the template itself")
    exact = exact if exact is not None else template_cells(review.csv)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.stem + ".writing" + out.suffix)
    shutil.copyfile(template, tmp)
    wb = openpyxl.load_workbook(tmp)
    try:
        for (sheet, (r, c)), cands in cell_candidates(review.items, exact, part).items():
            ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
            it = cands[0]
            cell = ws.cell(row=r, column=c)
            cell.value = _cell_value(it.value)
            tint = confidence_tint(it)[1:]
            cell.fill = PatternFill("solid", start_color=tint, end_color=tint)
            others = [o.value for o in cands[1:] if o.value != it.value]
            cell.comment = Comment(
                f"{it.fields[0]}: {it.value}\nfrom {it.where}\nscore {it.score:g}"
                + (f" (accepted)" if it.accept == "Y" else "")
                + (f"\nalso offered: {', '.join(others[:5])}" if others else ""), "Council")
        wb.save(tmp)
    finally:
        wb.close()
    os.replace(tmp, out)
    return str(out)


def filled_dir(review_path) -> Path:
    """Where a review's filled templates go: beside it, '<review>_filled'."""
    p = Path(review_path)
    return p.with_name(p.stem + "_filled")


def _safe_name(part: str) -> str:
    return re.sub(r"[^\w\-. ]+", "_", part).strip() or "part"


def write_all_filled(review: Review, progress=None) -> List[str]:
    """One filled template per part (write_filled), in filled_dir. A part
    that cannot be written is skipped."""
    if review.kind != "samples" or not review.csv:
        return []
    exact = template_cells(review.csv)
    if not exact:
        return []
    out = []
    folder = filled_dir(review.path)
    names = parts(review)
    for n, part in enumerate(names, 1):
        if progress is not None:
            progress(f"filling the template for {part} ({n}/{len(names)})")
        try:
            out.append(write_filled(review, part, folder / f"{_safe_name(part)}.xlsx", exact))
        except Exception:                                   # noqa: BLE001
            continue
    return out
