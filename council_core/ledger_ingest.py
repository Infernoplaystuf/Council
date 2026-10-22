"""
ledger_ingest.py — the spreadsheet-ingestion pipeline as three calls a front
end (the Ledger tab, council_qt/tabs/ledger.py) or an agent tool can make, each returning
a summary instead of printing one.

    create_parent(template, csv)          template -> parent CSV (+ schema)
    run_extract(csv, sources, review)     sources  -> review workbook
    run_apply(review, fill_dir)           review   -> parent CSV (+ filled templates)

Nothing here knows about a toolkit: no widget, no thread. The tab runs these
on a worker thread; the tests call them directly. The work itself is
sample_ledger.py and source_extract.py — this file only gathers their results
into the numbers the user is shown (values found / pre-accepted / to review /
images read / problems).

ledger_tools(policy) builds the same two steps as safe_agent Tools for a
HUMAN to register in a ToolRegistry before it is frozen. Nothing in this file
registers anything: the model can never add a tool.
"""
from __future__ import annotations

import os
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from sample_ledger import LedgerError, ParentCSV, read_template
from source_extract import (REVIEW_TOO_LONG, ImageReader, apply_review, decide,
                            default_image_reader, extract, load_aliases, load_taught,
                            path_too_long, too_long_to_write, workbook_files,
                            write_review)

#: listener(line, fraction or None) - called from the worker thread with every
#: progress line (also printed in the terminal); the caller hops to its UI.
Listener = Callable[[str, Optional[float]], None]


def default_csv_for(template) -> Path:
    """The parent CSV sample_ledger.py's command line would make."""
    t = Path(template)
    return t.with_name(t.stem.replace(" ", "_") + "_parent.csv")


def default_review_for(csv_path) -> Path:
    return Path(csv_path).with_name(f"review_{datetime.now():%Y%m%d_%H%M%S}.xlsx")


def is_ledger_output(p: Path, template_path: Optional[str] = None) -> bool:
    """The Ledger's own files: a review workbook, a filled template, the
    template itself. A recursive scan of a folder that also holds the
    Ledger read them back as sources - old reviews' values outranked the
    real ones (review, 2026-10-19)."""
    return bool(_ledger_output_kind(p, template_path))


def _ledger_output_kind(p: Path, template_path: Optional[str] = None) -> str:
    """'the template' / 'a Ledger review' / 'a filled template', or ''."""
    from council_core import kg_excel
    from sample_ledger import is_filled_template
    if template_path:
        try:
            if p.resolve() == Path(template_path).resolve():
                return "the template"
        except OSError:
            pass
    if kg_excel.is_review_workbook(p):
        return "a Ledger review"
    return "a filled template" if is_filled_template(p) else ""


def collect_sources(paths: Sequence,
                    ignore_patterns: Optional[Sequence[str]] = None,
                    template_path: Optional[str] = None,
                    reasons: Optional[Dict[Path, str]] = None
                    ) -> Tuple[List[Path], List[Path]]:
    """Expand paths (files kept as-is, folders scanned recursively) into
    the Ledger's source list.

    ``ignore_patterns`` is a sequence of ``fnmatch`` glob patterns
    (``cover*.xlsx``, ``agenda*.xlsx``). Each matched filename is kept
    out of ``kept`` and recorded in ``ignored``, so the Ledger log can
    say exactly which files were skipped and why. Case-insensitive;
    matched against the base name only.

    Returns ``(kept, ignored)``. Legacy single-list callers can
    discard the ignored list. ``reasons``, when given, gets why each
    ignored file was left out ("Skip pattern 'cover*'", "a Ledger
    review")."""
    from source_extract import file_matches_ignore
    patterns = list(ignore_patterns or [])
    kept: List[Path] = []
    ignored: List[Path] = []
    for s in paths:
        p = Path(s)
        scanned = p.is_dir()
        found = workbook_files(p) if scanned else [p]
        for x in found:
            if x in kept or x in ignored:
                continue
            # a folder's scan leaves the Ledger's own outputs out; a file
            # the user picked by name is read as asked
            pat = next((q for q in patterns if file_matches_ignore(x.name, [q])), None)
            why = f"Skip pattern {pat!r}" if pat is not None else (
                _ledger_output_kind(x, template_path) if scanned else "")
            if why:
                ignored.append(x)
                if reasons is not None:
                    reasons[x] = why
            else:
                kept.append(x)
    return kept, ignored


# ── step 1: template -> parent CSV ──────────────────────────────────────
@dataclass
class ParentSummary:
    csv: str
    created: bool
    fields: int
    items: int
    samples: int
    rows: int
    description: str
    #: a table template (council_core.ledger_records): one row per record
    records_mode: bool = False

    def describe(self) -> str:
        if self.records_mode:
            return self.description
        verb = "created" if self.created else "extended"
        return (f"Parent CSV {verb}: {self.csv}\n{self.fields} fields, {self.items} items; "
                f"{self.samples} samples / {self.rows} rows so far.")


def create_parent(template, csv_path=None, id_field: Optional[str] = None,
                  mode: str = "auto") -> ParentSummary:
    """``mode``: 'samples' - one workbook is one sample (sample_ledger);
    'records' - the template is a table and each of its rows is a record,
    many per workbook (council_core.ledger_records); 'auto' - a table Ledger
    stays one, and a template with a table but no sample-ID field (the BTB
    example, user 2026-10-20) makes one."""
    from council_core import ledger_records
    csv_path = Path(csv_path) if csv_path else default_csv_for(template)
    if mode == "records" or (mode == "auto" and ledger_records.is_record_ledger(csv_path)):
        return _create_records(template, csv_path)
    try:
        schema = read_template(template, id_field)
    except LedgerError as e:
        if mode == "auto" and not id_field and "no sample-ID field" in str(e) \
                and ledger_records.is_record_template(template):
            return _create_records(template, csv_path)
        raise
    existed = csv_path.exists()
    ledger = ParentCSV.create(csv_path, schema)
    return ParentSummary(str(ledger.path), not existed, len(schema.fields),
                         len(schema.items), len(ledger.sample_ids()), len(ledger.rows),
                         schema.describe())


def _create_records(template, csv_path: Path) -> ParentSummary:
    from council_core import ledger_records
    r = ledger_records.create_ledger(template, csv_path)
    _cols, rows = ledger_records._read_rows(Path(r.csv))
    return ParentSummary(r.csv, r.created, len(r.fields) + len(r.forms), 0,
                         len(rows), len(rows), r.describe(), records_mode=True)


# ── step 2: sources -> review workbook ──────────────────────────────────
@dataclass
class ExtractSummary:
    review: str
    files: int
    values: int
    pre_accepted: int
    to_review: int
    unmatched: int
    images: int
    images_read: int
    problems: List[str] = field(default_factory=list)
    #: One ``(filename, kind)`` per workbook Excel refused to open, so the
    #: Ledger log can show a breakdown at the end of a batch.
    open_failures: List[Any] = field(default_factory=list)
    #: Per-workbook ``(name, seconds)`` wall-clock timings. The log
    #: appends the slowest few so a user can pinpoint which workbook
    #: ate the hour.
    timings: List[Any] = field(default_factory=list)
    #: True when the extract exited early via the Ledger tab's Cancel
    #: event. The caller shows a "Cancelled" status instead of "Done".
    cancelled: bool = False
    #: Cancelled after every workbook was read, before or while Council's
    #: model was asked about the labels left over: every value the rules
    #: found is in the review - it is not a partial one (second review,
    #: 2026-10-20).
    read_all: bool = False
    #: What Council's model was asked and found (empty when not asked).
    assist: str = ""
    #: One ``(name, reason)`` per file left out before reading: a Skip
    #: pattern, or the Ledger's own review / filled template / template.
    #: Only a progress line said so, and it went by (review, 2026-10-19).
    skipped: List[Any] = field(default_factory=list)
    #: The part-name pattern the run used, and whether it was the Ledger's
    #: saved one (Teach saved it; the pattern field was empty).
    pattern: str = ""
    pattern_saved: bool = False
    #: a table Ledger's run: records found, the guesses waiting for a Y on
    #: its Decisions sheet, and what it learned (the user's highlights)
    records: int = -1
    decisions: int = 0
    learned: List[str] = field(default_factory=list)

    def describe(self) -> str:
        if self.records >= 0:
            return self._describe_records()
        s = (f"{self.files} files: {self.values} values found, {self.pre_accepted} "
             f"pre-accepted, {self.to_review} to review, {self.unmatched} values under "
             f"labels that matched no field, {self.images_read} of {self.images} "
             f"images read")
        if self.cancelled:
            s = ("CANCELLED after every workbook was read — " if self.read_all else
                 "CANCELLED — ") + s
        s += "".join(f"\n{line}" for line in self.learned)
        if self.pattern_saved:
            s += (f"\nPart-name pattern {self.pattern!r}: this Ledger's saved one (the "
                  "pattern field was empty) — a hint for file and sheet names and "
                  "titles; a part of another shape is still pre-accepted")
        if self.skipped:
            s += "\n" + summarize_skipped(self.skipped)
        if self.assist:
            s += "\n" + self.assist
        if self.open_failures:
            s += "\n" + summarize_open_failures(self.open_failures)
        if self.timings:
            s += "\n" + summarize_slowest(self.timings)
        return s + "".join(f"\n  ! {p}" for p in self.problems) + f"\nreview: {self.review}"


    def _describe_records(self) -> str:
        s = (f"{self.files} files: {self.records} records, {self.values} values found, "
             f"{self.pre_accepted} pre-accepted, {self.to_review} to review"
             + (f", {self.decisions} guesses to decide on the Decisions sheet"
                if self.decisions else ""))
        if self.cancelled:
            s = "CANCELLED — " + s
        s += "".join(f"\n{line}" for line in self.learned)
        if self.skipped:
            s += "\n" + summarize_skipped(self.skipped)
        if self.assist:
            s += "\n" + self.assist
        if self.open_failures:
            s += "\n" + summarize_open_failures(self.open_failures)
        if self.timings:
            s += "\n" + summarize_slowest(self.timings)
        return s + "".join(f"\n  ! {p}" for p in self.problems) + f"\nreview: {self.review}"


_OPEN_FAILURE_LABELS = {
    "password": "password-protected",
    "rights": "rights-protected (IRM)",
    "corrupt": "corrupt / wrong format / blocked",
    "timeout": "timed out (sign-in prompt?)",
    "access": "access denied",
    # (review, 2026-10-20) a path past Excel's 255 characters or Windows'
    # MAX_PATH was 'other error' or 'corrupt'; a 0-byte file was 'read'
    "path": "path too long",
    "empty": "empty file (0 bytes)",
    "unavailable": "Excel COM not available",
    # council_core.file_guard (user 2026-10-22): took too long or stalled
    # the app before - skipped until it changes
    "slow": "too slow / stalled the app (skipped until it changes)",
    "unknown": "unknown reason",
    "other": "other error",
}


def summarize_open_failures(open_failures) -> str:
    """A one-line roll-up for the Ledger log."""
    from collections import Counter
    counts = Counter(k for _n, k in open_failures)
    parts = [f"{counts[k]} {label}" for k, label in _OPEN_FAILURE_LABELS.items()
             if counts.get(k)]
    return f"{len(open_failures)} files failed to open — " + ", ".join(parts)


def summarize_skipped(skipped, top: int = 10) -> str:
    """'left out 3 file(s): cover_sheet.xlsx (Skip pattern 'cover*'), …'"""
    shown = ", ".join(f"{n} ({why})" for n, why in list(skipped)[:max(1, int(top))])
    more = len(skipped) - min(len(skipped), max(1, int(top)))
    return (f"left out {len(skipped)} file(s): {shown}"
            + (f", and {more} more" if more > 0 else ""))


def summarize_slowest(timings, top: int = 5) -> str:
    """A one-line roll-up of the slowest workbooks in a batch — the
    "What's slow?" line the Ledger log appends so a 20-minute extract's
    hot spots are obvious without scrolling the terminal.

    ``timings`` is ``[(name, seconds), ...]``; the ``top`` slowest are
    rendered in descending order."""
    if not timings:
        return ""
    ranked = sorted(timings, key=lambda t: -float(t[1]))[:max(1, int(top))]
    total = sum(float(s) for _n, s in timings)
    parts = ", ".join(f"{n} {s:.1f}s" for n, s in ranked)
    return (f"Slowest {len(ranked)} of {len(timings)} workbooks "
            f"({total:.1f}s total): {parts}")


def _images_read(notes) -> int:
    return sum(1 for n in notes if n.reader not in ("", "none installed")
               and not n.text.startswith(("reader failed", "(no text found)")))


def run_extract(csv_path, sources: Sequence, review_path=None, part: str = "",
                use_images: bool = True, image_reader: Optional[ImageReader] = None,
                progress: Optional[Listener] = None,
                pattern_mask: Optional[str] = None,
                cancel: Optional[Any] = None,
                log: Optional[Callable[[str], None]] = None,
                assist: Any = False) -> ExtractSummary:
    """``log``: where a stall dump goes (the Ledger tab's log). ``assist``:
    ask Council's model about the labels and part numbers the rules could
    not place (council_core.ledger_llm); its answers are suggestions, and
    the user's answers in Teach go in as its examples.

    ``pattern_mask`` empty: the Ledger's saved part-name pattern (Teach
    saves one; source_extract.load_patterns) is used as a hint - a part of
    another shape is still pre-accepted (extract's pattern_strict) - and the
    summary says so (ExtractSummary.pattern_saved)."""
    from progress_report import Progress
    with Progress("ledger extract", listener=progress) as prog:
        s = _extract(csv_path, sources, review_path, part, use_images, image_reader,
                     prog, pattern_mask, cancel, log, assist)
        tail = "cancelled" if s.cancelled else f"{len(s.problems)} problems"
        prog.finish(f"{s.files} workbooks: {s.values} values found, {s.to_review} to "
                    f"review, {tail}")
        return s


def _extract(csv_path, sources, review_path, part, use_images, image_reader,
             prog, pattern_mask: Optional[str] = None,
             cancel: Optional[Any] = None,
             log: Optional[Callable[[str], None]] = None,
             assist: Any = False, chat=None) -> ExtractSummary:
    say = prog.update
    from council_core import ledger_records
    if ledger_records.is_record_ledger(csv_path):
        return _extract_records(csv_path, sources, review_path, prog, cancel, assist, chat)
    ledger = ParentCSV(csv_path)
    if ledger.schema is None:
        raise LedgerError(f"{Path(csv_path).name} has no schema beside it; "
                          "create it from a template first")
    from source_extract import load_ignore_files
    ignore_patterns = load_ignore_files(csv_path)
    why_skipped: Dict[Path, str] = {}
    files, skipped_files = collect_sources(sources,
                                           ignore_patterns=ignore_patterns,
                                           template_path=ledger.schema.template_path,
                                           reasons=why_skipped)
    if not files:
        if skipped_files and not sources:
            raise LedgerError("no source workbooks: add .xlsx files or a folder of them")
        if skipped_files:
            raise LedgerError(
                f"no source workbooks left after the Skip patterns and the Ledger's own "
                f"outputs were left out "
                f"({len(skipped_files)} skipped: "
                + ", ".join(sorted({p.name for p in skipped_files}))[:200]
                + "). Edit the Skip these files field or clear it.")
        raise LedgerError("no source workbooks: add .xlsx files or a folder of them")
    # one whose path is only too long for Windows is there: the run says so
    # for it (source_extract.load_workbook_grids) and reads the rest
    missing = [f.name for f in files if not f.exists() and not path_too_long(f)]
    if missing:
        raise LedgerError(f"source workbooks not found: {missing}")
    out = Path(review_path) if review_path else default_review_for(csv_path)
    n = too_long_to_write(out)
    if n:
        # said now, not as a raw FileNotFoundError once every workbook was
        # read (review, 2026-10-19)
        raise LedgerError(REVIEW_TOO_LONG.format(n=n))
    mapping, ignore = load_aliases(csv_path)
    # what the Ledger tab's Teach learned (council_core.ledger_teach): its
    # inspectors' labels, flaw-code columns - and its part-name pattern,
    # used when the field is empty, and said so (user, 2026-10-20). Only as
    # a hint: one sample's shape held back every part of another shape
    # from pre-acceptance (review, 2026-10-20)
    taught = load_taught(csv_path)
    pattern_saved = False
    if not (pattern_mask or "").strip() and taught.patterns:
        pattern_mask, pattern_saved = taught.patterns[0], True
        say(f"using this Ledger's saved part-name pattern {pattern_mask!r} as a hint (the "
            "pattern field is empty)")
    if use_images and image_reader is None:
        say("starting the local image reader…")
        image_reader = default_image_reader()
    if skipped_files:
        say(f"leaving out {len(skipped_files)} file(s) (Skip patterns, or the Ledger's "
            "own reviews / filled templates): "
            + ", ".join(sorted({p.name for p in skipped_files}))[:200])
    say(f"reading {len(files)} workbooks…")
    report = part_finder = None
    label_examples: List[Any] = []
    if assist:
        from council_core import ledger_llm, ledger_teach
        report = ledger_llm.AssistReport()
        stop = cancel.is_set if cancel is not None else None
        # ``assist`` is the model the user picked ('ollama:<name>' or a GGUF
        # path); True alone means Council's own choice. Asked only while it
        # answers: one failure, one problem line (ledger_llm.one_shot).
        chat = ledger_llm.one_shot(chat or (ledger_llm.chat_for(assist, should_stop=stop)
                                            if isinstance(assist, str) else
                                            ledger_llm.council_chat(should_stop=stop)),
                                   report, stop)
        # the user's answers in Teach go in as examples (few-shot)
        label_examples = ledger_teach.label_examples(csv_path, ledger.schema)
        part_examples = ledger_teach.part_examples(csv_path)

        def part_finder(name, grids):
            return ledger_llm.find_part(name, grids, pattern_mask, chat, report,
                                        examples=part_examples)
    import pptx_open
    with pptx_open.user_pairings(csv_path):           # your text box <-> table links
        run = extract(ledger.schema, files, mapping, ignore,
                      image_dir=out.with_name(out.stem + "_images"),
                      image_reader=image_reader, part=part, use_images=use_images,
                      progress=prog, pattern_mask=pattern_mask, cancel=cancel, log=log,
                      part_finder=part_finder, taught=taught,
                      pattern_strict=not pattern_saved,
                      # the records already accepted are worked answers too
                      examples=ledger.rows)
    read_all = False
    if report is not None and not run.cancelled:
        rules = "every value the rules found is in this review"
        if cancel is not None and cancel.is_set():
            # pressed after the last workbook was read: every value is in
            # the review; the model is not asked about the labels after a
            # Cancel (review, 2026-10-19). Said as the Cancel below is: one
            # said 'Review written…', the other 'Cancelled… partial review'
            # (second review, 2026-10-20)
            read_all = run.cancelled = True
            report.problems.append("Cancelled before Council's model was asked about the "
                                   f"labels left over; {rules}.")
        else:
            prog.stage("asking Council's model about the labels left over")
            ledger_llm.place_labels(run, ledger.schema, chat, report,
                                    examples=label_examples)
            if cancel is not None and cancel.is_set():
                # pressed while the model was asked: a cancelled run - it
                # said 'Review written…' and that the model did not answer
                # ('stopped by the caller', review 2026-10-20)
                read_all = run.cancelled = True
                if report.cancelled or not report.values_placed:
                    run.problems.append("Cancelled by user while Council's model was asked "
                                        f"about the labels left over; {rules}.")
                else:
                    # a model Cancel cannot stop (a GGUF in this process)
                    # answered anyway: its suggestions are in the review
                    run.problems.append(
                        "Cancelled by user while Council's model was asked about the labels "
                        f"left over, but it answered first; {rules}, and the "
                        f"{report.values_placed} value(s) it placed are too (amber: check "
                        "them).")
        run.problems.extend(report.problems)
    prog.stage("writing the review workbook")
    dec = decide(run, ledger.schema)
    # the mask and Part override go with the review: Apply reads the
    # sources again and must give its values the same parts (review,
    # 2026-10-19: mapped labels went to a phantom part or were dropped)
    write_review(out, run, dec, ledger.schema, str(csv_path),
                 pattern_mask=pattern_mask or "", part=part or "",
                 pattern_hint=pattern_saved)
    ok = sum(d.accept for d in dec)
    return ExtractSummary(str(out), len(files), len(dec), ok, len(dec) - ok,
                          len(run.unmatched), len(run.images), _images_read(run.images),
                          list(run.problems), list(run.open_failures),
                          list(run.timings), bool(run.cancelled), read_all=read_all,
                          assist=report.describe() if report is not None else "",
                          skipped=[(p.name, why_skipped.get(p, "left out"))
                                   for p in skipped_files],
                          pattern=pattern_mask or "", pattern_saved=pattern_saved,
                          learned=list(getattr(run, "learned", []) or []))


def _extract_records(csv_path, sources, review_path, prog, cancel, assist,
                     chat=None) -> ExtractSummary:
    """A table Ledger's Extract (council_core.ledger_records): the same
    sources, Skip patterns and model choice as the sample Ledger's."""
    from council_core import ledger_records
    from source_extract import load_ignore_files
    recipe = ledger_records.load_recipe(csv_path)
    why_skipped: Dict[Path, str] = {}
    files, skipped_files = collect_sources(sources,
                                           ignore_patterns=load_ignore_files(csv_path),
                                           template_path=recipe.template.get("path"),
                                           reasons=why_skipped)
    if not files:
        raise LedgerError("no source workbooks: add .xlsx files or a folder of them"
                          + (f" ({len(skipped_files)} left out by the Skip patterns or as "
                             "the Ledger's own files)" if skipped_files else ""))
    missing = [f.name for f in files if not f.exists() and not path_too_long(f)]
    if missing:
        raise LedgerError(f"source workbooks not found: {missing}")
    out = Path(review_path) if review_path else default_review_for(csv_path)
    n = too_long_to_write(out)
    if n:
        raise LedgerError(REVIEW_TOO_LONG.format(n=n))
    report = None
    if assist:
        from council_core import ledger_llm
        report = ledger_llm.AssistReport()
        stop = cancel.is_set if cancel is not None else None
        chat = ledger_llm.one_shot(chat or (ledger_llm.chat_for(assist, should_stop=stop)
                                            if isinstance(assist, str) else
                                            ledger_llm.council_chat(should_stop=stop)),
                                   report, stop)
    run = ledger_records.extract(csv_path, files, chat=chat, progress=prog, cancel=cancel,
                                 report=report)
    if report is not None:
        run.problems.extend(report.problems)
    prog.stage("writing the review workbook")
    ledger_records.write_review(out, run, str(csv_path), [str(f) for f in files])
    cells, pre = run.cells, run.cells_pre_accepted
    assist_line = ""
    if assist:
        assist_line = (f"Council's model ({run.model or 'not answering'}) was asked about "
                       f"{run.model_asked} header layout(s); its answers are amber")
    return ExtractSummary(str(out), len(files), cells, pre, cells - pre,
                          sum(1 for u in run.tables if u.use == "skipped"), 0, 0,
                          list(run.problems), list(run.open_failures), list(run.timings),
                          bool(run.cancelled), assist=assist_line,
                          skipped=[(p.name, why_skipped.get(p, "left out"))
                                   for p in skipped_files],
                          records=len(run.records), decisions=len(run.decisions),
                          learned=list(run.learned))


# ── step 3: review -> parent CSV (+ filled templates) ───────────────────
@dataclass
class ApplySummary:
    accepted: int
    rows: int
    csv_result: str
    learned: Dict[str, str]
    ignored: List[str]
    filled: List[str]
    problems: List[str] = field(default_factory=list)
    #: initials -> full name saved from the review's People sheet
    people: Dict[str, str] = field(default_factory=dict)

    def describe(self) -> str:
        s = ""
        if self.learned:
            s += f"learned {len(self.learned)} label mappings\n"
        if self.people:
            s += ("saved " + ", ".join(f"{k} = {v}" for k, v in sorted(self.people.items()))
                  + " to the people list\n")
        s += f"{self.accepted} values accepted -> {self.rows} rows: {self.csv_result}"
        s += "".join(f"\n  filled {f}" for f in self.filled)
        return s + "".join(f"\n  ! {w}" for w in self.problems[:30])


def run_apply(review_path, fill_dir=None, use_images: bool = True,
              image_reader: Optional[ImageReader] = None,
              progress: Optional[Listener] = None) -> ApplySummary:
    from progress_report import Progress
    from council_core import ledger_records
    if ledger_records.is_record_review(review_path):
        with Progress("ledger apply", listener=progress) as prog:
            r = ledger_records.apply_review(review_path, fill_dir or None, progress=prog)
            prog.finish(f"{r.records} records accepted -> {r.rows} rows: {r.csv_result}")
            return ApplySummary(r.accepted, r.rows, r.csv_result, dict(r.learned),
                                list(r.ignored), list(r.filled), list(r.warnings))
    with Progress("ledger apply", listener=progress) as prog:
        if use_images and image_reader is None:
            prog.update("starting the local image reader…")
            image_reader = default_image_reader()
        r = apply_review(review_path, fill_dir or None, image_reader=image_reader,
                         use_images=use_images, progress=prog)
        prog.finish(f"{r.accepted} values accepted -> {r.rows} rows: {r.csv_result}")
        return ApplySummary(r.accepted, r.rows, r.csv_result, dict(r.learned),
                            list(r.ignored), list(r.filled), list(r.warnings),
                            people=dict(getattr(r, "people", {}) or {}))


def open_path(path) -> None:
    """Open a workbook or folder in its default app (Excel for a review)."""
    p = str(path)
    if sys.platform.startswith("win"):
        os.startfile(p)                                   # noqa: S606
    elif sys.platform == "darwin":
        subprocess.Popen(["open", p])
    else:
        subprocess.Popen(["xdg-open", p])


# ── agent tools (for a human to register) ───────────────────────────────
def _inside(root: Path, raw: Any, what: str) -> Path:
    if not isinstance(raw, str) or not raw.strip():
        raise ValueError(f"'{what}' must be a non-empty string")
    p = Path(raw.strip())
    full = (p if p.is_absolute() else root / p).resolve()
    try:
        full.relative_to(root)
    except ValueError:
        raise PermissionError(f"{what}: {raw!r} is outside {root}") from None
    return full


def ledger_tools(policy) -> Dict[str, Any]:
    """`ingest_extract` and `ingest_apply` as safe_agent Tools.

    Sources are read only from policy.file_root; the review workbook, the
    parent CSV it applies to, and filled templates must all be inside
    policy.output_dir — the agent's one writeable area. Values reach the CSV
    only through a review a person saved, never on the model's say-so.
    Images are not read by these tools (no OCR inside an agent step's time
    limit); the review lists them for the person.

    Returned, not registered: process-start wiring calls
    ``registry.register(tool)`` for the ones it wants, then ``freeze()``.
    """
    from safe_agent import Tool

    def extract_tool(args: Dict[str, Any], pol) -> Dict[str, Any]:
        out_root, in_root = Path(pol.output_dir), Path(pol.file_root)
        csv = _inside(out_root, args.get("csv"), "csv")
        raw = args.get("sources")
        if not isinstance(raw, list) or not raw:
            return {"error": "'sources' must be a non-empty list of paths"}
        sources = [_inside(in_root, s, "sources") for s in raw]
        review = out_root / f"review_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        part = args.get("part", "")
        s = run_extract(csv, sources, review, part=part if isinstance(part, str) else "",
                        use_images=False)
        return {"review": str(review), "summary": s.describe()}

    def apply_tool(args: Dict[str, Any], pol) -> Dict[str, Any]:
        import json
        import openpyxl
        out_root, in_root = Path(pol.output_dir), Path(pol.file_root)
        review = _inside(out_root, args.get("review"), "review")
        wb = openpyxl.load_workbook(review, read_only=True)
        try:
            meta = json.loads(wb["_run"]["A1"].value)
        except Exception:
            return {"error": f"{review.name} is not a review workbook"}
        finally:
            wb.close()
        _inside(out_root, meta.get("csv"), "the review's parent CSV")
        for s in meta.get("sources", []):
            _inside(in_root, s, "the review's sources")
        fill = args.get("fill_dir")
        fill_dir = _inside(out_root, fill, "fill_dir") if fill else None
        s = run_apply(review, fill_dir, use_images=False)
        return {"summary": s.describe()}

    return {
        "ingest_extract": Tool(
            name="ingest_extract", fn=extract_tool, timeout_s=300,
            description="Read source Excel workbooks and write a review workbook of the "
                        "values found, for a person to check. Writes nothing else.",
            schema={"type": "object", "required": ["csv", "sources"], "properties": {
                "csv": {"type": "string", "description": "parent CSV (in the output area)"},
                "sources": {"type": "array", "items": {"type": "string"},
                            "description": ".xlsx files or folders (in the file root)"},
                "part": {"type": "string", "description": "part number, if not in them"}}}),
        "ingest_apply": Tool(
            name="ingest_apply", fn=apply_tool, timeout_s=300,
            description="Put the values a person accepted in a saved review workbook "
                        "into its parent CSV, and optionally write filled templates.",
            schema={"type": "object", "required": ["review"], "properties": {
                "review": {"type": "string"},
                "fill_dir": {"type": "string", "description": "optional, in the output area"}}}),
    }
