"""Standalone Granite Q&A pass on the mfg corpus — closes the loop with
row-level citations attached to every CSV/XLSX/JSON source the model used.

Runs as its own process so a CLIP-or-chromadb shutdown segfault in the
big eval doesn't take the QA results with it.
"""
from __future__ import annotations

import _windll_bootstrap  # noqa: F401

import time
from pathlib import Path

VAULT = Path.home() / ".council" / "vault"
CORPUS = VAULT / "mfg_eval"
GGUF = Path(__file__).parent / "models" / "granite-3.1-8b-instruct-Q4_K_M.gguf"


def _extract_for_llm(p: Path) -> str:
    suf = p.suffix.lower()
    try:
        if suf == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(p))
            return "\n".join((pg.extract_text() or "") for pg in r.pages[:5])
        if suf == ".docx":
            from docx import Document
            d = Document(str(p))
            return "\n".join(par.text for par in d.paragraphs if par.text)
        if suf in (".xlsx", ".xlsm"):
            import openpyxl as _xl
            wb = _xl.load_workbook(p, read_only=True, data_only=True)
            parts = []
            for sh in wb.worksheets:
                parts.append(f"## {sh.title}")
                for row in sh.iter_rows(max_row=30, values_only=True):
                    parts.append("\t".join("" if v is None else str(v) for v in row))
            return "\n".join(parts)
        return p.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""


def main() -> int:
    from vault_index import VaultIndex
    from llama_cpp import Llama
    from row_citations import cite_rows

    print(f"Vault:  {VAULT}")
    print(f"Corpus: {CORPUS}")
    print(f"Model:  {GGUF.name}")

    idx = VaultIndex(VAULT)
    print(f"  vault records: {len(idx.records)}")

    print("loading Granite...")
    t0 = time.time()
    llm = Llama(model_path=str(GGUF), n_gpu_layers=99, n_ctx=8192, verbose=False)
    print(f"  loaded in {time.time()-t0:.1f}s")

    questions = [
        "Why was ECN-2026-017 raised and which part does it cover?",
        "What was the disposition for NCR-99001 and what work order is it linked to?",
        "Summarise the BOM relationships for PART-3001 from parts_master.xlsx.",
        "What PPE does the MSDS require when handling PART-3001?",
    ]

    for q in questions:
        print("\n" + "=" * 80)
        print(f"Q: {q}")
        print("=" * 80)
        t0 = time.time()
        hits, _ = idx.search(q, k=4, folder="mfg_eval")
        dt_s = (time.time() - t0) * 1000
        print(f"retrieval: {len(hits)} hits in {dt_s:.0f} ms")
        ctx_blocks = []
        cited_paths = []
        for rank, (sc, rec) in enumerate(hits[:4], 1):
            p = Path(rec.get("path") or "")
            try:
                rel = p.relative_to(VAULT)
            except Exception:
                rel = p
            print(f"  {rank}. [{sc:6.2f}] {rel}")
            if not p.exists():
                continue
            cited_paths.append(p)
            txt = _extract_for_llm(p)
            if len(txt) > 1800:
                txt = txt[:1800] + "\n...[truncated]..."
            ctx_blocks.append(f"[SOURCE #{rank}] {p.name}\n{txt}")

        prompt = (
            "Use ONLY the sources below. Cite [SOURCE #N] inline.\n\n"
            + "\n---\n".join(ctx_blocks)
            + f"\n\nQUESTION: {q}\nANSWER:"
        )
        t0 = time.time()
        out = llm.create_chat_completion(
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2, max_tokens=320,
        )
        dt_l = time.time() - t0
        ans = out["choices"][0]["message"]["content"].strip()
        print(f"\nGranite ({dt_l:.1f}s):")
        for line in ans.splitlines():
            print(f"  {line}")
        rc = cite_rows(ans, cited_paths)
        if rc:
            print("\nRow-level citations:")
            for line in rc.splitlines():
                print(f"  {line}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
