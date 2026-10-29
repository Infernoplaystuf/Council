"""
tools/council_fulltest.py — the whole app on this PC, with real local models,
in one command, ending in one report.

    conda activate wizardCouncil
    python tools\\council_fulltest.py --model ollama:gemma3:12b
    python tools\\council_fulltest.py --model ollama:gemma3:12b --long

Each step is a separate process with its own time limit and log, so one that
hangs or crashes costs that step, not the run. Run it again with the same
--out to carry on: steps that finished are not run again (--redo ID,ID to
run some again). Everything is written under --out; every step runs against
a scratch vault there (COUNCIL_VAULT_ROOT), never your own vault or data.

STEPS (id: what it covers; [model] = needs one; [long] = only with --long)
  env          Python, Qt, Ollama models, GPU: what this run ran on
  suite        tests/run_all.py: every test file, incl. every tab building
               and being shown offscreen (no model)
  stubs        the model benches' own harnesses with canned answers: Typhon
               rebuild (oracle), Designer code writer (stub), Docs (oracle).
               Anything under 100% here is this PC, not a model
  ab           tools/council_ab.py: Council chat before/after PR #51,
               Connections links before/after #50, Ledger long-log speed
               before/after #48 [model]
  ledger       Ledger accuracy: ledger_bench's eleven synthetic workbooks
               (stacked, side by side, by sheet, long log, pictures, messy)
               read and scored against their answer key
  connections  Connections end to end on the synthetic Ironbridge vault:
               build the graph (timed), then the model reads the sentences
               and tables; links scored against the answer key [model]
  llm          council_core.llm_bench: GUI Designer layouts, code, docs
               answers, as the Code and Designer tabs make them [model]
  describe     Designer "Describe it": wireframe, generate, run [model]
  codebehind   Designer "Write it with the model…" [model]
  docs         Docs tab: answers from an invented package's docs [model]
  typhon       Rebuild the Typhon app: layout, wiring, widgets [model][long]
  code         Code tab agent redoes 3 small past commits of this repo
               until their tests pass [model][long]
  council      the whole council (Deliberation on) on council_bench's 20
               questions [model][long]

    <out>/report.md, report.html    one line a step, then each step's summary
    <out>/<step>.log                everything the step printed
    <out>/ab/report.html            the before/after comparison
"""
from __future__ import annotations

import argparse
import html
import json
import os
import platform
import shutil
import subprocess
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(HERE))

import council_ab as ab  # noqa: E402

PY = sys.executable


# ============================================================
# Steps that run in this file (as child processes)
# ============================================================

def env_report(out: Path) -> int:
    lines = [f"Python {sys.version.split()[0]} ({PY})",
             f"OS {platform.platform()}"]
    for mod in ("PySide6", "openpyxl", "pandas", "llama_cpp", "pytesseract"):
        try:
            m = __import__(mod)
            lines.append(f"{mod} {getattr(m, '__version__', 'installed')}")
        except Exception as exc:                          # noqa: BLE001
            lines.append(f"{mod}: not importable ({type(exc).__name__})")
    for cmd in (["git", "-C", str(REPO), "log", "-1", "--oneline"],
                ["ollama", "list"], ["tesseract", "--version"],
                ["nvidia-smi", "--query-gpu=name,memory.total,driver_version",
                 "--format=csv,noheader"]):
        try:
            p = subprocess.run(cmd, capture_output=True, text=True,
                               timeout=30)
            text = (p.stdout or p.stderr).strip().splitlines()
            lines.append(f"$ {' '.join(cmd[:2])}: " + (
                "\n  ".join(text[:12]) if text else "(nothing)"))
        except Exception as exc:                          # noqa: BLE001
            lines.append(f"$ {cmd[0]}: not available ({type(exc).__name__})")
    for k in sorted(os.environ):
        if k.startswith(("COUNCIL_", "DI_")):
            lines.append(f"{k}={os.environ[k]}")
    text = "\n".join(lines)
    print(text)
    (out / "env.txt").write_text(text + "\n", encoding="utf-8")
    return 0


def make_lmm_template(path: Path) -> Path:
    """The stand-in LMM template (the real one stays on your PC): the
    sections and field names ledger_bench writes for."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LMM"
    big = Font(size=16, bold=True)

    def head(cell, text):
        ws[cell] = text
        ws[cell].font = big

    def grid(top, rows, extras=()):
        r = top
        for b in "ABC":
            ws.cell(r, 2, "Flaw #")
            for i in range(10):
                ws.cell(r, 3 + i, f"{i + 1}{b}")
            for j, e in enumerate(extras):
                ws.cell(r, 13 + j, e)
            for k, label in enumerate(rows):
                ws.cell(r + 1 + k, 2, label.format(b=b))
            r += len(rows) + 2
        return r

    head("A1", "Part Attributes")
    for col, label in enumerate(
            ["Part Number", "Frequency (kHz)", "%pec", "Power (W)",
             "Scan Speed (mm/s)", "Jump Speed (mm/s)", "Beam Expander",
             "Spot Size (mm)", "Date", "Roughness"], start=1):
        ws.cell(2, col, label)
    head("A5", "Flaw Attributes")
    head("R5", "Tables for Presentation")
    end = grid(6, ["Laser Team FOD {b} (um)", "Flaw Orientation (deg)",
                   "Flaw Length (in)"])
    for rr in range(6, end):
        ws.cell(rr, 19, f"=C{rr}")
        ws.cell(rr, 20, f"=AVERAGE(C{rr}:L{rr})")
    head(f"A{end + 1}", "NDE Inspection")
    end = grid(end + 2, ["Eddy Current Response Pre-Conceal",
                         "Eddy Current Response Mid-Conceal",
                         "Eddy Current Response Post-Conceal", "FPI"],
               extras=["Inspector Initials", "Date"])
    head(f"A{end + 1}", "Concealment")
    for col, label in enumerate(["Media", "Media Run Time", "Media Size",
                                 "Tub Status", "Amplitude", "Amount", "Notes"],
                                start=2):
        ws.cell(end + 2, col, label)
    vis = end + 10
    head(f"A{vis}", "Visual Inspection")
    grid(vis + 1, ["Inspector 1", "Inspector 2", "Inspector 3"],
         extras=["Inspector Initials"])
    wb.save(path)
    return path


def ledger_accuracy(work: Path, images: bool) -> int:
    work.mkdir(parents=True, exist_ok=True)
    tpl = make_lmm_template(work / "LMM_standin_template.xlsx")
    src = work / "sources"
    if src.exists():
        shutil.rmtree(src)
    cmd = [PY, str(REPO / "ledger_bench.py")]
    for args in (["make", str(tpl), str(src)],
                 ["score", str(src)] + ([] if images else ["--no-images"])):
        t0 = time.monotonic()
        code = subprocess.call(cmd + args, cwd=REPO)
        print(f"  ledger_bench {args[0]}: {time.monotonic() - t0:.1f} s",
              flush=True)
        if code:
            return code
    return 0


def connections(vault: Path, model: str) -> int:
    """Build the graph over the Ironbridge vault, then (with a model) read
    its sentences and tables; score links against the answer key."""
    sys.path.insert(0, str(REPO))
    os.chdir(REPO)
    from council_core import kg_corpus, kg_extract as kx
    from council_core import knowledge_graph as kgm
    if vault.exists():
        shutil.rmtree(vault)
    vault.mkdir(parents=True)
    key = kg_corpus.build(vault)
    out: Dict[str, Any] = {}
    kg = kgm.KnowledgeGraph(vault)
    t0 = time.monotonic()
    kg.seed()
    out["seed_s"] = round(time.monotonic() - t0, 1)
    out["after_seed"] = kg.counts()
    out["found_after_seed"] = score_links(kg, key)
    print(f"  graph built in {out['seed_s']} s: {out['after_seed']}",
          flush=True)
    print(f"  answer-key links found: {out['found_after_seed']['found']}/"
          f"{out['found_after_seed']['total']}", flush=True)
    if model:
        t0 = time.monotonic()
        out["text_stats"] = kg.extract_from_text(
            {f"this PC ({model})": (kx.ollama_chat(model), model)},
            model=model)
        out["text_s"] = round(time.monotonic() - t0, 1)
        t0 = time.monotonic()
        out["table_stats"] = kg.read_tables(kx.ollama_chat(model),
                                            model=model)
        out["tables_s"] = round(time.monotonic() - t0, 1)
        out["after_model"] = kg.counts()
        out["found_after_model"] = score_links(kg, key)
        print(f"  model read sentences in {out['text_s']} s, tables in "
              f"{out['tables_s']} s", flush=True)
        print(f"  answer-key links found: {out['found_after_model']['found']}"
              f"/{out['found_after_model']['total']} "
              f"(missed: {', '.join(out['found_after_model']['missed'][:8])})",
              flush=True)
    kg.close()
    (vault.parent / "connections.json").write_text(
        json.dumps(out, indent=1, default=str), encoding="utf-8")
    return 0


def score_links(kg, key: Dict[str, Any]) -> Dict[str, Any]:
    """Answer-key relations the graph holds (any status but rejected)."""
    ents = key["entities"]

    def ids(k: str) -> set:
        e = ents.get(k) or {}
        got = set()
        for name in [e.get("name", "")] + list(e.get("aliases") or []):
            if name:
                got |= {r["id"] for r in kg.search(name, exact=True)}
        return {kg.resolve(i) for i in got}

    found, missed = 0, []
    for r in key["relations"]:
        s, o = ids(r["subject"]), ids(r["object"])
        hit = False
        for si in s:
            for oi in o:
                row = kg.db.execute(
                    "SELECT 1 FROM relations WHERE subject_id=? AND "
                    "predicate=? AND object_id=? AND status != 'rejected'",
                    (si, r["predicate"], oi)).fetchone()
                if row:
                    hit = True
        if hit:
            found += 1
        else:
            missed.append(f"{r['subject']} {r['predicate']} {r['object']}")
    return {"found": found, "total": len(key["relations"]),
            "missed": missed}


# ============================================================
# The run
# ============================================================

def steps(a, out: Path, vault: Path) -> List[Dict[str, Any]]:
    me = str(Path(__file__).resolve())
    model = a.model
    gguf = model if model and not model.startswith("ollama:") else ""
    ollama = model[len("ollama:"):] if model.startswith("ollama:") else ""
    scratch = str(out / "scratch-vault")
    ab_args = [PY, "-u", str(HERE / "council_ab.py"), "--out",
               str(out / "ab"), "--repeat", str(a.repeat)]
    if model:
        ab_args += ["--model", model]
    if a.from_vault:
        ab_args += ["--from-vault", a.from_vault]
    S = [
        {"id": "env", "cmd": [PY, "-u", me, "_env", "--out", str(out)],
         "limit": 120},
        {"id": "suite", "cmd": [PY, "-u", str(REPO / "tests" / "run_all.py"),
                                "--logs", str(out / "suite-logs")],
         "limit": 120 * 60},
        {"id": "stubs", "cmd": [PY, "-u", me, "_stubs", "--out", str(out)],
         "limit": 60 * 60},
        {"id": "ab", "cmd": ab_args, "limit": 8 * 3600, "model": True},
        {"id": "ledger", "cmd": [PY, "-u", me, "_ledger", "--out", str(out)]
         + (["--images"] if a.ledger_images else []), "limit": 30 * 60},
        {"id": "connections", "cmd": [PY, "-u", me, "_connections", "--out",
                                      str(out), "--kg-model",
                                      a.kg_model or ollama],
         "limit": 3 * 3600},
        {"id": "llm", "cmd": [PY, "-u", "-m", "council_core.llm_bench",
                              "--suites", "gui,code,docs", "--pipeline",
                              "new", "--vault", str(out / "llm-vault"),
                              "--out", str(out / "llm_bench.json"),
                              "--max-minutes", "180"]
         + (["--model", model] if model else []),
         "limit": 4 * 3600, "model": True},
        {"id": "describe", "cmd": [PY, "-u", "run_describe_prompts.py",
                                   "--vault", scratch, "--out",
                                   str(out / "describe.jsonl")]
         + (["--model", gguf] if gguf else []),
         "limit": 3 * 3600, "model": True},
        {"id": "codebehind", "cmd": [PY, "-u", "run_codebehind_prompts.py",
                                     "--vault", scratch, "--out",
                                     str(out / "codebehind.jsonl")]
         + (["--model", model] if model else []),
         "limit": 3 * 3600, "model": True},
        {"id": "docs", "cmd": [PY, "-u", "-m", "council_core.docs_bench",
                               "--out", str(out / "docs_bench.json")],
         "limit": 2 * 3600, "model": True},
        {"id": "typhon", "cmd": [PY, "-u", "run_typhon_recreate.py",
                                 "--vault", scratch, "--out",
                                 str(out / "typhon")]
         + (["--model", gguf] if gguf else []),
         "limit": 6 * 3600, "model": True, "long": True},
        {"id": "code", "cmd": [PY, "-u", "-m", "council_core.code_bench",
                               str(REPO), "--count", "3", "--label",
                               "fulltest", "--vault", scratch],
         "limit": 6 * 3600, "model": True, "long": True},
        {"id": "council", "cmd": [PY, "-u", "-m",
                                  "council_core.council_bench", "run",
                                  "--label", "fulltest", "--vault",
                                  str(vault)],
         "limit": 6 * 3600, "model": True, "long": True},
    ]
    return S


def stubs(out: Path) -> int:
    scratch = out / "stub-vault"
    bad = 0
    for name, cmd in (
            ("Typhon rebuild, oracle answers",
             [PY, "-u", "run_typhon_recreate.py", "--vault", str(scratch),
              "--stub", "oracle", "--out", str(out / "typhon-oracle")]),
            ("Designer code writer, canned answers",
             [PY, "-u", "run_codebehind_prompts.py", "--vault",
              str(scratch), "--stub"]),
            ("Docs bench, answers from the key",
             [PY, "-u", "-m", "council_core.docs_bench", "--oracle"])):
        print(f"--- {name}", flush=True)
        t0 = time.monotonic()
        code = subprocess.call(cmd, cwd=REPO)
        print(f"--- {name}: exit {code}, {time.monotonic() - t0:.0f} s",
              flush=True)
        bad += code != 0
    return 1 if bad else 0


def tail(path: Path, n: int = 25) -> str:
    try:
        lines = path.read_text(encoding="utf-8",
                               errors="replace").splitlines()
    except OSError:
        return ""
    lines = [x for x in lines if not x.startswith("[DEBUG")]
    return "\n".join(lines[-n:])


def run_step(s: Dict[str, Any], out: Path, env: Dict[str, str]
             ) -> Dict[str, Any]:
    log = out / f"{s['id']}.log"
    t0 = time.monotonic()
    rec = {"id": s["id"], "started": time.strftime("%H:%M:%S")}
    with log.open("w", encoding="utf-8") as fh:
        fh.write("$ " + " ".join(s["cmd"]) + "\n\n")
        fh.flush()
        try:
            p = subprocess.run(s["cmd"], cwd=REPO, env=env, stdout=fh,
                               stderr=subprocess.STDOUT, timeout=s["limit"])
            rec["status"] = "ok" if p.returncode == 0 else \
                f"failed (exit {p.returncode})"
        except subprocess.TimeoutExpired:
            rec["status"] = f"stopped after {s['limit'] // 60} min"
        except OSError as exc:
            rec["status"] = f"could not start: {exc}"
    rec["seconds"] = round(time.monotonic() - t0, 1)
    return rec


def write_report(out: Path, results: List[Dict[str, Any]],
                 meta: Dict[str, Any]) -> str:
    md = [f"# Council full test — {meta['when']}", "",
          f"Model: {meta['model'] or 'your saved setup'}. "
          f"Build: {meta['build']}.", "",
          "| step | result | minutes |", "|---|---|---|"]
    for r in results:
        md.append(f"| {r['id']} | {r['status']} | "
                  f"{r.get('seconds', 0) / 60:.1f} |")
    h = ["<!doctype html><meta charset='utf-8'><title>Council full test"
         "</title><style>body{font:14px/1.45 system-ui,sans-serif;margin:24px;"
         "max-width:1100px;color:#1d2330;background:#fff}table{border-"
         "collapse:collapse}td,th{border:1px solid #d5d9e0;padding:4px 8px;"
         "text-align:left}th{background:#f1f3f6}pre{background:#f6f7f9;"
         "padding:8px;white-space:pre-wrap;font-size:12px}.ok{color:#0b7a3e}"
         ".bad{color:#b3261e}</style>",
         f"<h1>Council full test</h1><p>{html.escape(md[2])}</p>",
         "<table><tr><th>step</th><th>result</th><th>minutes</th></tr>"]
    for r in results:
        cls = "ok" if r["status"] in ("ok", "skipped: no model",
                                      "skipped: needs --long") else "bad"
        h.append(f"<tr><td><a href='#{r['id']}'>{r['id']}</a></td>"
                 f"<td class='{cls}'>{html.escape(r['status'])}</td>"
                 f"<td>{r.get('seconds', 0) / 60:.1f}</td></tr>")
    h.append("</table>")
    if (out / "ab" / "report.html").is_file():
        md += ["", "Before/after comparison: ab/report.html"]
        h.append("<p><a href='ab/report.html'>Before/after comparison"
                 "</a></p>")
    for r in results:
        t = tail(out / f"{r['id']}.log")
        if not t:
            continue
        md += ["", f"## {r['id']} — {r['status']}", "", "```", t, "```"]
        h.append(f"<h2 id='{r['id']}'>{r['id']} — "
                 f"{html.escape(r['status'])}</h2><pre>{html.escape(t)}"
                 "</pre>")
    text = "\n".join(md) + "\n"
    (out / "report.md").write_text(text, encoding="utf-8")
    (out / "report.html").write_text("\n".join(h), encoding="utf-8")
    return text


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(prog="council_fulltest", description=__doc__,
                                 formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("internal", nargs="?", default="",
                    help=argparse.SUPPRESS)
    ap.add_argument("--model", default="",
                    help="ollama:NAME or a .gguf path (default: your saved "
                         "setup, copied from --from-vault)")
    ap.add_argument("--out", default="")
    ap.add_argument("--long", action="store_true",
                    help="also the slow steps (typhon, code, council)")
    ap.add_argument("--only", default="", help="step ids, comma separated")
    ap.add_argument("--skip", default="", help="step ids, comma separated")
    ap.add_argument("--redo", default="",
                    help="run these again though they finished")
    ap.add_argument("--repeat", type=int, default=1,
                    help="ask the before/after questions this many times")
    ap.add_argument("--from-vault", default="")
    ap.add_argument("--kg-model", default="",
                    help="Ollama model for Connections (default: --model's)")
    ap.add_argument("--images", dest="ledger_images", action="store_true",
                    help="ledger: read the pictures too (needs Tesseract)")
    a = ap.parse_args(argv)
    out = Path(a.out) if a.out else (
        Path.home() / "council_fulltest" / time.strftime("%Y%m%d-%H%M%S"))
    out.mkdir(parents=True, exist_ok=True)

    if a.internal == "_env":
        return env_report(out)
    if a.internal == "_stubs":
        return stubs(out)
    if a.internal == "_ledger":
        return ledger_accuracy(out / "ledger", a.ledger_images)
    if a.internal == "_connections":
        return connections(out / "connections-vault", a.kg_model)

    # The scratch vault every step runs in: synthetic files, and your model
    # settings when no --model is given.
    vault = out / "vault"
    if not vault.exists():
        from_vault = Path(a.from_vault) if a.from_vault else (
            Path(os.environ["COUNCIL_VAULT_ROOT"])
            if os.environ.get("COUNCIL_VAULT_ROOT") else REPO / "vault")
        ab.make_vault(vault, None, None if a.model else from_vault)
    env = ab.model_env(a.model)
    env["COUNCIL_VAULT_ROOT"] = str(vault)
    env.setdefault("PYTHONIOENCODING", "utf-8")
    env.pop("COUNCIL_SAMPLE_DATA", None)

    state_path = out / "steps.json"
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        state = {}
    only = {x for x in a.only.split(",") if x}
    skip = {x for x in a.skip.split(",") if x}
    redo = {x for x in a.redo.split(",") if x}
    sha = subprocess.run(["git", "-C", str(REPO), "log", "-1",
                          "--format=%h %s"], capture_output=True, text=True
                         ).stdout.strip()
    meta = {"when": time.strftime("%Y-%m-%d %H:%M"), "model": a.model,
            "build": sha}
    print(f"Output: {out}", flush=True)
    results = []
    for s in steps(a, out, vault):
        sid = s["id"]
        if (only and sid not in only) or sid in skip:
            continue
        if s.get("long") and not a.long and sid not in only:
            rec = {"id": sid, "status": "skipped: needs --long"}
        elif state.get(sid, {}).get("status") == "ok" and sid not in redo:
            rec = state[sid]
            print(f"== {sid}: done before ({rec['status']})", flush=True)
        else:
            print(f"== {sid} (log: {out / (sid + '.log')})", flush=True)
            rec = run_step(s, out, env)
            print(f"   {rec['status']}, {rec['seconds'] / 60:.1f} min",
                  flush=True)
            state[sid] = rec
            state_path.write_text(json.dumps(state, indent=1),
                                  encoding="utf-8")
        results.append(rec)
        write_report(out, results, meta)
    print(write_report(out, results, meta))
    print(f"Report: {out / 'report.html'}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
