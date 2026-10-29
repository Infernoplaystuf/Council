"""The Ledger extract no longer freezes or crawls (review, 2026-10-19).

What the stall hunt measured, and what each test pins:
- one 2,000-row sheet took 276 s (every 'what is on row r' scanned the
  whole sheet); 1,000 rows now take a few seconds;
- one formatted cell at row 1,048,576 made openpyxl build 21 million
  cells (4 GB, multi-second GC pauses holding the GIL);
- a space-padded cell made the initials-heading pattern backtrack for
  54 s with the GIL held;
- the review writer and 'Other values offered' were quadratic;
- Cancel did nothing inside a workbook, a Cancel that killed Excel lost
  the whole run, and Cancel during the last workbook was not recorded -
  and then a Cancel pressed after the last workbook was read made a whole
  review a cancelled, 'partial' one;
- Apply re-read the sources without the part-name mask;
- the stall dump could not be written while the app was frozen.
"""
from __future__ import annotations

import json
import sys
import threading
import time
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

sys.path.insert(0, str(Path(__file__).parent))
import ledger_fixtures as lf  # noqa: E402
import source_extract as se  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import stall_watchdog  # noqa: E402
from sample_ledger import ParentCSV  # noqa: E402


def _log_book(path: Path, rows: int, cols: int = 20) -> Path:
    """A long inspection log: header + rows of mixed text and numbers."""
    head = ["Part Number", "Flaw ID", "Inspector", "Date", "Stage", "Response",
            "Notes", "Laser Power (W)", "Frequency (kHz)", "Scan Speed (mm/s)",
            "Comment", "Operator", "Batch", "Station", "Shift", "Lot",
            "Length (mm)", "Orientation (deg)", "Passes", "Remarks"][:cols]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inspection Log"
    ws.append(head)
    words = ["ok", "see note", "rework", "pass", "visual check", "porosity"]
    for r in range(rows):
        row = []
        for c, h in enumerate(head):
            if h == "Part Number":
                row.append(f"PN-{1000 + r // 50}")
            elif h == "Flaw ID":
                row.append(lf.ITEMS[r % len(lf.ITEMS)])
            elif h in ("Inspector", "Operator"):
                row.append(["JB", "MK", "TR"][r % 3])
            elif h == "Date":
                row.append(f"2026-09-{1 + r % 28:02d}")
            elif h == "Stage":
                row.append(["EC Pre", "EC Post", "Pre-Conceal"][r % 3])
            elif h in ("Notes", "Comment", "Remarks", "Station", "Shift"):
                row.append(words[(r + c) % len(words)])
            else:
                row.append(round(10 + (r * 7 + c) % 90 + 0.5, 2))
        ws.append(row)
    wb.save(path)
    return path


@pytest.fixture()
def ledger(tmp_path):
    tpl = lf.make_template(tmp_path / "template.xlsx")
    ps = li.create_parent(tpl)
    return Path(ps.csv), ParentCSV(ps.csv).schema


# ---------------------------------------------- speed

def test_a_1000_row_sheet_is_read_in_seconds_not_minutes(tmp_path, ledger):
    csv, schema = ledger
    src = _log_book(tmp_path / "log.xlsx", 1000)
    t0 = time.perf_counter()
    run = se.extract(schema, [src], use_images=False)
    took = time.perf_counter() - t0
    assert run.obs, "the log gave values"
    # was ~80 s on this shape; generous for a slow CI box
    assert took < 30, f"took {took:.1f}s"


def test_one_far_away_formatted_cell_does_not_fill_the_sheet(tmp_path):
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 6):
        ws.cell(r, 1, f"Label {r}")
        ws.cell(r, 2, r * 10)
    ws.cell(1048576, 20).font = Font(bold=True)           # styled, no value
    wb.save(tmp_path / "stray.xlsx")
    t0 = time.perf_counter()
    grids, _images = se.load_workbook_grids(tmp_path / "stray.xlsx")
    assert time.perf_counter() - t0 < 5
    assert len(grids[0].cells) == 10


def test_the_initials_heading_pattern_is_linear_on_padded_cells():
    cell = "NOTE:" + " " * 4000 + "x"
    t0 = time.perf_counter()
    assert se._WHO_HEADING.match(cell.strip()) is None
    assert time.perf_counter() - t0 < 0.5                  # was 54 s
    for ok in ("JB", "JB — 2026-09-18", "MK - 09/18", "TR (2026-09-18)"):
        assert se._WHO_HEADING.match(ok), ok
    for no in ("Notes", "JB was here", "1A"):
        assert not se._WHO_HEADING.match(no), no


def test_other_values_offered_is_capped(ledger):
    _csv, schema = ledger
    f = next(x for x in schema.fields if not x.per_item and x.name != schema.id_field)
    obs = [se.Obs("101", "", f.name, str(v), 0.9, "match", "x", f"s.xlsx!B{v}")
           for v in range(1, 101)]
    decs = se.decide(se.Extraction(["s.xlsx"], obs=obs), schema)
    text = decs[0].conflict
    assert text.endswith(", and 79 more")
    assert text.count(",") == se.OTHERS_SHOWN


def test_write_review_with_many_rows_is_fast(tmp_path, ledger):
    """It grew faster than the rows did (~10 s at 6k, 75 s at 16k). Checked
    by how it scales - four times the rows may take at most about six times
    as long - not by a number of seconds, which depends on the machine: a
    CI runner took 23 s at 6k where a desktop takes a few."""
    csv, schema = ledger
    f = next(x for x in schema.fields if not x.per_item and x.name != schema.id_field)

    def write(n: int, name: str) -> float:
        obs = [se.Obs(str(1000 + i), "", f.name, str(i), 0.5, "ambiguous", "x",
                      f"s.xlsx!B{i}") for i in range(n)]
        run = se.Extraction(["s.xlsx"], obs=obs)
        decs = se.decide(run, schema)
        t0 = time.perf_counter()
        se.write_review(tmp_path / name, run, decs, schema, str(csv))
        return time.perf_counter() - t0

    small = write(1500, "small.xlsx")
    big = write(6000, "r.xlsx")
    assert big < 6 * small + 2, (small, big)              # linear is 4x
    assert big < 120                                       # and never a stall
    ws = openpyxl.load_workbook(tmp_path / "r.xlsx")["Values"]
    assert ws.cell(2, 1).fill.fgColor.rgb.endswith("FFF2CC")   # amber row kept


# ---------------------------------------------- cancel

class _Prog:
    """A progress stand-in that presses Cancel once a sheet is being read."""
    def __init__(self, cancel, when):
        self.cancel, self.when = cancel, when

    def stage(self, *a, **k):
        pass

    def update(self, what="", **k):
        if self.when in str(what):
            self.cancel.set()

    def advance(self, *a, **k):
        pass


def test_cancel_inside_a_big_workbook_stops_it_and_keeps_the_earlier_one(tmp_path, ledger):
    _csv, schema = ledger
    small = lf.make_laser_log(tmp_path / "a_laser.xlsx")
    big = _log_book(tmp_path / "b_big.xlsx", 1500)
    cancel = threading.Event()
    run = se.extract(schema, [small, big], use_images=False, cancel=cancel,
                     progress=_Prog(cancel, "b_big.xlsx: finding values"))
    assert run.cancelled
    assert run.obs and all("a_laser" in o.where for o in run.obs)
    assert any("while reading 'b_big.xlsx'" in p for p in run.problems)


def test_cancel_that_kills_excel_mid_read_ends_cleanly(tmp_path, ledger, monkeypatch):
    _csv, schema = ledger
    a = lf.make_laser_log(tmp_path / "a.xlsx")
    b = tmp_path / "b.xls"
    b.write_bytes(b"x")
    cancel = threading.Event()
    real = se.load_workbook_grids

    def fake(path, *args, **kw):
        if Path(path).name == "b.xls":
            cancel.set()                                   # Cancel killed Excel...
            raise RuntimeError("com_error(-2147023174, 'The RPC server is unavailable')")
        return real(path, *args, **kw)
    monkeypatch.setattr(se, "load_workbook_grids", fake)
    run = se.extract(schema, [a, b], use_images=False, cancel=cancel)
    assert run.cancelled and run.obs                       # a's values kept
    assert not run.open_failures                           # b is not blamed


def test_a_workbook_that_fails_oddly_does_not_end_the_run(tmp_path, ledger, monkeypatch):
    _csv, schema = ledger
    a = tmp_path / "a.xlsx"
    a.write_bytes(b"x")
    b = lf.make_laser_log(tmp_path / "b.xlsx")
    real = se.load_workbook_grids

    def fake(path, *args, **kw):
        if Path(path).name == "a.xlsx":
            raise RuntimeError("something unexpected")
        return real(path, *args, **kw)
    monkeypatch.setattr(se, "load_workbook_grids", fake)
    run = se.extract(schema, [a, b], use_images=False)
    assert not run.cancelled
    assert run.open_failures == [("a.xlsx", "other")]
    assert run.obs and any("something unexpected" in p for p in run.problems)


def test_cancel_during_the_last_workbook_is_recorded(tmp_path, ledger):
    _csv, schema = ledger
    big = _log_book(tmp_path / "only.xlsx", 1500)
    cancel = threading.Event()
    run = se.extract(schema, [big], use_images=False, cancel=cancel,
                     progress=_Prog(cancel, "finding values"))
    assert run.cancelled and not run.obs


class _CancelWhenRead(_Prog):
    """Presses Cancel once a workbook has been read to the end."""
    def update(self, what="", **k):
        pass

    def advance(self, *a, what="", **k):
        if self.when in str(what):
            self.cancel.set()


def _partless(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Build"
    ws["A1"], ws["B1"] = "Frequency (kHz)", 55
    ws["A2"], ws["B2"] = "Zorp rate", 7
    wb.save(path)
    return path


def test_a_cancel_after_the_last_workbook_was_read_is_a_whole_review(tmp_path, ledger):
    """Pressed after the last workbook was read, it marked a whole review
    'Cancelled... partial review' (review, 2026-10-19). With Council's
    model ticked, the Cancel stops its question: cancelled, and said as one
    pressed while it is asked is - a review with every value the rules
    found, not a partial one (second review, 2026-10-20)."""
    csv, schema = ledger
    a = lf.make_laser_log(tmp_path / "a_laser.xlsx")
    b = _partless(tmp_path / "b_nopart.xlsx")
    cancel = threading.Event()
    run = se.extract(schema, [a, b], use_images=False, cancel=cancel,
                     progress=_CancelWhenRead(cancel, "b_nopart.xlsx read in"))
    assert cancel.is_set() and not run.cancelled
    assert not any("ancel" in p for p in run.problems)
    assert {o.where.split(" › ")[0] for o in run.obs} == {"a_laser.xlsx", "b_nopart.xlsx"}

    # the Ledger tab's call, Council's model ticked: a whole review, and the
    # model is not asked about the labels after the Cancel
    asked = []

    def chat(prompt, _schema):
        asked.append(prompt)
        return json.dumps({"part": "", "evidence": ""} if "Which part" in prompt
                          else {"mappings": []}), "fake-granite"
    cancel = threading.Event()
    s = li._extract(str(csv), [str(a), str(b)], tmp_path / "review.xlsx", "", False, None,
                    _CancelWhenRead(cancel, "b_nopart.xlsx read in"), None, cancel, None,
                    True, chat)
    assert cancel.is_set() and s.cancelled and s.read_all and s.values
    assert s.describe().startswith("CANCELLED after every workbook was read — ")
    assert asked and all("Which part" in p for p in asked)
    assert ("Cancelled before Council's model was asked about the labels left over; every "
            "value the rules found is in this review.") in s.problems


def test_a_cancel_that_cut_the_last_workbook_short_is_still_a_cancelled_run(
        tmp_path, ledger, monkeypatch):
    """Cut short where the read swallowed it - Excel killed while it copied
    a picture out, the model stopped mid-answer: that workbook is not in
    the review, and the run is cancelled."""
    _csv, schema = ledger
    a = lf.make_laser_log(tmp_path / "a_laser.xlsx")
    b = _partless(tmp_path / "b_nopart.xlsx")
    cancel = threading.Event()

    def stopped_model(name, grids):
        cancel.set()                         # Cancel pressed: the model stops
        return "", ""
    run = se.extract(schema, [a, b], use_images=False, cancel=cancel,
                     part_finder=stopped_model)
    assert run.cancelled
    assert run.obs and all("a_laser" in o.where for o in run.obs)
    assert any("while reading 'b_nopart.xlsx'" in p for p in run.problems)

    real = se.load_workbook_grids

    def killed_while_copying(path, *args, **kw):
        out = real(path, *args, **kw)
        if Path(path).name == "b_nopart.xlsx":
            cancel.set()                     # what Excel gave back is cut short
        return out
    monkeypatch.setattr(se, "load_workbook_grids", killed_while_copying)
    cancel = threading.Event()
    run = se.extract(schema, [a, b], use_images=False, cancel=cancel)
    assert run.cancelled
    assert run.obs and all("a_laser" in o.where for o in run.obs)


# ---------------------------------------------- Apply keeps the mask

def test_apply_reads_the_sources_with_the_extracts_mask_and_part(tmp_path, ledger, monkeypatch):
    csv, schema = ledger
    src = lf.make_laser_log(tmp_path / "insp_ABC1DE234_run3.xlsx")
    s = li.run_extract(str(csv), [str(src)], use_images=False, pattern_mask="@@@#@@###",
                       part="")
    meta = json.loads(openpyxl.load_workbook(s.review)["_run"]["A1"].value)
    assert meta["pattern_mask"] == "@@@#@@###"
    seen = {}
    real = se.extract

    def spy(*a, **kw):
        seen.update(kw)
        return real(*a, **kw)
    monkeypatch.setattr(se, "extract", spy)
    li.run_apply(s.review, use_images=False)
    assert seen.get("pattern_mask") == "@@@#@@###"


# ---------------------------------------------- the stall dump

def test_a_frozen_app_gets_a_faulthandler_dump_in_the_file(tmp_path):
    """A C call that holds the GIL freezes the whole app and the Python
    watchdog with it; faulthandler's own thread still writes the stacks."""
    import re as _re
    out = tmp_path / "logs" / "stalls.txt"
    with stall_watchdog.watch("frozen step", lambda _t: None, stall_s=1.0, sample_s=0.5,
                              dump_file=str(out)):
        time.sleep(0.6)
        _re.match(r"(a+)+$", "a" * 28 + "b")          # ~5 s holding the GIL
    text = out.read_text(encoding="utf-8")
    assert "frozen step: watching" in text
    assert "most recent call first" in text                # faulthandler's dump
    assert "test_a_frozen_app_gets_a_faulthandler_dump" in text


def test_a_slow_but_running_app_is_reported_without_faulthandler(tmp_path):
    """Nothing moves but the app runs: the Python watchdog reports it, and
    faulthandler - which reads other threads unlocked, and crashed the
    process once while a thread exited - stays quiet."""
    out = tmp_path / "logs" / "stalls.txt"
    said = []
    with stall_watchdog.watch("slow step", said.append, stall_s=1.0, sample_s=0.5,
                              dump_file=str(out)):
        time.sleep(2.5)                                    # GIL free: not frozen
    assert said and str(out) in said[0]
    assert "most recent call first" not in out.read_text(encoding="utf-8")


def test_cancel_is_only_offered_for_the_extract(tmp_path):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    from council_qt.tabs.ledger import LedgerActions, LedgerTab
    QApplication.instance() or QApplication([])
    gate = threading.Event()
    seen = {}

    class Acts(LedgerActions):
        def apply(self, review, fill_dir, use_images, progress):
            seen["cancel_enabled"] = tab.cancel_btn.isEnabled()
            gate.wait(5)
            raise se.LedgerError("stop")

    tab = LedgerTab(actions=Acts())
    try:
        tab.review.setText(str(tmp_path / "r.xlsx"))
        tab.on_apply()
        time.sleep(0.2)
        QApplication.processEvents()
        assert tab.busy and not tab.cancel_btn.isEnabled()
        gate.set()
        assert tab.wait(10)
    finally:
        gate.set()
        tab.deleteLater()


# ---------------------------------------------- Excel pictures off the GIL

class _Shape:
    def __init__(self):
        self.copied = 0

    def CopyPicture(self, **kw):
        self.copied += 1


def _dib(color=(200, 30, 90)) -> bytes:
    import io
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (9, 5), color).save(b, "BMP")
    return b.getvalue()[14:]                               # CF_DIB: no file header


def test_a_picture_is_read_through_pywin32_not_imagegrab(monkeypatch):
    import io
    import excel_open
    from PIL import Image, ImageGrab
    monkeypatch.setattr(excel_open, "_clipboard_dib", lambda timeout=5.0: _dib())
    monkeypatch.setattr(ImageGrab, "grabclipboard",
                        lambda: pytest.fail("ImageGrab holds the GIL; not used"))
    shape = _Shape()
    png = excel_open._copy_picture(shape)
    assert shape.copied == 1
    img = Image.open(io.BytesIO(png))
    assert img.size == (9, 5) and img.convert("RGB").getpixel((0, 0)) == (200, 30, 90)


def test_imagegrab_is_still_the_fallback(monkeypatch):
    import excel_open
    from PIL import Image, ImageGrab

    def broken(timeout=5.0):
        raise OSError("no pywin32 clipboard")
    monkeypatch.setattr(excel_open, "_clipboard_dib", broken)
    monkeypatch.setattr(ImageGrab, "grabclipboard", lambda: Image.new("RGB", (4, 4)))
    assert excel_open._copy_picture(_Shape())


# ---------------------------------------------- what a folder scan reads

def _accept_all(review: str) -> None:
    wb = openpyxl.load_workbook(review)
    v = wb["Values"]
    for r in range(2, v.max_row + 1):
        v.cell(r, 1, "Y")
    wb.save(review)


def test_same_named_workbooks_in_sibling_folders_are_both_applied(tmp_path, ledger):
    """A folder per sample, each with its own laser_log.xlsx: Apply matched
    the reviewer's choices by file name, so one sample was lost."""
    csv, _schema = ledger
    root = tmp_path / "samples"
    for part, power in (("201", 280), ("202", 300)):
        d = root / part
        d.mkdir(parents=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Part {part}"
        ws["A1"], ws["B1"] = "Laser Power (W)", power
        wb.save(d / "laser_log.xlsx")
    s = li.run_extract(str(csv), [str(root)], use_images=False)
    wheres = {r[8] for r in openpyxl.load_workbook(s.review)["Values"].iter_rows(
        min_row=2, values_only=True)}
    assert any(w.startswith("201/laser_log.xlsx") for w in wheres)
    assert any(w.startswith("202/laser_log.xlsx") for w in wheres)
    _accept_all(s.review)
    li.run_apply(s.review, use_images=False)
    parts = {r[ParentCSV(csv).schema.id_field] for r in ParentCSV(csv).rows}
    assert {"201", "202"} <= parts


def test_a_folder_scan_leaves_the_ledgers_own_outputs_out(tmp_path, ledger):
    csv, schema = ledger
    root = tmp_path / "inspections"
    root.mkdir()
    lf.make_laser_log(root / "laser_101.xlsx")
    s = li.run_extract(str(csv), [str(root)], use_images=False)
    review = Path(s.review)
    (root / "old_review.xlsx").write_bytes(review.read_bytes())   # a review kept here
    from sample_ledger import fill_template
    fill_template(schema, [{schema.id_field: "101"}] if not schema.item_field else
                  [{schema.id_field: "101", schema.item_field: schema.items[0]}],
                  root / "filled_101.xlsx")
    kept, skipped = li.collect_sources([str(root)], template_path=schema.template_path)
    assert [p.name for p in kept] == ["laser_101.xlsx"]
    assert {p.name for p in skipped} == {"old_review.xlsx", "filled_101.xlsx"}
    # picked by name, a file is read as asked
    kept2, _ = li.collect_sources([str(root / "filled_101.xlsx")],
                                  template_path=schema.template_path)
    assert [p.name for p in kept2] == ["filled_101.xlsx"]


def test_a_folder_scan_does_not_follow_a_junction(tmp_path):
    _winapi = pytest.importorskip("_winapi")
    if not hasattr(_winapi, "CreateJunction"):
        pytest.skip("no junctions here")
    root = tmp_path / "root"
    (root / "real").mkdir(parents=True)
    lf.make_laser_log(root / "real" / "a.xlsx")
    _winapi.CreateJunction(str(root / "real"), str(root / "loop"))
    assert [p.name for p in se.workbook_files(root)] == ["a.xlsx"]


# ---------------------------------------------- huge padded cells (end-to-end test, 2026-10-19)

@pytest.mark.parametrize("cell", ["NOTE:" + " " * 32000 + "x", "MK" + " " * 32000,
                                  "  EC  Pre " + " " * 20000 + "(%FSH)"])
def test_normalizing_a_cell_padded_to_excels_limit_is_instant(cell):
    import sample_ledger as sl
    t0 = time.perf_counter()
    sl.normalize_label(cell)
    sl._split_unit(cell)
    assert time.perf_counter() - t0 < 0.2                  # was 2.5-3.8 s holding the GIL


@pytest.mark.parametrize("label,want", [("Flaw Length (in)", "flaw length"),
                                        ("  EC   Pre  (%FSH) ", "ec pre"),
                                        ("Power (W)", "power"), ("Notes:", "notes")])
def test_normalize_label_gives_the_same_names(label, want):
    import sample_ledger as sl
    assert sl.normalize_label(label) == want


def test_a_sheet_of_padded_cells_extracts_in_seconds(tmp_path, ledger):
    _csv, schema = ledger
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 41):
        ws.cell(r, 1, "NOTE:" + " " * 8000 + f"x{r}")
        ws.cell(r, 2, "MK" + " " * 8000)
    ws["D1"], ws["E1"] = "Power (W)", 280
    wb.save(tmp_path / "padded.xlsx")
    t0 = time.perf_counter()
    run = se.extract(schema, [tmp_path / "padded.xlsx"], use_images=False)
    assert time.perf_counter() - t0 < 10                  # was minutes, GUI frozen
    assert any(o.field.startswith("Power") for o in run.obs)


# ---------------------------------------------- worst-case padded cells (review, 2026-10-20)

_MOST = 32767                                   # Excel's most characters in a cell


def _worst_cells():
    """Cells that made a pattern backtrack with the GIL held: runs of
    spaces (leading, trailing, inside), tabs and line breaks, and runs of
    ':', '*' and '(' - each as long as Excel lets a cell be."""
    n, h = _MOST, _MOST // 2
    return {
        "leading spaces": " " * (n - 4) + "NOTE",
        "trailing spaces": "NOTE" + " " * (n - 4),
        "spaces x spaces y": " " * h + "x" + " " * (n - h - 2) + "y",
        "spaces inspected by spaces x": " " * h + "inspected by" + " " * (n - h - 13) + "x",
        "Inspector spaces x": "Inspector" + " " * (n - 10) + "x",
        "Inspector: spaces x line y": "Inspector: " + " " * (n - 14) + "x\ny",
        "Run spaces x": "Run" + " " * (n - 4) + "x",
        "Response spaces x": "Response" + " " * (n - 9) + "x",
        "tabs": "\t" * (n - 4) + "NOTE",
        "line breaks": "\n" * (n - 4) + "NOTE",
        "mixed whitespace": (" \t\r\n" * n)[:n - 4] + "NOTE",
        "colons then x": ":" * (n - 1) + "x",
        "colon space pairs": (": " * n)[:n - 1] + "x",
        "stars then x": "*" * (n - 1) + "x",
        "open brackets": "(" * (n - 2) + "x)",
        "open bracket pairs": ("( " * n)[:n],
    }


def test_split_inline_reads_as_the_pattern_did():
    """_split_inline is read without a pattern over the whole cell now: the
    same labels and values as '^\\s*([^:=]{1,60}?)\\s*[:=]\\s*(.+)$'."""
    assert se._split_inline("Operator: AK") == ("Operator", "AK")
    assert se._split_inline("  Inspector  =  JB  \n") == ("Inspector", "JB")
    assert se._split_inline("Operator: ") == ("Operator", "")
    assert se._split_inline("Operator:") == ("Operator:", "")
    assert se._split_inline("Notes: a\nb") == ("Notes: a\nb", "")     # a value on one line
    assert se._split_inline("Path: /x") == ("Path: /x", "")
    assert se._split_inline("x" * 61 + ": y") == ("x" * 61 + ": y", "")
    assert se._split_inline("x" * 60 + ": y") == ("x" * 60, "y")
    assert se._split_inline("10:30") == ("10:30", "")


@pytest.mark.parametrize("name", sorted(_worst_cells()))
def test_every_pattern_reads_a_worst_case_cell_at_once(name, ledger):
    """Each pattern the extract and the knowledge graph try on a cell's
    text, on cells padded to Excel's limit: _split_inline took 2 s on 4,000
    leading spaces and 63 s on 20,000, 'Inspector' and spaces held the GIL
    for seconds, and nothing could be pressed (review, 2026-10-20)."""
    import sample_ledger as sl
    from council_core import kg_excel as kx
    from council_core import knowledge_graph as kgm
    _csv, schema = ledger
    cell = _worst_cells()[name]
    m = se.Matcher(schema)
    part = next(r for r in kgm.DEFAULT_FIELD_RULES if r.label == "Part Number")
    calls = {
        "_split_inline": se._split_inline, "_WHO_INLINE": se._WHO_INLINE.match,
        "_WHO_BY": se._WHO_BY.match, "_WHO_HEAD": se._WHO_HEAD.match,
        "_RUN_HEAD": se._RUN_HEAD.match, "_VALUE_HEAD": se._VALUE_HEAD.match,
        "match": m.match, "is_label": m.is_label, "is_item_header": m.is_item_header,
        "is_placeholder": se.is_placeholder, "_initials": se._initials,
        "normalize_label": sl.normalize_label,
        "kg _label_text": kx._label_text, "kg _who_head": kx._who_head,
        "kg _who_inline": kx._who_inline, "kg split_people": kx.split_people,
        "kg _split": lambda v: kx._split(v, part),
        "kg plausible": lambda v: kx.plausible(part, v),
    }
    slow = {}
    for what, fn in calls.items():
        t0 = time.perf_counter()
        fn(cell)
        took = time.perf_counter() - t0
        if took >= 0.05:
            slow[what] = round(took, 3)
    assert not slow, slow


def _place_everywhere(cell):
    """A Grid's cells with ``cell`` where a label, a value, a header, a
    heading, an item grid's label and a long log's stage go."""
    c = {(1, 1): cell, (1, 2): 12.5,                         # a label beside a number
         (3, 1): "Notes", (3, 2): cell, (4, 1): "Inspector", (4, 2): cell,
         (6, 1): "Flaw", (6, 2): cell, (6, 3): "EC Pre"}      # a header
    for i, it in enumerate(("1A", "1B")):
        c[(7 + i, 1)], c[(7 + i, 2)], c[(7 + i, 3)] = it, 1.0 + i, cell
    c[(10, 2)] = cell                                         # a heading over a table
    c[(11, 2)], c[(11, 3)] = "Flaw", "EC Pre"
    c[(12, 2)], c[(12, 3)] = "1A", 50
    c[(14, 1)], c[(14, 2)], c[(14, 3)], c[(14, 4)] = "Flaw #", "1A", "1B", "2A"
    c[(15, 1)], c[(15, 2)], c[(15, 3)], c[(15, 4)] = cell, 1, 2, 3   # an item grid
    c[(16, 1)], c[(16, 2)], c[(16, 3)], c[(16, 4)] = "EC Pre", cell, 4, 5
    c[(18, 1)], c[(18, 2)], c[(18, 3)] = "Flaw", "Stage", "Response"  # a long log
    c[(19, 1)], c[(19, 2)], c[(19, 3)] = "1A", cell, cell
    c[(20, 1)], c[(20, 2)], c[(20, 3)] = "1A", "EC Pre", 6
    return c


@pytest.mark.parametrize("name", sorted(_worst_cells()))
def test_a_worst_case_cell_is_read_in_milliseconds_wherever_it_is(name, ledger):
    _csv, schema = ledger
    m, codes = se.Matcher(schema), se.ItemCodes(schema.items)
    cell = _worst_cells()[name]
    cells = _place_everywhere(cell)
    places = sum(1 for v in cells.values() if v is cell)
    ex = se._Extractor(se.Grid("w.xlsx", "Notes", cells), m, codes, "", None)
    t0 = time.perf_counter()
    ex.find_part()
    ex.find_inspectors()
    ex.item_grids()
    ex.tables()
    ex.forms()
    each = (time.perf_counter() - t0) / places
    assert each < 0.05, f"{each * 1000:.0f} ms a cell"


def _worst_book(path: Path) -> Path:
    """40 worst-case cells: alone on their rows, labels with a value beside
    them, values beside a label, in a header row and under it."""
    shapes = list(_worst_cells().values())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws["A1"], ws["B1"] = "Laser Power (W)", 280
    r = 3
    for k in range(8):
        ws.cell(r, 1, shapes[k])
        r += 1
    for k in range(14):
        ws.cell(r, 1, shapes[(k + 8) % len(shapes)])
        ws.cell(r, 2, 10 + k)
        r += 1
    for k in range(14):
        ws.cell(r, 1, ("Notes", "Inspector", "Frequency (kHz)", "Operator")[k % 4])
        ws.cell(r, 2, shapes[(k + 6) % len(shapes)])
        r += 1
    r += 1
    ws.cell(r, 1, "Flaw")
    ws.cell(r, 2, shapes[0])
    ws.cell(r, 3, "EC Pre")
    ws.cell(r, 4, shapes[5])
    ws.cell(r + 1, 1, "1A")
    ws.cell(r + 1, 2, shapes[2])
    ws.cell(r + 1, 3, 61)
    ws.cell(r + 1, 4, shapes[11])
    wb.save(path)
    return path


def test_a_sheet_of_40_worst_case_cells_extracts_in_seconds(tmp_path, ledger):
    """The Ledger tab's whole Extract (the review written) and the knowledge
    graph's read of the same workbook: minutes before, the GUI frozen."""
    from council_core import kg_excel as kx
    from council_core import knowledge_graph as kgm
    csv, _schema = ledger
    src = _worst_book(tmp_path / "worst.xlsx")
    t0 = time.perf_counter()
    s = li.run_extract(str(csv), [str(src)], tmp_path / "review.xlsx", use_images=False)
    took = time.perf_counter() - t0
    assert took < 5, f"extract took {took:.1f}s"
    assert s.values and Path(s.review).exists()
    t0 = time.perf_counter()
    kx.read_workbook(src, "worst.xlsx", kgm.DEFAULT_FIELD_RULES)
    took = time.perf_counter() - t0
    assert took < 5, f"knowledge graph read took {took:.1f}s"


def test_cancel_during_a_sheet_of_padded_cells_ends_the_run_at_once(tmp_path, ledger):
    """Cancel is seen every few hundred cells - in the part search too -
    and no cell holds the GIL for long: pressed at +1 s, it was seen at
    +9.9 s, once the sheet was read, and the review was written anyway
    (review, 2026-10-20)."""
    _csv, schema = ledger
    wb = openpyxl.Workbook()
    ws = wb.active
    # named for a template field: a sheet with no word of the template is
    # not read cell by cell at all (synthetic test, 2026-10-27)
    ws.title = "Power notes"
    for r in range(1, 601):
        ws.cell(r, 1, " " * 4000 + f"NOTE {r}")
        ws.cell(r, 2, " " * 2000 + "inspected by" + " " * 2000 + "x")
    wb.save(tmp_path / "padded.xlsx")
    cancel = threading.Event()
    pressed = {}

    class _PressOnce(_Prog):
        def update(self, what="", **k):
            if "reading sheet" in str(what) and not cancel.is_set():
                pressed["at"] = time.perf_counter()
                cancel.set()
    run = se.extract(schema, [tmp_path / "padded.xlsx"], use_images=False, cancel=cancel,
                     progress=_PressOnce(cancel, ""))
    assert run.cancelled and not run.obs
    assert time.perf_counter() - pressed["at"] < 2


# ---------------------------------------------- a who cell typed on two lines (review, 2026-10-20)

def _who_read(cell, schema):
    """Who the Ledger (find_inspectors) and the knowledge graph read from
    ``cell`` above a reading."""
    from council_core import kg_excel as kx
    m, codes = se.Matcher(schema), se.ItemCodes(schema.items)
    g = se.Grid("w.xlsx", "Sheet1", {(1, 1): cell, (2, 1): "Laser Power (W)", (2, 2): 280})
    ex = se._Extractor(g, m, codes, "", None)
    ex.find_inspectors()
    return [w for _r, _c, w in ex.markers], kx._who_inline(cell)


@pytest.mark.parametrize("cell", ["Inspector: JB\nMK", "Inspector: JB\r\nMK",
                                  "Inspector: J\nB", "Inspector: John\nBlake",
                                  "Inspector: JB  \n  MK", "Inspected by JB\nMK",
                                  "Inspector: JB\rMK"])
def test_a_who_cell_on_two_lines_is_nobody(cell, ledger):
    """The who patterns read a cell's words with its spaces collapsed, line
    breaks too: 'Inspector: JB' Alt+Enter 'MK' became one person 'Jb Mk' -
    'JBMK' in the graph - who signed and pre-accepted every reading under
    it. The pattern as typed kept a who on one line, and still does (review,
    2026-10-20); a lone '\\r' is a line break too."""
    _csv, schema = ledger
    assert _who_read(cell, schema) == ([], [])


@pytest.mark.parametrize("cell, who", [("Inspector: JB", "JB"), ("Inspector:\nJB", "JB"),
                                       ("Inspector: JB\n", "JB"), ("Checked by\nJB", "JB"),
                                       ("Inspector:\tJB", "JB"), ("Operator - MK", "MK"),
                                       ("  Inspector:  " + " " * 3000 + "JB  ", "JB")])
def test_a_who_on_one_line_is_still_read(cell, who, ledger):
    _csv, schema = ledger
    assert _who_read(cell, schema) == ([who], [who])


def test_readings_under_a_two_line_who_cell_are_not_signed_by_a_phantom(tmp_path, ledger):
    """End to end: an EC table under 'Inspector: JB' Alt+Enter 'MK'. Every
    reading was pre-accepted with Inspector 'Jb Mk' and the graph had a
    person 'JBMK'; as typed, nobody signed: the readings are to review."""
    from council_core import kg_excel as kx
    from council_core import knowledge_graph as kgm
    csv, schema = ledger
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EC"
    ws["A1"], ws["B1"] = "Part Number", "LMM1AB101"
    ws["A3"] = "Inspector: JB\nMK"
    ws["B5"] = "Flaw #"
    for i, it in enumerate(lf.ITEMS[:3]):
        ws.cell(5, 3 + i, it)
        ws.cell(6, 3 + i, 60 + i)
    ws["B6"] = "Eddy Current Response Pre-Conceal"
    src = tmp_path / "ec.xlsx"
    wb.save(src)
    run = se.extract(schema, [src], use_images=False)
    assert run.obs and not any(o.inspector for o in run.obs)
    s = li.run_extract(str(csv), [str(src)], tmp_path / "review.xlsx", use_images=False)
    assert s.values and s.pre_accepted == 0
    reading = kx.read_workbook(src, "ec.xlsx", kgm.DEFAULT_FIELD_RULES)
    assert not [v.raw for rec in reading.records for v in rec.values
                if v.rule.type == "PERSON"]
