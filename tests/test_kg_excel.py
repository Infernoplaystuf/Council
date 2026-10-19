"""The knowledge graph reading Excel the Ledger's way (council_core/kg_excel.py):
parent CSVs (Ledgers) as firm sources, raw workbooks through the Ledger's
readers, inspector initials through the Ledger's people list and name
questions.

Every vault is a tmp_path; workbooks are synthetic (tests/ledger_fixtures.py
and the small builders below). Excel is never started: the .xls / .xlsb
tests drive a stand-in for Excel.Application, as tests/test_excel_open.py
does. Pictures are read by a stand-in reader, never by Tesseract.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import sys
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

sys.path.insert(0, str(Path(__file__).parent))
import excel_open  # noqa: E402
import ledger_fixtures as lf  # noqa: E402
import source_extract as se  # noqa: E402
from sample_ledger import INITIALS_SUFFIX, ParentCSV, read_template  # noqa: E402

from council_core import kg_excel  # noqa: E402
from council_core import knowledge_graph as kgm  # noqa: E402


# ── helpers ──────────────────────────────────────────────────────────────
def _vault(tmp_path, name="v"):
    v = tmp_path / name
    (v / "data_in").mkdir(parents=True)
    return v


def _kg(v, reader=None, confirm=True):
    kg = kgm.KnowledgeGraph(v, image_reader=reader)
    if confirm:
        kg.confirm_all_rules()
    return kg


def _rels(kg):
    return {(r["subject"]["name"], r["predicate"], r["object"]["name"]): r["status"]
            for r in kg.all_relations() if r["predicate"] != "DOCUMENTED_IN"}


def _names(kg, etype):
    return sorted(r[0] for r in kg.db.execute(
        "SELECT name FROM entities WHERE type=? AND merged_into IS NULL", (etype,)))


def _hashes(root: Path):
    return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(root.rglob("*")) if p.is_file()}


def _schema(tmp_path):
    t = tmp_path / "templates" / "lmm.xlsx"
    t.parent.mkdir(parents=True, exist_ok=True)
    if not t.exists():
        lf.make_template(t)
    return read_template(t)


def _ec(schema):
    """The template's per-flaw EC reading that takes inspectors."""
    return next(f.name for f in schema.fields if schema.takes_inspectors(f.name))


def _ledger(tmp_path, folder: Path, rows, *, people=None, name="lmm_parent.csv",
            source="nde.xlsx"):
    schema = _schema(tmp_path)
    folder.mkdir(parents=True, exist_ok=True)
    led = ParentCSV.create(folder / name, schema)
    led.add_samples(rows, source=source)
    if people is not None:
        se.save_aliases(folder / name, {}, [])
        p = se.aliases_path_for(folder / name)
        data = json.loads(p.read_text(encoding="utf-8"))
        data["people"] = people
        p.write_text(json.dumps(data), encoding="utf-8")
    return folder / name, schema


def _row(schema, part, item, **cols):
    r = {schema.id_field: part, schema.item_field: item}
    r.update(cols)
    return r


def _book(path: Path, sheets):
    """``sheets``: {name: {"A1": value, ...}}; ``merge``: list of ranges."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for ref, v in cells.items():
            if ref == "merge":
                for rng in v:
                    ws.merge_cells(rng)
            else:
                ws[ref] = v
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ── Ledgers (parent CSVs) ────────────────────────────────────────────────
def test_a_ledger_row_links_its_inspector_to_its_part_firmly(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in" / "ledger", [
        _row(schema, "101", "1A", **{f"{ec} [JB]": "63", ec: "63",
                                     ec + INITIALS_SUFFIX: "JB"}),
        _row(schema, "101", "2A", **{f"{ec} [MK]": "70"}),
    ], people={"JB": "Jordan Blake", "mk": "Mary King"})
    kg = _kg(v)
    stats = kg.seed()
    assert stats["unreadable"] == [] and stats["ledgers"] == 1
    assert _rels(kg) == {("Jordan Blake", "INSPECTED", "101"): "seeded",
                         ("Mary King", "INSPECTED", "101"): "seeded"}
    jb = kg.search("Jordan Blake", "PERSON")[0]
    (n,) = kg.neighbors(jb["id"])
    (ev,) = n["evidence"]
    assert ev["path"] == "ledger/lmm_parent.csv" and ev["method"] == "ledger"
    loc = ev["locator"]
    assert loc["row"] == 2 and loc["item"] == "Flaw # 1A"
    # the person's two columns (their own reading, the initials) and the ID
    assert set(loc["columns"]) == {f"{ec} [JB]", ec + INITIALS_SUFFIX, schema.id_field}
    assert loc["sources"] == "nde.xlsx"
    assert ev["where"].startswith("row 2, Flaw # 1A (")
    assert "source workbook(s): nde.xlsx" in ev["where"]
    assert f"{ec} [JB] = 63" in ev["quote"]
    # 'JB' is how Jordan signs: an alias, so a search for it finds him
    assert [h["name"] for h in kg.search("JB", "PERSON")] == ["Jordan Blake"]
    assert kg.reviews() == []


def test_ledger_sidecars_and_review_workbooks_are_not_documents(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    csv, _ = _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A")], people={})
    src = _book(v / "data_in" / "src.xlsx", {"Part 101": {"A1": "Inspector: JB"}})
    from council_core.ledger_ingest import run_extract
    run_extract(csv, [src], v / "data_in" / "review_x.xlsx", use_images=False)
    kg = _kg(v)
    docs = [p.name for p in kg.document_files()]
    assert docs == ["lmm_parent.csv", "src.xlsx"]
    assert kg_excel.is_review_workbook(v / "data_in" / "review_x.xlsx")
    assert kg_excel.is_ledger_sidecar(v / "data_in" / "lmm_parent.schema.json")


def test_values_still_waiting_in_a_review_are_never_read(tmp_path):
    # The review holds a reading by 'ZZ' nobody accepted; only the CSV is
    # the user's word, so ZZ is in no link and no question.
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    csv, _ = _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A")])
    src = tmp_path / "outside" / "src.xlsx"
    cells = {"A1": "Inspector: ZZ", "B3": "Flaw", "B4": "EC Pre (%FSH)"}
    for i, item in enumerate(lf.ITEMS):
        cells[f"{chr(67 + i)}3"] = item
        cells[f"{chr(67 + i)}4"] = 60 + i
    _book(src, {"Part 101": cells})
    from council_core.ledger_ingest import run_extract
    s = run_extract(csv, [src], v / "data_in" / "review_x.xlsx", use_images=False)
    assert s.values > 0
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {}
    assert [r["surface"] for r in kg.reviews()] == []
    assert "ZZ" not in json.dumps(kg.export_json())


def _src_with_inspector(path, who="JB", sheet="Part 101"):
    cells = {"A1": f"Inspector: {who}", "B3": "Flaw", "B4": "EC Pre (%FSH)"}
    for i, item in enumerate(lf.ITEMS):
        col = openpyxl.utils.get_column_letter(3 + i)
        cells[f"{col}3"] = item
        cells[f"{col}4"] = 60 + i
    return _book(path, {sheet: cells})


def _accept_all(review):
    wb = openpyxl.load_workbook(review)
    ws = wb["Values"]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1, "Y")
    wb.save(review)


def test_ledger_evidence_points_at_the_workbook_cell_it_came_from(tmp_path):
    """The whole Ledger loop: extract -> review (a person accepts) -> apply;
    the graph then cites the CSV row AND the source cell the review kept."""
    from council_core.ledger_ingest import run_apply, run_extract
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    folder = v / "data_in" / "ledger"
    csv, _ = _ledger(tmp_path, folder, [], people={"JB": "Jordan Blake"})
    src = _src_with_inspector(folder / "nde_101.xlsx")
    review = folder / "review_1.xlsx"
    run_extract(csv, [src], review, use_images=False)
    _accept_all(review)
    run_apply(review, use_images=False)
    assert "JB" in ParentCSV(csv).rows[0].get(_ec(schema) + INITIALS_SUFFIX, "")
    kg = _kg(v)
    kg.seed()
    rels = _rels(kg)
    assert rels[("Jordan Blake", "INSPECTED", "101")] == "seeded"
    jb = kg.search("Jordan Blake", "PERSON")[0]
    (n,) = kg.neighbors(jb["id"])
    led = [e for e in n["evidence"] if e["method"] == "ledger"]
    assert len(led) == len(lf.ITEMS)                 # one per flaw row
    first = min(led, key=lambda e: e["locator"]["row"])
    assert first["locator"]["item"] == "Flaw # 1A"
    assert first["locator"]["from"] == "nde_101.xlsx › Part 101!C4"
    assert first["where"].endswith("— from nde_101.xlsx › Part 101!C4")
    # the source workbook itself says the same, as the Ledger's reading: a
    # suggestion, folded into the firm link
    raw = [e for e in n["evidence"] if e["path"].endswith("nde_101.xlsx")]
    assert raw and {e["method"] for e in raw} == {"sheet"}
    # the in-app preview of the CSV row names the cell too
    ctx = kgm.source_context(kg.root, first["path"], first["locator"])
    assert ("from", "nde_101.xlsx › Part 101!C4", False) in ctx
    summary = next(s for s in kg.sources() if s["kind"] == "ledger")
    assert summary["summary"]["located"] >= len(lf.ITEMS)
    assert "traced to their workbook cell" in summary["text"]


def test_ledger_columns_are_read_by_the_schema_not_guessed(tmp_path):
    """A Notes field holding 'JB' is not a person; a part-level 'Operator'
    is; a 'Project' field is a project once its rule is confirmed, giving
    inspector - part - project."""
    t = tmp_path / "templates" / "proj.xlsx"
    t.parent.mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(["Part Number", "Project", "Operator", "Notes"], start=1):
        ws.cell(1, c, h)
    wb.save(t)
    schema = read_template(t)
    v = _vault(tmp_path)
    led = ParentCSV.create(v / "data_in" / "parts_parent.csv", schema)
    led.add_samples([{"Part Number": "2041", "Project": "PRJ-7", "Operator": "Dana Whitfield",
                      "Notes": "JB"}], source="x.xlsx")
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "INSPECTED", "2041"): "seeded",
                         ("PRJ-7", "USES_PART", "2041"): "seeded"}
    assert _names(kg, "PERSON") == ["Dana Whitfield"]
    assert kg.reviews() == []
    # Not confirmed: 'Project' is just a column.
    v2 = _vault(tmp_path, "v2")
    led2 = ParentCSV.create(v2 / "data_in" / "parts_parent.csv", schema)
    led2.add_samples([{"Part Number": "2041", "Project": "PRJ-7"}])
    kg2 = _kg(v2, confirm=False)
    kg2.seed()
    assert _names(kg2, "PROJECT") == [] and _names(kg2, "PART") == ["2041"]


def test_a_damaged_ledger_schema_is_reported_and_not_read_as_a_table(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    csv, _ = _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A")])
    kgm_schema = kg_excel.schema_path(csv)
    kgm_schema.write_text("{ not json", encoding="utf-8")
    kg = _kg(v)
    stats = kg.seed()
    assert stats["unreadable"] == ["lmm_parent.csv: its Ledger schema lmm_parent.schema.json"
                                   " could not be read (JSONDecodeError), so the Ledger was"
                                   " not read"]
    assert _names(kg, "PART") == []
    assert any(s["text"].startswith("Not read: its Ledger schema") for s in kg.sources())


def test_a_ledger_of_many_rows_is_streamed_not_held(tmp_path):
    # 3,000 parts x 1 flaw: read row by row; every row cited once.
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    rows = [_row(schema, str(1000 + i), "1A", **{f"{ec} [JB]": str(i)}) for i in range(3000)]
    _ledger(tmp_path, v / "data_in", rows, people={"JB": "Jordan Blake"})
    kg = _kg(v)
    stats = kg.seed()
    assert stats["records"] == 3000
    jb = kg.search("Jordan Blake", "PERSON")[0]
    ns = kg.neighbors(jb["id"], evidence_limit=5)
    assert len(ns) == 3000 and all(n["evidence_total"] == 1 for n in ns)


# ── initials: the people list, the name questions, never a guess ────────
def test_unknown_initials_are_asked_once_and_never_guessed(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in", [_row(schema, "101", it, **{f"{ec} [MK]": "1"})
                                      for it in ("1A", "2A", "3A")])
    _ledger(tmp_path, v / "data_in" / "b", [_row(schema, "102", "1A",
                                                 **{ec + INITIALS_SUFFIX: "MK"})])
    # Mary King is known from a labelled table: she has those initials, so
    # she is OFFERED - but the rows are not linked to her on that alone.
    (v / "data_in" / "team.csv").write_text("Project,Owner\nPRJ-1,Mary King\n",
                                            encoding="utf-8")
    kg = _kg(v)
    kg.seed()
    assert ("Mary King", "INSPECTED", "101") not in _rels(kg)
    assert not [k for k in _rels(kg) if k[1] == "INSPECTED"]
    qs = [r for r in kg.reviews() if r["kind"] == "initials"]
    assert len(qs) == 1
    (q,) = qs
    assert q["surface"] == "MK"
    assert [c["name"] for c in q["candidates"]] == ["Mary King"]
    assert "signed 4 value(s) in 2 documents" in q["snippet"]
    assert "MK" not in _names(kg, "PERSON")


def test_an_answer_everywhere_links_every_row_and_lasts(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A", **{f"{ec} [MK]": "1"})])
    _ledger(tmp_path, v / "data_in" / "b", [_row(schema, "102", "1A",
                                                 **{f"{ec} [MK]": "2"})])
    (v / "data_in" / "team.csv").write_text("Project,Owner\nPRJ-1,Mary King\n",
                                            encoding="utf-8")
    kg = _kg(v)
    kg.seed()
    (q,) = kg.reviews()
    kg.answer_review(q["id"], q["candidates"][0]["id"], everywhere=True)
    kg.seed()
    rels = _rels(kg)
    assert rels[("Mary King", "INSPECTED", "101")] == "seeded"
    assert rels[("Mary King", "INSPECTED", "102")] == "seeded"
    assert kg.reviews() == []
    assert [a["surface"] for a in kg.answers()] == ["MK"]
    kg.seed()                                   # and again: still answered
    assert kg.reviews() == [] and ("Mary King", "INSPECTED", "102") in _rels(kg)


def test_initials_answered_with_a_typed_name_and_with_neither(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A", **{f"{ec} [JB]": "1",
                                                                   f"{ec} [QX]": "2"})])
    kg = _kg(v)
    kg.seed()
    qs = {r["surface"]: r for r in kg.reviews()}
    assert set(qs) == {"JB", "QX"} and qs["JB"]["candidates"] == []
    eid = kg.add_person("Jordan Blake")           # someone no document names
    kg.answer_review(qs["JB"]["id"], eid, everywhere=True)
    kg.answer_review(qs["QX"]["id"], None, everywhere=True)   # a person of their own
    kg.seed()
    rels = _rels(kg)
    assert rels[("Jordan Blake", "INSPECTED", "101")] == "seeded"
    assert rels[("QX", "INSPECTED", "101")] == "seeded"
    assert kg.reviews() == []
    kg.seed()
    assert ("Jordan Blake", "INSPECTED", "101") in _rels(kg)


def test_two_ledgers_that_disagree_about_initials(tmp_path):
    """Each Ledger's rows use its own people list; a workbook outside both
    gets no pick between them - it is asked."""
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in" / "a", [_row(schema, "101", "1A", **{f"{ec} [JB]": "1"})],
            people={"JB": "Jordan Blake"})
    _ledger(tmp_path, v / "data_in" / "b", [_row(schema, "202", "1A", **{f"{ec} [JB]": "2"})],
            people={"JB": "Jill Brown"})
    _book(v / "data_in" / "c.xlsx", {"Part 303": {"A1": "Inspector: JB"}})
    kg = _kg(v)
    kg.seed()
    rels = _rels(kg)
    assert rels[("Jordan Blake", "INSPECTED", "101")] == "seeded"
    assert rels[("Jill Brown", "INSPECTED", "202")] == "seeded"
    assert not [k for k in rels if k[2] == "303"]
    (q,) = kg.reviews()
    assert q["surface"] == "JB" and q["path"] == "c.xlsx"
    assert sorted(c["name"] for c in q["candidates"]) == ["Jill Brown", "Jordan Blake"]


def test_the_people_list_survives_the_ledgers_own_saves(tmp_path):
    csv = tmp_path / "x_parent.csv"
    csv.write_text("a\n", encoding="utf-8")
    se.aliases_path_for(csv).write_text(json.dumps(
        {"map": {"pn": "Part Number"}, "ignore": [], "people": {"j.b.": "Jordan  Blake"}}),
        encoding="utf-8")
    se.save_aliases(csv, {"pn": "Part Number", "freq": "Frequency"}, ["notes"])
    data = json.loads(se.aliases_path_for(csv).read_text(encoding="utf-8"))
    assert data["people"] == {"j.b.": "Jordan  Blake"}
    assert data["map"] == {"freq": "Frequency", "pn": "Part Number"}
    assert se.load_people(csv) == {"JB": "Jordan Blake"}
    assert se.load_aliases(csv) == ({"freq": "Frequency", "pn": "Part Number"}, ["notes"])


# ── raw workbooks through the Ledger's reader ───────────────────────────
def test_a_header_below_a_title_and_two_tables_on_one_sheet(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "tracker.xlsx", {"Tracker": {
        "A1": "Program tracker — Q3", "merge": ["A1:D1"],
        "A4": "Part Number", "B4": "Owner", "C4": "Status",
        "A5": "PN-1001", "B5": "Dana Whitfield", "C5": "open",
        "A6": "PN-1002", "B6": "Carol Lee", "C6": "closed",
        # a second table to the right, after an empty column
        "F4": "Project", "G4": "Program Lead",
        "F5": "PRJ-0915", "G5": "Hana Sato",
        # and a third under the first, after a blank row
        "A9": "Project ID", "B9": "Point of Contact",
        "A10": "PRJ-0420", "B10": "Tomas Reyes"}})
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "OWNS", "PN-1001"): "seeded",
                         ("Carol Lee", "OWNS", "PN-1002"): "seeded",
                         ("Hana Sato", "LEADS", "PRJ-0915"): "seeded",
                         ("Tomas Reyes", "CONTACT_FOR", "PRJ-0420"): "seeded"}
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    (n,) = kg.neighbors(dana["id"])
    (ev,) = n["evidence"]
    assert ev["method"] == "row"
    assert ev["locator"] == {"sheet": "Tracker", "row": 5, "header": 4,
                             "columns": ["Owner", "Part Number"], "cells": ["B5", "A5"]}
    assert ev["where"] == "sheet Tracker, row 5 (Owner, Part Number)"
    ctx = kgm.source_context(kg.root, "tracker.xlsx", ev["locator"])
    assert ("B5", "Owner: Dana Whitfield", True) in ctx and ("A5", "Part Number: PN-1001", True) in ctx
    assert ("C5", "Status: open", False) in ctx


def test_labelled_cells_beside_inline_and_below(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "form.xlsx", {"Build": {
        "A1": "Build record",
        "A3": "P/N: PN-2041",                       # inline
        "A4": "Program Lead", "B4": "Dana Whitfield",   # beside
        "A6": "Project:", "A7": "PRJ-0915",          # below
        "A9": "Freq (kHz)", "B9": 50}})
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "LEADS", "PRJ-0915"): "suggested",
                         ("PRJ-0915", "USES_PART", "PN-2041"): "suggested"}
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    ev = kg.neighbors(dana["id"])[0]["evidence"][0]
    assert ev["where"] == "sheet Build, cell B4 (Program Lead)" and ev["method"] == "sheet"
    # beside (B4) and inline values are firm; the one read from below the
    # label (A7) is not, so its links wait for the user
    kg2 = _kg(_vault(tmp_path, "v2"))
    _book(kg2.root / "form.xlsx", {"Build": {"A3": "P/N: PN-2041", "A4": "Program Lead",
                                             "B4": "Dana Whitfield", "A5": "Project",
                                             "B5": "PRJ-0915"}})
    kg2.seed()
    assert set(_rels(kg2).values()) == {"seeded"}


def test_a_row_of_labelled_values_is_not_read_as_a_table(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "f.xlsx", {"S": {
        "A1": "Project", "B1": "PRJ-1", "C1": "Program Lead", "D1": "Dana Whitfield",
        "A2": "Owner", "B2": "Bob Smith", "C2": "Part", "D2": "PN-7"}})
    kg = _kg(v)
    kg.seed()
    assert _names(kg, "PROJECT") == ["PRJ-1"]
    assert _names(kg, "PART") == ["PN-7"]
    assert _names(kg, "PERSON") == ["Bob Smith", "Dana Whitfield"]
    assert _rels(kg) == {("Dana Whitfield", "LEADS", "PRJ-1"): "seeded",
                         ("Bob Smith", "OWNS", "PN-7"): "seeded",
                         ("PRJ-1", "USES_PART", "PN-7"): "seeded"}


def test_merged_cells_are_one_label_and_one_value(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "m.xlsx", {"S": {
        "A1": "Project", "B1": "Part Number", "C1": "Owner", "merge": ["B1:B1", "A2:A3"],
        "A2": "PRJ-5", "B2": "PN-1", "C2": "Dana Whitfield",
        "B3": "PN-2", "C3": "Carol Lee"}})
    kg = _kg(v)
    kg.seed()
    rels = _rels(kg)
    # the project merged down two rows covers both
    assert rels[("PRJ-5", "USES_PART", "PN-1")] == "seeded"
    assert rels[("PRJ-5", "USES_PART", "PN-2")] == "seeded"
    assert rels[("Carol Lee", "OWNS", "PN-2")] == "seeded"


def test_ledger_style_workbooks_give_suggestions(tmp_path):
    """The Ledger's own fixtures: 'Operator: JB' under a flaw grid on sheet
    'Part 101', and 'P/N: 101' on a laser log. Only the Ledger's reading says
    JB worked on 101 - a suggestion - and only once JB is known."""
    v = _vault(tmp_path)
    lf.make_nde_book(v / "data_in" / "nde.xlsx")
    lf.make_laser_log(v / "data_in" / "laser.xlsx")
    (v / "data_in" / "ledger").mkdir()
    schema = _schema(tmp_path)
    _ledger(tmp_path, v / "data_in" / "ledger", [_row(schema, "999", "1A")],
            people={"JB": "Jordan Blake"})
    kg = _kg(v, reader=lf.fake_reader(lf.CONCEALMENT_IMAGE_ROWS))
    stats = kg.seed()
    assert stats["unreadable"] == []
    assert _rels(kg) == {("Jordan Blake", "INSPECTED", "101"): "suggested"}
    jb = kg.search("Jordan Blake", "PERSON")[0]
    ev = {e["where"] for e in kg.neighbors(jb["id"])[0]["evidence"]}
    assert ev == {"sheet Part 101, cell B8 (who) | sheet 'Part 101'".split(" | ")[0]} or \
        any("sheet Part 101" in w for w in ev)
    assert "101" in _names(kg, "PART")            # P/N: 101 on the laser log
    assert stats["pictures"] == 1 and stats["pictures_read"] == 1


def test_text_boxes_and_charts_are_read(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    book = _book(v / "data_in" / "nde.xlsx", {"Part 101": {"A1": "NDE"}})
    real = excel_open.zip_extras

    def extras(path):
        # the workbook is parsed from its bytes (a BytesIO) or, in the
        # preview, its path
        if isinstance(path, (str, Path)) and Path(path) != book:
            return real(path)
        return [excel_open.Extra("Part 101", "Part 101 › text box 1", excel_open.text_cells(
                    "Program Lead: Dana Whitfield\nProject: PRJ-0915\n"
                    "Dana Whitfield checked the shim stack with Carol Lee on site.")),
                excel_open.Extra("Part 101", "Part 101 › chart 2", excel_open.chart_cells(
                    "EC Pre-Conceal", "", ["1A", "2A"], [("JB", [70, 71]), ("MK", [72, None])]))]
    monkeypatch.setattr(excel_open, "zip_extras", extras)
    schema = _schema(tmp_path)
    _ledger(tmp_path, v / "data_in" / "led", [_row(schema, "999", "1A")],
            people={"JB": "Jordan Blake", "MK": "Mary King"})
    kg = _kg(v)
    stats = kg.seed()
    rels = _rels(kg)
    assert rels[("Dana Whitfield", "LEADS", "PRJ-0915")] == "seeded"   # the text box
    # the chart's series are initials: they inspected the sheet's part
    assert rels[("Jordan Blake", "INSPECTED", "101")] == "suggested"
    assert rels[("Mary King", "INSPECTED", "101")] == "suggested"
    srcs = {s["path"]: s for s in kg.sources()}
    assert srcs["nde.xlsx"]["summary"]["text_boxes"] == 1
    assert srcs["nde.xlsx"]["summary"]["charts"] == 1
    # the text box's sentence is free text, cited by its box and line
    carol = [m for m in kg.mentions(kg.search("Carol Lee")[0]["id"])] if kg.search("Carol Lee") \
        else []
    assert not carol                             # one-word-safe: Carol is not known yet
    lines = kg.db.execute("SELECT lines FROM doc_cache c JOIN documents d ON d.id=c.document_id"
                          " WHERE d.path='nde.xlsx'").fetchone()[0]
    assert ["Part 101 › text box 1", 3,
            "Dana Whitfield checked the shim stack with Carol Lee on site."] in json.loads(lines)
    assert stats["workbooks"] == 1


def test_pictures_are_read_by_ocr_and_their_links_are_suggested(tmp_path):
    from openpyxl.drawing.image import Image as XLImage
    v = _vault(tmp_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visual"
    ws["A1"] = "Visual inspection"
    ws.add_image(XLImage(io.BytesIO(lf._png([["x"]]))), "F2")
    wb.save(v / "data_in" / "vis.xlsx")
    rows = [["Part Number", "Owner"], ["PN-77", "Dana Whitfield"]]
    kg = _kg(v, reader=lf.fake_reader(rows))
    stats = kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "OWNS", "PN-77"): "suggested"}
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    ev = kg.neighbors(dana["id"])[0]["evidence"][0]
    assert ev["method"] == "picture" and ev["locator"]["image"] == "F2"
    assert ev["where"] == ("sheet Visual, picture at F2, row 2 (Owner, Part Number) — read "
                           "from a picture, how sure unknown — check it")
    assert stats["pictures_read"] == 1
    assert kgm.source_context(kg.root, "vis.xlsx", ev["locator"], quote=ev["quote"])[0][1] \
        == "a picture at F2 on sheet Visual, read by OCR"
    # no reader: a plain note, not a crash, and the picture is not saved
    kg2 = _kg(_vault(tmp_path, "v2"))
    (kg2.root / "vis.xlsx").write_bytes((v / "data_in" / "vis.xlsx").read_bytes())
    st2 = kg2.seed()
    assert st2["notes"] == ["vis.xlsx: 1 picture(s) not read: install Tesseract (local OCR)"
                            " to read them"]
    assert _rels(kg2) == {}
    assert sorted(p.name for p in kg2.root.rglob("*")) == ["vis.xlsx"]


def test_a_picture_the_reader_fails_on_is_a_note(tmp_path):
    from openpyxl.drawing.image import Image as XLImage
    v = _vault(tmp_path)
    wb = openpyxl.Workbook()
    wb.active.add_image(XLImage(io.BytesIO(lf._png([["x"]]))), "B2")
    wb.save(v / "data_in" / "p.xlsx")

    def broken(_data):
        raise RuntimeError("tesseract exited 1")
    broken.name = "Tesseract (local)"
    stats = _kg(v, reader=broken).seed()
    assert stats["unreadable"] == []
    assert stats["notes"] == ["p.xlsx: the picture at Sheet!B2 could not be read (RuntimeError)"]


# ── Excel-only workbooks: a stand-in for Excel.Application ──────────────
class _Range:
    def __init__(self, row=1, col=1, value=None):
        self.Row, self.Column, self.Value = row, col, value
        self.MergeCells = False


class _Coll:
    def __init__(self, items):
        self._items, self.Count = items, len(items)

    def Item(self, i):
        return self._items[i - 1]


class _Sheet:
    def __init__(self, name, rows):
        self.Name = name
        self.UsedRange = _Range(1, 1, tuple(tuple(r) for r in rows))
        self.Shapes = _Coll([])

    def Cells(self, r, c):
        return _Range(r, c)


class _App:
    """Excel.Application as excel_open drives it; logs what it was asked."""

    def __init__(self, log, sheets, fail=None):
        self.log = log
        app = self

        class Books:
            def Open(self, path, **kw):
                app.log.append(("open", Path(path).name, kw))
                if fail:
                    raise Exception(fail)
                return app._book
        self.Workbooks = Books()

        class Book:
            Worksheets = sheets

            def Close(self, SaveChanges=True):
                app.log.append(("close", SaveChanges))

            def Save(self):
                app.log.append(("SAVED",))
        self._book = Book()

    def __setattr__(self, k, v):
        if k[0].isupper() and k != "Workbooks":
            self.log.append(("set", k, v))
        object.__setattr__(self, k, v)

    def Quit(self):
        self.log.append(("quit",))


def _fake_excel(monkeypatch, sheets, fail=None):
    log = []
    real = excel_open.read_with_excel
    monkeypatch.setattr(excel_open, "excel_available", lambda: True)
    monkeypatch.setattr(excel_open, "read_with_excel",
                        lambda path, want_pictures=True, **kw: real(
                            path, want_pictures=want_pictures,
                            app_factory=lambda: _App(log, sheets, fail)))
    return log


def test_an_xls_is_read_through_excel_read_only_and_closed_unsaved(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    xls = v / "data_in" / "old.xls"
    xls.write_bytes(excel_open.OLE_MAGIC + b"\0" * 504 + "Workbook".encode("utf-16-le"))
    before = _hashes(v / "data_in")
    log = _fake_excel(monkeypatch, [_Sheet("Parts", [
        ("Parts list", None, None),
        (None, None, None),
        ("Part Number", "Owner", "Qty"),
        ("PN-9", "Dana Whitfield", 4.0)])])
    kg = _kg(v)
    stats = kg.seed()
    assert stats["unreadable"] == []
    assert _rels(kg) == {("Dana Whitfield", "OWNS", "PN-9"): "seeded"}
    opened = next(e for e in log if e[0] == "open")[2]
    assert opened["ReadOnly"] is True and opened["UpdateLinks"] == 0
    sets = {e[1]: e[2] for e in log if e[0] == "set"}
    assert sets["Visible"] is False and sets["AutomationSecurity"] == 3
    assert ("close", False) in log and log[-1] == ("quit",)
    assert ("SAVED",) not in log
    assert _hashes(v / "data_in") == before
    (src,) = kg.sources()
    assert src["summary"]["opened_with"] == "Excel (hidden, read-only)"
    # unchanged: not opened again
    log.clear()
    kg.seed()
    assert log == []


def test_an_xls_with_no_excel_and_no_xlrd_is_a_plain_note(tmp_path, monkeypatch):
    import importlib.util
    v = _vault(tmp_path)
    (v / "data_in" / "old.xls").write_bytes(excel_open.OLE_MAGIC + b"\0" * 504)
    (v / "data_in" / "bin.xlsb").write_bytes(b"PK\x03\x04")
    monkeypatch.setattr(excel_open, "excel_available", lambda: False)
    real = importlib.util.find_spec
    monkeypatch.setattr(importlib.util, "find_spec",
                        lambda name, *a: None if name == "xlrd" else real(name, *a))
    stats = _kg(v).seed()
    assert sorted(stats["unreadable"]) == [
        "bin.xlsb: a .xlsb workbook opens only through Microsoft Excel (with the pywin32"
        " package), which is not available here",
        "old.xls: a .xls workbook opens only through Microsoft Excel (with the pywin32"
        " package), which is not available here"]


def test_a_password_workbook_is_not_read_nor_retried(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    pw = v / "data_in" / "pw.xlsx"
    pw.write_bytes(excel_open.OLE_MAGIC + b"\0" * 504 + b"".join(
        n.encode("utf-16-le") for n in ("EncryptedPackage", "StrongEncryptionDataSpace")))
    asked = []
    monkeypatch.setattr(excel_open, "read_with_excel",
                        lambda *a, **k: asked.append(a) or pytest.fail("Excel started"))
    kg = _kg(v)
    stats = kg.seed()
    assert stats["unreadable"] == ["pw.xlsx: has a password to open, so it was not read (open"
                                   " it in Excel and save a copy without the password)"]
    calls = []
    monkeypatch.setattr(kg_excel, "read_workbook", lambda *a, **k: calls.append(a))
    stats = kg.seed()
    assert calls == [] and len(stats["unreadable"]) == 1     # remembered, not re-read
    assert asked == []


def test_an_excel_failure_is_tried_again_next_time(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    (v / "data_in" / "b.xlsb").write_bytes(b"PK\x03\x04")
    log = _fake_excel(monkeypatch, [], fail="Excel is busy")
    kg = _kg(v)
    stats = kg.seed()
    assert stats["unreadable"] == ["b.xlsb: Excel could not open b.xlsb: Excel is busy"]
    assert log[-1] == ("quit",)                     # Excel closed after a failure too
    log.clear()
    kg.seed()
    assert any(e[0] == "open" for e in log)


# ── free text, incremental rebuilds, safety ─────────────────────────────
def test_workbook_notes_feed_the_gazetteer_and_the_free_text_pass(tmp_path):
    v = _vault(tmp_path)
    (v / "data_in" / "team.csv").write_text(
        "Project,Program Lead,Part\nPRJ-0915,Dana Whitfield,PN-1001\n", encoding="utf-8")
    _book(v / "data_in" / "notes.xlsx", {"Log": {
        "A1": "Site log",
        "A3": "Nothing to report today at the cell.",
        "A9": "Dana Whitfield fitted PN-1001 to the rig on Tuesday."}})
    kg = _kg(v)
    kg.seed()
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    gaz = [m for m in kg.mentions(dana["id"]) if m["method"] == "gazetteer"]
    assert [(m["path"], m["where"]) for m in gaz] == [("notes.xlsx", "sheet Log, row 9")]
    chunks = [c for c in kg.text_chunks() if c["path"] == "notes.xlsx"]
    assert len(chunks) == 1 and chunks[0]["line_nos"] == [3, 9]
    # a model's link from the workbook is cited at the sheet row it is on
    # (rows 3 and 9 are lines 1 and 2 of the passage)
    quote = "Dana Whitfield fitted PN-1001 to the rig on Tuesday."

    def chat(messages, **kw):
        if quote not in messages[-1]["content"]:
            return json.dumps({"links": []})
        return json.dumps({"links": [{"subject": "Dana Whitfield", "predicate": "OWNS",
                                      "object": "PN-1001", "quote": quote}]})
    stats = kg.suggest_from_text({"this PC": chat}, model="stand-in")
    assert stats["errors"] == 0 and stats["links"] == 1
    owns = [n for n in kg.neighbors(dana["id"]) if n["predicate"] == "OWNS"]
    (ev,) = [e for n in owns for e in n["evidence"] if e["method"] == "model"]
    assert ev["path"] == "notes.xlsx" and ev["where"] == "sheet Log, row 9"
    assert owns[0]["status"] == "suggested"


def test_an_unchanged_workbook_is_skipped_and_a_changed_one_replaces_its_own(tmp_path,
                                                                            monkeypatch):
    v = _vault(tmp_path)
    for i in range(5):
        _book(v / "data_in" / f"t{i}.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                                   "A2": f"PN-{i}", "B2": "Dana Whitfield"}})
    kg = _kg(v)
    kg.seed()
    first = _rels(kg)
    reads = []
    real = se.load_workbook_grids
    monkeypatch.setattr(se, "load_workbook_grids",
                        lambda p, *a, **k: reads.append(Path(p).name) or real(p, *a, **k))
    stats = kg.seed()
    assert reads == [] and stats["documents_unchanged"] == 5 and _rels(kg) == first
    _book(v / "data_in" / "t2.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                            "A2": "PN-2", "B2": "Carol Lee"}})
    stats = kg.seed()
    assert reads == ["t2.xlsx"] and stats["documents_read"] == 1
    now = _rels(kg)
    assert ("Dana Whitfield", "OWNS", "PN-2") not in now
    assert now[("Carol Lee", "OWNS", "PN-2")] == "seeded"
    assert {k: s for k, s in now.items() if k[2] != "PN-2"} == \
        {k: s for k, s in first.items() if k[2] != "PN-2"}


def test_seeding_never_changes_a_source_file(tmp_path):
    from council_core.ledger_ingest import run_apply, run_extract
    v = _vault(tmp_path)
    folder = v / "data_in" / "ledger"
    csv, _ = _ledger(tmp_path, folder, [], people={"JB": "Jordan Blake"})
    src = _src_with_inspector(folder / "nde_101.xlsx")
    review = folder / "review_1.xlsx"
    run_extract(csv, [src], review, use_images=False)
    _accept_all(review)
    run_apply(review, use_images=False)
    lf.make_nde_book(v / "data_in" / "nde.xlsx")
    before = _hashes(v / "data_in")
    kg = _kg(v, reader=lf.fake_reader(lf.CONCEALMENT_IMAGE_ROWS))
    kg.seed()
    kg.seed()
    kgm.source_context(kg.root, "nde.xlsx", {"sheet": "Part 101", "cell": "B8"})
    kg.write_exports()
    assert _hashes(v / "data_in") == before


def test_a_big_workbook_is_streamed_with_the_same_answers(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    cells = {"A1": "Parts", "A3": "Part Number", "B3": "Owner"}
    for i in range(200):
        cells[f"A{4 + i}"] = f"PN-{i}"
        cells[f"B{4 + i}"] = "Dana Whitfield" if i % 2 else "Carol Lee"
    _book(v / "data_in" / "big.xlsx", {"Data": cells})
    loaded = _kg(v)
    loaded.seed()
    want = _rels(loaded)
    loaded.close()
    monkeypatch.setattr(kg_excel, "STREAM_XML_BYTES", 10)
    called = []
    monkeypatch.setattr(se, "load_workbook_grids", lambda *a, **k: called.append(1))
    v2 = _vault(tmp_path, "v2")
    (v2 / "data_in" / "big.xlsx").write_bytes((v / "data_in" / "big.xlsx").read_bytes())
    streamed = _kg(v2)
    stats = streamed.seed()
    assert called == [] and _rels(streamed) == want and len(want) == 200
    assert stats["notes"] == ["big.xlsx: a big workbook: read one row at a time, without "
                              "merged cells or pictures"]


def test_the_corpus_tracker_workbook_gives_the_answer_keys_links(tmp_path):
    """The Ironbridge tracker (Read Me sheet, then a Projects table) read by
    the Ledger's reader: every link the key gives it, none other, cited as
    before ('sheet Projects, row 2 (Program Lead, Project ID)')."""
    from council_core import kg_corpus
    v = tmp_path / "corpus"
    key = kg_corpus.build(v)
    kg = _kg(v)
    kg.add_label_synonym("POC", "Point of Contact")
    kg.seed()

    def k_of(ent):
        for k, e in key["entities"].items():
            if e["type"] == ent["type"] and kgm.entity_key(e["type"], e["name"]) == ent["key"]:
                return k
        return "?" + ent["name"]
    from_tracker = set()
    for r in kg.all_relations():
        for e in kg.evidence(r["id"]):
            if e["path"].endswith("program_tracker.xlsx"):
                from_tracker.add((k_of(r["subject"]), r["predicate"], k_of(r["object"])))
    truth = {(r["subject"], r["predicate"], r["object"]) for r in key["relations"]
             if r["file"] == "trackers/program_tracker.xlsx"}
    assert from_tracker == truth
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    where = {e["where"] for n in kg.neighbors(dana["id"]) if n["predicate"] == "LEADS"
             for e in n["evidence"] if e["path"].endswith(".xlsx")}
    assert where == {"sheet Projects, row 2 (Program Lead, Project ID)"}


def test_initials_are_never_a_name_and_names_are_not_initials():
    assert kg_excel.is_initials("JB") and kg_excel.is_initials("j.b.")
    assert kg_excel.is_initials("JMB") and not kg_excel.is_initials("Bob")
    assert kg_excel.person_ok("Jordan Blake") and kg_excel.person_ok("Lee, Carol")
    assert not kg_excel.person_ok("EC Pre-Conceal") and not kg_excel.person_ok("seen")
    assert not kg_excel.person_ok("63") and not kg_excel.person_ok("Date")
    assert kg_excel.split_people("JB/MK") == ["JB", "MK"]
    assert kg_excel.split_people("Lee, Carol") == ["Lee, Carol"]
    assert kg_excel.split_people("JB, MK") == ["JB", "MK"]


# ── folders outside the vault (read only) ───────────────────────────────
def _outside(tmp_path, schema_people=None):
    """A folder outside the vault: a Ledger and a workbook, plus a dot-folder."""
    out = tmp_path / "Q3 Ledgers"
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, out, [_row(schema, "101", "1A", **{f"{ec} [JB]": "63"})],
            people={"JB": "Jordan Blake"})
    _book(out / "tracker.xlsx", {"T": {"A2": "Part Number", "B2": "Owner",
                                       "A3": "PN-1", "B3": "Dana Whitfield"}})
    _book(out / ".hidden" / "skip.xlsx", {"T": {"A1": "Part Number", "B1": "Owner",
                                                "A2": "PN-X", "B2": "Nobody Here"}})
    return out


def test_an_added_folder_is_read_and_never_written(tmp_path):
    v = _vault(tmp_path)
    out = _outside(tmp_path)
    before = _hashes(out)
    kg = _kg(v)
    entry = kg.add_root(out)
    assert entry["name"] == "Q3 Ledgers" and entry["exists"]
    kg.seed()
    rels = _rels(kg)
    assert rels == {("Jordan Blake", "INSPECTED", "101"): "seeded",
                    ("Dana Whitfield", "OWNS", "PN-1"): "seeded"}
    paths = sorted(r[0] for r in kg.db.execute("SELECT path FROM documents"))
    rid = entry["id"]
    assert paths == [f"<{rid}>/lmm_parent.csv", f"<{rid}>/tracker.xlsx"]
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    (ev,) = kg.neighbors(dana["id"])[0]["evidence"]
    assert ev["file"] == "Q3 Ledgers/tracker.xlsx"
    assert kg.abs_path(ev["path"]) == out.resolve() / "tracker.xlsx"
    ctx = kgm.source_context(kg.root, str(kg.abs_path(ev["path"])), ev["locator"])
    assert ("B3", "Owner: Dana Whitfield", True) in ctx
    assert [s["file"] for s in kg.sources()] == ["Q3 Ledgers/lmm_parent.csv",
                                                 "Q3 Ledgers/tracker.xlsx"]
    kg.write_exports()
    assert _hashes(out) == before                     # read, never written
    # kept per vault, in the store
    kg.close()
    again = kgm.KnowledgeGraph(v, image_reader=None)
    assert [r["path"] for r in again.extra_roots()] == [str(out.resolve())]
    assert kgm.KnowledgeGraph(_vault(tmp_path, "other"), image_reader=None).extra_roots() == []


def test_folders_that_overlap_are_refused(tmp_path):
    v = _vault(tmp_path)
    out = _outside(tmp_path)
    kg = _kg(v)
    for bad, why in ((v, "overlaps the vault"), (v / "data_in", "overlaps the vault"),
                     (tmp_path, "overlaps the vault"),
                     (tmp_path / "nope", "is not a folder")):
        with pytest.raises(ValueError, match=why):
            kg.add_root(bad)
    kg.add_root(out)
    for bad in (out, out / ".hidden"):
        with pytest.raises(ValueError, match="overlaps a folder already added"):
            kg.add_root(bad)
    assert len(kg.extra_roots()) == 1
    assert [d["action"] for d in kg.decisions() if "folder" in d["action"]] == ["add_folder"]


def test_a_removed_folder_leaves_the_graph_at_the_next_rebuild(tmp_path):
    v = _vault(tmp_path)
    out = _outside(tmp_path)
    before = _hashes(out)
    kg = _kg(v)
    rid = kg.add_root(out)["id"]
    kg.seed()
    assert _rels(kg)
    kg.remove_root(rid)
    assert _rels(kg)                                   # until the rebuild
    kg.seed()
    assert _rels(kg) == {}
    assert {r[0] for r in kg.db.execute("SELECT status FROM documents")} == {"missing"}
    assert kg.db.execute("SELECT COUNT(*) FROM doc_cache").fetchone()[0] == 0
    assert _names(kg, "PERSON") == []
    assert _hashes(out) == before
    with pytest.raises(KeyError):
        kg.remove_root(rid)


def test_an_added_folder_that_is_not_there_is_a_note(tmp_path):
    v = _vault(tmp_path)
    out = _outside(tmp_path)
    kg = _kg(v)
    kg.add_root(out)
    kg.seed()
    away = tmp_path / "unplugged"
    out.rename(away)
    stats = kg.seed()
    assert stats["notes"] == [f"{out.resolve()}: this added folder is not there now; its "
                              "documents count as gone until it is back"]
    assert _rels(kg) == {}
    away.rename(out)
    kg.seed()
    assert ("Dana Whitfield", "OWNS", "PN-1") in _rels(kg)


# ── progress in the terminal ─────────────────────────────────────────────
def _capture_progress(monkeypatch, heartbeat=0.0):
    import progress_report
    out = io.StringIO()
    real = progress_report.Progress

    class Captured(real):
        def __init__(self, task, **kw):
            kw.update(stream=out, heartbeat=heartbeat, every=0)
            super().__init__(task, **kw)
    monkeypatch.setattr(progress_report, "Progress", Captured)
    return out


def test_a_rebuild_prints_its_progress_with_percentages(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    for i in range(4):
        _book(v / "data_in" / f"t{i}.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                                   "A2": f"PN-{i}", "B2": "Dana Whitfield"}})
    out = _capture_progress(monkeypatch)
    _kg(v).seed()
    lines = out.getvalue().splitlines()
    assert lines[0] == "[kg rebuild] listing documents · 0:00"
    assert "[kg rebuild] reading documents: 0/4 files (0%) · 0:00" in lines
    assert any(ln.startswith("[kg rebuild] reading documents: 2/4 files (50%) · t1.xlsx read")
               for ln in lines)
    assert any(ln.startswith("[kg rebuild] building the graph: 4/4 documents (100%)")
               for ln in lines)
    assert lines[-1].startswith("[kg rebuild] done in 0:0") and \
        lines[-1].endswith("4 documents (4 read, 0 unchanged, 0 not read); 4 links, "
                           "0 suggested, 0 question(s)")


def test_a_slow_workbook_shows_a_still_working_line(tmp_path, monkeypatch):
    import time as _t
    v = _vault(tmp_path)
    _book(v / "data_in" / "slow.xlsx", {"S": {"A1": "x"}})
    out = _capture_progress(monkeypatch, heartbeat=0.05)
    real = kg_excel.read_workbook
    monkeypatch.setattr(kg_excel, "read_workbook",
                        lambda *a, **k: _t.sleep(0.4) or real(*a, **k))
    _kg(v).seed()
    assert any("still working, nothing new for" in ln and "reading slow.xlsx" in ln
               for ln in out.getvalue().splitlines())


def test_a_failed_rebuild_says_so_in_the_terminal(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    _book(v / "data_in" / "a.xlsx", {"S": {"A1": "x"}})
    out = _capture_progress(monkeypatch)
    kg = _kg(v)

    def stop(n, total):
        raise KeyboardInterrupt("the app was closed")
    with pytest.raises(KeyboardInterrupt):
        kg.seed(on_progress=stop)
    assert out.getvalue().splitlines()[-1].startswith(
        "[kg rebuild] stopped after 0:00: KeyboardInterrupt: the app was closed")


def test_suggesting_links_prints_its_progress(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    (v / "data_in" / "team.csv").write_text(
        "Project,Program Lead,Part\nPRJ-0915,Dana Whitfield,PN-1001\n", encoding="utf-8")
    (v / "data_in" / "memo.txt").write_text(
        "Dana Whitfield fitted PN-1001 to the rig on Tuesday.\n", encoding="utf-8")
    kg = _kg(v)
    kg.seed()
    out = _capture_progress(monkeypatch)
    kg.suggest_from_text({"this PC": lambda m, **k: json.dumps({"links": []})},
                         model="stand-in")
    lines = out.getvalue().splitlines()
    assert "[kg suggest links (stand-in)] finding passages: 0/1 documents (0%) · 0:00" in lines
    assert any(ln.startswith("[kg suggest links (stand-in)] asking the model: 1/1 passages"
                             " (100%) · memo.txt read by this PC") for ln in lines)
    assert lines[-1].endswith("1 of 1 passages read, 0 link(s) suggested")


def test_the_readers_are_toolkit_free_and_offline():
    src = (Path(__file__).resolve().parent.parent / "council_core" / "kg_excel.py").read_text(
        encoding="utf-8")
    for word in ("tkinter", "PySide6", "PyQt5", "urllib", "requests", "socket", "http"):
        assert word not in src, word


def test_a_huge_neighbourhood_is_read_in_part(tmp_path, monkeypatch):
    """MEASURED (review, 2026-10-16): neighbors() of a person with 4,000
    links took 3.4 s on the GUI thread, though the tab draws 300 of them. Now
    only the links drawn are looked up; the rest are counted."""
    v = _vault(tmp_path)
    rows = "\n".join(f"PN-{i},Dana Whitfield" for i in range(600))
    (v / "data_in" / "t.csv").write_text("Part,Owner\n" + rows + "\n", encoding="utf-8")
    kg = _kg(v)
    kg.seed()
    dana = kg.search("Dana Whitfield", "PERSON")[0]["id"]
    calls = []
    real = kg.evidence
    monkeypatch.setattr(kg, "evidence", lambda rid, limit=None: calls.append(rid)
                        or real(rid, limit))
    ns = kg.neighbors(dana, evidence_limit=5, relation_limit=300)
    assert len(ns) == 300 and len(calls) == 300
    assert kg.neighbor_counts(dana) == {("OWNS", "out"): 600}


# ── review (i), 2026-10-16: correctness findings, each failing before ────
def test_r1_an_everywhere_answer_never_overrides_a_ledgers_own_list(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in" / "a", [_row(schema, "101", "1A", **{f"{ec} [JB]": "1"})],
            people={"JB": "Jordan Blake"})
    _ledger(tmp_path, v / "data_in" / "b", [_row(schema, "202", "1A", **{f"{ec} [JB]": "2"})])
    (v / "data_in" / "team.csv").write_text("Project,Owner\nPRJ-1,Jill Brown\n", encoding="utf-8")
    kg = _kg(v)
    kg.seed()
    (q,) = [r for r in kg.reviews() if r["kind"] == "initials"]
    assert q["path"] == "b/lmm_parent.csv"
    kg.answer_review(q["id"], kg.search("Jill Brown", "PERSON")[0]["id"], everywhere=True)
    kg.seed()
    rels = {k: s for k, s in _rels(kg).items() if k[1] == "INSPECTED"}
    assert rels == {("Jordan Blake", "INSPECTED", "101"): "seeded",
                    ("Jill Brown", "INSPECTED", "202"): "seeded"}


@pytest.mark.parametrize("heads", [
    ["Part Number", "Owner", "Q1", "Q2"],          # titles that look like codes
    ["Part Number", "Owner", 2025, 2026],          # titles that are numbers
    ["Part Number", None, "Owner"],                # a title cell left empty
    ["Part Number", "Owner", "Lot 3"],
])
def test_r1_a_header_with_code_number_or_empty_titles_is_still_a_header(tmp_path, heads):
    v = _vault(tmp_path)
    cells = {}
    for i, h in enumerate(heads):
        if h is not None:
            cells[f"{openpyxl.utils.get_column_letter(i + 1)}1"] = h
    oc = openpyxl.utils.get_column_letter(heads.index("Owner") + 1)
    for r, (pn, who) in enumerate((("PN-1001", "Dana Whitfield"), ("PN-1002", "Carol Lee"),
                                   ("PN-1003", "Hana Sato")), start=2):
        cells[f"A{r}"], cells[f"{oc}{r}"] = pn, who
        for i in range(len(heads)):
            col = openpyxl.utils.get_column_letter(i + 1)
            if heads[i] not in ("Part Number", "Owner", None):
                cells[f"{col}{r}"] = r
    _book(v / "data_in" / "plan.xlsx", {"Plan": cells})
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "OWNS", "PN-1001"): "seeded",
                         ("Carol Lee", "OWNS", "PN-1002"): "seeded",
                         ("Hana Sato", "OWNS", "PN-1003"): "seeded"}


def test_r1_side_by_side_tables_one_empty_column_apart_stay_apart(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "t.xlsx", {"S": {
        "A1": "Part Number", "B1": "Owner", "D1": "Project", "E1": "Program Lead",
        "A2": "PN-1", "B2": "Dana Whitfield", "D2": "PRJ-9", "E2": "Hana Sato"}})
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "OWNS", "PN-1"): "seeded",
                         ("Hana Sato", "LEADS", "PRJ-9"): "seeded"}


def test_r1_a_csv_beside_a_json_table_schema_is_a_plain_table(tmp_path):
    v = _vault(tmp_path)
    (v / "data_in" / "team.csv").write_text("Project,Program Lead\nPRJ-1,Mary King\n",
                                            encoding="utf-8")
    (v / "data_in" / "team.schema.json").write_text(json.dumps(
        {"fields": [{"name": "Project", "type": "string"},
                    {"name": "Program Lead", "type": "string"}]}), encoding="utf-8")
    kg = _kg(v)
    stats = kg.seed()
    assert stats["unreadable"] == []
    assert _rels(kg) == {("Mary King", "LEADS", "PRJ-1"): "seeded"}
    assert sorted(p.name for p in kg.document_files()) == ["team.csv", "team.schema.json"]


def test_r1_label_synonyms_are_still_found_in_workbooks(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "p.xlsx", {"S": {"A1": "Proj Code", "B1": "POC", "C1": "Qty",
                                           "A2": "PRJ-1", "B2": "Dana Whitfield", "C2": 3}})
    kg = _kg(v)
    kg.seed()
    assert {"Proj Code", "POC", "Qty"} <= set(kg.harvest_labels())
    assert ("POC", "Point of Contact") in {(s["label"], s["means"])
                                           for s in kg.suggest_label_synonyms()}


def test_r1_initials_under_any_person_column_go_through_the_list(tmp_path):
    t = tmp_path / "templates" / "own.xlsx"
    t.parent.mkdir(parents=True)
    wb = openpyxl.Workbook()
    for c, h in enumerate(["Part Number", "Owner", "Notes"], start=1):
        wb.active.cell(1, c, h)
    wb.save(t)
    schema = read_template(t)
    v = _vault(tmp_path)
    led = ParentCSV.create(v / "data_in" / "own_parent.csv", schema)
    led.add_samples([{"Part Number": "2041", "Owner": "JB"}])
    p = se.aliases_path_for(v / "data_in" / "own_parent.csv")
    p.write_text(json.dumps({"map": {}, "ignore": [], "people": {"JB": "Jordan Blake"}}),
                 encoding="utf-8")
    _book(v / "data_in" / "w.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                           "A2": "PN-9", "B2": "MK"}})
    kg = _kg(v)
    kg.seed()
    rels = _rels(kg)
    assert rels[("Jordan Blake", "OWNS", "2041")] == "seeded"
    assert not [k for k in rels if k[0] in ("JB", "MK")]
    assert [r["surface"] for r in kg.reviews()] == ["MK"]


def test_r1_placeholders_are_nobody():
    for v in ("N/A", "n/a", "TBD", "OK", "NA", "tbd", "n.a.", "-", "?"):
        assert kg_excel.split_people(v) == [], v
    assert kg_excel.split_people("EC") == ["EC"]        # Eva Chen signs as EC
    assert kg_excel.split_people("JB/MK") == ["JB", "MK"]


def test_r1_a_blank_row_ends_a_table(tmp_path):
    v = _vault(tmp_path)
    _book(v / "data_in" / "tracker.xlsx", {"Tracker": {
        "A1": "Part Number", "B1": "Owner",
        "A2": "PN-1001", "B2": "Dana Whitfield",
        "A3": "PN-1002", "B3": "Carol Lee",
        "A5": "Name", "B5": "Role",
        "A6": "Hana Sato", "B6": "Engineer",
        "A7": "Tomas Reyes", "B7": "Technician",
        "A9": "Note: values from the Q3 audit"}})
    kg = _kg(v)
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "OWNS", "PN-1001"): "seeded",
                         ("Carol Lee", "OWNS", "PN-1002"): "seeded"}
    assert _names(kg, "PART") == ["PN-1001", "PN-1002"]


def test_r1_a_marker_above_a_table_is_cited_at_the_row(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    _ledger(tmp_path, v / "data_in" / "led", [_row(schema, "999", "1A")],
            people={"JB": "Jordan Blake"})
    _book(v / "data_in" / "insp.xlsx", {"Results": {
        "A1": "Inspector: JB",
        "A3": "Part Number", "B3": "Description", "C3": "Qty",
        "A4": "PN-1001", "B4": "bracket", "C4": 2,
        "A5": "PN-1002", "B5": "shim", "C5": 4}})
    kg = _kg(v)
    kg.seed()
    jb = kg.search("Jordan Blake", "PERSON")[0]
    wheres = sorted(e["where"] for n in kg.neighbors(jb["id"]) for e in n["evidence"]
                    if e["path"] == "insp.xlsx")
    assert wheres == ["sheet Results, row 4 (Part Number)", "sheet Results, row 5 (Part Number)"]


def _src_who(folder, name, who):
    return _src_with_inspector(folder / name, who=who)


def test_r1_from_cites_only_the_review_that_was_applied(tmp_path):
    import time as _t
    from council_core.ledger_ingest import run_apply, run_extract
    v = _vault(tmp_path)
    folder = v / "data_in" / "ledger"
    csv, _ = _ledger(tmp_path, folder, [], people={"JB": "Jordan Blake"})
    a = _src_who(folder, "nde_first.xlsx", "JB")
    r1 = folder / "review_1.xlsx"
    run_extract(csv, [a], r1, use_images=False)
    _accept_all(r1)
    run_apply(r1, use_images=False)
    _t.sleep(1.1)
    b = _src_who(folder, "nde_resend.xlsx", "JB")
    r2 = folder / "review_2.xlsx"
    run_extract(csv, [b], r2, use_images=False)
    _accept_all(r2)                                     # saved, never applied
    kg = _kg(v)
    kg.seed()
    jb = kg.search("Jordan Blake", "PERSON")[0]["id"]
    froms = {e["locator"].get("from") for n in kg.neighbors(jb) for e in n["evidence"]
             if e["method"] == "ledger" and e["locator"].get("item") == "Flaw # 1A"}
    assert froms == {"nde_first.xlsx › Part 101!C4"}


def test_r1_agreeing_inspectors_are_each_cited_at_their_own_cell(tmp_path):
    from council_core.ledger_ingest import run_apply, run_extract
    v = _vault(tmp_path)
    folder = v / "data_in" / "ledger"
    csv, _ = _ledger(tmp_path, folder, [], people={"JB": "Jordan Blake", "MK": "Mary King"})
    a = _src_who(folder, "nde_jb.xlsx", "JB")
    b = _src_who(folder, "nde_mk.xlsx", "MK")
    review = folder / "review_1.xlsx"
    run_extract(csv, [a, b], review, use_images=False)
    _accept_all(review)
    run_apply(review, use_images=False)
    kg = _kg(v)
    kg.seed()
    for who, book in (("Mary King", "nde_mk.xlsx"), ("Jordan Blake", "nde_jb.xlsx")):
        eid = kg.search(who, "PERSON")[0]["id"]
        froms = {e["locator"].get("from") for n in kg.neighbors(eid) for e in n["evidence"]
                 if e["method"] == "ledger" and e["locator"].get("item") == "Flaw # 1A"}
        assert froms == {f"{book} › Part 101!C4"}, (who, froms)


def test_r1_part_level_values_are_traced_too(tmp_path):
    from sample_ledger import TemplateField, TemplateSchema
    from source_extract import Decision, Extraction, Obs, rows_from, write_review
    v = _vault(tmp_path)
    base = _schema(tmp_path)
    fields = list(base.fields)
    fields.insert(1, TemplateField(name="Project", label="Project", sheet="Sheet",
                                   section="Part Attributes", layout="table", cells={"": "F3"}))
    schema = TemplateSchema(base.template_path, fields, base.id_field, base.item_field, base.items)
    ec = _ec(schema)
    folder = v / "data_in" / "led"
    folder.mkdir(parents=True)
    csv = folder / "lmm_parent.csv"
    led = ParentCSV.create(csv, schema)
    obs = [Obs("101", "", "Project", "PRJ-7", 1.0, "alias", "Project", "src.xlsx › Part 101!B2"),
           Obs("101", "1A", f"{ec} [JB]", "63", 1.0, "alias", "EC Pre",
               "src.xlsx › Part 101!C4", inspector="JB")]
    decs = [Decision(o, True) for o in obs]
    run = Extraction(["src.xlsx"], obs=obs)
    write_review(folder / "review_1.xlsx", run, decs, schema, str(csv))
    led.add_samples(rows_from(decs, schema), source="src.xlsx")
    reading = kg_excel.read_ledger(csv, "led/lmm_parent.csv", list(kgm.DEFAULT_FIELD_RULES))
    projects = [val for rec in reading.records for val in rec.values
                if val.rule.label == "Project"]
    assert projects and all(val.locator.get("from") == "src.xlsx › Part 101!B2"
                            for val in projects)


def test_r1_a_picture_text_line_previews(tmp_path):
    from openpyxl.drawing.image import Image as XLImage
    v = _vault(tmp_path)
    wb = openpyxl.Workbook()
    wb.active.title = "Visual"
    wb.active.add_image(XLImage(io.BytesIO(lf._png([["x"]]))), "F2")
    wb.save(v / "data_in" / "vis.xlsx")
    kg = _kg(v, reader=lf.fake_reader([["Dana Whitfield checked the shim stack today"]]))
    kg.seed()
    (line,) = json.loads(kg.db.execute("SELECT lines FROM doc_cache").fetchone()[0])
    loc = {"page": line[0], "line": line[1]}
    ctx = kgm.source_context(kg.root, "vis.xlsx", loc, quote=line[2])
    assert ctx and ctx[0][1] == "a picture at F2 on sheet Visual, read by OCR"
    assert kgm.locator_text(loc) == "sheet Visual › picture F2, row 1"


def test_r1_sheet_names_that_are_years_or_codes(tmp_path):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    _ledger(tmp_path, v / "data_in" / "led", [_row(schema, "999", "1A")],
            people={"JB": "Jordan Blake"})
    _book(v / "data_in" / "y.xlsx", {"2025": {"A1": "Inspector: JB"},
                                     "PN-1001": {"A1": "Inspector: JB"},
                                     "Part 101": {"A1": "Inspector: JB"}})
    kg = _kg(v)
    kg.seed()
    assert {k[2] for k in _rels(kg) if k[1] == "INSPECTED"} == {"PN-1001", "101"}


# ── review (ii), 2026-10-16: safety and scale findings, each failing before ──
def _owner_rows(v, n, who="Dana Whitfield", name="t.csv"):
    rows = "\n".join(f"PN-{i},{who}" for i in range(n))
    (v / "data_in" / name).write_text("Part,Owner\n" + rows + "\n", encoding="utf-8")


def test_r2_the_agent_tool_asks_for_a_bounded_neighbourhood(tmp_path, monkeypatch):
    from council_core import agent_profiles
    v = _vault(tmp_path)
    _owner_rows(v, 5)
    _kg(v).seed()
    seen = []
    real = kgm.KnowledgeGraph.neighbors
    monkeypatch.setattr(kgm.KnowledgeGraph, "neighbors",
                        lambda self, eid, **k: seen.append(k) or real(self, eid, **k))
    out = agent_profiles._graph_tools(v)["graph_neighbors"].fn({"name": "Dana Whitfield"}, None)
    assert len(out["links"]) == 5
    assert seen and seen[0].get("evidence_limit") and seen[0].get("relation_limit")


def test_r2_a_hidden_excel_is_tied_to_the_council_and_dies_with_it(tmp_path):
    """Quitting the Council while a rebuild's daemon worker is inside Excel
    never ran read_with_excel's finally: the hidden Excel stayed. Now every
    Excel the Council starts is in a Windows job that closes with it."""
    import subprocess
    if os.name != "nt":
        pytest.skip("Windows job objects")
    pytest.importorskip("win32job")
    parent = tmp_path / "parent.py"
    parent.write_text(
        "import subprocess, sys\n"
        f"sys.path.insert(0, {str(Path(__file__).resolve().parent.parent)!r})\n"
        "import excel_open\n"
        "child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'],\n"
        "                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)\n"
        "print(child.pid, excel_open.tie_to_this_process(child.pid), flush=True)\n",
        encoding="utf-8")
    out = subprocess.run([sys.executable, str(parent)], capture_output=True, text=True,
                         timeout=60)
    pid, tied = out.stdout.split()
    assert tied == "True", out.stderr
    import time as _t
    import psutil
    for _ in range(50):
        if not psutil.pid_exists(int(pid)):
            break
        _t.sleep(0.1)
    assert not psutil.pid_exists(int(pid))


def test_r2_read_with_excel_ties_the_excel_it_starts(tmp_path, monkeypatch):
    tied = []
    monkeypatch.setattr(excel_open, "_tie_excel", lambda app: tied.append(app))
    f = tmp_path / "x.xls"
    f.write_bytes(excel_open.OLE_MAGIC + b"\0" * 504)
    log = []
    excel_open.read_with_excel(f, want_pictures=False,
                               app_factory=lambda: _App(log, [_Sheet("S", [("a",)])]))
    assert len(tied) == 1


def test_r2_a_password_xls_is_not_retried_nor_harvested(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    (v / "data_in" / "locked.xls").write_bytes(excel_open.OLE_MAGIC + b"\0" * 504)
    log = _fake_excel(monkeypatch, [], fail="The password you supplied is not correct. "
                                            "Verify that the CAPS LOCK key is off")
    kg = _kg(v)
    for _ in range(3):
        stats = kg.seed()
    assert sum(1 for e in log if e[0] == "open") == 1
    assert len(stats["unreadable"]) == 1 and "password" in stats["unreadable"][0]
    kg.suggest_label_synonyms()
    assert sum(1 for e in log if e[0] == "open") == 1


def test_r2_the_harvest_does_not_reopen_a_workbook_that_failed(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    (v / "data_in" / "busy.xls").write_bytes(excel_open.OLE_MAGIC + b"\0" * 504)
    log = _fake_excel(monkeypatch, [], fail="Excel is busy")
    kg = _kg(v)
    kg.seed()
    n = sum(1 for e in log if e[0] == "open")
    kg.suggest_label_synonyms()
    assert sum(1 for e in log if e[0] == "open") == n


def test_r2_folder_overlaps_cannot_be_spelt_around(tmp_path):
    v = _vault(tmp_path)
    (v / "data_in" / "sub").mkdir()
    kg = _kg(v)
    with pytest.raises(ValueError, match="overlaps the vault"):
        kg.add_root("\\\\?\\" + str((v / "data_in" / "sub").resolve()))
    assert kg.extra_roots() == []


def test_r2_a_junction_in_an_added_folder_is_not_followed(tmp_path):
    if os.name != "nt":
        pytest.skip("Windows junctions")
    import _winapi
    v = _vault(tmp_path)
    (v / "data_out").mkdir()
    (v / "data_in" / "team.csv").write_text("Project,Program Lead\nPRJ-1,Dana Whitfield\n",
                                            encoding="utf-8")
    (v / "data_out" / "answer.txt").write_text("Program Lead: Secret Person\nProject: PRJ-9\n",
                                               encoding="utf-8")
    out = tmp_path / "outside"
    out.mkdir()
    (out / "own.txt").write_text("Program Lead: Lee Carol\nProject: PRJ-2\n", encoding="utf-8")
    _winapi.CreateJunction(str(v), str(out / "j_vault"))
    try:
        kg = _kg(v)
        kg.add_root(out)
        kg.seed()
        docs = sorted(kg.display_path(kg._rel(p)) for p in kg.document_files())
        assert docs == ["outside/own.txt", "team.csv"]
        assert "Secret Person" not in _names(kg, "PERSON")
    finally:
        os.rmdir(out / "j_vault")


def test_r2_system_folders_of_an_added_drive_are_skipped(tmp_path):
    v = _vault(tmp_path)
    out = tmp_path / "drive"
    for sub in ("$RECYCLE.BIN/S-1-5", "System Volume Information", "work"):
        (out / sub).mkdir(parents=True)
    (out / "$RECYCLE.BIN" / "S-1-5" / "$R4XK2.txt").write_text(
        "Program Lead: Deleted Person\nProject: PRJ-DEL\n", encoding="utf-8")
    (out / "System Volume Information" / "x.txt").write_text("Project: PRJ-SYS\n",
                                                             encoding="utf-8")
    (out / "work" / "a.txt").write_text("Program Lead: Dana Whitfield\nProject: PRJ-1\n",
                                        encoding="utf-8")
    kg = _kg(v)
    kg.add_root(out)
    assert [p.name for p in kg.document_files()] == ["a.txt"]


def test_r2_the_spreadsheets_list_reads_only_each_readings_summary(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in", [_row(schema, str(1000 + i), "1A", **{f"{ec} [JB]": "1"})
                                      for i in range(400)], people={"JB": "Jordan Blake"})
    kg = _kg(v)
    kg.seed()
    full = kg.db.execute("SELECT length(records) FROM doc_cache").fetchone()[0]
    assert full > 100_000
    sizes = []
    real = kgm._reading_load
    monkeypatch.setattr(kgm, "_reading_load",
                        lambda text: sizes.append(len(text or "")) or real(text))
    (src,) = kg.sources()
    assert src["summary"]["rows"] == 400
    assert sizes and max(sizes) < 10_000


def test_r2_an_unchanged_rebuild_does_not_rebuild_the_graph(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A", **{f"{ec} [MK]": "1"})])
    _owner_rows(v, 3)
    kg = _kg(v)
    kg.seed()
    before = _rels(kg), [r["surface"] for r in kg.reviews()]
    calls = []
    real = kg._seed_records
    monkeypatch.setattr(kg, "_seed_records", lambda *a, **k: calls.append(1) or real(*a, **k))
    stats = kg.seed()
    assert calls == [] and stats["graph_unchanged"] is True
    assert (_rels(kg), [r["surface"] for r in kg.reviews()]) == before
    # ... but anything the graph is worked out from makes it build again:
    (q,) = kg.reviews()
    kg.answer_review(q["id"], kg.add_person("Mary King"), everywhere=True)
    kg.seed()
    assert calls == [1] and ("Mary King", "INSPECTED", "101") in _rels(kg)
    p = se.aliases_path_for(v / "data_in" / "lmm_parent.csv")
    p.write_text(json.dumps({"map": {}, "ignore": [], "people": {"MK": "Mike Kent"}}),
                 encoding="utf-8")
    kg.seed()
    assert calls == [1, 1]
    dana = kg.search("Dana Whitfield", "PERSON")[0]["id"]
    kg.merge(kg.search("Mary King", "PERSON")[0]["id"], dana)
    kg.seed()
    assert calls == [1, 1, 1]
    kg.set_rule_status("Owner", "rejected")
    kg.seed()
    assert calls == [1, 1, 1, 1]


def test_r2_a_graph_being_rebuilt_is_busy_not_damaged(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    _owner_rows(v, 3)
    kg = _kg(v)
    kg.seed()
    monkeypatch.setattr(kgm.KnowledgeGraph, "BUSY_TIMEOUT_S", 0.2)
    kg.db.execute("BEGIN EXCLUSIVE")
    try:
        with pytest.raises(kgm.KnowledgeGraphBusy, match="being rebuilt"):
            kgm.KnowledgeGraph(v, image_reader=None)
        assert not issubclass(kgm.KnowledgeGraphBusy, kgm.KnowledgeGraphDamaged)
    finally:
        kg.db.rollback()
    kgm.KnowledgeGraph(v, image_reader=None).close()       # free again


def test_r2_a_ledger_reads_with_no_openpyxl(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    ec = _ec(schema)
    csv, _ = _ledger(tmp_path, v / "data_in", [_row(schema, "101", "1A",
                                                    **{f"{ec} [JB]": "1"})],
                     people={"JB": "Jordan Blake"})
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    reading = kg_excel.read_ledger(csv, "lmm_parent.csv", [])
    assert reading.summary["rows"] == 1 and reading.notes == []
    (csv.parent / "review_1.xlsx").write_bytes(b"PK")                 # a review is there
    monkeypatch.setattr(kg_excel, "_reviews_of", lambda p: [csv.parent / "review_1.xlsx"])
    reading = kg_excel.read_ledger(csv, "lmm_parent.csv", [])
    assert reading.notes == ["review workbooks were not read (they need the openpyxl "
                             "package, which is not installed): no workbook cells for "
                             "these values"]


def test_r2_a_new_picture_reader_re_reads_only_workbooks(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    _owner_rows(v, 2)
    (v / "data_in" / "memo.txt").write_text("Program Lead: Dana Whitfield\nProject: P-1\n",
                                            encoding="utf-8")
    _book(v / "data_in" / "w.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                           "A2": "PN-9", "B2": "Carol Lee"}})
    kgm.KnowledgeGraph(v, image_reader=None).seed()
    import collections
    calls = collections.Counter()
    real = kgm.read_records
    monkeypatch.setattr(kgm, "read_records",
                        lambda p, rel, rules: calls.update([Path(p).name]) or
                        real(p, rel, rules))
    kg = kgm.KnowledgeGraph(v, image_reader=lf.fake_reader([["x"]]))
    stats = kg.seed()
    assert set(calls) == {"w.xlsx"} and stats["documents_read"] == 1


def test_r2_excel_failing_half_way_is_said_in_words(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    (v / "data_in" / "half.xls").write_bytes(excel_open.OLE_MAGIC + b"\0" * 504)

    class Bad:
        def __iter__(self):
            raise RuntimeError("the RPC server is unavailable")
    log = _fake_excel(monkeypatch, Bad())
    stats = _kg(v).seed()
    assert stats["unreadable"] == ["half.xls: Excel stopped while it was being read "
                                   "(RuntimeError: the RPC server is unavailable)"]
    assert ("close", False) in log and log[-1] == ("quit",)


class _PicShape:
    Type, HasChart = 13, False
    TopLeftCell = type("c", (), {"Address": lambda self, a=False, b=False: "E2",
                                 "GetAddress": lambda self, a=False, b=False: "E2"})()
    TextFrame2 = type("tf", (), {"HasText": False})()


def test_r2_a_workbook_is_read_from_memory_not_held_open(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    _book(v / "data_in" / "w.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                           "A2": "PN-9", "B2": "Carol Lee"}})
    seen = []
    real = openpyxl.load_workbook
    monkeypatch.setattr(openpyxl, "load_workbook",
                        lambda f, *a, **k: seen.append(type(f).__name__) or real(f, *a, **k))
    _kg(v).seed()
    assert seen and set(seen) == {"BytesIO"}
    assert _rels(_kg(v)) == {("Carol Lee", "OWNS", "PN-9"): "seeded"}


def test_r2_free_text_is_capped_by_characters_too(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    cells = {f"A{r}": "a long note about the shim stack and the bracket " * 4
             for r in range(1, 60)}
    _book(v / "data_in" / "n.xlsx", {"S": cells})
    monkeypatch.setattr(kg_excel, "MAX_TEXT_CHARS", 1000)
    kg = _kg(v)
    stats = kg.seed()
    lines = json.loads(kg.db.execute("SELECT lines FROM doc_cache").fetchone()[0])
    assert sum(len(t) for _p, _l, t in lines) <= 1000
    assert "n.xlsx: only the first 1000 characters of free text were kept" in stats["notes"]


def test_a_rebuild_takes_records_as_cache_lines_not_objects(tmp_path, monkeypatch):
    """MEASURED (2026-10-18): a 100k-row workbook's records, held as objects
    and then turned into text, peaked at 274 MB; their lines are 46 MB."""
    v = _vault(tmp_path)
    _book(v / "data_in" / "w.xlsx", {"S": {"A1": "Part Number", "B1": "Owner",
                                           "A2": "PN-9", "B2": "Carol Lee"}})
    got = []
    real = kg_excel.read_workbook
    monkeypatch.setattr(kg_excel, "read_workbook",
                        lambda *a, **k: got.append(real(*a, **k)) or got[-1])
    kg = _kg(v)
    kg.seed()
    (reading,) = got
    assert reading.records == [] and len(reading.record_lines) == 1
    assert _rels(kg) == {("Carol Lee", "OWNS", "PN-9"): "seeded"}
    direct = real(v / "data_in" / "w.xlsx", "w.xlsx", [r for r, _s in kg.field_rules()])
    assert direct.record_lines is None and len(direct.records) == 1


# ── pictures: read, but trusted only when confirmed (user, 2026-10-18) ──
# "Lets have the knowledge graphs try to read the images but maybe not have
# the data from the images be 100% trusted unless we are positive that the
# data pulled from the images is accurate."
def _conf_reader(rows, words):
    """A stand-in OCR that also says how sure it is of each cell (by
    position, as source_extract's Tesseract reader does): a cell's least
    sure word, 0 for a word not in ``words``."""
    def read(_data):
        return [list(r) for r in rows]

    def matrix():
        return [[min((words.get(t, 0) for t in str(c).split()), default=0) for c in r]
                for r in rows]
    read.name = "test OCR"
    read.with_confidence = lambda _data: ([list(r) for r in rows], matrix())
    return read


#: The header words, read sure.
HEAD = {"Part": 99, "Number": 99, "Owner": 99}


def _picture_book(path, sheet="Visual", anchor="F2"):
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    wb.active.title = sheet
    wb.active["A1"] = "Visual inspection"
    wb.active.add_image(XLImage(io.BytesIO(lf._png([["x"]]))), anchor)
    wb.save(path)
    return path


def _known_elsewhere(v):
    """PN-77 and Dana Whitfield are known from tables - but nothing says Dana
    owns PN-77 except the picture."""
    (v / "data_in" / "parts.csv").write_text("Part,Project\nPN-77,PRJ-1\n", encoding="utf-8")
    (v / "data_in" / "leads.csv").write_text("Project,Program Lead\nPRJ-1,Dana Whitfield\n",
                                             encoding="utf-8")


ROWS = [["Part Number", "Owner"], ["PN-77", "Dana Whitfield"]]


def test_word_confidences_come_from_tesseracts_tsv():
    tsv = ("level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight"
           "\tconf\ttext\n"
           "5\t1\t1\t1\t1\t1\t10\t10\t40\t12\t96.5\tPN-77\n"
           "5\t1\t1\t1\t1\t2\t80\t10\t40\t12\t91\tDana\n"
           "5\t1\t1\t1\t1\t3\t125\t10\t60\t12\t88.2\tWhitfield\n"
           "5\t1\t1\t1\t2\t1\t80\t30\t40\t12\t99\tDana\n"
           "4\t1\t1\t1\t1\t0\t10\t10\t200\t12\t-1\t\n")
    assert se.word_confidences(tsv) == {"PN-77": 96.5, "Dana": 91.0, "Whitfield": 88.2}
    reader = se.tesseract_reader(exe=Path(sys.executable))
    assert callable(reader.with_confidence)          # built, not run


def test_a_picture_link_is_firm_only_when_ocr_was_sure_and_every_name_is_known(tmp_path):
    v = _vault(tmp_path)
    _known_elsewhere(v)
    _picture_book(v / "data_in" / "vis.xlsx")
    kg = _kg(v, reader=_conf_reader(ROWS, {"PN-77": 97, "Dana": 95, "Whitfield": 93,
                                           "Part": 99, "Number": 99, "Owner": 99}))
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "seeded"
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    (ev,) = [e for n in kg.neighbors(dana["id"]) if n["predicate"] == "OWNS"
             for e in n["evidence"]]
    assert ev["method"] == "picture"
    assert ev["locator"]["ocr"] == 93 and ev["locator"]["checked"] is True
    assert ev["where"] == ("sheet Visual, picture at F2, row 2 (Owner, Part Number) — read "
                           "from a picture, 93% sure, every name found elsewhere")


@pytest.mark.parametrize("words, known, why", [
    ({"PN-77": 97, "Dana": 95, "Whitfield": 62, **HEAD}, True, "62% sure"),    # OCR unsure
    ({"PN-77": 97, "Dana": 95, "Whitfield": 93, **HEAD}, False, "93% sure"),   # PN-77 unknown
    ({}, True, "how sure unknown"),                                            # no confidence
])
def test_a_picture_link_otherwise_waits_for_the_user(tmp_path, words, known, why):
    v = _vault(tmp_path)
    if known:
        _known_elsewhere(v)
    else:
        (v / "data_in" / "leads.csv").write_text("Project,Program Lead\nPRJ-1,Dana Whitfield\n",
                                                 encoding="utf-8")
    _picture_book(v / "data_in" / "vis.xlsx")
    reader = _conf_reader(ROWS, words) if words else lf.fake_reader(ROWS)
    kg = _kg(v, reader=reader)
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "suggested"
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    (ev,) = [e for n in kg.neighbors(dana["id"]) if n["predicate"] == "OWNS"
             for e in n["evidence"]]
    assert ev["locator"]["checked"] is False
    assert ev["where"].endswith(f"— read from a picture, {why} — check it")


def test_a_link_another_source_states_is_firm_whatever_the_picture(tmp_path):
    v = _vault(tmp_path)
    (v / "data_in" / "own.csv").write_text("Part,Owner\nPN-77,Dana Whitfield\n", encoding="utf-8")
    _picture_book(v / "data_in" / "vis.xlsx")
    kg = _kg(v, reader=_conf_reader(ROWS, {"PN-77": 40, "Dana": 40, "Whitfield": 40}))
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "seeded"


def _xls_with_picture(v):
    (v / "data_in" / "pics.xls").write_bytes(excel_open.OLE_MAGIC + b"\0" * 504)
    sheet = _Sheet("Visual", [("Inspection",)])
    sheet.Shapes = _Coll([_PicShape()])
    return sheet


def _excel_with_pictures(monkeypatch, sheet):
    asked = []
    real = excel_open.read_with_excel
    monkeypatch.setattr(excel_open, "excel_available", lambda: True)
    # the stand-in never touches the real clipboard
    monkeypatch.setattr(excel_open, "_copy_picture", lambda shape: lf._png([["x"]]))
    monkeypatch.setattr(excel_open, "_clip_text", lambda: None)
    monkeypatch.setattr(excel_open, "_clip_restore", lambda text: None)
    monkeypatch.setattr(excel_open, "_clip_seq", lambda: 1)
    monkeypatch.setattr(excel_open, "read_with_excel",
                        lambda path, want_pictures=True, careful=False, **kw:
                        asked.append(want_pictures) or (want_pictures and not careful and
                                                        pytest.fail("not careful")) or
                        real(path, want_pictures=want_pictures, careful=careful,
                             app_factory=lambda: _App([], [sheet])))
    return asked


def test_pictures_in_an_excel_only_workbook_are_read_when_the_clipboard_is_safe(
        tmp_path, monkeypatch):
    v = _vault(tmp_path)
    _known_elsewhere(v)
    sheet = _xls_with_picture(v)
    asked = _excel_with_pictures(monkeypatch, sheet)
    monkeypatch.setattr(excel_open, "clipboard_safe", lambda: (True, ""))
    kg = _kg(v, reader=_conf_reader(ROWS, {"PN-77": 97, "Dana": 95, "Whitfield": 93,
                                           **HEAD}))
    stats = kg.seed()
    assert asked == [True] and stats["pictures_read"] == 1
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "seeded"


@pytest.mark.parametrize("why", [
    "the clipboard holds something besides plain text, which copying would replace",
    "Windows syncs this PC's clipboard to your other devices",
])
def test_pictures_in_an_excel_only_workbook_wait_when_the_clipboard_is_not_safe(
        tmp_path, monkeypatch, why):
    v = _vault(tmp_path)
    sheet = _xls_with_picture(v)
    asked = _excel_with_pictures(monkeypatch, sheet)
    monkeypatch.setattr(excel_open, "clipboard_safe", lambda: (False, why))
    stats = _kg(v, reader=_conf_reader(ROWS, {})).seed()
    assert asked == [False]
    assert stats["notes"] == [f"pics.xls: 1 picture(s) not read: {why} (rebuild again "
                              "later to read them)"]


def test_an_excel_only_workbooks_pictures_without_tesseract(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    sheet = _xls_with_picture(v)
    asked = _excel_with_pictures(monkeypatch, sheet)
    called = []
    monkeypatch.setattr(excel_open, "clipboard_safe", lambda: called.append(1) or (True, ""))
    stats = _kg(v, reader=None).seed()
    assert asked == [False] and called == []          # no reader: the clipboard untouched
    assert stats["notes"] == ["pics.xls: 1 picture(s) not read: install Tesseract (local "
                              "OCR) to read them"]


def test_the_clipboard_check_names_what_would_be_lost(monkeypatch):
    if os.name != "nt":
        pytest.skip("the Windows clipboard")
    monkeypatch.setattr(excel_open, "_clipboard_formats", lambda: {13, 1, 16})
    monkeypatch.setattr(excel_open, "cloud_clipboard_on", lambda: False)
    monkeypatch.setattr(excel_open, "_remote_session", lambda: False)
    assert excel_open.clipboard_safe() == (True, "")
    for held in ({13, 8}, {13, "html format"}, {"rich text format"}, {"image/jpeg"}, {12}):
        monkeypatch.setattr(excel_open, "_clipboard_formats", lambda held=held: held)
        assert excel_open.clipboard_safe() == (
            False, "the clipboard holds something besides plain text, which copying "
                   "would replace"), held
    monkeypatch.setattr(excel_open, "_clipboard_formats", lambda: set())
    monkeypatch.setattr(excel_open, "cloud_clipboard_on", lambda: True)
    assert excel_open.clipboard_safe() == (
        False, "Windows syncs this PC's clipboard to your other devices")


def test_a_name_known_only_from_a_picture_does_not_confirm_another_picture(tmp_path):
    """'Inspector: JB' read from one picture made Jordan Blake 'known from a
    source that is not a picture' (the marker was not marked as picture
    data), so a second picture's link to him was taken as firm."""
    from openpyxl.drawing.image import Image as XLImage
    v = _vault(tmp_path)
    schema = _schema(tmp_path)
    _ledger(tmp_path, v / "data_in" / "led", [_row(schema, "999", "1A")],
            people={"JB": "Jordan Blake"})
    (v / "data_in" / "parts.csv").write_text("Part,Project\nPN-5,PRJ-1\n", encoding="utf-8")
    pics = {}
    for name, rows in (("a.xlsx", [["Inspector: JB"]]),
                       ("b.xlsx", [["Part Number", "Owner"], ["PN-5", "Jordan Blake"]])):
        png = lf._png([[name]])
        pics[png] = rows
        wb = openpyxl.Workbook()
        wb.active.title = "Visual"
        wb.active.add_image(XLImage(io.BytesIO(png)), "F2")
        wb.save(v / "data_in" / name)
    def read(data):
        return [list(r) for r in pics.get(data, [])] or None
    read.name = "test OCR"
    read.with_confidence = lambda data: (read(data), [[99] * len(r) for r in read(data)])
    kg = _kg(v, reader=read)
    kg.seed()
    assert _rels(kg)[("Jordan Blake", "OWNS", "PN-5")] == "suggested"


# ── review (iii), 2026-10-18: the picture rule, each failing before ───────
def _picture_with(v, name, rows, sheet="Visual"):
    from openpyxl.drawing.image import Image as XLImage
    png = lf._png([[name]])
    wb = openpyxl.Workbook()
    wb.active.title = sheet
    wb.active.add_image(XLImage(io.BytesIO(png)), "F2")
    wb.save(v / "data_in" / name)
    return png


def _by_picture(pics, confs):
    """A stand-in OCR that reads each picture (by its bytes) as its rows, with
    the confidence matrix given for it."""
    def read(data):
        return [list(r) for r in pics.get(data, [])] or None
    read.name = "test OCR"
    read.with_confidence = lambda data: (read(data), confs.get(data))
    return read


def test_r3_a_pictures_id_and_name_row_never_renames_other_sources(tmp_path):
    v = _vault(tmp_path)
    (v / "data_in" / "leads.csv").write_text("Project,Program Lead\nHelios,Dana Whitfield\n",
                                             encoding="utf-8")
    rows = [["Project ID", "Project Name"], ["PRJ-9", "Helios"]]
    png = _picture_with(v, "p.xlsx", rows)
    kg = _kg(v, reader=_by_picture({png: rows}, {png: [[40, 40], [40, 40]]}))
    kg.seed()
    assert _rels(kg) == {("Dana Whitfield", "LEADS", "Helios"): "seeded"}


@pytest.mark.parametrize("rows, confs", [
    ([["Part Number", "Owner"], ["PN-77", "Dana Whitfield"]], [[99, 12], [97, 95]]),
    ([["Owner", "Dana Whitfield"], ["Part Number", "PN-77"]], [[12, 95], [99, 97]]),
])
def test_r3_the_title_or_label_a_picture_value_hangs_on_must_be_read_sure_too(
        tmp_path, rows, confs):
    v = _vault(tmp_path)
    _known_elsewhere(v)
    png = _picture_with(v, "p.xlsx", rows)
    kg = _kg(v, reader=_by_picture({png: rows}, {png: confs}))
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "suggested"


def test_r3_an_initial_guessed_from_a_picture_name_confirms_nothing(tmp_path):
    v = _vault(tmp_path)
    (v / "data_in" / "parts.csv").write_text("Part,Project\nPN-77,PRJ-1\n", encoding="utf-8")
    (v / "data_in" / "leads.csv").write_text("Project,Program Lead\nPRJ-1,D. Whitfield\n",
                                             encoding="utf-8")
    png = _picture_with(v, "p.xlsx", ROWS)
    kg = _kg(v, reader=_by_picture({png: ROWS}, {png: [[99, 99], [97, 95]]}))
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "suggested"


def test_r3_a_picture_link_suggested_for_any_reason_is_worded_check_it(tmp_path):
    v = _vault(tmp_path)
    _known_elsewhere(v)
    rows = [["Part Number", "Owner"], ["PN-77", "D. Whitfield"]]
    png = _picture_with(v, "p.xlsx", rows)
    kg = _kg(v, reader=_by_picture({png: rows}, {png: [[99, 99], [97, 93]]}))
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "suggested"
    dana = kg.search("Dana Whitfield", "PERSON")[0]
    (ev,) = [e for n in kg.neighbors(dana["id"]) if n["predicate"] == "OWNS"
             for e in n["evidence"]]
    assert ev["where"].endswith("— read from a picture, 93% sure — check it")


def test_r3_a_merge_the_user_made_counts_as_known(tmp_path):
    v = _vault(tmp_path)
    _known_elsewhere(v)
    rows = [["Part Number", "Owner"], ["PN-77", "Dana Ross"]]
    png = _picture_with(v, "p.xlsx", rows)
    kg = _kg(v, reader=_by_picture({png: rows}, {png: [[99, 99], [97, 95]]}))
    kg.seed()
    assert _rels(kg)[("Dana Ross", "OWNS", "PN-77")] == "suggested"
    kg.merge(kg.search("Dana Whitfield", "PERSON")[0]["id"],
             kg.search("Dana Ross", "PERSON")[0]["id"])
    kg.seed()
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "seeded"


def test_r3_ocr_confidence_follows_each_box_not_the_word(tmp_path):
    head = ("level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth"
            "\theight\tconf\ttext\n")
    tsv = head + ("5\t1\t1\t1\t1\t1\t10\t10\t40\t12\t35\tPN-7\n"     # glued to the next
                  "5\t1\t1\t1\t2\t1\t50\t10\t10\t12\t30\t7\n"
                  "5\t1\t1\t1\t3\t1\t10\t60\t50\t12\t97\tPN-77\n")   # elsewhere, sure
    rows, confs = se.rows_with_confidence(tsv)
    assert rows == [["PN-77"], ["PN-77"]]
    assert confs == [[30], [97]]


def test_r3_pictures_not_copied_are_read_at_a_later_rebuild(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    _known_elsewhere(v)
    sheet = _xls_with_picture(v)
    asked = _excel_with_pictures(monkeypatch, sheet)
    state = {"safe": (False, "Windows syncs this PC's clipboard to your other devices")}
    monkeypatch.setattr(excel_open, "clipboard_safe", lambda: state["safe"])
    kg = _kg(v, reader=_conf_reader(ROWS, {"PN-77": 97, "Dana": 95, "Whitfield": 93,
                                           **HEAD}))
    kg.seed()
    state["safe"] = (True, "")
    stats = kg.seed()
    assert asked == [False, True] and stats["pictures_read"] == 1
    assert _rels(kg)[("Dana Whitfield", "OWNS", "PN-77")] == "seeded"


def test_r3_a_remote_desktop_session_never_copies_pictures(monkeypatch):
    if os.name != "nt":
        pytest.skip("Windows")
    monkeypatch.setattr(excel_open, "cloud_clipboard_on", lambda: False)
    monkeypatch.setattr(excel_open, "_clipboard_formats", lambda: set())
    monkeypatch.setattr(excel_open, "_remote_session", lambda: True)
    assert excel_open.clipboard_safe() == (
        False, "this is a Remote Desktop session, whose clipboard is shared with the "
               "other computer")


def _careful_read(monkeypatch, seqs, copies):
    """read_with_excel in careful mode over a sheet with a picture; the
    clipboard's change counter gives ``seqs`` in turn."""
    seq = iter(seqs)
    restored = []
    monkeypatch.setattr(excel_open, "_clip_seq", lambda: next(seq))
    monkeypatch.setattr(excel_open, "_clip_text", lambda: "the user's text")
    monkeypatch.setattr(excel_open, "_clip_restore", lambda text: restored.append(text))
    monkeypatch.setattr(excel_open, "_copy_picture",
                        lambda shape: copies.append(excel_open._CLIPBOARD_LOCK.locked())
                        or lf._png([["x"]]))
    sheet = _Sheet("Visual", [("Inspection",)])
    sheet.Shapes = _Coll([_PicShape()])
    book = excel_open.read_with_excel(Path(__file__), want_pictures=True, careful=True,
                                      app_factory=lambda: _App([], [sheet]))
    return book, restored


def test_r3_a_clipboard_the_user_changed_meanwhile_is_left_alone(monkeypatch):
    copies = []
    book, restored = _careful_read(monkeypatch, [5, 9], copies)    # 5 at the start, 9 later
    assert copies == [] and restored == []
    assert [p.data for p in book.pictures] == [b""]
    assert book.notes == ["Part 101: picture(s) not copied: the clipboard changed while the "
                          "workbook was being read"] or any(
        "the clipboard changed while the workbook was being read" in n for n in book.notes)


def test_r3_pictures_are_copied_one_reader_at_a_time(monkeypatch):
    copies = []
    book, restored = _careful_read(monkeypatch, [5, 5, 6, 6], copies)
    assert copies == [True]                       # under the lock
    assert len(book.pictures) == 1 and book.pictures[0].data
    assert restored == ["the user's text"]



def test_r3_a_careful_read_does_not_wait_for_another_reading(monkeypatch):
    copies = []
    assert excel_open._CLIPBOARD_LOCK.acquire(timeout=5)
    try:
        book, restored = _careful_read(monkeypatch, [5], copies)
    finally:
        excel_open._CLIPBOARD_LOCK.release()
    assert copies == [] and restored == [] and [p.data for p in book.pictures] == [b""]
    assert any(n.endswith("another reading was using the clipboard") for n in book.notes)


def test_r3_pictures_left_uncopied_by_a_changed_clipboard_say_so_once(tmp_path, monkeypatch):
    v = _vault(tmp_path)
    sheet = _xls_with_picture(v)
    _excel_with_pictures(monkeypatch, sheet)
    seq = iter(range(1, 100))
    monkeypatch.setattr(excel_open, "_clip_seq", lambda: next(seq))     # the user copying
    monkeypatch.setattr(excel_open, "clipboard_safe", lambda: (True, ""))
    stats = _kg(v, reader=_conf_reader(ROWS, {})).seed()
    assert stats["pictures_read"] == 0
    assert stats["notes"] == ["pics.xls: Visual: picture(s) not copied: the clipboard "
                              "changed while the workbook was being read"]
