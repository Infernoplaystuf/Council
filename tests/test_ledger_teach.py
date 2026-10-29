"""Teach from a sample workbook (council_core/ledger_teach.py).

The user's request (2026-10-20): "taking a sample workbook from the folder and
I line up where they match the templates, or say data found in the images
and what they are in the template. First I want an LLM guess to be made, then
I give an answer — hopefully this will help train the Ledger and knowledge
graph to work."

What is pinned: the guess lists what the rules placed, what they left, the
part number, the inspectors and a picture's text, with where each guess came
from; Council's model is asked only where the rules stop, with the user's
past answers as examples; learning writes the aliases, ignores, inspectors'
labels, people, the part-name pattern and the examples file, and ticks the
knowledge graph's rules; the next guess places by rule what the model had to
guess. Synthetic workbooks only; a fake model; no network.
"""
from __future__ import annotations

import json
import sys
import threading
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import Font  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
import ledger_fixtures as lf  # noqa: E402
import source_extract as se  # noqa: E402
from council_core import kg_excel  # noqa: E402
from council_core import knowledge_graph as kgm  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_llm  # noqa: E402
from council_core import ledger_teach as lt  # noqa: E402
from sample_ledger import LedgerError, ParentCSV  # noqa: E402

INSPECTOR = "[an inspector's initials or name]"
NOT_DATA = "[not data: a title, heading or note]"


class FakeModel:
    """Answers Teach's two questions from fixed answers; keeps the prompts."""

    def __init__(self, mapping=None, part="XYZ9AB123-R2"):
        self.mapping, self.part, self.prompts = dict(mapping or {}), part, []

    def __call__(self, prompt, schema):
        self.prompts.append(prompt)
        if "Which part" in prompt:
            return json.dumps({"part": self.part, "evidence": "Serial"}), "fake-granite"
        labels = schema["properties"]["mappings"]["items"]["properties"]["label"]["enum"]
        return json.dumps({"mappings": [{"label": lab, "field": self.mapping.get(lab)}
                                        for lab in labels]}), "fake-granite"

    def label_prompts(self):
        return [p for p in self.prompts if "template field it means" in p]

    def part_prompts(self):
        return [p for p in self.prompts if "Which part" in p]


@pytest.fixture()
def ledger(tmp_path):
    ps = li.create_parent(lf.make_template(tmp_path / "template.xlsx"))
    return Path(ps.csv), ParentCSV(ps.csv).schema


def _field(schema, word):
    return next(f.name for f in schema.fields if word in f.name)


def _unlabelled_book(path: Path) -> Path:
    """No 'Part Number' label, an odd serial, labels no rule knows, an
    inspector's label no rule knows, and a note."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Serial: XYZ9AB123-R2"
    ws["A3"], ws["B3"] = "Power (W)", 280
    ws["A4"], ws["B4"] = "Zorp rate", 55
    ws["A5"], ws["B5"] = "Tech Sig", "JB"
    ws["A6"], ws["B6"] = "Remarks to self", "call MK"
    wb.save(path)
    return path


def _labelled_book(path: Path, with_picture: bool = False) -> Path:
    """A part number the rules read, a label they match, one they do not,
    'Inspector: JB', and a flaw table with a 'Flaw ID' column."""
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Build"
    ws["A1"] = "Part Number: LMM1AB101"
    ws["A2"] = "Inspector: JB"
    ws["A3"], ws["B3"] = "Power (W)", 280
    ws["A4"], ws["B4"] = "Zorp rate", 55
    fl = wb.create_sheet("Flaws")
    fl.append(["Flaw ID", "Passes"])
    for item in lf.ITEMS[:3]:
        fl.append([item, 300])
    if with_picture:
        ws.add_image(XLImage(__import__("io").BytesIO(lf._png(lf.CONCEALMENT_IMAGE_ROWS))), "F2")
    wb.save(path)
    return path


def _row(session, label):
    return next(r for r in session.rows if r.label == label)


# ── the dropdown's choices ──────────────────────────────────────────────
def test_choices_hold_every_role_and_field_but_the_part_numbers(ledger):
    _csv, schema = ledger
    keys = [k for k, _t in lt.choices(schema)]
    assert keys[:5] == [lt.ROLE_PART, lt.ROLE_INSPECTOR, lt.ROLE_ITEM, lt.ROLE_IGNORE,
                        lt.ROLE_UNSURE]
    assert "field:Power (W)" in keys and "field:Part Number" not in keys
    texts = dict(lt.choices(schema))
    assert texts[lt.ROLE_ITEM] == "Flaw # code"
    assert "(per Flaw #)" in texts["field:" + _field(schema, "Passes")]
    for key in keys:
        role, fname = lt.parse_choice(key)
        assert lt.choice_key(role, fname) == key
    assert lt.parse_choice("nonsense") == (lt.ROLE_UNSURE, "")
    assert lt.parse_choice(lt.Guess(lt.ROLE_FIELD, "Power (W)")) == (lt.ROLE_FIELD, "Power (W)")
    assert lt.parse_choice(("inspector", "x")) == (lt.ROLE_INSPECTOR, "")


def test_a_part_numbers_shape_becomes_a_pattern():
    assert se.mask_from_part("ABC1DE234") == "@@@#@@###"
    assert se.mask_from_part("XYZ9AB123-R2") == "@@@#@@###-@#"
    assert se.mask_to_regex(se.mask_from_part("P/N-0042")).fullmatch("Q/R-1234")


# ── the guess ───────────────────────────────────────────────────────────
def test_the_rules_guess_lists_matched_unmatched_part_inspector_item_and_picture(
        tmp_path, ledger):
    csv, schema = ledger
    se.save_people(csv, {"JB": "Jordan Blake"})
    src = _labelled_book(tmp_path / "sample.xlsx", with_picture=True)
    s = lt.guess(csv, src, use_images=True,
                 image_reader=lf.fake_reader(lf.CONCEALMENT_IMAGE_ROWS))
    assert not s.cancelled and s.model == "" and not s.model_picked
    # a label the rules matched: shown, so a wrong match can be corrected
    power = _row(s, "Power (W)")
    assert (power.kind, power.guess.key, power.guess.via) == ("label", "field:Power (W)",
                                                              lt.VIA_RULE)
    assert power.where == "Build!A3" and power.example == "280" and power.guess.score >= 0.85
    # a label nothing placed
    zorp = _row(s, "Zorp rate")
    assert (zorp.guess.role, zorp.guess.via, zorp.example) == (lt.ROLE_UNSURE, lt.VIA_NONE,
                                                               "55")
    # the part number's label and the part
    pn = _row(s, "Part Number")
    assert pn.kind == "part" and pn.guess.role == lt.ROLE_PART and pn.example == "LMM1AB101"
    assert s.part.value == "LMM1AB101" and s.part.via == lt.VIA_RULE
    assert s.part.where == "Build!A1"
    # the inspector's label, and who it is from the people list
    insp = _row(s, "Inspector")
    assert (insp.kind, insp.guess.role, insp.example, insp.initials) == (
        "inspector", lt.ROLE_INSPECTOR, "JB", ["JB"])
    assert [(p.initials, p.full_name) for p in s.inspectors] == [("JB", "Jordan Blake")]
    # the flaw-code column
    flaw = _row(s, "Flaw ID")
    assert flaw.kind == "item" and flaw.guess.role == lt.ROLE_ITEM and flaw.example == "1A"
    passes = _row(s, "Passes")
    assert passes.count == 3 and "Flaw # codes (1A … 3A)" in passes.note
    # what the picture's reader read, with the picture named
    media = next(r for r in s.rows if r.picture and r.label == "Media")
    assert media.picture == "picture 1 on Build"
    assert media.where == "picture 1 on Build, row 1 column 1"
    assert "ceramic" in media.picture_text and media.guess.role == lt.ROLE_FIELD
    assert "read from a picture" in media.note
    assert (s.pictures, s.pictures_read) == (1, 1)
    # keys are stable: the same workbook gives the same keys
    again = lt.guess(csv, src, use_images=True,
                     image_reader=lf.fake_reader(lf.CONCEALMENT_IMAGE_ROWS))
    assert [r.key for r in again.rows] == [r.key for r in s.rows]
    assert len({r.key for r in s.rows}) == len(s.rows)
    # what needs an answer first: the part's, the inspector's, the flaw
    # codes', then the unplaced labels before the sure rule matches
    kinds = [r.kind for r in s.rows]
    assert kinds.index("part") < kinds.index("inspector") < kinds.index("item") \
        < kinds.index("label")
    labels = [r for r in s.rows if r.kind == "label"]
    assert labels.index(zorp) < labels.index(power)
    text = s.describe()
    assert "sample.xlsx" in text and "LMM1AB101" in text and "JB (Jordan Blake)" in text


def test_the_model_guesses_where_the_rules_stop(tmp_path, ledger):
    csv, schema = ledger
    freq = _field(schema, "Frequency")
    model = FakeModel({"Zorp rate": freq, "Tech Sig": INSPECTOR,
                       "Remarks to self": NOT_DATA})
    src = _unlabelled_book(tmp_path / "untitled.xlsx")
    # exact: loose, the rules read 'XYZ9AB123-R2' themselves
    # (test_a_loose_pattern_takes_a_dash_tail)
    s = lt.guess(csv, src, chat=model, use_images=False, pattern_mask="=@@@#@@###")
    assert s.model == "fake-granite"
    zorp, tech, note = _row(s, "Zorp rate"), _row(s, "Tech Sig"), _row(s, "Remarks to self")
    assert (zorp.guess.key, zorp.guess.via, zorp.guess.score) == (
        f"field:{freq}", lt.VIA_MODEL, ledger_llm.MODEL_SCORE)
    assert (tech.guess.role, tech.guess.via, tech.initials) == (lt.ROLE_INSPECTOR,
                                                                lt.VIA_MODEL, ["JB"])
    assert note.guess.role == lt.ROLE_IGNORE
    # what the rules placed is not asked about; only labels are sent
    asked = model.label_prompts()
    assert len(asked) == 1
    assert json.loads(asked[0].split("Labels:\n", 1)[1]) == [
        "Remarks to self", "Serial", "Tech Sig", "Zorp rate"]
    assert "55" not in asked[0] and "call MK" not in asked[0]
    # the roles are offered beside the fields, never the part number's field
    offered = json.loads(asked[0].split("Template fields:\n", 1)[1].split("\n\n")[0])
    assert INSPECTOR in offered and NOT_DATA in offered and "Part Number" not in offered
    # no part number by the rules: the model's, written in the workbook,
    # with the user's pattern as an example - and its label is the part's
    assert s.part.value == "XYZ9AB123-R2" and s.part.via == lt.VIA_MODEL
    assert s.part.label == "Serial" and "check it" in s.part.how
    serial = _row(s, "Serial")
    assert (serial.guess.role, serial.guess.via) == (lt.ROLE_PART, lt.VIA_MODEL)
    assert "'@@@#@@###'" in model.part_prompts()[0]
    # the inspector the model found is in the People section, to be named
    assert [p.initials for p in s.inspectors] == ["JB"]


def test_a_model_that_is_not_us_origin_is_refused_before_reading(tmp_path, ledger,
                                                                 monkeypatch):
    csv, _schema = ledger
    from council_core import local_models
    monkeypatch.setattr(local_models, "maker_and_origin", lambda name: ("Someone", "non-US"))
    with pytest.raises(LedgerError, match="not a US-origin model"):
        lt.guess(csv, _unlabelled_book(tmp_path / "u.xlsx"), model_id="ollama:qwen2.5:7b",
                 use_images=False)


def test_cancel_stops_the_guess(tmp_path, ledger):
    csv, _schema = ledger
    stop = threading.Event()
    stop.set()
    s = lt.guess(csv, _labelled_book(tmp_path / "s.xlsx"), use_images=False, cancel=stop)
    assert s.cancelled and not s.rows and "cancelled" in s.describe()


def test_cancel_while_the_model_is_asked_keeps_the_rules_guesses(tmp_path, ledger):
    csv, _schema = ledger
    stop = threading.Event()
    lines = []

    def chat(prompt, schema):
        stop.set()                                  # pressed while it was asked
        raise RuntimeError("stopped by the caller")
    s = lt.guess(csv, _unlabelled_book(tmp_path / "u.xlsx"), chat=chat, use_images=False,
                 cancel=stop, progress=lambda line, frac=None: lines.append(line))
    assert s.cancelled and s.rows
    assert _row(s, "Zorp rate").guess.via == lt.VIA_NONE
    assert not any("did not answer" in p for p in s.problems)   # Cancel is no failure
    assert "Cancel stopped Council's model" in s.describe()
    assert any("asking Council's model" in line for line in lines)


def test_a_model_that_does_not_answer_is_said_once(tmp_path, ledger):
    csv, _schema = ledger

    def chat(prompt, schema):
        raise ConnectionRefusedError("nothing listens there")
    s = lt.guess(csv, _unlabelled_book(tmp_path / "u.xlsx"), chat=chat, use_images=False)
    assert s.model.endswith("(did not answer)") and not s.cancelled
    assert sum("did not answer" in p for p in s.problems) == 1
    assert _row(s, "Zorp rate").guess.via == lt.VIA_NONE and s.part is None
    assert "part number: not found" in s.describe()


def test_a_long_list_is_capped_with_a_note(tmp_path, ledger):
    csv, _schema = ledger
    s = lt.guess(csv, _labelled_book(tmp_path / "s.xlsx"), use_images=False, max_rows=2)
    assert len(s.rows) == 2 and s.hidden > 0
    assert s.rows[0].kind == "part"
    assert any("not shown" in p for p in s.problems)


def test_an_unknown_workbook_or_a_ledger_without_schema_is_said(tmp_path, ledger):
    csv, _schema = ledger
    with pytest.raises(LedgerError, match="not found"):
        lt.guess(csv, tmp_path / "nowhere.xlsx", use_images=False)
    plain = tmp_path / "plain.csv"
    plain.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(LedgerError, match="no schema"):
        lt.guess(plain, _labelled_book(tmp_path / "s.xlsx"), use_images=False)


# ── learning ────────────────────────────────────────────────────────────
def _taught(tmp_path, ledger, vault=None):
    """Guess the unlabelled book with the model, answer every guess as it
    is, name JB, confirm the part and keep its shape."""
    csv, schema = ledger
    freq = _field(schema, "Frequency")
    model = FakeModel({"Zorp rate": freq, "Tech Sig": INSPECTOR,
                       "Remarks to self": NOT_DATA})
    src = _unlabelled_book(tmp_path / "untitled.xlsx")
    s = lt.guess(csv, src, chat=model, use_images=False)
    rep = lt.learn(csv, s, lt.answers_from_guesses(s), {"JB": "Jordan Blake", "MK": ""},
                   part="XYZ9AB123-R2", save_pattern=True, connections_vault=vault)
    return csv, schema, src, s, rep, model


def test_learning_writes_aliases_ignores_inspectors_people_pattern_and_examples(
        tmp_path, ledger):
    csv, schema, _src, _s, rep, _m = _taught(tmp_path, ledger)
    freq = _field(schema, "Frequency")
    data = json.loads(se.aliases_path_for(csv).read_text(encoding="utf-8"))
    assert data["map"] == {"serial": "Part Number", "zorp rate": freq}
    assert data["ignore"] == ["remarks to self"]
    assert data["who_labels"] == ["Tech Sig"]
    assert data["people"] == {"JB": "Jordan Blake"}            # a blank name is no answer
    assert data["patterns"] == ["@@@#@@###-@#"]
    assert se.load_patterns(csv) == ["@@@#@@###-@#"]
    assert rep.mapped == {"Zorp rate": freq} and rep.who_labels == ["Tech Sig"]
    assert rep.part_labels == ["Serial"] and rep.ignored == ["Remarks to self"]
    assert rep.pattern == "@@@#@@###-@#" and rep.people == {"JB": "Jordan Blake"}
    assert rep.unchanged == 1                                  # Power (W): the rules knew
    # the examples: labels and roles; a value only for the part number
    ex = lt.examples(csv)
    by = {(e.label, e.role) for e in ex}
    assert {("Zorp rate", "field"), ("Tech Sig", "inspector"), ("Serial", "part"),
            ("Remarks to self", "ignore"), ("Power (W)", "field")} <= by
    assert {e.value for e in ex if e.role != lt.ROLE_PART} == {""}
    assert {e.value for e in ex if e.role == lt.ROLE_PART} == {"XYZ9AB123-R2"}
    # the stored text, less each line's time stamp (19:19:55 has a 55 in it)
    raw = "".join(json.dumps({k: v for k, v in json.loads(line).items() if k != "ts"})
                  for line in lt.examples_path_for(csv).read_text(encoding="utf-8")
                  .splitlines() if line.strip())
    assert "call MK" not in raw and "55" not in raw
    assert all(e.workbook == "untitled.xlsx" for e in ex)
    text = rep.describe()
    assert "Zorp rate → " in text and "Tech Sig" in text and "@@@#@@###-@#" in text


def test_the_same_answers_again_add_no_examples(tmp_path, ledger):
    csv, _schema, src, s, rep, _m = _taught(tmp_path, ledger)
    n = len(lt.examples(csv))
    rep2 = lt.learn(csv, s, lt.answers_from_guesses(s))
    assert rep2.examples == 0 and len(lt.examples(csv)) == n
    assert rep.examples == n


def test_a_second_guess_places_by_rule_what_the_model_had_to_guess(tmp_path, ledger):
    csv, schema, src, before, _rep, _m = _taught(tmp_path, ledger)
    model = FakeModel()
    after = lt.guess(csv, src, chat=model, use_images=False)
    zorp, tech = _row(after, "Zorp rate"), _row(after, "Tech Sig")
    assert (zorp.guess.key, zorp.guess.via) == (f"field:{_field(schema, 'Frequency')}",
                                                lt.VIA_ALIAS)
    assert (tech.kind, tech.guess.role, tech.guess.via) == ("inspector", lt.ROLE_INSPECTOR,
                                                            lt.VIA_ALIAS)
    assert _row(after, "Remarks to self").guess.key == lt.ROLE_IGNORE
    # the part number is read by the rules now (the label is the part's)
    assert after.part.value == "XYZ9AB123-R2" and after.part.via == lt.VIA_ALIAS
    assert after.pattern_saved and after.pattern_mask == "@@@#@@###-@#"
    # nothing left for the model: it was not asked
    assert model.prompts == []
    assert "not needed" in after.describe()
    assert [(p.initials, p.full_name) for p in after.inspectors] == [("JB", "Jordan Blake")]
    said = lt.improvement(before, after)
    assert said.startswith("4 thing(s)") and "Zorp rate → " in said
    assert "XYZ9AB123-R2 is now read by the rules" in said and "nothing left unsure" in said


def test_past_answers_go_in_the_models_prompts(tmp_path, ledger):
    csv, schema, _src, _s, _rep, _m = _taught(tmp_path, ledger)
    other = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Note: no id on this one"
    ws["A2"], ws["B2"] = "Zorp rates", 57
    ws["A3"], ws["B3"] = "Blip count", 3
    wb.save(other)
    model = FakeModel(part="")
    lt.guess(csv, other, chat=model, use_images=False)
    asked = model.label_prompts()[0]
    freq = _field(schema, "Frequency")
    assert "The user answered these labels before" in asked
    assert f'"Zorp rate" -> "{freq}"' in asked                 # most like 'Zorp rates'
    assert f'"Tech Sig" -> "{INSPECTOR}"' in asked
    # the part question shows the part numbers the user confirmed
    assert '"XYZ9AB123-R2" (labelled "Serial")' in model.part_prompts()[0]


def test_extracts_model_gets_the_examples_too(tmp_path, ledger):
    csv, schema, _src, _s, _rep, _m = _taught(tmp_path, ledger)
    other = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Blip count", 3
    wb.save(other)
    model = FakeModel(part="")
    li._extract(str(csv), [str(other)], None, "", False, None, _NoProg(), None, None,
                None, True, model)
    asked = model.label_prompts()[0]
    assert f'"Zorp rate" -> "{_field(schema, "Frequency")}"' in asked
    # Extract's question offers fields only: a role's answer is null, said why
    assert '"Tech Sig" -> null (the user said: Inspector (initials or name))' in asked


def test_extract_uses_the_saved_pattern_when_the_field_is_empty(tmp_path, ledger):
    csv, _schema = ledger
    se.save_taught(csv, se.Taught(patterns=["@@@#@@###"]))
    src = lf.make_laser_log(tmp_path / "laser_101.xlsx")
    s = li.run_extract(csv, [str(src)], tmp_path / "review.xlsx", use_images=False)
    assert s.pattern_saved and s.pattern == "@@@#@@###"
    assert "this Ledger's saved one (the pattern field was empty)" in s.describe()
    meta = json.loads(openpyxl.load_workbook(s.review)["_run"]["A1"].value)
    assert meta["pattern_mask"] == "@@@#@@###"                 # Apply reads it the same way
    assert meta["pattern_hint"] is True                        # ... as a hint
    typed = li.run_extract(csv, [str(src)], tmp_path / "review2.xlsx", use_images=False,
                           pattern_mask="###")
    assert not typed.pattern_saved and typed.pattern == "###"
    meta = json.loads(openpyxl.load_workbook(typed.review)["_run"]["A1"].value)
    assert "pattern_hint" not in meta


def _lmm_book(path: Path, title: str = "", label: str = "") -> Path:
    """A workbook of one part: a title naming it, or 'Part Number: ...'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Build"
    if title:
        ws["A1"] = title
    if label:
        ws["A2"] = label
    ws["A3"], ws["B3"] = "Power (W)", 280
    ws["A4"], ws["B4"] = "Scan Speed", 1200
    wb.save(path)
    return path


def test_the_saved_pattern_is_a_hint_and_can_be_forgotten(tmp_path, ledger):
    """A pattern Teach saved from one sample ('XYZ9AB123-R2' ->
    '@@@#@@###-@#') held back every part of another shape from
    pre-acceptance, and nothing could take it back (review, 2026-10-20).
    Saved, it is a hint; typed, it still holds a misfit back."""
    csv, _schema = ledger
    books = [str(lf.make_laser_log(tmp_path / "laser_101.xlsx")),
             str(_lmm_book(tmp_path / "lmm.xlsx", label="Part Number: LMM1AB101"))]
    plain = li.run_extract(csv, books, tmp_path / "r0.xlsx", use_images=False)
    assert plain.pre_accepted > 0
    se.save_taught(csv, se.Taught(patterns=["@@@#@@###-@#"]))
    hinted = li.run_extract(csv, books, tmp_path / "r1.xlsx", use_images=False)
    assert hinted.pattern_saved and hinted.pre_accepted == plain.pre_accepted
    typed = li.run_extract(csv, books, tmp_path / "r2.xlsx", use_images=False,
                           pattern_mask="@@@#@@###-@#")
    assert typed.pre_accepted == 0                             # a typed one still holds back
    # Apply reads the review as Extract did: a hint, not a filter
    run = se.extract(ParentCSV(csv).schema, books, use_images=False,
                     pattern_mask="@@@#@@###-@#", pattern_strict=False)
    assert not any(se.MISFIT_PART_HOW in o.part_how for o in run.obs)
    # Teach's own read takes it as a hint too
    s = lt.guess(csv, books[0], use_images=False)
    assert s.pattern_saved and "a hint" in s.describe()
    assert se.forget_patterns(csv) == ["@@@#@@###-@#"]
    assert se.load_patterns(csv) == [] and se.forget_patterns(csv) == []
    assert se.aliases_path_for(csv).exists()                   # the rest of the file is kept


def test_the_pattern_finds_a_part_written_only_in_a_title(tmp_path, ledger):
    """Teach confirmed 'LMM1AB101' from the title 'LMM1AB101 laser build
    record' and saved its shape, but the pattern was tried on sheet and
    file names only: a rules-only Extract found no part (review,
    2026-10-20). It is looked for in short text cells now - a whole word,
    one part only - and such a part is never pre-accepted."""
    csv, schema = ledger
    a = _lmm_book(tmp_path / "a.xlsx", title="LMM1AB101 laser build record")
    b = _lmm_book(tmp_path / "b.xlsx", title="LMM1AB102 visual inspection")
    s = li.run_extract(csv, [str(a), str(b)], tmp_path / "r0.xlsx", use_images=False)
    assert sum("have no Part Number" in p for p in s.problems) == 2   # no pattern: no part
    # Teach on a.xlsx: the model found the part in the title; the user
    # confirms it and keeps its shape
    model = FakeModel(part="LMM1AB101")
    first = lt.guess(csv, a, chat=model, use_images=False)
    assert first.part.value == "LMM1AB101" and first.part.via == lt.VIA_MODEL
    rep = lt.learn(csv, first, {}, part="LMM1AB101", save_pattern=True)
    assert rep.pattern == "@@@#@@###"
    # the next rules-only Extract finds both parts in their titles
    run = se.extract(schema, [a, b], use_images=False, pattern_mask="@@@#@@###",
                     pattern_strict=False)
    parts = {o.where.split(" › ")[0]: (o.part, o.part_how) for o in run.obs}
    assert parts["a.xlsx"][0] == "LMM1AB101" and parts["b.xlsx"][0] == "LMM1AB102"
    assert parts["a.xlsx"][1] == f"{se.TEXT_PART_HOW} Build!A1 — check it"
    assert not any(d.accept for d in se.decide(run, schema))  # never pre-accepted
    s = li.run_extract(csv, [str(a), str(b)], tmp_path / "r1.xlsx", use_images=False)
    assert not any("have no Part Number" in p for p in s.problems)
    # Guess again: the rules read it now, without the model
    again = lt.guess(csv, a, chat=FakeModel(part=""), use_images=False)
    assert (again.part.value, again.part.via) == ("LMM1AB101", lt.VIA_ALIAS)
    assert again.part.where == "Build!A1" and "your part-name pattern" in again.part.how
    assert "LMM1AB101 is now read by the rules" in lt.improvement(first, again)
    # only a whole word: 'XYZ9AB123' of 'Serial: XYZ9AB123-R2' is no part for
    # the exact shape; the loose one takes it with its tail (user,
    # 2026-10-22: "have a -@ or -#"); two parts in titles: no guess
    def parts(mask):
        return se._Extractor(se.Grid("x", "S", {(1, 1): "Serial: XYZ9AB123-R2"}),
                             se.Matcher(schema), se.ItemCodes(schema.items), "",
                             se.mask_to_regex(mask)).parts_in_text()
    assert parts("=@@@#@@###") == []
    assert parts("@@@#@@###") == [(1, 1, "XYZ9AB123-R2")]
    two = se._Extractor(se.Grid("x", "S", {(1, 1): "LMM1AB101 build", (2, 1): "LMM1AB102"}),
                        se.Matcher(schema), se.ItemCodes(schema.items), "",
                        se.mask_to_regex("@@@#@@###"))
    assert se._text_part([two]) == ("", "")


def test_ignored_labels_are_not_offered_to_map_or_sent_to_the_model(tmp_path, ledger):
    """Labels answered 'Ignore (not data)' were listed on the next review's
    'Labels to map' and sent to Council's model (review, 2026-10-20)."""
    csv, schema = ledger
    book = _unlabelled_book(tmp_path / "u.xlsx")
    s = lt.guess(csv, book, use_images=False)
    rep = lt.learn(csv, s, {_row(s, "Zorp rate").key: lt.ROLE_IGNORE})
    assert rep.ignored == ["Zorp rate"]
    run = se.extract(schema, [book], *se.load_aliases(csv), use_images=False)
    assert run.unmatched and not any(o.label == "Zorp rate" for o in run.unmatched)
    model = FakeModel(part="")
    summary = li._extract(str(csv), [str(book)], tmp_path / "review.xlsx", "", False, None,
                          _NoProg(), None, None, None, True, model)
    asked = model.label_prompts()[0].split("Labels:\n", 1)[1]
    assert "Zorp rate" not in asked and "Remarks to self" in asked
    to_map = openpyxl.load_workbook(summary.review)["Labels to map"]
    assert not any(c.value == "Zorp rate" for row in to_map.iter_rows() for c in row)
    # Teach still shows it, as your earlier answer
    again = _row(lt.guess(csv, book, use_images=False), "Zorp rate")
    assert (again.guess.key, again.guess.via) == (lt.ROLE_IGNORE, lt.VIA_ALIAS)


def test_guesses_left_unconfirmed_are_not_learned(tmp_path, ledger):
    """Save wrote every guess left as it was, a 50%-sure model guess too, as
    an alias that wins outright (review, 2026-10-20): the model's guesses
    and the rules' unsure matches are kept only when confirmed."""
    csv, schema = ledger
    freq = _field(schema, "Frequency")
    model = FakeModel({"Zorp rate": freq, "Tech Sig": INSPECTOR})
    s = lt.guess(csv, _unlabelled_book(tmp_path / "u.xlsx"), chat=model, use_images=False)
    zorp, tech, serial = _row(s, "Zorp rate"), _row(s, "Tech Sig"), _row(s, "Serial")
    power = _row(s, "Power (W)")
    assert lt.needs_confirm(zorp) and lt.needs_confirm(tech) and lt.needs_confirm(serial)
    assert not lt.needs_confirm(power)                         # a sure rule match
    weak = lt.TeachRow(key="k", kind="label", label="Media", where="", sheet="",
                       guess=lt.Guess(lt.ROLE_FIELD, "Concealment 1: Media", lt.VIA_RULE, 0.8))
    assert lt.needs_confirm(weak)
    sure = lt.confirmed_by_default(s)
    assert power.key in sure and zorp.key not in sure
    answers = lt.answers_from_guesses(s)
    rep = lt.learn(csv, s, answers, confirmed=sure)
    assert sorted(rep.unconfirmed) == ["Serial", "Tech Sig", "Zorp rate"]
    assert "3 guess(es) you did not confirm, so not learned" in rep.describe()
    assert se.load_aliases(csv)[0] == {} and se.load_taught(csv).who == []
    assert not any(e.role == lt.ROLE_FIELD and e.label == "Zorp rate" for e in lt.examples(csv))
    # confirmed (the box ticked, or another answer picked): learned
    rep = lt.learn(csv, s, answers, confirmed=sure | {zorp.key, tech.key})
    assert rep.mapped == {"Zorp rate": freq} and rep.who_labels == ["Tech Sig"]
    assert rep.unconfirmed == ["Serial"]


def test_inspector_labels_sign_readings_in_extract(tmp_path, ledger):
    """'Tech Sig' answered as an inspector's label: the next Extract reads
    the initials beside it as who logged the readings under it, and never
    maps the label to a field."""
    csv, schema = ledger
    book = tmp_path / "ec.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part 101"
    ws["A1"], ws["B1"] = "Tech Sig", "MK"
    ws["B3"] = "Flaw"
    ws["B4"] = "EC Pre (%FSH)"
    for i, item in enumerate(lf.ITEMS):
        ws.cell(3, 3 + i, item)
        ws.cell(4, 3 + i, 60 + i)
    wb.save(book)
    before = se.extract(schema, [book], use_images=False)
    assert not any(o.inspector == "MK" for o in before.obs)
    s = lt.guess(csv, book, use_images=False)
    tech = _row(s, "Tech Sig")
    rep = lt.learn(csv, s, {tech.key: lt.ROLE_INSPECTOR})
    assert rep.who_labels == ["Tech Sig"]
    taught = se.load_taught(csv)
    after = se.extract(schema, [book], *se.load_aliases(csv), use_images=False, taught=taught)
    ec = [o for o in after.obs if "Pre-Conceal" in o.field]
    assert len(ec) == 9 and {o.inspector for o in ec} == {"MK"}
    assert not any(o.label == "Tech Sig" for o in after.obs + after.unmatched)
    assert se.Matcher(schema, taught=taught).match("Tech Sig").target is None


def _headed_book(path: Path, heading: str, sheet: str = "Part 101") -> Path:
    """A heading over a small EC table: 'NDT — 2026-09-18', 'Inspector: JB'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = heading
    ws["A2"], ws["B2"] = "Flaw", "EC Pre (%FSH)"
    for i, item in enumerate(lf.ITEMS[:3]):
        ws.cell(3 + i, 1, item)
        ws.cell(3 + i, 2, 60 + i)
    wb.save(path)
    return path


def _signed(csv, schema, book):
    """(reading, who signed it) of the EC readings the next Extract gives."""
    run = se.extract(schema, [book], *se.load_aliases(csv), use_images=False,
                     taught=se.load_taught(csv))
    return sorted((o.value, o.inspector) for o in run.obs
                  if "Pre-Conceal" in o.field and not lt.initials_field(o.field))


def test_a_heading_that_is_no_one_is_not_read_as_initials_again(tmp_path, ledger):
    """'NDT — 2026-09-18' over a table is read as an inspector's heading;
    answered 'Not an inspector', NDT is no one's initials in headings from
    then on. The answer is offered on that row only."""
    csv, schema = ledger
    book = _headed_book(tmp_path / "ndt.xlsx", "NDT — 2026-09-18")
    s = lt.guess(csv, book, use_images=False)
    head = next(r for r in s.rows if r.kind == "inspector")
    assert head.shape == "heading" and head.initials == ["NDT"]
    offered = [k for k, _t in lt.row_choices(s, head)]
    assert offered.index(lt.ROLE_NO_ONE) == offered.index(lt.ROLE_INSPECTOR) + 1
    assert lt.ROLE_NO_ONE not in [k for k, _t in lt.row_choices(s, _row(s, "Flaw"))]
    rep = lt.learn(csv, s, {head.key: lt.ROLE_NO_ONE})
    assert rep.not_who == ["NDT"] and se.load_taught(csv).not_who == ["NDT"]
    assert not rep.ignored                                     # a value, not a label
    assert "names no one, in every workbook" in rep.describe()
    again = lt.guess(csv, book, use_images=False)
    assert not any(r.kind == "inspector" for r in again.rows)
    assert again.inspectors == []
    assert _signed(csv, schema, book) == [("60", ""), ("61", ""), ("62", "")]


def test_ignore_on_a_heading_of_initials_changes_nobodys_readings(tmp_path, ledger):
    """'Ignore (not data)' on a heading 'JB — 2026-09-18' made JB no one in
    every workbook, even beside 'Inspector:' (review, 2026-10-20). Ignore
    now learns nothing from a heading, and says so; 'Not an inspector'
    makes JB no one's in headings only."""
    csv, schema = ledger
    sample = _headed_book(tmp_path / "sample.xlsx", "JB — 2026-09-18")
    labelled = _headed_book(tmp_path / "labelled.xlsx", "Inspector: JB", "Part 102")
    signed = [("60", "JB"), ("61", "JB"), ("62", "JB")]
    assert _signed(csv, schema, labelled) == signed
    s = lt.guess(csv, sample, use_images=False)
    head = next(r for r in s.rows if r.kind == "inspector")
    assert (head.shape, head.initials) == ("heading", ["JB"])
    rep = lt.learn(csv, s, {head.key: lt.ROLE_IGNORE})
    assert not rep.not_who and se.load_taught(csv).not_who == []
    assert not rep.learned and "JB still signed the readings under it" in rep.notes[0]
    assert "'Not an inspector (names no one)'" in rep.describe()
    assert _signed(csv, schema, sample) == signed and _signed(csv, schema, labelled) == signed
    # 'Not an inspector': no one in a heading, still JB beside 'Inspector:'
    rep = lt.learn(csv, s, {head.key: lt.ROLE_NO_ONE})
    assert rep.not_who == ["JB"]
    assert _signed(csv, schema, sample) == [("60", ""), ("61", ""), ("62", "")]
    assert _signed(csv, schema, labelled) == signed
    other = _headed_book(tmp_path / "other.xlsx", "Initials: JB", "Part 103")
    assert _signed(csv, schema, other) == signed
    # answered Inspector again, JB is someone in headings again
    lt.learn(csv, s, {head.key: lt.ROLE_INSPECTOR})
    assert se.load_taught(csv).not_who == []
    assert _signed(csv, schema, sample) == signed


def _initials_column_book(path: Path, who: str = "JB", head: str = "Inspector Initials"
                          ) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part 101"
    for j, h in enumerate(["Flaw ID", "EC Pre (%FSH)", head]):
        ws.cell(1, 1 + j, h)
    for i, it in enumerate(lf.ITEMS[:3]):
        ws.cell(2 + i, 1, it)
        ws.cell(2 + i, 2, 60 + i)
        ws.cell(2 + i, 3, who)
    wb.save(path)
    return path


def test_an_inspector_initials_field_answer_is_the_inspector(tmp_path, ledger):
    """An 'Inspector Initials' column answered as the template's own '... —
    Inspector Initials' field (the dropdown offered it) made 'Inspector
    Initials' no one's for the whole Ledger and put the initials in one
    reading's field (review, 2026-10-20). Those fields are not offered;
    such an answer is 'Inspector'; 'Ignore' on an inspector's column turns
    off nobody's signatures either."""
    csv, schema = ledger
    companion = next(f.name for f in schema.fields if lt.initials_field(f.name))
    assert lt.choice_key(lt.ROLE_FIELD, companion) not in dict(lt.choices(schema))
    sample = _initials_column_book(tmp_path / "sample.xlsx")
    other = _initials_column_book(tmp_path / "other.xlsx", who="MK")
    signed = [("60", "MK"), ("61", "MK"), ("62", "MK")]
    assert _signed(csv, schema, other) == signed
    s = lt.guess(csv, sample, use_images=False)
    col = _row(s, "Inspector Initials")
    assert (col.kind, col.shape, col.guess.role, col.initials) == (
        "inspector", "column", lt.ROLE_INSPECTOR, ["JB"])
    rep = lt.learn(csv, s, {col.key: lt.choice_key(lt.ROLE_FIELD, companion)})
    assert rep.unchanged == 1 and not rep.not_who and not rep.mapped
    taught = se.load_taught(csv)
    assert taught.not_who == [] and se.load_aliases(csv)[0] == {}
    assert _signed(csv, schema, other) == signed
    run = se.extract(schema, [other], *se.load_aliases(csv), use_images=False, taught=taught)
    assert not any(o.field == companion for o in run.obs)        # no one reading's field
    # 'Ignore' on it: not data either way - it still signs, and says so
    rep = lt.learn(csv, s, {col.key: lt.ROLE_IGNORE})
    assert not rep.not_who and not rep.ignored and "still says who signed" in rep.notes[0]
    assert _signed(csv, schema, other) == signed
    # 'Not an inspector' is the answer that makes the label no one's
    rep = lt.learn(csv, s, {col.key: lt.ROLE_NO_ONE})
    assert rep.not_who == ["Inspector Initials"]
    assert _signed(csv, schema, other) == [("60", ""), ("61", ""), ("62", "")]


def test_labels_taught_as_no_ones_never_undo_the_templates_person_columns(tmp_path):
    """kg_excel.read_ledger: a template field whose label says who signed
    ('Inspector Initials') stays a person column whatever a workbook's
    label was taught to be; the fingerprint follows only the inspectors'
    labels (review, 2026-10-20)."""
    tpl = tmp_path / "tpl.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Part Attributes"
    ws["A1"].font = Font(size=16, bold=True)
    for col, label in enumerate(["Part Number", "Power (W)", "Inspector Initials"], start=1):
        ws.cell(2, col, label)
    wb.save(tpl)
    csv = Path(li.create_parent(tpl).csv)
    ParentCSV(csv).add_samples([{"Part Number": "101", "Power (W)": "280",
                                 "Inspector Initials": "JB"}])
    before = kg_excel.read_ledger(csv, "ledger.csv", [])
    fp = kg_excel.ledger_fingerprint(csv)
    assert "Inspector Initials" in before.summary["person_columns"]
    se.save_taught(csv, se.Taught(not_who=["Inspector Initials"]))
    after = kg_excel.read_ledger(csv, "ledger.csv", [])
    assert "Inspector Initials" in after.summary["person_columns"]
    assert after.summary["people"] == ["JB"]
    assert kg_excel.ledger_fingerprint(csv) == fp


@pytest.mark.parametrize("variant", ["title", "colon", "column"])
def test_initials_in_brackets_are_read_and_taught(tmp_path, ledger, variant):
    """The 'initials in parentheses' layout: 'Inspector (JB)' over a table,
    'Inspector: (JB)', a column 'Insp.' of '(JB)'. None signed a reading,
    and Teach's 'Inspector' answer was said to be learned while the next
    Extract still signed nothing (review, 2026-10-20)."""
    csv, schema = ledger
    book = tmp_path / f"{variant}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part 101"
    head = ["Flaw ID", "EC Pre (%FSH)"] + (["Insp."] if variant == "column" else [])
    if variant != "column":
        ws["A1"] = "Inspector (JB)" if variant == "title" else "Inspector: (JB)"
    for j, h in enumerate(head):
        ws.cell(4, 1 + j, h)
    for i, it in enumerate(lf.ITEMS[:3]):
        ws.cell(5 + i, 1, it)
        ws.cell(5 + i, 2, 60 + i)
        if variant == "column":
            ws.cell(5 + i, 3, "(JB)")
    wb.save(book)
    signed = [("60", "JB"), ("61", "JB"), ("62", "JB")]
    s = lt.guess(csv, book, use_images=False)
    if variant != "column":
        # read by the rules: a row to answer, its initials in People
        row = next(r for r in s.rows if r.kind == "inspector")
        assert (row.label, row.shape, row.initials) == ("Inspector", "label", ["JB"])
        assert [p.initials for p in s.inspectors] == ["JB"]
        assert _signed(csv, schema, book) == signed
        return
    insp = _row(s, "Insp.")
    assert insp.example == "(JB)" and lt.initials_in(insp.example) == ["JB"]
    rep = lt.learn(csv, s, {insp.key: lt.ROLE_INSPECTOR}, {"JB": "Jordan Blake"})
    assert rep.who_labels == ["Insp."]
    assert _signed(csv, schema, book) == signed
    again = lt.guess(csv, book, use_images=False)
    col = _row(again, "Insp.")
    assert (col.kind, col.guess.via, col.initials) == ("inspector", lt.VIA_ALIAS, ["JB"])
    assert [(p.initials, p.full_name) for p in again.inspectors] == [("JB", "Jordan Blake")]


def test_bracketed_initials_and_their_look_alikes():
    assert se._initials("(JB)") == "JB" and se._initials("[j.b.]") == "JB"
    assert se._initials("(12)") == "" and se._initials("(N/A)") == ""
    assert lt.initials_in("(JB/MK)") == ["JB", "MK"] and lt.initials_in("(ok)") == []
    assert se._who_paren("Inspector (JB)") == ("Inspector", "JB")
    assert se._who_paren("Tech Sig [M.K.]") == ("Tech Sig", "MK")
    # a column's title, a unit, a word: no one
    for cell in ("Operator (ID)", "Inspector (Name)", "Inspector (mm)", "Speed (IPM)",
                 "Inspector (Level II)", "Notes (see JB)"):
        assert se._who_paren(cell) == ("", ""), cell
    assert kg_excel.split_people("(JB)") == ["JB"]


def test_a_flaw_code_column_is_learned(tmp_path, ledger):
    csv, schema = ledger
    book = tmp_path / "site.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Part Number", "Loc", "Passes"])
    ws.append(["LMM1AB101", "Spot one", 300])
    wb.save(book)
    s = lt.guess(csv, book, use_images=False)
    loc = _row(s, "Loc")
    rep = lt.learn(csv, s, {loc.key: lt.ROLE_ITEM})
    assert rep.item_labels == ["Loc"]
    taught = se.load_taught(csv)
    assert taught.items == ["Loc"]
    m = se.Matcher(schema, taught=taught)
    assert m.is_item_header("Loc") and m.match("Loc").target is None


def test_an_answer_overrides_a_wrong_rule_match(tmp_path, ledger):
    """The rules matched 'Laser Power (W)' to Power (W); the user says it
    is the Frequency: the alias wins on the next guess."""
    csv, schema = ledger
    src = lf.make_laser_log(tmp_path / "laser_101.xlsx")
    s = lt.guess(csv, src, use_images=False)
    lp = _row(s, "Laser Power (W)")
    assert lp.guess.key == "field:Power (W)" and lp.guess.via == lt.VIA_RULE
    freq = _field(schema, "Frequency")
    rep = lt.learn(csv, s, {lp.key: lt.choice_key(lt.ROLE_FIELD, freq)})
    assert rep.mapped == {"Laser Power (W)": freq}
    again = _row(lt.guess(csv, src, use_images=False), "Laser Power (W)")
    assert (again.guess.key, again.guess.via) == (f"field:{freq}", lt.VIA_ALIAS)


def test_two_answers_for_one_label_keep_the_first(tmp_path, ledger):
    csv, schema = ledger
    book = tmp_path / "two.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "One"
    wb.active["A1"], wb.active["B1"] = "Zorp rate", 1
    two = wb.create_sheet("Two")
    two["A1"], two["B1"] = "Zorp rate", 2
    wb.save(book)
    s = lt.guess(csv, book, use_images=False)
    one, other = [r for r in s.rows if r.label == "Zorp rate"]
    freq = _field(schema, "Frequency")
    rep = lt.learn(csv, s, {one.key: f"field:{freq}", other.key: lt.ROLE_IGNORE})
    assert rep.mapped == {"Zorp rate": freq} and not rep.ignored
    assert rep.conflicts and "answered two ways" in rep.conflicts[0]


def test_unsure_and_unknown_fields_teach_nothing(tmp_path, ledger):
    csv, _schema = ledger
    s = lt.guess(csv, _unlabelled_book(tmp_path / "u.xlsx"), use_images=False)
    zorp = _row(s, "Zorp rate")
    rep = lt.learn(csv, s, {zorp.key: lt.ROLE_UNSURE})
    assert rep.unsure == 1 and not rep.learned
    rep = lt.learn(csv, s, {zorp.key: "field:No Such Field"})
    assert "not a field" in rep.problems[0]
    assert not se.aliases_path_for(csv).exists()


def test_a_damaged_aliases_file_is_not_overwritten(tmp_path, ledger):
    csv, _schema = ledger
    se.aliases_path_for(csv).write_text("{not json", encoding="utf-8")
    session = lt.TeachSession(csv=str(csv), workbook="x.xlsx", rows=[lt.TeachRow(
        key="S|zorp rate", kind="label", label="Zorp rate", where="S!A1", sheet="S")])
    with pytest.raises(LedgerError, match="could not be read"):
        lt.learn(csv, session, {"S|zorp rate": lt.ROLE_IGNORE})
    assert se.aliases_path_for(csv).read_text(encoding="utf-8") == "{not json"
    with pytest.raises(LedgerError):
        se.save_taught(csv, se.Taught(who=["Tech Sig"]))


# ── the knowledge graph ─────────────────────────────────────────────────
def test_the_graphs_part_rules_are_ticked(tmp_path, ledger):
    vault = tmp_path / "vault"
    (vault / "data_in").mkdir(parents=True)
    _csv, _schema, _src, _s, rep, _m = _taught(tmp_path, ledger, vault=vault)
    with kgm.KnowledgeGraph(vault) as kg:
        rules = {r.label: (r, st) for r, st in kg.field_rules()}
    assert rules["Serial"][1] == "confirmed" and rules["Serial"][0].type == "PART"
    assert rules["Part Number"][1] == "confirmed"
    assert any("added 'Serial' as another name for 'Part Number'" in k for k in rep.kg)
    # an existing rule is ticked, not added again
    book = tmp_path / "sn.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "S/N: ABC1DE234"
    wb.save(book)
    csv = _csv
    s = lt.guess(csv, book, use_images=False)
    sn = _row(s, "S/N")
    assert sn.kind == "part" and sn.guess.via == lt.VIA_RULE
    # the default rules start ticked (2026-10-21): one the user unticked
    with kgm.KnowledgeGraph(vault) as kg:
        kg.set_rule_status("S/N", "proposed")
    rep2 = lt.learn(csv, s, {sn.key: lt.ROLE_PART}, connections_vault=vault)
    assert rep2.unchanged == 1 and rep2.part_labels == ["S/N"]
    with kgm.KnowledgeGraph(vault) as kg:
        assert dict((r.label, st) for r, st in kg.field_rules())["S/N"] == "confirmed"
    assert any("ticked 'S/N'" in k for k in rep2.kg)


def test_the_templates_own_part_rule_is_ticked_once(tmp_path, ledger):
    """'Part Number' answered as the part number is the graph's own rule:
    ticked, and said, once (the Ledger log said it twice). The default rules
    start ticked (2026-10-21): here the user had unticked it; ticked
    already, nothing is said (below)."""
    csv, _schema = ledger
    vault = tmp_path / "vault"
    (vault / "data_in").mkdir(parents=True)
    with kgm.KnowledgeGraph(vault) as kg:
        kg.set_rule_status("Part Number", "proposed")
    s = lt.guess(csv, _labelled_book(tmp_path / "s.xlsx"), use_images=False)
    pn = _row(s, "Part Number")
    rep = lt.learn(csv, s, {pn.key: lt.ROLE_PART}, connections_vault=vault)
    assert rep.kg == ["ticked 'Part Number' (a part number)"]
    with kgm.KnowledgeGraph(vault) as kg:
        assert dict((r.label, st) for r, st in kg.field_rules())["Part Number"] == "confirmed"
    again = lt.learn(csv, s, {pn.key: lt.ROLE_PART}, connections_vault=vault)
    assert again.kg == []


def test_a_field_named_like_a_confirmed_rule_gets_the_label_as_a_synonym(tmp_path):
    """A template field 'Project' (a graph rule the user ticked): a label
    mapped to it is added as that rule's other name."""
    tpl = tmp_path / "tpl.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Part Attributes"
    ws["A1"].font = Font(size=16, bold=True)
    for col, label in enumerate(["Part Number", "Project", "Power (W)"], start=1):
        ws.cell(2, col, label)
    wb.save(tpl)
    csv = Path(li.create_parent(tpl).csv)
    vault = tmp_path / "vault"
    with kgm.KnowledgeGraph(vault) as kg:
        kg.set_rule_status("Project", "confirmed")
    book = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "Prog. Code", "ORION"
    wb.save(book)
    s = lt.guess(csv, book, use_images=False)
    row = _row(s, "Prog. Code")
    rep = lt.learn(csv, s, {row.key: "field:Project"}, connections_vault=vault)
    assert rep.mapped == {"Prog. Code": "Project"}
    with kgm.KnowledgeGraph(vault) as kg:
        rules = {r.label: (r, st) for r, st in kg.field_rules()}
    assert rules["Prog. Code"] == (rules["Project"][0].renamed("Prog. Code"), "confirmed")


def test_a_busy_graph_is_said_and_the_rest_is_kept(tmp_path, ledger):
    class Busy:
        def open_graph(self):
            raise kgm.KnowledgeGraphBusy("being rebuilt")
    _csv, _schema, _src, _s, rep, _m = _taught(tmp_path, ledger, vault=Busy())
    assert rep.kg and "being rebuilt elsewhere" in rep.kg[0]
    assert rep.mapped and rep.who_labels                        # the answers are kept


def test_the_ledgers_reading_follows_its_inspector_labels(tmp_path):
    """kg_excel.read_ledger: a template field whose label the Ledger's Teach
    learned is an inspector's is a person column, and the fingerprint
    changes so the graph reads the Ledger again."""
    tpl = tmp_path / "tpl.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Part Attributes"
    ws["A1"].font = Font(size=16, bold=True)
    for col, label in enumerate(["Part Number", "Power (W)", "Tech Sig"], start=1):
        ws.cell(2, col, label)
    wb.save(tpl)
    csv = Path(li.create_parent(tpl).csv)
    ParentCSV(csv).add_samples([{"Part Number": "101", "Power (W)": "280",
                                 "Tech Sig": "JB"}])
    before = kg_excel.read_ledger(csv, "ledger.csv", [])
    fp = kg_excel.ledger_fingerprint(csv)
    assert "Tech Sig" not in before.summary["person_columns"]
    se.save_taught(csv, se.Taught(who=["Tech Sig"]))
    after = kg_excel.read_ledger(csv, "ledger.csv", [])
    assert "Tech Sig" in after.summary["person_columns"]
    assert after.summary["people"] == ["JB"]
    assert kg_excel.ledger_fingerprint(csv) != fp


# ── ledger_llm's examples ───────────────────────────────────────────────
def test_examples_most_like_the_labels_come_first():
    ex = [("Zorp rate", "Frequency (kHz)", ""), ("Power", "Power (W)", ""),
          ("Tech Sig", None, "inspector"), ("Zorp rate", "Date", "")]   # newest wins
    got = ledger_llm.pick_examples(ex, ["Zorp rates", "Blip"], limit=2)
    assert got[0] == ("Zorp rate", "Date", "") and len(got) == 2
    assert ledger_llm.pick_examples(ex, ["Power"], limit=1) == [("Power", "Power (W)", "")]
    text = ledger_llm._examples_text(ex, ["Power (W)"])
    assert '"Power" -> "Power (W)"' in text and '"Tech Sig" -> null (inspector)' in text
    assert "Frequency" not in text                              # not a field offered


class _NoProg:
    def __getattr__(self, _name):
        return lambda *a, **k: None
