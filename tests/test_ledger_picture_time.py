"""A workbook of many pictures no longer holds the Ledger's Extract (user,
2026-10-30: "the ledger still has a stall in the extract -> review
workbook"; "an excel file with many images in it ... at least 10 minutes";
"the issue is image grab and multi thread").

- the pictures of one file have a time budget; past it the rest are saved
  and listed on the Images sheet, not read, and the file says so once;
- a photo with no plain page or screen in it is not given to the reader;
- a camera-size picture is made smaller before Tesseract reads it;
- each helper's Tesseract keeps to one core, and the file the Ledger waits
  for says which picture it is on.
"""
from __future__ import annotations

import io
import sys
import time
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
np = pytest.importorskip("numpy")
from PIL import Image, ImageDraw  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import ledger_fixtures as lf  # noqa: E402
import source_extract as se  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_workers as lw  # noqa: E402
from sample_ledger import ParentCSV  # noqa: E402


@pytest.fixture(autouse=True)
def app_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("COUNCIL_APP_DIR", str(tmp_path / "app"))
    monkeypatch.delenv("COUNCIL_LEDGER_PICTURE_BUDGET_S", raising=False)
    monkeypatch.delenv("COUNCIL_LEDGER_READ_PHOTOS", raising=False)


@pytest.fixture()
def schema(tmp_path):
    tpl = lf.make_template(tmp_path / "template.xlsx")
    return ParentCSV(li.create_parent(tpl).csv).schema


def _jpeg(img) -> bytes:
    buf = io.BytesIO()
    img.save(buf, "JPEG")
    return buf.getvalue()


def _photo(w=1600, h=1200, seed=0) -> bytes:
    """Shades and noise everywhere, nothing like a page: a camera's photo."""
    rng = np.random.default_rng(seed)
    y, x = np.mgrid[0:h, 0:w]
    a = (127 + 60 * np.sin(x / 90.0) + 50 * np.cos(y / 70.0)
         + rng.normal(0, 25, (h, w))).clip(0, 255).astype("uint8")
    return _jpeg(Image.fromarray(np.stack([a, (a * 0.8).astype("uint8"), 255 - a], -1)))


def _page(w=1600, h=1200) -> bytes:
    """A page with lines of text on a desk, photographed."""
    img = Image.new("L", (w, h), 90)
    d = ImageDraw.Draw(img)
    d.rectangle([w // 5, h // 12, w * 4 // 5, h * 11 // 12], fill=240)
    for r in range(20):
        d.text((w // 5 + 40, h // 12 + 30 + r * 45), "Flaw 1A  Response 0.245 mm  JB", fill=20)
    return _jpeg(img)


def _book(path: Path, pictures) -> Path:
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"], ws["B1"] = "Part Number", "2041"
    for i, data in enumerate(pictures):
        ws.add_image(XLImage(io.BytesIO(data)), f"D{2 + 30 * i}")
    wb.save(path)
    return path


class _Reader:
    """A picture reader that takes its time and counts what it was given."""
    name = "slow test reader"

    def __init__(self, seconds=0.0):
        self.seconds, self.calls = seconds, 0

    def __call__(self, data):
        self.calls += 1
        time.sleep(self.seconds)
        return [["Power (W)", "255"]]


# ---------------------------------------------- photos

def test_a_photo_is_told_from_a_page_screenshot_or_small_picture():
    assert se.looks_like_photo(_photo())
    assert not se.looks_like_photo(_page())
    assert not se.looks_like_photo(_photo(300, 200))          # small: read
    assert not se.looks_like_photo(b"not a picture at all")   # unsure: read
    shot = Image.new("RGB", (1280, 800), (255, 255, 255))
    ImageDraw.Draw(shot).text((40, 40), "Laser Power (W) 255", fill=(0, 0, 0))
    assert not se.looks_like_photo(_jpeg(shot))


def test_telling_a_photo_takes_milliseconds_on_a_camera_picture():
    big = _photo(4000, 3000)
    t0 = time.perf_counter()
    assert se.looks_like_photo(big)
    assert time.perf_counter() - t0 < 1.0


def test_photos_are_not_given_to_the_reader_and_count_as_no_text(tmp_path, schema):
    src = _book(tmp_path / "photos.xlsx", [_photo(seed=k) for k in range(3)] + [_page()])
    reader = _Reader()
    run = se.extract(schema, [src], image_reader=reader, use_images=True,
                     image_dir=tmp_path / "imgs")
    assert reader.calls == 1                                   # only the page
    notes = sorted(n.text for n in run.images)
    assert sum(t == se.PHOTO_NO_TEXT for t in notes) == 3
    assert li._images_blank(run.images) == 3 and li._images_read(run.images) == 1
    assert all(n.saved_as for n in run.images)                 # every one saved


def test_every_picture_is_read_when_photos_are_asked_for(tmp_path, schema, monkeypatch):
    monkeypatch.setenv("COUNCIL_LEDGER_READ_PHOTOS", "1")
    src = _book(tmp_path / "photos.xlsx", [_photo(seed=k) for k in range(2)])
    reader = _Reader()
    se.extract(schema, [src], image_reader=reader, use_images=True)
    assert reader.calls == 2


# ---------------------------------------------- the time budget

def test_a_files_pictures_stop_being_read_past_their_time(tmp_path, schema, monkeypatch):
    monkeypatch.setenv("COUNCIL_LEDGER_PICTURE_BUDGET_S", "0.5")
    src = _book(tmp_path / "many.xlsx", [_page(800, 600) for _ in range(12)])
    reader = _Reader(0.2)
    t0 = time.perf_counter()
    run = se.extract(schema, [src], image_reader=reader, use_images=True,
                     image_dir=tmp_path / "imgs")
    took = time.perf_counter() - t0
    assert 2 <= reader.calls <= 5 and took < 12 * 0.2
    late = [n for n in run.images if n.text.startswith(se.PICTURES_OUT_OF_TIME)]
    assert len(late) == 12 - reader.calls and all(n.saved_as for n in late)
    line = next(p for p in run.problems if "pictures not read" in p)
    assert line.startswith(f"many.xlsx: {len(late)} of 12 pictures not read")
    assert li._images_read(run.images) == reader.calls
    assert "over the time a file's pictures may take" in li._images_unread(run.images)


def test_no_budget_reads_every_picture(tmp_path, schema, monkeypatch):
    monkeypatch.setenv("COUNCIL_LEDGER_PICTURE_BUDGET_S", "0")
    src = _book(tmp_path / "many.xlsx", [_page(800, 600) for _ in range(4)])
    reader = _Reader(0.2)
    run = se.extract(schema, [src], image_reader=reader, use_images=True)
    assert reader.calls == 4 and not any("pictures not read" in p for p in run.problems)


def test_a_camera_size_picture_is_made_smaller_for_tesseract():
    out = Image.open(io.BytesIO(se._for_ocr(_page(4000, 3000))))
    assert max(out.size) == se.OCR_MAX_SIDE


# ---------------------------------------------- the helpers

def test_each_helpers_tesseract_keeps_to_one_core(monkeypatch):
    seen = {}

    class Proc:
        pid = 1

        def poll(self):
            return 1

        def kill(self):
            pass

    def popen(cmd, env=None, **_k):
        seen.update(env)
        return Proc()
    monkeypatch.delenv("OMP_THREAD_LIMIT", raising=False)
    monkeypatch.setattr(lw.subprocess, "Popen", popen)
    pool = lw.Pool({}, [("a.xlsx", "a.xlsx")], 1, start_timeout=0.2)
    pool.close()
    assert seen["OMP_THREAD_LIMIT"] == "1"


def test_the_file_waited_for_says_which_picture_it_is_on(monkeypatch):
    """A helper's 'picture 3 of 40' reaches the Ledger's progress while it
    waits for that file - and another file's does not."""
    pool = lw.Pool.__new__(lw.Pool)
    pool.jobs = [("a.xlsx", "a"), ("b.xlsx", "b")]
    pool.results, pool.steps = {}, {}
    pool._next, pool._want = 2, 0
    import queue

    class H:
        alive, job = True, 0

        def stop(self):
            pass
    h = H()
    pool.helpers = [h]
    pool._inbox = queue.Queue()
    for msg in ({"beat": 1, "what": "b.xlsx: picture 1 of 9"},
                {"beat": 0, "what": "a.xlsx: picture 3 of 40"}, {"beat": 0},
                {"done": 0, "part": {"x": 1}, "steps": []}):
        pool._inbox.put((h, msg))
    said = []
    assert pool.result(0, say=said.append) == {"x": 1}
    assert said == ["a.xlsx: picture 3 of 40"]


def test_a_helper_sends_which_picture_it_is_on(tmp_path, schema):
    """_job's beat: a plain beat just before must not swallow the line."""
    src = _book(tmp_path / "pics.xlsx", [_page(800, 600) for _ in range(3)])
    setup = lw._Setup(lw.config_json(schema, {}, [], None, [], "", "", True, "", "none"))
    sent = []
    reply = lw._job(setup, str(src), src.name, sent.append, 7)
    assert reply["done"] == 7
    said = [m["what"] for m in sent if m.get("beat") == 7 and m.get("what")]
    assert "pics.xlsx: picture 1 of 3" in said
