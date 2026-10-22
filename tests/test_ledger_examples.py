"""Learning by example from filled copies of the template (and the records
already accepted), council_core.ledger_examples - user 2026-10-21: "Lets do
a learn by example from filled copies addition"."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
import sample_ledger as sl  # noqa: E402
import source_extract as se  # noqa: E402
from council_core import ledger_examples as lx  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_records as lr  # noqa: E402
from ledger_fixtures import make_template  # noqa: E402
from test_ledger_records import book, creature_rows, table  # noqa: E402
from test_ledger_records import make_template as btb_template  # noqa: E402


def test_values_compare_by_what_they_are():
    assert lx.value_key(6300) == lx.value_key(6300.0) == lx.value_key("6,300")
    assert lx.value_key("Greg #1") == lx.value_key("greg  #1")
    assert lx.distinctive(1200) and lx.distinctive("Greg #1")
    assert not any(lx.distinctive(v) for v in (0, 1, "Yes", "OK", "M", "", None))


# ── the sample Ledger ───────────────────────────────────────────────────
@pytest.fixture
def lmm(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    tpl = make_template(tmp_path / "LMM Template.xlsx")
    csv = Path(li.create_parent(tpl, tmp_path / "lmm.csv").csv)
    schema = sl.TemplateSchema.load(sl.schema_path_for(csv))
    return tmp_path, src, csv, schema


def filled_copy(schema, path, part, values, sheet_title=None):
    """A copy of the template filled in by hand (no Council mark)."""
    rows = [dict(values, **{schema.id_field: part, schema.item_field: it})
            for it in schema.items]
    sl.fill_template(schema, rows, path)
    wb = openpyxl.load_workbook(path)
    wb.properties.keywords = ""
    if sheet_title:
        wb.active.title = sheet_title
    wb.save(path)
    return path


def messy(path, part, labels):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Run"
    ws["A1"], ws["B1"] = "P/N", part
    for i, (label, value) in enumerate(labels, start=3):
        ws.cell(i, 1, label)
        ws.cell(i, 2, value)
    wb.save(path)
    return path


def test_a_filled_copy_is_told_by_the_template_s_words_in_place(lmm):
    tmp, src, csv, schema = lmm
    layout = lx.template_layout(schema.template_path)
    copy = filled_copy(schema, src / "f.xlsx", "101", {"Power (W)": "285"})
    cells = {k: v for g in se.load_workbook_grids(copy)[0] for k, v in g.cells.items()}
    assert lx.is_filled_copy(cells, layout) == "Sheet"
    other = messy(src / "m.xlsx", 101, [("Traverse rate", 1200)])
    cells = se.load_workbook_grids(other)[0][0].cells
    assert lx.is_filled_copy(cells, layout) is None
    rows = lx.filled_rows(next(g for g in se.load_workbook_grids(copy)[0]).cells, schema,
                          "Sheet")
    assert len(rows) == 9 and rows[0][schema.id_field] == "101"
    assert rows[0]["Power (W)"] == "285"


def test_a_filled_copy_teaches_the_labels_of_every_messy_workbook(lmm):
    tmp, src, csv, schema = lmm
    filled_copy(schema, src / "filled_101.xlsx", "101",
                {"Power (W)": "285", "Scan Speed (mm/s)": "1200", "Frequency (kHz)": "50"})
    # part 101's messy log, beside its filled copy: where the labels are learned
    messy(src / "messy_101.xlsx", 101, [("Traverse rate", 1200), ("Zap level", 285),
                                       ("Operator note", "fine")])
    # part 202 has no filled copy: what was learned reads it
    messy(src / "messy_202.xlsx", 202, [("Traverse rate", 900), ("Zap level", 310)])
    ex = li.run_extract(csv, [src], use_images=False)
    assert any("'traverse rate' is Scan Speed (mm/s)" in x for x in ex.learned)
    wb = openpyxl.load_workbook(ex.review)
    rows = {r[3]: r for r in wb["Values"].iter_rows(min_row=2, values_only=True)
            if "messy_202" in str(r[8])}
    speed = rows["Scan Speed (mm/s)"]
    assert (speed[1], speed[4]) == ("202", "900") and rows["Power (W)"][4] == "310"
    assert speed[0] is None                                # a suggestion: amber
    assert "learned from your filled copies (filled_101.xlsx › Sheet)" in speed[10]
    labels = [r[2] for r in wb["Labels to map"].iter_rows(min_row=2, values_only=True)]
    assert "Operator note" in labels and "Traverse rate" not in labels
    # part 101's same values are the filled copy's own rows: nothing is
    # taught by them - only a Y on a suggestion teaches the label
    for row in wb["Values"].iter_rows(min_row=2):
        if "messy_202" in str(row[8].value) and row[3].value == "Scan Speed (mm/s)":
            row[0].value = "Y"
    wb.save(ex.review)
    li.run_apply(ex.review, use_images=False)
    taught = se.load_aliases(csv)[0]
    assert taught["traverse rate"] == "Scan Speed (mm/s)" and "zap level" not in taught


def test_the_parent_csv_is_a_worked_answer_too(lmm):
    tmp, src, csv, schema = lmm
    ledger = sl.ParentCSV(csv)
    ledger.add_samples([{schema.id_field: "101", schema.item_field: it, "Power (W)": "285"}
                        for it in schema.items], source="earlier.xlsx")
    ledger.save()
    messy(src / "messy_101.xlsx", 101, [("Zap level", 285)])
    ex = li.run_extract(csv, [src], use_images=False)
    assert ex.learned == ["Learned by example: 'zap level' is Power (W) (1 of 1 values agree)"]


def test_coincidences_teach_nothing(lmm):
    tmp, src, csv, schema = lmm
    # two fields hold 50: which one is 'Knob'? - neither
    filled_copy(schema, src / "f.xlsx", "101", {"Frequency (kHz)": "50", "Power (W)": "50",
                                                "Scan Speed (mm/s)": "1"})
    messy(src / "m.xlsx", 101, [("Knob", 50), ("Flag", 1)])  # 1 says nothing
    messy(src / "other_part.xlsx", 202, [("Dial", 50)])       # no answer for part 202
    ex = li.run_extract(csv, [src], use_images=False)
    assert ex.learned == []


# ── the table Ledger ────────────────────────────────────────────────────
HEADS = ["Tame", "Name", "Health", "Stam", "Oxygen", "Food", "Weight", "Melee", "Gender",
         "Saddle BP"]


@pytest.fixture
def btb(tmp_path):
    csv = tmp_path / "T.csv"
    lr.create_ledger(btb_template(tmp_path / "T.xlsx"), csv)
    src = tmp_path / "s"
    src.mkdir()
    greg = creature_rows("Greg", 6)
    book(src / "filled.xlsx", {"Copy": table(1, 1, HEADS, [
        ["Spider", r["Name"], r["Health"], r["Stam"], None, None, r["Weight"], r["Melee"],
         r["Gender"], "Greg"] for r in greg[:4]])})
    return csv, src, greg


def test_a_messy_table_is_read_by_its_filled_copy(btb):
    csv, src, greg = btb
    book(src / "messy.xlsx", {"Log": table(1, 1, ["Creature", "HP", "Mass", "Sex", "Remarks"],
                                           [[r["Name"], r["Health"], r["Weight"],
                                             r["Gender"], "ok"] for r in greg])})
    run = lr.extract(csv, sorted(src.glob("*.xlsx")))
    assert run.learned == ["Learned by example in messy.xlsx › Log: 'Creature' is Name, "
                           "'HP' is Health, 'Mass' is Weight, 'Sex' is Gender"]
    recs = [r for r in run.records if r.file == "messy.xlsx"]
    assert len(recs) == 6                                   # Greg #5, #6 are not in the copy
    last = recs[-1]
    assert (last.values["Name"].value, last.values["Health"].value) == ("Greg #6", 7700.0)
    assert last.values["Health"].how == "example" and last.accept == ""
    assert "learned from your filled copies" in run.decisions["C|Health|hp"].detail
    # a Y teaches it: the next Extract reads the column by its header
    out = csv.with_name("r.xlsx")
    lr.write_review(out, run, csv, [])
    wb = openpyxl.load_workbook(out)
    for row in wb["Decisions"].iter_rows(min_row=2):
        if str(row[0].value).startswith("C|"):
            row[1].value = "Y"
    wb.save(out)
    lr.apply_review(out)
    again = [r for r in lr.extract(csv, sorted(src.glob("*.xlsx"))).records
             if r.file == "messy.xlsx"]
    assert again[-1].values["Health"].how == "taught" and again[-1].accept == "Y"   # Greg #6
    # Greg #1 adds nothing to its filled copy's record: left out, said
    assert again[0].accept == "N" and "duplicate" in again[0].flags


def test_a_column_that_agrees_once_in_four_teaches_nothing(btb):
    csv, src, greg = btb
    book(src / "messy.xlsx", {"Log": table(1, 1, ["Name", "Score"], [
        [r["Name"], r["Health"] if i == 0 else 1 + i] for i, r in enumerate(greg[:4])])})
    run = lr.extract(csv, sorted(src.glob("*.xlsx")))
    assert run.learned == []
