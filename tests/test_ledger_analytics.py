"""
Analysis of what a Ledger holds (council_core.ledger_analytics) and its
Analyze window - user, 2026-10-30: "add a data analytics function for
example the ledger output file we can do some analysis on it"; a window
inside the Ledger, no inputs / results split, no spec limits.

Run:  python -m pytest tests/test_ledger_analytics.py -q
"""
from __future__ import annotations

import os
import random
import sys
from pathlib import Path

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).parent))

from council_core import ledger_analytics as la  # noqa: E402
from council_core import ledger_groups as lg  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from ledger_fixtures import make_template  # noqa: E402
from sample_ledger import FILLED_MARK, LedgerError, ParentCSV  # noqa: E402

FOD = "Flaw Attributes: Laser Team FOD (um)"
LENGTH = "Flaw Attributes: Flaw Length (in)"
MEDIA = "Concealment 1: Media"
JB = FOD + " [JB]"


@pytest.fixture
def app_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("COUNCIL_APP_DIR", str(tmp_path / "app"))
    return tmp_path / "app"


@pytest.fixture
def csv(tmp_path, app_dir):
    """20 parts S-1..S-20, 3 flaws each: Power per part, FOD per flaw (S-7's
    flaws all 500 um, a misread), a date per part, Media on even parts."""
    path = li.create_parent(make_template(tmp_path / "T.xlsx"), tmp_path / "l.csv").csv
    rnd = random.Random(4)
    rows = []
    for p in range(1, 21):
        for item in ("1A", "2A", "3A"):
            rows.append({"Part Number": f"S-{p}", "Flaw #": item, "Power (W)": str(300 + p),
                         "Date": f"2025-{(p - 1) // 2 + 1:02d}-15",
                         FOD: "500" if p == 7 else f"{rnd.gauss(50, 3):.1f}",
                         MEDIA: "ceramic" if p % 2 == 0 else "",
                         "source_file": f"book{p}.xlsx"})
    ParentCSV(path).add_samples(rows)
    return Path(path)


def test_fields_get_their_kind_unit_and_level(csv):
    d = la.load(csv)
    assert d.fields["Power (W)"].kind == la.NUMBER and d.fields["Power (W)"].unit == "W"
    assert d.fields[FOD].kind == la.NUMBER and d.fields[FOD].per_item
    assert not d.fields["Power (W)"].per_item
    assert d.fields["Date"].kind == la.DATE
    assert d.fields[MEDIA].kind == la.TEXT
    assert d.fields[LENGTH].kind == la.EMPTY
    assert d.parts[:3] == ["S-1", "S-2", "S-3"] and len(d.parts) == 20
    assert la.to_number("1,200") == 1200 and la.to_number("12.5 um") == 12.5
    assert la.to_number("n/a") is None and la.to_number("") is None
    assert la.to_number("12-34") is None


def test_coverage_counts_part_fields_per_part_and_flaw_fields_per_flaw(csv):
    cov = la.coverage(la.load(csv)).set_index("Field")
    assert (cov.loc["Power (W)", "Filled"], cov.loc["Power (W)", "Could hold"]) == (20, 20)
    assert (cov.loc[FOD, "Filled"], cov.loc[FOD, "Could hold"]) == (60, 60)
    assert cov.loc[MEDIA, "Filled %"] == 50.0
    assert cov.loc[LENGTH, "Filled"] == 0
    missing = la.missing_by_part(la.load(csv))
    # odd parts have no Media, so they are blank on one more field
    assert missing.iloc[0]["Part"] == "S-1"
    assert missing.iloc[0]["Fields blank"] == missing.iloc[-1]["Fields blank"] + 1


def test_summary_counts_a_part_level_field_once_per_part(csv):
    d = la.load(csv)
    s = la.summary(d).set_index("Field")
    assert s.loc["Power (W)", "Count"] == 20                # not 60: one per part
    assert s.loc["Power (W)", "Min"] == 301 and s.loc["Power (W)", "Max"] == 320
    assert s.loc[FOD, "Count"] == 60 and s.loc[FOD, "Max"] == 500
    assert s.loc["Date", "Min"] == "2025-01-15" and s.loc["Date", "Max"] == "2025-10-15"
    assert s.loc[MEDIA, "Most common"] == "ceramic (10)"
    by_part = la.summary(d, per=la.PER_PART).set_index("Field")
    assert by_part.loc[FOD, "Count"] == 20                  # each part's average


def test_trend_has_control_limits_and_flags_the_misread_part(csv):
    t = la.trend(la.load(csv), FOD, by="Date")
    assert len(t.points) == 20 and t.x_label == "Date"
    assert t.lcl < t.center < t.ucl
    assert list(t.points["X"]) == sorted(t.points["X"])
    flagged = t.flagged
    assert "S-7" in set(flagged["Part"])
    assert "1" in flagged.set_index("Part").loc["S-7", "Rules"]
    by_order = la.trend(la.load(csv), "Power (W)", per=la.PER_ROW)
    # a part-level field once per part even per flaw
    assert list(by_order.points["Part"])[:3] == ["S-1", "S-2", "S-3"]
    assert len(by_order.points) == 20
    with pytest.raises(LedgerError):
        la.trend(la.load(csv), "nope")


def test_outliers_find_the_misread_values_and_say_where_they_came_from(csv):
    out = la.outliers(la.load(csv))
    assert set(out["Part"]) == {"S-7"} and set(out["Field"]) == {FOD}
    assert sorted(out["Item"]) == ["1A", "2A", "3A"]
    assert set(out["Source"]) == {"book7.xlsx"}
    per_part = la.outliers(la.load(csv), per=la.PER_PART)
    assert list(per_part["Part"]) == ["S-7"] and per_part.iloc[0]["Item"] == ""


@pytest.fixture
def related(csv):
    """Scan Speed = 1000 - 2 x Power on every part (a part-level pair); Flaw
    Length follows FOD per flaw (S-7's 500 um misread included)."""
    ledger = ParentCSV(csv)
    ledger.columns.insert(ledger.columns.index(FOD) + 1, JB)
    for r in ledger.rows:
        p = int(r["Part Number"].split("-")[1])
        r["Scan Speed (mm/s)"] = str(1000 - 2 * (300 + p))
        r[LENGTH] = f"{float(r[FOD]) / 100 + 0.01 * (p % 3):.3f}"
        r[JB] = f"{float(r[FOD]) + 1:.1f}" if p <= 10 else ""   # one inspector's reading
    ledger.save()
    return csv


def test_the_correlation_matrix_finds_the_pairs_that_move_together(related):
    """user, 2026-10-30: "we need a correlation matrix so that we can better
    determine which variables correlate to each other"."""
    d = la.load(related)
    c = la.correlation(d)
    assert set(c.fields) == {"Power (W)", "Scan Speed (mm/s)", FOD, LENGTH, JB}
    # an inspector's own reading is per flaw, with its field's unit
    assert d.fields[JB].per_item and d.fields[JB].unit == "um"
    assert c.n.loc[FOD, JB] == 30
    assert c.r.loc["Power (W)", "Scan Speed (mm/s)"] == pytest.approx(-1.0)
    assert c.r.loc[FOD, LENGTH] > 0.99
    # a part-level pair is counted once per part, not once per flaw
    assert c.n.loc["Power (W)", "Scan Speed (mm/s)"] == 20
    assert c.n.loc[FOD, LENGTH] == 60
    assert c.n.loc["Power (W)", FOD] == 60
    pairs = c.pairs()
    assert {tuple(sorted(pairs.iloc[i][["Field A", "Field B"]])) for i in range(2)} == \
        {("Power (W)", "Scan Speed (mm/s)"), tuple(sorted((FOD, LENGTH)))}
    names = {frozenset((a, b)) for a, b in zip(pairs["Field A"], pairs["Field B"])}
    assert frozenset((FOD, JB)) not in names                 # agreement, not a relation
    twins = c.pairs(twins=True)
    assert frozenset((FOD, JB)) in {frozenset((a, b)) for a, b in
                                    zip(twins["Field A"], twins["Field B"])}
    assert pairs.iloc[0]["Sure"] == "yes" and "very strong" in pairs.iloc[0]["Strength"]
    assert "strongly correlated" in c.note
    cov = la.coverage(d).set_index("Field")
    assert cov.loc[JB, "Could hold"] == 60 and cov.loc[JB, "Filled"] == 30
    assert la.label(FOD, "um") == "Laser Team FOD (um)"
    assert la.label(JB, "um", 20).endswith("[JB]")
    by_part = la.correlation(d, per=la.PER_PART)
    assert by_part.n.loc[FOD, LENGTH] == 20
    spear = la.correlation(d, method="spearman")
    assert spear.r.loc["Power (W)", "Scan Speed (mm/s)"] == pytest.approx(-1.0)
    la.correlation_figure(c)
    la.scatter_figure(d, FOD, LENGTH)


def test_a_correlation_needs_two_fields_with_different_values(csv):
    lg.save_groups(csv, [lg.Group(name="One", by="part", pattern="S-1")])
    c = la.correlation(la.load(csv), "One")
    assert c.pairs().empty and "nothing to correlate" in c.note


def test_a_group_takes_its_parts_and_leaves_out_fields_it_does_not_track(csv):
    g = lg.Group(name="Early", by="part", pattern="S-1, S-2, S-3", off=["Power (W)"],
                 yes_no={MEDIA: ""})
    lg.save_groups(csv, [g])
    d = la.load(csv)
    assert d.group_names() == ["Early"]
    assert set(la.rows(d, "Early")["Part Number"]) == {"S-1", "S-2", "S-3"}
    assert "Power (W)" not in la.fields_of(d, "Early")
    s = la.summary(d, "Early").set_index("Field")
    assert "Power (W)" not in s.index
    assert s.loc[MEDIA, "Most common"] == "No (2), Yes (1)"


def test_export_writes_the_analysis_workbook_beside_the_ledger(csv):
    d = la.load(csv)
    out = la.export(d)
    assert out == csv.with_name("l_analysis.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Coverage", "Missing", "Summary", "Trends", "Correlation",
                             "Outliers", "Charts"]
    assert FILLED_MARK in wb.properties.keywords             # no scan reads it back
    assert len(wb["Charts"]._images) == len(d.number_fields) + 1      # + the matrix
    assert any(c.value == "S-7" for row in wb["Outliers"].iter_rows() for c in row)
    assert la.export(d) == out                               # its own: written again
    other = csv.with_name("mine.xlsx")
    openpyxl.Workbook().save(other)
    with pytest.raises(LedgerError):
        la.export(d, other)                                  # never over someone's file


def test_the_analyze_window(csv):
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from council_qt.tabs.ledger_analytics_dialog import AnalyticsDialog
    lg.save_groups(csv, [lg.Group(name="Early", by="part", pattern="S-1, S-2, S-3")])
    shown, opened = [], []
    dlg = AnalyticsDialog(la.load(csv), on_part=shown.append, opener=opened.append)
    assert dlg.coverage_table.rowCount() == len(la.load(csv).fields)
    assert dlg.outlier_table.rowCount() == 3
    dlg.field_choice.setCurrentText(FOD)
    assert dlg.canvas is not None and "run rule" in dlg.trend_note.text()
    dlg.outlier_table.cellDoubleClicked.emit(0, 0)
    assert shown == ["S-7"]
    dlg.group_choice.setCurrentText("Early")
    assert "3 part(s)" in dlg.facts.text() and dlg.outlier_table.rowCount() == 0
    dlg.group_choice.setCurrentText("All parts")
    assert dlg.pairs_table.rowCount() >= 1 and dlg.scatter_canvas is not None
    dlg.method_choice.setCurrentText("Spearman")
    assert dlg.current_corr.method == "spearman"
    out = dlg.on_export()
    assert out and out.exists() and opened == [str(out)]
    dlg.close()


def test_analyze_opens_from_the_ledger_tab_and_a_part_opens_its_filled_template(tmp_path,
                                                                              app_dir):
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    from ledger_fixtures import make_laser_log, make_nde_book, make_summary
    from test_ledger_tab import make_tab
    qapp = QApplication.instance() or QApplication([])
    src = tmp_path / "sources"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    make_summary(src / "summary.xlsx")
    tab, _ = make_tab(qapp)
    assert not tab.analyze_btn.isEnabled()
    tab.template.setText(str(make_template(tmp_path / "LMM Template.xlsx")))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.on_extract()
    assert tab.wait(120)
    assert tab.analyze_btn.isEnabled()
    tab.analyze_btn.click()
    assert tab.wait(60)
    dlg = tab.analytics_dialog
    assert dlg.data.csv == tab.csv.text()
    part = dlg.data.parts[0] if dlg.data.parts else "101"
    dlg.on_part(part)
    assert tab.wait(60)
    assert tab.filled_dialog.part == part
