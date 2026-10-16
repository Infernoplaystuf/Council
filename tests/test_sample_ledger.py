"""
sample_ledger — template workbook -> field list -> parent CSV in sample order.

The template is built in tmp_path with both layouts a real one uses: a form
sheet (merged title, "Label:" cells, an instruction paragraph) and a table
sheet (a header row with a unit in brackets and a dropdown under one header).

Run:  python -m pytest tests/test_sample_ledger.py -q
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

from openpyxl.styles import Font

from sample_ledger import (LedgerError, ParentCSV, TemplateSchema, fill_template,
                           read_filled, read_template, sample_sort_key)


@pytest.fixture
def template(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Build Data Entry Template"
    ws.merge_cells("A1:F1")
    ws["B3"] = "Sample ID:"
    ws["B4"] = "Operator:"
    ws["B5"] = "Build Date:"
    ws["A8"] = ("Fill in one sheet per build. Values in grey cells are calculated "
                "and must not be typed over.")
    m = wb.create_sheet("Measurements")
    for col, label in enumerate(["Sample ID", "Laser Power (W)", "Scan Speed [mm/s]",
                                 "Density (g/cc)", "Pass/Fail"], start=1):
        m.cell(2, col, label)
    dv = DataValidation(type="list", formula1='"Pass,Fail,Retest"')
    dv.add("E3:E50")
    m.add_data_validation(dv)
    p = tmp_path / "template.xlsx"
    wb.save(p)
    return p


def _names(schema):
    return [f.name for f in schema.fields]


def test_reads_form_and_table_fields(template):
    s = read_template(template)
    assert s.id_field == "Sample ID"
    assert sorted(_names(s)) == sorted(["Sample ID", "Operator", "Build Date",
                                        "Laser Power (W)", "Scan Speed [mm/s]",
                                        "Density (g/cc)", "Pass/Fail"])
    by = {f.name: f for f in s.fields}
    assert by["Operator"].layout == "form" and by["Operator"].cells == {"": "C4"}
    assert by["Laser Power (W)"].layout == "table"
    assert by["Laser Power (W)"].unit == "W"
    assert by["Scan Speed [mm/s]"].unit == "mm/s"
    assert by["Pass/Fail"].choices == ["Pass", "Fail", "Retest"]
    # title and instruction paragraph are not fields
    assert not any("Template" in n or "Fill in" in n for n in _names(s))


def test_id_column_first_and_schema_roundtrip(template, tmp_path):
    s = read_template(template)
    assert s.columns[0] == "Sample ID"
    s.save(tmp_path / "s.json")
    assert TemplateSchema.load(tmp_path / "s.json").columns == s.columns


def test_explicit_id_field_and_bad_one(template):
    assert read_template(template, id_field="operator").id_field == "Operator"
    with pytest.raises(LedgerError, match="not a field"):
        read_template(template, id_field="Nope")


def test_no_id_field_is_an_actionable_error(tmp_path):
    wb = openpyxl.Workbook()
    for col, label in enumerate(["Alpha", "Beta", "Gamma"], start=1):
        wb.active.cell(1, col, label)
    wb.save(tmp_path / "t.xlsx")
    with pytest.raises(LedgerError, match="id_field"):
        read_template(tmp_path / "t.xlsx")


def test_natural_sort():
    ids = ["S-10", "S-2", "S-1", "12", "7", "S-2a"]
    assert sorted(ids, key=sample_sort_key) == ["7", "12", "S-1", "S-2", "S-2a", "S-10"]


def test_parent_csv_rows_in_numeric_order(template, tmp_path):
    ledger = ParentCSV.create(tmp_path / "parent.csv", read_template(template))
    ledger.add_samples([{"Sample ID": 10, "Operator": "AK"}], source="b.xlsx")
    ledger.add_samples([{"sample id": 2.0, "laser power": 285},
                        {"Sample ID": "9", "Density (g/cc)": 7.9}], source="a.xlsx")
    again = ParentCSV(tmp_path / "parent.csv")
    assert again.sample_ids() == ["2", "9", "10"]
    assert again.get(2)["Laser Power (W)"] == "285"
    assert again.get(10)["source_file"] == "b.xlsx"
    with open(tmp_path / "parent.csv", encoding="utf-8-sig", newline="") as fh:
        header = next(csv.reader(fh))
    assert header[0] == "Sample ID" and header[-2:] == ["source_file", "updated_at"]


def test_merge_fills_blanks_and_reports_conflicts(template, tmp_path):
    ledger = ParentCSV.create(tmp_path / "p.csv", read_template(template))
    ledger.add_samples([{"Sample ID": 1, "Operator": "AK"}])
    r = ledger.add_samples([{"Sample ID": 1, "Operator": "JB", "Pass/Fail": "Pass"}])
    assert r.updated == ["1"]
    assert r.conflicts == [("1", "Operator", "AK", "JB")]
    assert ledger.get(1)["Operator"] == "AK" and ledger.get(1)["Pass/Fail"] == "Pass"
    r = ledger.add_samples([{"Sample ID": 1, "Operator": "JB"}], overwrite=True)
    assert ledger.get(1)["Operator"] == "JB" and not r.conflicts


def test_unknown_columns_and_missing_ids_are_reported(template, tmp_path):
    ledger = ParentCSV.create(tmp_path / "p.csv", read_template(template))
    r = ledger.add_samples([{"Sample ID": 3, "Hatch": 0.1}, {"Operator": "AK"}])
    assert r.unknown_columns == ["Hatch"] and r.skipped_no_key == 1
    assert ledger.sample_ids() == ["3"]


def test_newer_template_adds_columns_keeps_rows(template, tmp_path):
    p = tmp_path / "p.csv"
    ParentCSV.create(p, read_template(template)).add_samples([{"Sample ID": 5, "Operator": "AK"}])
    wb = openpyxl.load_workbook(template)
    wb["Measurements"].cell(2, 6, "Hatch Spacing (mm)")
    wb.save(template)
    ledger = ParentCSV.create(p, read_template(template))
    assert "Hatch Spacing (mm)" in ledger.columns
    assert ledger.columns[-2:] == ["source_file", "updated_at"]
    assert ledger.get(5)["Operator"] == "AK"


def test_cli_creates_csv_and_schema(template, capsys):
    from sample_ledger import main
    assert main([str(template)]) == 0
    assert template.with_name("template_parent.csv").exists()
    assert template.with_name("template_parent.schema.json").exists()
    assert "7 part-level fields" in capsys.readouterr().out


# ── grid templates (one row per sample per item) ───────────────────────
# Laid out like the LMM template: big-type section headings, a part table,
# a flaw grid in three blocks (1A.., 1B.., 1C..) whose row labels carry the
# block letter and one typo, per-row extras, a repeating table, a formula-
# only presentation area to the right, and a trailing notes heading.
CODES = {b: [f"{i}{b}" for i in range(1, 4)] for b in "ABC"}


@pytest.fixture
def grid_template(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    big = Font(size=16, bold=True)

    def head(cell, text):
        ws[cell] = text
        ws[cell].font = big
    head("A1", "Part Attributes")
    for col, label in enumerate(["Part Number", "Power (W)", "At Focus"], start=1):
        ws.cell(2, col, label)
    ws["C3"] = "Y"
    head("A5", "Flaw Attributes")
    head("H5", "Tables for Presentation")
    r = 6
    for b in "ABC":
        ws.cell(r, 2, "Flaw #")
        for i, code in enumerate(CODES[b]):
            ws.cell(r, 3 + i, code)
        ws.cell(r + 1, 2, f"Laser FOD {b} (um)")
        ws.cell(r + 2, 2, "Passes")
        for i in range(3):
            ws.cell(r + 2, 3 + i, 300)
        r += 4
    for rr in range(6, 18):                      # presentation: formulas only
        ws.cell(rr, 9, f"=C{rr}")
        ws.cell(rr, 10, f"=AVERAGE(C{rr}:E{rr})")
    head("A19", "NDE Inspection")
    r = 20
    for b in "ABC":
        ws.cell(r, 2, "Flaw #")
        for i, code in enumerate(CODES[b]):
            ws.cell(r, 3 + i, code)
        ws.cell(r, 6, "Inspector Initials")
        ws.cell(r + 1, 2, "EC Repsonse Pre" if b == "A" else "EC Response Pre")
        ws.cell(r + 2, 2, "EC Response Post")
        r += 4
    head("A33", "Concealment")
    for col, label in enumerate(["Media", "Run Time", "Amplitude"], start=2):
        ws.cell(34, col, label)
    head("A38", "Additional Notes")
    p = tmp_path / "grid.xlsx"
    wb.save(p)
    return p


def test_grid_fields_items_and_cells(grid_template):
    s = read_template(grid_template)
    assert s.key_columns == ["Part Number", "Flaw #"]
    assert s.items == CODES["A"] + CODES["B"] + CODES["C"]
    per_item = [f.name for f in s.fields if f.per_item]
    assert per_item == [
        "Flaw Attributes: Laser FOD (um)", "Flaw Attributes: Passes",
        "NDE Inspection: EC Response Pre",
        "NDE Inspection: EC Response Pre — Inspector Initials",
        "NDE Inspection: EC Response Post",
        "NDE Inspection: EC Response Post — Inspector Initials"]
    fod = s.field("Flaw Attributes: Laser FOD (um)")
    assert (fod.cells["1A"], fod.cells["2B"], fod.cells["3C"]) == ("C7", "D11", "E15")
    assert s.field("Flaw Attributes: Passes").defaults["1A"] == "300"
    ini = s.field("NDE Inspection: EC Response Pre — Inspector Initials")
    assert ini.cells["1A"] == ini.cells["3A"] == "F21" and ini.cells["1B"] == "F25"
    assert s.field("At Focus").defaults == {"": "Y"}
    # the presentation area is derived, so nothing from it is a field
    assert not any(f.section == "Tables for Presentation" for f in s.fields)


def test_repeating_table_and_note(grid_template):
    s = read_template(grid_template)
    runs = [f.name for f in s.fields if f.layout == "repeat"]
    assert "Concealment 1: Media" in runs and "Concealment 3: Amplitude" in runs
    assert "Concealment 4: Media" not in runs     # row 38 is the next heading
    assert s.field("Concealment 2: Media").cells[""] == "B36"
    note = s.field("Additional Notes")
    assert note.layout == "note" and note.cells[""] == "A39"


def test_fill_then_read_back_round_trip(grid_template, tmp_path):
    s = read_template(grid_template)
    ini = "NDE Inspection: EC Response Pre — Inspector Initials"
    rows = [{"Part Number": "P-7", "Flaw #": it, "Power (W)": "285",
             "Flaw Attributes: Laser FOD (um)": f"4{it[0]}0",
             ini: "AK" if it.endswith("A") else "JB",
             "Concealment 2: Media": "ceramic"} for it in s.items]
    out = tmp_path / "filled.xlsx"
    assert fill_template(s, rows, out) == []
    ws = openpyxl.load_workbook(out).active
    assert ws["D11"].value == 420 and ws["F21"].value == "AK"   # numbers stay numbers
    back = {r["Flaw #"]: r for r in read_filled(s, out)}
    assert back["2B"]["Flaw Attributes: Laser FOD (um)"] == "420"
    assert back["3C"][ini] == "JB"
    assert back["1A"]["Concealment 2: Media"] == "ceramic"
    assert back["1A"]["Flaw Attributes: Passes"] == "300"      # template default


def test_fill_warns_when_items_disagree_on_a_shared_cell(grid_template, tmp_path):
    s = read_template(grid_template)
    col = "NDE Inspection: EC Response Pre — Inspector Initials"
    rows = [{"Part Number": "1", "Flaw #": "1A", col: "AK"},
            {"Part Number": "1", "Flaw #": "2A", col: "ZZ"}]
    warn = fill_template(s, rows, tmp_path / "o.xlsx")
    assert len(warn) == 1 and "'AK' kept" in warn[0]
    with pytest.raises(LedgerError, match="over the template"):
        fill_template(s, rows, grid_template)


def test_grid_parent_csv_sorts_parts_then_items(grid_template, tmp_path):
    s = read_template(grid_template)
    ledger = ParentCSV.create(tmp_path / "parent.csv", s)
    for part in ("12", "2"):
        out = tmp_path / f"p{part}.xlsx"
        fill_template(s, [{"Part Number": part, "Flaw #": "1A"}], out)
        r = ledger.add_samples(read_filled(s, out), source=out.name)
        assert len(r.added) == 9
    again = ParentCSV(tmp_path / "parent.csv")
    assert again.keys == ["Part Number", "Flaw #"]
    assert again.sample_ids() == ["2", "12"]
    assert [r["Flaw #"] for r in again.rows_for("2")] == s.items
    assert again.get("12", "3C")["source_file"] == "p12.xlsx"
    # re-adding the same file changes nothing
    assert ledger.add_samples(read_filled(s, tmp_path / "p2.xlsx")).added == []


REAL = Path.home() / "Downloads" / "LMM Template.xlsx"


@pytest.mark.skipif(not REAL.exists(), reason="the LMM template is not on this machine")
def test_real_lmm_template():
    s = read_template(REAL)
    assert s.key_columns == ["Part Number", "Flaw #"] and len(s.items) == 30
    pre = s.field("NDE Inspection: Eddy Current Response Pre-Conceal")
    assert (pre.cells["1A"], pre.cells["1B"], pre.cells["10C"]) == ("C34", "C41", "L48")
    assert s.field("Visual Inspection: Inspector 4 — Date").cells["10C"] == "N80"
    assert len([f for f in s.fields if f.per_item]) == 36
