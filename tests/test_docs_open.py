"""Text documents as Ledger sources (docs_open, user 2026-10-21: "some labels
may be properly explained in the text documents (word, .txt etc) or they
could have data in them"). Word documents are built here with python-docx;
Council reads them with the standard library."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
import docs_open  # noqa: E402
import source_extract as se  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_records as lr  # noqa: E402
from ledger_fixtures import make_laser_log, make_template  # noqa: E402


# ── a line: data or an explanation ──────────────────────────────────────
@pytest.mark.parametrize("line,kind", [
    ("Part Number: 101", "data"),
    ("Inspected by JB", "data"),
    ("Tank — North tank", "data"),
    ("Time: 10:30", "data"),
    ("HP — health points at full growth", "explanation"),
    ("OX: the oxygen reading taken after the dive", "explanation"),
    ("EC means eddy current", "explanation"),
    ("• FOD = foreign object debris found on the part", "explanation"),
])
def test_data_is_told_from_an_explanation(line, kind):
    assert docs_open.classify(line)[0] == kind


def test_explanations_point_to_fields():
    exs = [docs_open.Explanation("HP", "the health points of the creature", "n.txt"),
           docs_open.Explanation("Traverse rate", "scan speed of the head", "n.txt"),
           docs_open.Explanation("Weight", 'written "Wt" or "Mass" in old logs', "n.txt"),
           docs_open.Explanation("Q", "the health or weight, whichever is lower", "n.txt"),
           docs_open.Explanation("Health", "the creature's hit points", "n.txt")]
    fields = [("Health", ["health"]), ("Weight", ["weight"]),
              ("Scan Speed (mm/s)", ["scan speed"])]
    got = {k: v[0] for k, v in docs_open.explanations_for(exs, fields).items()}
    # Q names two fields equally: none; Health explained is no new word
    assert got == {"hp": "Health", "traverse rate": "Scan Speed (mm/s)", "wt": "Weight",
                   "mass": "Weight"}


# ── reading ─────────────────────────────────────────────────────────────
def notes_docx(path: Path) -> Path:
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_heading("Part 101 build notes", 1)
    d.add_paragraph("Part Number: 101")
    d.add_paragraph("Traverse rate — the scan speed of the laser head, in mm/s")
    d.add_paragraph("The part was built on the second machine after the filter change, "
                    "and the operators noted nothing unusual during the run.")
    d.add_heading("Concealment", 2)
    d.add_paragraph("Media and amplitude used for each run:")
    t = d.add_table(rows=3, cols=2)
    for i, row in enumerate([["Media", "Amplitude"], ["ceramic", "40"], ["steel", "35"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    g = d.add_table(rows=3, cols=2)
    for i, row in enumerate([["Term", "Meaning"], ["FOD", "laser team foreign object debris"],
                             ["Amp", "the amplitude"]]):
        for j, v in enumerate(row):
            g.cell(i, j).text = v
    d.save(path)
    return path


def test_a_word_document_is_text_tables_and_explanations(tmp_path):
    doc = docs_open.read_document(notes_docx(tmp_path / "notes.docx"))
    names = [s.name for s in doc.sheets]
    assert names == ["Text", "Table 1"]                     # the Term|Meaning table is no data
    text = [v for _k, v in sorted(doc.sheets[0].cells.items())]
    assert "Part Number: 101" in text and not any("operators noted" in t for t in text)
    assert not any(t.startswith("Traverse rate") for t in text)      # an explanation
    table = doc.sheets[1].cells
    lines = [v for (r, c), v in sorted(table.items()) if c == 1]
    assert lines[:2] == ["Part 101 build notes", "Concealment"]
    assert "Media and amplitude used for each run:" in lines
    assert ("ceramic", "40") in {(table.get((r, 1)), table.get((r, 2))) for r, _ in table}
    terms = {e.term: e.text for e in doc.explanations}
    assert terms["Traverse rate"].startswith("the scan speed") and terms["FOD"]


def test_a_text_file_with_markdown_and_pasted_tables(tmp_path):
    p = tmp_path / "log.md"
    p.write_text("# Visual check\nInspector: MK\n\n| Defect | Result |\n|---|---|\n"
                 "| 1A | seen |\n| 2A | seen |\n\nWidth\tLength\n2.5\t10\n3.0\t12\n",
                 encoding="utf-8")
    doc = docs_open.read_document(p)
    assert [s.name for s in doc.sheets] == ["Text", "Table 1", "Table 2"]
    t1 = doc.sheets[1].cells
    assert [v for (r, c), v in sorted(t1.items()) if c == 1][:2] == ["Visual check",
                                                                    "Inspector: MK"]
    assert ("1A", "seen") in {(t1.get((r, 1)), t1.get((r, 2))) for r, _ in t1}
    t2 = doc.sheets[2].cells
    assert ("2.5", "10") in {(t2.get((r, 1)), t2.get((r, 2))) for r, _ in t2}


def test_a_pdf_and_an_old_doc(tmp_path):
    old = tmp_path / "old.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(docs_open.DocError, match="save it as .docx"):
        docs_open.read_document(old)
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter
    pdf = tmp_path / "blank.pdf"
    w = PdfWriter()
    w.add_blank_page(100, 100)
    with open(pdf, "wb") as fh:
        w.write(fh)
    assert docs_open.read_document(pdf).sheets == []


def test_a_folder_scan_finds_documents(tmp_path):
    (tmp_path / "a.txt").write_text("x")
    (tmp_path / "b.docx").write_bytes(b"")
    (tmp_path / "~$b.docx").write_bytes(b"")
    assert [p.name for p in se.workbook_files(tmp_path)] == ["a.txt", "b.docx"]


# ── the sample Ledger ───────────────────────────────────────────────────
def test_a_document_explains_a_label_and_holds_data(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    log = make_laser_log(src / "laser.xlsx")
    wb = openpyxl.load_workbook(log)
    wb["Build"]["A6"] = "Traverse rate"                    # no rule knows it
    wb.save(log)
    notes_docx(src / "notes.docx")
    tpl = make_template(tmp_path / "LMM Template.xlsx")
    csv = Path(li.create_parent(tpl, tmp_path / "lmm.csv").csv)
    ex = li.run_extract(csv, [src], use_images=False)
    wb = openpyxl.load_workbook(ex.review)
    rows = list(wb["Values"].iter_rows(min_row=2, values_only=True))
    speed = next(r for r in rows if r[3] == "Scan Speed (mm/s)")
    assert speed[4] == "1200" and speed[0] is None          # a suggestion: amber
    assert "explained in notes.docx: Traverse rate — the scan speed" in speed[10]
    media = [r for r in rows if str(r[3]).startswith("Concealment") and "Media" in r[3]]
    assert {r[4] for r in media} == {"ceramic", "steel"}    # the Word table's data
    assert all(r[0] is None and "read from a text document" in r[10] for r in media)
    # accepting it teaches it
    for row in wb["Values"].iter_rows(min_row=2):
        if row[3].value == "Scan Speed (mm/s)":
            row[0].value = "Y"
    wb.save(ex.review)
    li.run_apply(ex.review, use_images=False)
    assert se.load_aliases(csv)[0]["traverse rate"] == "Scan Speed (mm/s)"


def test_council_model_is_told_what_the_documents_say(tmp_path):
    from council_core import ledger_llm
    asked = []

    def chat(prompt, schema):
        asked.append(prompt)
        return '{"mappings": []}', "fake"
    exs = [docs_open.Explanation("Zq", "the flaw orientation in degrees", "n.txt",
                                 "Zq — the flaw orientation in degrees")]
    ledger_llm.map_labels(["Zq"], ["Flaw Orientation (deg)"], chat,
                          notes=docs_open.notes_text(exs, ["Zq"]))
    assert "n.txt: Zq — the flaw orientation in degrees" in asked[0]


# ── the table Ledger ────────────────────────────────────────────────────
def test_a_note_explains_a_header_for_the_table_ledger(tmp_path):
    from test_ledger_records import simple_ledger, book, table
    tpl, csv = simple_ledger(tmp_path)
    src = tmp_path / "s"
    src.mkdir()
    book(src / "a.xlsx", {"S": table(1, 1, ["Name", "HP", "Weight"], [["A", 1, 2],
                                                                      ["B", 3, 4]])})
    (src / "notes.txt").write_text("HP — the health points of the creature\n",
                                   encoding="utf-8")
    run = lr.extract(csv, sorted(src.iterdir()))
    rec = run.records[0]
    assert rec.values["Health"].how == "explained" and rec.values["Health"].check
    assert "notes.txt: HP — the health points" in run.decisions["C|Health|hp"].detail
    out = tmp_path / "r.xlsx"
    lr.write_review(out, run, csv, [])
    wb = openpyxl.load_workbook(out)
    for row in wb["Decisions"].iter_rows(min_row=2):
        row[1].value = "Y"
    wb.save(out)
    lr.apply_review(out)
    assert lr.extract(csv, sorted(src.iterdir())).records[0].values["Health"].how == "taught"


def test_a_document_of_records_for_the_table_ledger(tmp_path):
    from test_ledger_records import simple_ledger
    tpl, csv = simple_ledger(tmp_path)
    p = tmp_path / "creatures.md"
    p.write_text("# Creatures\n| Name | Health | Weight |\n|---|---|---|\n"
                 "| Rex 1 | 4480 | 574 |\n| Rex 2 | 7700 | 532 |\n", encoding="utf-8")
    run = lr.extract(csv, [p])
    assert [(r.values["Name"].value, r.values["Health"].value) for r in run.records] == [
        ("Rex 1", 4480), ("Rex 2", 7700)]
    # a text document's records wait for one Y on the Decisions sheet
    assert all(r.accept == "" and r.checks == ["D|creatures.md"] for r in run.records)
    assert run.decisions["D|creatures.md"].records == 2
