"""
The Ledger's pictures in an .xlsx (user, 2026-10-26: "I am only getting about
1/3 to 1/2 of the images read"): read from the file's zip
(excel_open.zip_pictures), not openpyxl's, which dropped pictures in a group,
in a drawing wrapped for newer Excel versions (mc:AlternateContent) - with
the rest of that sheet's -, at a fixed position, and pictures placed in a
cell (Excel's "Place in Cell"). And the count: a picture read that has no
text in it (a photo) is not one 'not read'.

Run:  python -m pytest tests/test_ledger_pictures.py -q
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import excel_open  # noqa: E402

RED, BLUE, GREEN = (255, 0, 0), (0, 0, 255), (0, 128, 0)

_X = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ' \
     'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def _png(colour) -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), colour).save(buf, "PNG")
    return buf.getvalue()


def _pic(rid: str) -> str:
    return (f'<pic><nvPicPr><cNvPr id="1" name="p"/><cNvPicPr/></nvPicPr><blipFill>'
            f'<a:blip {_X} r:embed="{rid}"/></blipFill><spPr/></pic>')


def _from(col: int, row: int) -> str:
    return (f"<from><col>{col}</col><colOff>0</colOff><row>{row}</row><rowOff>0</rowOff>"
            "</from>")


def every_kind_of_picture(path: Path) -> Path:
    """A sheet 'EC' with a picture at B2; a group of two at D4; one wrapped
    in mc:AlternateContent at F6 (its fallback the same one); one at a fixed
    position; and one placed in cell C5."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    wb.active.title = "EC"
    wb.active["A1"] = "Part"
    wb.active.add_image(XLImage(io.BytesIO(_png(RED))), "B2")
    raw = io.BytesIO()
    wb.save(raw)
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    drawing = (
        '<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        f'<oneCellAnchor>{_from(1, 1)}<ext cx="1" cy="1"/>{_pic("rId1")}<clientData/>'
        '</oneCellAnchor>'
        f'<twoCellAnchor>{_from(3, 3)}{_from(4, 4).replace("from", "to")}'
        f'<grpSp><nvGrpSpPr><cNvPr id="2" name="g"/><cNvGrpSpPr/></nvGrpSpPr><grpSpPr/>'
        f'{_pic("rId1")}{_pic("rId2")}</grpSp><clientData/></twoCellAnchor>'
        '<mc:AlternateContent><mc:Choice Requires="a14">'
        f'<twoCellAnchor>{_from(5, 5)}{_from(6, 6).replace("from", "to")}{_pic("rId2")}'
        '<clientData/></twoCellAnchor></mc:Choice><mc:Fallback>'
        f'<twoCellAnchor>{_from(5, 5)}{_from(6, 6).replace("from", "to")}{_pic("rId2")}'
        '<clientData/></twoCellAnchor></mc:Fallback></mc:AlternateContent>'
        f'<absoluteAnchor><pos x="0" y="0"/><ext cx="1" cy="1"/>{_pic("rId1")}<clientData/>'
        '</absoluteAnchor></wsDr>')
    rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/image" Target="../media/image1.png"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/image" Target="../media/image2.png"/></Relationships>')
    rich = {
        "xl/metadata.xml":
            '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion='
            '"120000"/></metadataTypes><futureMetadata name="XLRICHVALUE" count="1"><bk>'
            '<extLst><ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}"><xlrd:rvb i="0"/>'
            '</ext></extLst></bk></futureMetadata><valueMetadata count="1"><bk>'
            '<rc t="1" v="0"/></bk></valueMetadata></metadata>',
        "xl/richData/rdrichvalue.xml":
            '<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" '
            'count="1"><rv s="0"><v>0</v><v>5</v></rv></rvData>',
        "xl/richData/rdrichvaluestructure.xml":
            '<rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/'
            'richdata" count="1"><s t="_localImage"><k n="_rvRel:LocalImageIdentifier" '
            't="i"/><k n="CalcOrigin" t="i"/></s></rvStructures>',
        "xl/richData/richValueRel.xml":
            '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/'
            'richvaluerel" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships"><rel r:id="rId1"/></richValueRels>',
        "xl/richData/_rels/richValueRel.xml.rels":
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/image" Target="../media/image3.png"/>'
            '</Relationships>',
    }
    with zipfile.ZipFile(path, "w") as out:
        for name in src.namelist():
            data = src.read(name)
            if name == "xl/drawings/drawing1.xml":
                data = drawing.encode()
            elif name == "xl/drawings/_rels/drawing1.xml.rels":
                data = rels.encode()
            elif name == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</row></sheetData>",
                                    b'</row><row r="5"><c r="C5" t="e" vm="1"><v>#VALUE!</v>'
                                    b'</c></row></sheetData>')
            out.writestr(name, data)
        out.writestr("xl/media/image2.png", _png(BLUE))
        out.writestr("xl/media/image3.png", _png(GREEN))
        for name, text in rich.items():
            out.writestr(name, text)
    return path


def _colour(data: bytes):
    from PIL import Image
    return Image.open(io.BytesIO(data)).convert("RGB").getpixel((0, 0))


def test_every_kind_of_picture_is_read_from_the_zip(tmp_path):
    pics = excel_open.zip_pictures(every_kind_of_picture(tmp_path / "pics.xlsx"))
    got = {k: [_colour(d) for d, _e in v] for k, v in pics.items()}
    assert got == {("EC", "B2"): [RED],
                   ("EC", "D4"): [RED, BLUE],               # both of the group
                   ("EC", "F6"): [BLUE],                    # once, not its fallback too
                   ("EC", "A1"): [RED],                     # at a fixed position
                   ("EC", "C5"): [GREEN]}                   # placed in the cell


def test_the_ledger_reads_them_all(tmp_path):
    import source_extract as se
    grids, images = se.load_workbook_grids(every_kind_of_picture(tmp_path / "pics.xlsx"))
    assert len(images) == 6
    assert sorted((i.sheet, i.anchor) for i in images) == sorted(
        [("EC", "B2"), ("EC", "D4"), ("EC", "D4"), ("EC", "F6"), ("EC", "A1"), ("EC", "C5")])
    assert any(g.sheet == "EC" and g.cells.get((1, 1)) == "Part" for g in grids)
    assert se.load_workbook_grids(tmp_path / "pics.xlsx", want_pictures=False)[1] == []


def test_a_picture_with_no_text_is_not_counted_as_not_read():
    from council_core import ledger_ingest as li
    from source_extract import ImageNote
    notes = [ImageNote("a.xlsx", "S", "A1", "", "OCR", "Power | 285"),
             ImageNote("a.xlsx", "S", "B1", "", "OCR", "(no text found)"),
             ImageNote("a.xlsx", "S", "C1", "", "OCR", "reader failed: timed out"),
             ImageNote("a.xlsx", "S", "D1", "", "", "not read: Excel could not copy it out "
                                                  "(see the problems)")]
    assert (li._images_read(notes), li._images_blank(notes)) == (1, 1)
    assert li._images_unread(notes) == "1 the reader failed, 1 Excel could not copy them out"
    s = li.ExtractSummary("r.xlsx", 1, 0, 0, 0, 0, 4, 1, images_blank=1,
                          images_unread=li._images_unread(notes))
    assert "1 of 4 images read (1 more had no text in them); not read: 1 the reader " \
           "failed, 1 Excel could not copy them out" in s.describe()


def test_an_emf_drawing_is_said_so_when_it_cannot_be_drawn(tmp_path):
    import source_extract as se
    data, ext, why = se._as_raster(b"not really an emf", "emf")
    assert (data, ext) == (b"not really an emf", "emf") and "EMF" in why
    data, ext, why = se._as_raster(_png(RED), "wmf")          # one Pillow can read
    assert ext == "png" and why == "" and _colour(data) == RED
