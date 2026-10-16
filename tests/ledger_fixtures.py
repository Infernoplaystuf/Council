"""
Synthetic workbooks for the ingestion tests: a small copy of the LMM template's
shape, and source workbooks laid out the ways real ones differ from it.

Nothing here is real data. Every workbook is built in a pytest tmp_path, so the
tests run on a fresh clone.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

BLOCKS = "ABC"
ITEMS = [f"{i}{b}" for b in BLOCKS for i in range(1, 4)]       # 1A..3A, 1B.., 1C..


def make_template(path: Path) -> Path:
    """Part table; Flaw / NDE / Visual grids in A-B-C blocks with per-row
    extras; a 3-run Concealment table; a formula-only presentation area."""
    wb = openpyxl.Workbook()
    ws = wb.active
    big = Font(size=16, bold=True)

    def head(cell, text):
        ws[cell] = text
        ws[cell].font = big

    def grid(top, rows, extras=(), letter_in=None):
        r = top
        for b in BLOCKS:
            ws.cell(r, 2, "Flaw #")
            for i in range(3):
                ws.cell(r, 3 + i, f"{i + 1}{b}")
            for j, e in enumerate(extras):
                ws.cell(r, 6 + j, e)
            for k, label in enumerate(rows):
                ws.cell(r + 1 + k, 2, label.format(b=b))
            r += len(rows) + 2
        return r

    head("A1", "Part Attributes")
    for col, label in enumerate(["Part Number", "Frequency (kHz)", "Power (W)",
                                 "Scan Speed (mm/s)", "Date"], start=1):
        ws.cell(2, col, label)
    head("A5", "Flaw Attributes")
    head("K5", "Tables for Presentation")
    end = grid(6, ["Laser Team FOD {b} (um)", "Flaw Orientation (deg)",
                   "Flaw Length (in)", "Passes"])
    for rr in range(6, end):
        ws.cell(rr, 12, f"=C{rr}")
        ws.cell(rr, 13, f"=AVERAGE(C{rr}:E{rr})")
    head(f"A{end + 1}", "NDE Inspection")
    end = grid(end + 2, ["Eddy Current Response Pre-Conceal",
                         "Eddy Current Response Post-Conceal", "FPI"],
               extras=["Inspector Initials", "Date"])
    head(f"A{end + 1}", "Concealment")
    for col, label in enumerate(["Media", "Amplitude", "Notes"], start=2):
        ws.cell(end + 2, col, label)
    vis = end + 6
    head(f"A{vis}", "Visual Inspection")
    grid(vis + 1, ["Inspector 1", "Inspector 2"], extras=["Inspector Initials"])
    wb.save(path)
    return path


def _png(text_rows) -> bytes:
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (320, 30 * len(text_rows) + 10), "white")
    d = ImageDraw.Draw(img)
    for i, row in enumerate(text_rows):
        d.text((10, 10 + 30 * i), "    ".join(row), fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


CONCEALMENT_IMAGE_ROWS = [["Media", "Amplitude"], ["ceramic", "40"], ["steel", "35"]]


def make_laser_log(path: Path) -> Path:
    """Part number written inline; flaws down the side with mm units."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Build"
    ws["A1"] = "Laser build log"
    ws["A3"] = "P/N: 101"
    ws["A4"], ws["B4"] = "Freq (kHz)", 50
    ws["A5"], ws["C5"] = "Laser Power (W)", 285          # value two cells over
    ws["A6"], ws["B6"] = "Scan Speed", 1200
    fl = wb.create_sheet("Flaws")
    fl.append(["Flaw ID", "FOD (mm)", "Orientation (deg)", "Length (mm)", "Passes"])
    for n, item in enumerate(ITEMS):
        fl.append([f"Flaw {item}", 0.4 + n / 1000, [0, 90, 45][n % 3], 2.0, 300])
    wb.save(path)
    return path


def make_nde_book(path: Path) -> Path:
    """Sheet named for the part; grid codes written 1-A; visual results
    transposed; concealment only as a picture; one unknown label."""
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part 101"
    ws["A1"] = "NDE"
    ws["B3"] = "Flaw"
    for i, item in enumerate(ITEMS):
        ws.cell(3, 3 + i, f"{item[0]}-{item[1]}")
    ws["B4"], ws["B5"] = "EC Pre (%FSH)", "EC Post (%FSH)"
    for i in range(len(ITEMS)):
        ws.cell(4, 3 + i, 60 + i)
        ws.cell(5, 3 + i, 20 + i)
    ws["B8"] = "Operator: JB"
    vis = wb.create_sheet("Visual")
    vis.append(["Defect", "Inspector 1", "Inspector 2"])
    for item in ITEMS:
        vis.append([item, "seen", "seen" if item != "3C" else "not seen"])
    img = XLImage(io.BytesIO(_png(CONCEALMENT_IMAGE_ROWS)))
    vis.add_image(img, "F2")
    wb.save(path)
    return path


def make_summary(path: Path) -> Path:
    """Several parts as rows; part 101's power disagrees with the laser log."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Part Number", "Power (W)", "Date"])
    ws.append([101, 290, "2026-10-01"])
    ws.append([102, 300, "2026-10-02"])
    wb.save(path)
    return path


def fake_reader(rows):
    """Stands in for OCR / a vision model: returns fixed rows for any image."""
    def read(_data):
        return [list(r) for r in rows]
    read.name = "test reader"
    return read
