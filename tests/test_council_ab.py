"""tools/council_ab.py — the before/after harness's scoring, vault and report
(no model: the runs themselves need a real one, on the user's PC)."""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))

import council_ab as ab  # noqa: E402


def test_checks_score_like_the_council_bench():
    assert ab.check("17 x 23 = 391.", [{"number": 391}]) == []
    assert ab.check("It is 1,391.", [{"number": 391}]) != []
    assert ab.check("```python\ndef f(:\n```", [{"python_parses": True}])
    assert ab.check("No code here", [{"python_parses": True}]) == [
        "no Python block"]
    assert ab.check("Buy from TradeWorks", [{"not_contains": ["tradeworks"]}])


def test_the_questions_load_and_have_checks():
    qs = ab.load_questions()
    assert len(qs) >= 20
    assert all(q.get("checks") for q in qs)
    assert {q.get("expect") for q in qs} <= {"one", "council", "data", None}


def test_the_vault_holds_every_planted_answer(tmp_path):
    vault = ab.make_vault(tmp_path / "v")
    inp = vault / "data_in"
    text = {p.name: p.read_text(encoding="utf-8") for p in inp.iterdir()}
    assert "EPDM" in text["pump_notes.txt"]
    assert "Rowan Ellis" in text["pump_notes.txt"]
    assert sum("Hastelloy" in t for t in text.values()) == 3
    runs = list(csv.DictReader((inp / "furnace_runs.csv").open()))
    dips = [r["run_id"] for r in runs if int(r["supply_voltage_v"]) < 200]
    assert dips == ["R-117"]
    rows = list(csv.DictReader((inp / "inspections.csv").open()))
    assert sum(r["result"] == "FAIL" for r in rows) == 7
    assert not any(m.lower() in t.lower() for t in text.values()
                   for m in ab.SAMPLE_MARKERS)


def test_the_vault_copies_only_the_model_settings(tmp_path):
    mine = tmp_path / "mine"
    mine.mkdir()
    (mine / "backend_settings.json").write_text('{"gguf_path": "x"}')
    (mine / "personality_backends.json").write_text("{}")
    (mine / "secret.txt").write_text("no")
    vault = ab.make_vault(tmp_path / "v", from_vault=mine)
    assert (vault / "backend_settings.json").is_file()
    assert (vault / "personality_backends.json").is_file()
    assert not (vault / "secret.txt").exists()


def _rec(qid, side, **kw):
    r = {"id": qid, "rep": 0, "kind": "math", "expect": "one", "side": side,
         "correct": True, "failed": [], "error": "", "path": "one",
         "handed_on": "", "calls": 1, "loads": 0, "wall_s": 2.0,
         "first_text_s": 0.5, "sample_leak": [], "answer": "391",
         "by_role": {"writer": 1}}
    r.update(kw)
    return r


def test_the_report_compares_the_sides(tmp_path):
    for side, recs in (("before-app", [
            _rec("a", "before-app", first_text_s=2.0),
            _rec("b", "before-app", expect="council", correct=False,
                 failed=["no 7"], sample_leak=["Power Strip"])]),
            ("after-app", [
            _rec("a", "after-app"),
            _rec("b", "after-app", expect="council", path="council",
                 handed_on="a coding question", calls=12,
                 by_role={"writer": 5, "judge": 7})])):
        with (tmp_path / f"chat-{side}.jsonl").open("w") as fh:
            for r in recs:
                fh.write(json.dumps(r) + "\n")
    (tmp_path / "ledger-before.json").write_text(json.dumps(
        [{"rows": 2000, "seconds": 0.5, "values": 1996}]))
    (tmp_path / "ledger-after.json").write_text(json.dumps(
        [{"rows": 2000, "timed_out": True, "seconds": 300}]))
    md = ab.report(tmp_path, ["before-app", "after-app"], {})
    s = ab.summarise(ab.read_jsonl(tmp_path / "chat-after-app.jsonl"))
    assert s["right"] == 2 and s["handed_on"] == 1 and s["routing_ok"] == 2
    assert "| Right answers | 1/2 (50%) | 2/2 (100%) (better) |" in md
    assert "| Answers with demo sample data | 1 | 0 (better) |" in md
    assert "over 300 s (stopped)" in md
    assert (tmp_path / "report.html").read_text().count("<details>") == 4


def test_the_long_log_has_the_rows_asked_for(tmp_path):
    import openpyxl
    tpl, log = ab.make_ledger_files(tmp_path, 50)
    ws = openpyxl.load_workbook(log).active
    assert ws.max_row == 51
    assert ws.cell(1, 3).value == "Part Number"
    assert tpl.is_file()


def test_a_model_choice_sets_the_backend():
    env = ab.model_env("ollama:gemma3:12b")
    assert env["COUNCIL_BACKEND"] == "ollama"
    assert env["COUNCIL_OLLAMA_MODEL"] == "gemma3:12b"
    env = ab.model_env("m.gguf")
    assert env["COUNCIL_BACKEND"] == "gguf"
    assert env["COUNCIL_GGUF_PATH"].endswith("m.gguf")


# ------------------------------------------------ tools/council_fulltest.py
import council_fulltest as ft  # noqa: E402


def _args(**kw):
    import argparse
    base = dict(model="ollama:gemma3:12b", repeat=1, from_vault="",
                kg_model="", ledger_images=False, long=False)
    base.update(kw)
    return argparse.Namespace(**base)


def test_every_step_has_an_id_a_command_and_a_limit(tmp_path):
    steps = ft.steps(_args(), tmp_path, tmp_path / "vault")
    ids = [s["id"] for s in steps]
    assert len(ids) == len(set(ids))
    assert {"suite", "ab", "ledger", "connections", "llm"} <= set(ids)
    assert all(s["cmd"] and s["limit"] > 0 for s in steps)
    assert {s["id"] for s in steps if s.get("long")} == {
        "typhon", "code", "council"}


def test_a_gguf_goes_only_to_the_steps_that_take_one(tmp_path):
    steps = {s["id"]: s["cmd"] for s in ft.steps(
        _args(model="C:/m.gguf"), tmp_path, tmp_path / "v")}
    assert "C:/m.gguf" in steps["describe"]
    steps = {s["id"]: s["cmd"] for s in ft.steps(_args(), tmp_path,
                                                 tmp_path / "v")}
    assert "--model" not in steps["describe"]
    assert "gemma3:12b" in steps["connections"]


def test_the_stand_in_template_reads_as_a_template(tmp_path):
    from sample_ledger import read_template
    schema = read_template(ft.make_lmm_template(tmp_path / "t.xlsx"))
    labels = " ".join(str(getattr(f, "label", f)) for f in schema.fields)
    assert "Part Number" in labels and "Eddy Current" in labels


def test_the_graph_alone_finds_most_answer_key_links(tmp_path):
    assert ft.connections(tmp_path / "kg-vault", "") == 0
    got = json.loads((tmp_path / "connections.json").read_text())
    found = got["found_after_seed"]
    assert found["total"] > 30 and found["found"] >= found["total"] - 10


def test_the_full_report_lists_every_step(tmp_path):
    (tmp_path / "suite.log").write_text("12 passed\n")
    text = ft.write_report(tmp_path, [
        {"id": "suite", "status": "ok", "seconds": 60},
        {"id": "typhon", "status": "skipped: needs --long"}],
        {"when": "now", "model": "", "build": "abc"})
    assert "| suite | ok | 1.0 |" in text and "12 passed" in text
    assert (tmp_path / "report.html").is_file()
