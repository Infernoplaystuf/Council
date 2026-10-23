"""
Saved Ledgers (council_core.ledger_saved) and a Ledger run again as its
folders gain files - user, 2026-10-23: "Lets say I create the filled
template once but then the folders it was run on get more data in them I
want to be able to rerun from that stage", "ledger runs should be saved,
catergorized, and named by the council", "both by template and or
project/program".

Run:  python -m pytest tests/test_ledger_saved.py -q
"""
from __future__ import annotations

import json
import os
import sys
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).parent))

from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_review as lv  # noqa: E402
from council_core import ledger_saved as ls  # noqa: E402
from ledger_fixtures import make_laser_log, make_nde_book, make_summary, make_template  # noqa: E402,E501
from sample_ledger import ParentCSV  # noqa: E402
from test_ledger_ingest import accept  # noqa: E402


@pytest.fixture
def app_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("COUNCIL_APP_DIR", str(tmp_path / "app"))
    return tmp_path / "app"


@pytest.fixture
def ledger(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    template = make_template(tmp_path / "LMM Template.xlsx")
    csv = li.create_parent(template, tmp_path / "l.csv").csv
    return Path(csv), src, template


def test_a_run_is_saved_beside_the_csv_and_listed(app_dir, ledger):
    csv, src, template = ledger
    s = li.run_extract(csv, [src], use_images=False)
    run = ls.record_run(csv, s, sources=[str(src)], template=str(template),
                        settings={"part": "", "prefer": "JB"})
    assert ls.path_for(csv).name == "l.csv.ledger.json" and ls.path_for(csv).is_file()
    now = datetime.now()
    assert run.name == f"LMM Template · src · {now.day} {now:%b %Y}"
    led = ls.load(csv)
    assert led.sources == [str(src)] and led.settings["prefer"] == "JB"
    assert len(led.files) == 2 and led.last.review == s.review
    assert (led.last.values, led.last.to_review) == (s.values, s.to_review)
    assert [x.csv for x in ls.listed()] == [str(csv.resolve())]
    ls.forget(csv)
    assert ls.listed() == [] and ls.path_for(csv).is_file()        # nothing deleted


def test_new_and_changed_files_are_told_apart(app_dir, ledger, tmp_path):
    csv, src, template = ledger
    s = li.run_extract(csv, [src], use_images=False)
    ls.record_run(csv, s, sources=[str(src)], template=str(template))
    files = li.source_files(csv, [src])
    assert ls.changed_files(ls.load(csv), files) == ([], [], [str(f) for f in files])
    make_summary(src / "summary.xlsx")
    log = src / "laser_log.xlsx"
    wb = openpyxl.load_workbook(log)
    wb.active["Z1"] = "edited"
    wb.save(log)
    os.utime(log, ns=(log.stat().st_atime_ns, log.stat().st_mtime_ns + 10**9))
    new, changed, same = ls.changed_files(ls.load(csv), li.source_files(csv, [src]))
    assert [Path(f).name for f in new] == ["summary.xlsx"]
    assert [Path(f).name for f in changed] == ["laser_log.xlsx"]
    assert [Path(f).name for f in same] == ["NDE results.xlsx"]


def test_ledgers_are_grouped_by_template_and_or_project(app_dir, tmp_path):
    def saved(name, template, project):
        csv = tmp_path / f"{name}.csv"
        csv.write_text("x\n", encoding="utf-8")
        ls.save(ls.SavedLedger(csv=str(csv), template=template, name=name, project=project))
    saved("a", "T1.xlsx", "Apollo")
    saved("b", "T2.xlsx", "Apollo")
    saved("c", "T1.xlsx", "")
    by = {k: [(h, [x.name for x in g]) for h, g in ls.group(ls.listed(), k)]
          for k in ls.GROUPS}
    assert by["template"] == [("T1", ["a", "c"]), ("T2", ["b"])]
    assert by["project"] == [("Apollo", ["a", "b"]), (ls.NO_PROJECT, ["c"])]
    assert by["project+template"] == [("Apollo › T1", ["a"]), ("Apollo › T2", ["b"]),
                                      (f"{ls.NO_PROJECT} › T1", ["c"])]
    assert ls.projects() == ["Apollo"]
    assert ls.set_details(tmp_path / "c.csv", project="Gemini")
    assert ls.load(tmp_path / "c.csv").project == "Gemini"


def test_a_run_is_named_by_the_model_or_the_rules_and_renamed(app_dir, ledger):
    csv, src, template = ledger
    s = li.run_extract(csv, [src], use_images=False)
    asked = []

    def chat(prompt, schema):
        asked.append(prompt)
        return json.dumps({"name": "  LMM laser and NDE, October "}), "test"
    run = ls.record_run(csv, s, sources=[str(src)], template=str(template),
                        namer=lambda facts, fb: ls.model_name(chat, facts, fb))
    assert run.name == "LMM laser and NDE, October"
    assert "template: LMM Template" in asked[0] and "files read: 2" in asked[0]
    assert "285" not in asked[0]                         # names and counts, no values

    def broken(prompt, schema):
        raise RuntimeError("no model")
    assert ls.model_name(broken, "facts", "the rules' name") == "the rules' name"
    assert ls.auto_name("T.xlsx", ["D:/a/NDE", "D:/a/Laser"], datetime(2026, 10, 23),
                        "update", 12) == "T · NDE + Laser · 23 Oct 2026 · 12 new files"
    assert ls.rename_run(csv, run.id, "October batch")
    assert ls.load(csv).last.name == "October batch"


def test_a_value_the_ledger_holds_is_accepted_again_and_another_flagged(ledger):
    """The Ledger has part 101's power as 285; a new summary says 290."""
    csv, src, template = ledger
    first = li.run_extract(csv, [src], use_images=False)
    accept(first.review, [(None, None, None, None)])
    li.run_apply(first.review, use_images=False)
    assert ParentCSV(csv).get("101", "2B")["Power (W)"] == "285"
    again = li.run_extract(csv, [src], use_images=False)
    held = [i for i in lv.load(again.review).items if i.accept != "Y"]
    assert again.pre_accepted > first.pre_accepted and not held   # all known now
    make_summary(src / "summary.xlsx")
    s = li.run_extract(csv, [src / "summary.xlsx"], use_images=False)
    items = {(i.part, i.fields[0]): i for i in lv.load(s.review).items}
    power = items[("101", "Power (W)")]
    assert power.accept == "" and "differs from the Ledger, which has 285" in power.detail
    assert "differs from the value already in the Ledger" in s.review_reasons
    assert items[("102", "Power (W)")].accept == "Y"     # a part the Ledger lacks


def test_a_flagged_value_replaces_the_old_one_only_when_accepted(ledger):
    csv, src, template = ledger
    first = li.run_extract(csv, [src], use_images=False)
    accept(first.review, [(None, None, None, None)])
    li.run_apply(first.review, use_images=False)
    make_summary(src / "summary.xlsx")
    s = li.run_extract(csv, [src / "summary.xlsx"], use_images=False)
    left = li.run_apply(s.review, use_images=False)          # not accepted: kept
    assert ParentCSV(csv).get("101", "2B")["Power (W)"] == "285"
    assert not any("written over" in p for p in left.problems)
    s = li.run_extract(csv, [src / "summary.xlsx"], use_images=False)
    accept(s.review, [("101", "", "Power (W)", "290")])
    took = li.run_apply(s.review, use_images=False)
    rows = ParentCSV(csv).rows_for("101")
    assert {r["Power (W)"] for r in rows} == {"290"}
    assert any("written over" in p for p in took.problems)
    assert not any("conflict kept" in p and "Power" in p for p in took.problems)


def test_the_filled_template_shows_the_whole_ledger(ledger):
    csv, src, template = ledger
    first = li.run_extract(csv, [src], use_images=False)
    accept(first.review, [(None, None, None, None)])
    li.run_apply(first.review, use_images=False)
    make_summary(src / "summary.xlsx")
    s = li.run_extract(csv, [src / "summary.xlsx"], use_images=False)
    review = lv.load(s.review)
    exact = lv.template_cells(csv)
    path = lv.write_filled(review, "101", Path(s.review).with_name("101.xlsx"), exact)
    wb = openpyxl.load_workbook(path)
    tints = {}
    for (fname, item), (sheet, (r, c)) in exact.items():
        cell = wb[sheet].cell(row=r, column=c)
        if cell.value is not None:
            tints[(fname, item)] = (str(cell.value), cell.fill.start_color.rgb[-6:])
    assert tints[("Power (W)", "")] == ("290", lv.TINT_CLASH[1:])
    fod = tints[("Flaw Attributes: Laser Team FOD (um)", "2B")]
    assert fod == ("404", lv.TINT_LEDGER[1:])            # in the Ledger, not read now
    assert "the Ledger has 285" in wb[exact[("Power (W)", "")][0]].cell(
        *exact[("Power (W)", "")][1]).comment.text


def test_settings_survive_and_a_namer_failure_never_loses_the_run(app_dir, ledger):
    csv, src, template = ledger
    s = li.run_extract(csv, [src], use_images=False)
    run = ls.record_run(csv, s, kind="update", new_files=1,
                        namer=lambda facts, fb: ls.model_name(lambda *_a: ("not json", ""),
                                                              facts, fb))
    assert run.name.endswith("1 new file") and run.kind == "update"
    assert ls.load(csv).sources                        # its files, when none were given
    assert ls.record_run(csv, SimpleNamespace(review="")) is None
