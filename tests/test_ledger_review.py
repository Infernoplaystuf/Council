"""
The Ledger's "Review by clicking…" (council_core.ledger_review and
council_qt/tabs/ledger_review_dialog.py; user 2026-10-22: "something similar
to the existing link by clicking ... Highlight the cells in question"): each
item of the review beside the template and source cells it links, the
answers written back into the review workbook, and Apply taking them.

Run:  python -m pytest tests/test_ledger_review.py -q
"""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("COUNCIL_NO_DIALOGS", "1")

sys.path.insert(0, str(Path(__file__).parent))
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_review as lv  # noqa: E402
from ledger_fixtures import make_laser_log, make_nde_book, make_template  # noqa: E402
from sample_ledger import ParentCSV  # noqa: E402


@pytest.fixture
def sample_review(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    csv = li.create_parent(make_template(tmp_path / "T.xlsx"), tmp_path / "l.csv").csv
    ex = li.run_extract(csv, [src], use_images=False)
    return Path(ex.review), Path(csv), src


@pytest.fixture
def records_review(tmp_path):
    from test_ledger_records import creature_rows, make_sample
    from test_ledger_records import make_template as make_table_template
    src = tmp_path / "s"
    src.mkdir()
    make_sample(src / "BTB_Spider.xlsx", creature_rows("Greg"),
                ["Name", "Health", "Stam", "Weight", "Melee", "Gender"], highlight=False)
    csv = li.create_parent(make_table_template(tmp_path / "BTB_Template.xlsx"),
                           tmp_path / "r.csv").csv
    ex = li.run_extract(csv, [src], use_images=False)
    return Path(ex.review), Path(csv), src


def test_where_is_parsed():
    assert lv.parse_where("laser log.xlsx › Build!B4") == ("laser log.xlsx", "Build", (4, 2))
    assert lv.parse_where("a.xlsx › My Sheet › x!C10") == ("a.xlsx", "My Sheet › x", (10, 3))
    f, s, cell = lv.parse_where("a.xlsx › Sheet1 › image 2 (row 3, col 4)")
    assert (f, cell) == ("a.xlsx", None)


def test_a_sample_review_is_items_with_their_cells(sample_review):
    review, csv, src = sample_review
    r = lv.load(review)
    assert r.kind == "samples" and Path(r.csv) == csv
    assert len(r.items) == 75 and sum(i.needs for i in r.items) == 9
    power = next(i for i in r.items if i.fields == ["Power (W)"])
    assert power.source_file == str(src / "laser_log.xlsx")
    assert (power.source_sheet, power.source_cells) == ("Build", [(5, 3)])
    assert "Laser Power (W)" in power.detail and power.value == "285"


def test_answers_go_into_the_review_and_apply_takes_them(sample_review):
    review, csv, _ = sample_review
    r = lv.load(review)
    power = next(i for i in r.items if i.fields == ["Power (W)"])
    first = next(i for i in r.items if i.needs)
    n = lv.save(review, {(power.sheet, power.row): {"accept": "Y", "value": "290",
                                                     "value_col": power.value_col},
                         (first.sheet, first.row): {"accept": "N"}})
    assert n == 2
    again = lv.load(review)
    p2 = next(i for i in again.items if i.row == power.row)
    assert (p2.accept, p2.value) == ("Y", "290")
    assert next(i for i in again.items if i.row == first.row).accept == "N"
    li.run_apply(review, use_images=False)
    rows = ParentCSV(csv).rows
    assert any(row.get("Power (W)") == "290" for row in rows)


def test_a_table_review_is_one_item_per_record(records_review):
    review, csv, src = records_review
    r = lv.load(review)
    assert r.kind == "records" and len(r.items) == 31
    it = r.items[0]
    assert it.source_file == str(src / "BTB_Spider.xlsx")
    assert it.values["Name"] == "Greg #1" and it.source_cells == [(3, 1)]
    from council_core import ledger_linking as ll
    view = ll.sample_view(it.source_file, r.csv)
    assert lv.record_row_cells(view, it) == [(3, c) for c in range(1, 7)]


# ── the window ─────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def qapp():
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def wait_for(qapp, cond, timeout=30.0):
    end = time.monotonic() + timeout
    while not cond() and time.monotonic() < end:
        qapp.processEvents()
        time.sleep(0.01)
    qapp.processEvents()
    return cond()


def test_the_window_highlights_the_template_and_source_cells(qapp, sample_review):
    from council_qt.tabs.ledger_review_dialog import (CELL_TINT, PICKED_TINT,
                                                      SHOW_ALL, ReviewDialog)
    review, _, src = sample_review
    dlg = ReviewDialog(lv.load(review))
    assert dlg.items.rowCount() == 9                         # what needs a decision first
    dlg.show_choice.setCurrentText(SHOW_ALL)
    assert dlg.items.rowCount() == 75
    power = next(i for i in dlg.review.items if i.fields == ["Power (W)"])
    dlg.select(power)
    assert wait_for(qapp, lambda: dlg.template is not None and dlg._source_models)
    assert wait_for(qapp, lambda: any(m.tints.get((5, 3)) == CELL_TINT
                                      for m in dlg._source_models))
    build = next(m for m in dlg._source_models if m.sheet.name == "Build")
    assert build.tints == {(5, 3): CELL_TINT}
    assert dlg.source_tabs.tabText(dlg.source_tabs.currentIndex()) == "Build"
    flaw = next(i for i in dlg.review.items if i.item == "2A" and "FOD" in i.fields[0])
    dlg.select(flaw)
    exact = [(m.sheet.name, rc) for m in dlg._template_models
             for rc, t in m.tints.items() if t == CELL_TINT]
    assert exact == [dlg._exact[(flaw.fields[0], "2A")]]   # the 2A cell, not the whole row
    dlg.select(power)
    picked = [rc for m in dlg._template_models for rc, t in m.tints.items()
              if t in (PICKED_TINT, CELL_TINT)]
    fields = {m.sheet.fields[rc] for m in dlg._template_models for rc in picked
              if rc in m.sheet.fields}
    assert picked and fields == {"Power (W)"}
    dlg.close()


def test_right_wrong_and_a_correction_are_saved(qapp, sample_review):
    from council_qt.tabs.ledger_review_dialog import ReviewDialog
    review, _, _ = sample_review
    saved = []
    dlg = ReviewDialog(lv.load(review), on_saved=saved.append)
    first, second = dlg.shown[0], dlg.shown[1]
    dlg.select(first)
    dlg.value_edit.setText("123")
    dlg.on_right()
    assert dlg.current is second                             # the next one comes up
    dlg.on_wrong()
    assert dlg.save_btn.isEnabled()
    assert dlg.on_save() and saved == [2]
    again = {(i.sheet, i.row): i for i in lv.load(review).items}
    assert (again[(first.sheet, first.row)].accept,
            again[(first.sheet, first.row)].value) == ("Y", "123")
    assert again[(second.sheet, second.row)].accept == "N"
    dlg.close()


def test_the_window_shows_a_records_row(qapp, records_review):
    from council_qt.tabs.ledger_review_dialog import ROW_TINT, SHOW_ALL, ReviewDialog
    review, _, _ = records_review
    dlg = ReviewDialog(lv.load(review))
    dlg.show_choice.setCurrentText(SHOW_ALL)
    dlg.select(dlg.review.items[2])
    assert wait_for(qapp, lambda: any(m.tints for m in dlg._source_models))
    m = next(m for m in dlg._source_models if m.tints)
    assert m.tints == {(5, c): ROW_TINT for c in range(1, 7)}
    assert dlg.value_edit.isHidden()
    dlg.close()


# ── why values are to review (user, 2026-10-23: 7312 of 7976 to review) ──
def test_the_run_says_why_values_are_to_review(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    csv = li.create_parent(make_template(tmp_path / "T.xlsx"), tmp_path / "l.csv").csv
    ex = li.run_extract(csv, [src], use_images=False)
    assert ex.to_review == 9
    assert ex.review_reasons.splitlines()[:2] == [
        "Why 9 values are to review (the first reason each was held back):",
        "       9  label only a loose match for the field"]
    assert "Fields with the most to review: " in ex.review_reasons
    assert ex.review_reasons in ex.describe()
    assert "285" not in ex.review_reasons                  # counts and names, no values


def test_each_hold_back_has_its_reason(tmp_path):
    from dataclasses import replace
    import source_extract as se
    schema = li.read_template(make_template(tmp_path / "T.xlsx"))
    base = se.Obs(part="P1", item="", field="Power (W)", value="285", score=0.95,
                  via="match", label="Laser Power (W)", where="a.xlsx › S!C5")
    cases = {
        "different workbooks disagree on the value": se.Decision(base, False, "290"),
        "no part number found": se.Decision(replace(base, part=""), False),
        "read from a picture": se.Decision(replace(base, from_image=True), False),
        "part taken from the file name": se.Decision(
            replace(base, part_how=se.FILE_PART_HOW + " a.xlsx"), False),
        "label placed by the model": se.Decision(replace(base, via="model"), False),
        "label only a loose match for the field": se.Decision(replace(base, score=0.6), False),
        "": se.Decision(base, True),
    }
    for why, d in cases.items():
        assert se.review_reason(d, schema) == why
    text = se.summarize_review_reasons(list(cases.values()), schema)
    assert text.startswith("Why 6 values are to review") and "Power (W) (6)" in text


def test_one_workbooks_readings_are_one_first_and_the_rest_set_aside(tmp_path):
    """user, 2026-10-23: 'multiple readings per sample in the source while
    in the template it only accepts one reading ... prioritize the first'.
    Two workbooks that disagree are still the reviewer's; '12' and '12.0'
    are one reading."""
    import source_extract as se
    schema = li.read_template(make_template(tmp_path / "T.xlsx"))
    def obs(value, where, inspector=""):
        return se.Obs(part="P1", item="", field="Power (W)", value=value, score=0.95,
                      via="match", label="Laser Power (W)", where=where,
                      inspector=inspector)
    run = se.Extraction(sources=[])
    run.obs = [obs("285", "a.xlsx › S!C5"), obs("290", "a.xlsx › S!C9", "AK")]
    first, second = sorted((d for d in se.decide(run, schema) if d.obs.field == "Power (W)"),
                           key=lambda d: d.obs.value)
    assert first.accept and not first.conflict and "the first reading" in first.obs.note
    assert not second.accept and second.alternate.startswith("set aside")
    text = se.summarize_review_reasons([first, second], schema)
    assert text.startswith("1 other readings set aside") and "AK (1)" in text
    assert se.review_reason(second, schema) == ""
    ak = {d.obs.value: d for d in se.decide(run, schema, prefer=se.parse_prefer("jb, ak"))
          if d.obs.field == "Power (W)"}
    assert ak["290"].accept and ak["285"].alternate
    run.obs = [obs("285", "a.xlsx › S!C5"), obs("290", "b.xlsx › S!C5")]
    assert all(d.conflict and not d.accept for d in se.decide(run, schema))
    run.obs = [obs("12", "a.xlsx › S!C5"), obs("12.0", "b.xlsx › S!C5")]
    (one,) = se.decide(run, schema)
    assert one.accept and not one.conflict


def test_an_operator_is_their_initials_or_their_name():
    """user, 2026-10-23: 'For operator sometimes its initials sometimes its
    a name'."""
    import source_extract as se
    assert se.parse_prefer("JB, Alex King; mk") == ["JB", "Alex King", "mk"]
    for a, b in (("JB", "Jane Brown"), ("Jane Brown", "J.B."), ("jb", "JB"),
                 ("Brown", "Jane Brown"), ("jane brown", "Jane Brown"),
                 ("Mary-Jo Smith", "MJS")):
        assert se.same_person(a, b) and se.same_person(b, a), (a, b)
    for a, b in (("JB", "Alex King"), ("JB", "AK"), ("Jane Brown", "John Brownlee"), ("", "JB")):
        assert not se.same_person(a, b), (a, b)
    assert se._prefer_rank("Jane Brown", ["AK", "JB"]) == 1
    assert se._prefer_rank("AK/JB", ["Jane Brown"]) == 0
    assert se._prefer_rank("MK", ["Jane Brown"]) is None


# ── the filled template (user, 2026-10-23) ──────────────────────────────
def test_extract_writes_the_template_filled_in_and_tinted_by_confidence(sample_review):
    import openpyxl
    review, csv, src = sample_review
    folder = lv.filled_dir(review)
    assert sorted(p.name for p in folder.iterdir()) == ["101.xlsx"]
    ws = openpyxl.load_workbook(folder / "101.xlsx").active
    exact = lv.template_cells(csv)
    sheet, (r, c) = exact[("Power (W)", "")]
    cell = ws.cell(r, c)
    assert cell.value == 285 and cell.fill.start_color.rgb.endswith(lv.TINT_HIGH[1:])
    assert "laser_log.xlsx › Build!C5" in cell.comment.text
    fod = next(i for i in lv.load(review).items if i.item == "2A" and "FOD" in i.fields[0])
    sheet, (r, c) = exact[(fod.fields[0], "2A")]
    assert ws.cell(r, c).fill.start_color.rgb.endswith(lv.confidence_tint(fod)[1:])
    assert lv.confidence_tint(fod) == lv.TINT_MID                # score 0.83


def test_cell_candidates_put_the_accepted_value_first():
    a = lv.Item("Values", 2, "", True, True, "a", "", fields=["F"], value="1", part="P",
                score=0.95)
    b = lv.Item("Values", 3, "Y", False, True, "b", "", fields=["F"], value="2", part="P",
                score=0.6)
    no = lv.Item("Values", 4, "N", False, False, "c", "", fields=["F"], value="3", part="P",
                 score=1.0)
    exact = {("F", ""): ("Sheet", (1, 1))}
    got = lv.cell_candidates([a, b, no], exact, "P")
    assert got == {("Sheet", (1, 1)): [b, a]}
    assert lv.cell_candidates([a, b], exact, "P", accept_of=lambda it: "N" if it is b
                              else it.accept) == {("Sheet", (1, 1)): [a]}
    assert lv.confidence_tint(a) == lv.TINT_CLASH


def test_the_filled_template_window_shows_a_cells_source(qapp, sample_review):
    from council_qt.tabs.ledger_review_dialog import CELL_TINT, FilledTemplateDialog
    review, csv, src = sample_review
    opened = []
    dlg = FilledTemplateDialog(lv.load(review), opener=opened.append)
    assert dlg.part == "101"
    assert wait_for(qapp, lambda: dlg.template is not None and dlg._exact)
    dlg._tint_template()
    sheet, rc = dlg._exact[("Power (W)", "")]
    m = next(m for m in dlg._template_models if m.sheet.name == sheet)
    assert m.texts[rc] == "285" and "Build!C5" in m.tips[rc]
    assert dlg.click_cell(sheet, rc)
    assert [i.fields for i in dlg.shown] == [["Power (W)"]]
    assert dlg.current.value == "285" and m.tints[rc] == CELL_TINT
    assert wait_for(qapp, lambda: any(sm.tints.get((5, 3)) == CELL_TINT
                                      for sm in dlg._source_models))
    dlg.on_wrong()                                          # said wrong: the cell empties
    assert rc not in m.texts
    assert dlg.on_save()
    assert next(i for i in lv.load(review).items if i.fields == ["Power (W)"]).accept == "N"
    assert not dlg.click_cell(sheet, (1, 1))
    path = dlg.on_export()
    assert path and opened == [path] and Path(path).name == "101.xlsx"
    dlg.show_whole_part()
    assert len(dlg.shown) == sum(1 for i in dlg.review.items if i.part == "101")
    dlg.close()


def test_answers_carry_into_a_new_review_of_the_same_sources(sample_review):
    """Re-run with my answers: what was said Right or Wrong is not asked
    again, and a correction stays."""
    import shutil
    review, csv, src = sample_review
    new = review.with_name("review_again.xlsx")
    shutil.copy(review, new)
    old = lv.load(review)
    held = [i for i in old.items if i.needs and not i.clash]
    green = next(i for i in old.items if i.accept == "Y" and not i.needs)
    right, fixed = held[0], held[1]
    lv.save(review, {(right.sheet, right.row): {"accept": "Y"},
                     (fixed.sheet, fixed.row): {"accept": "Y", "value": "999",
                                                "value_col": fixed.value_col},
                     (green.sheet, green.row): {"accept": "N"}})
    assert lv.carry_answers(lv.load(review), lv.load(new)) == 3
    got = {(i.sheet, i.row): i for i in lv.load(new).items}
    assert got[(right.sheet, right.row)].accept == "Y"
    assert (got[(fixed.sheet, fixed.row)].accept, got[(fixed.sheet, fixed.row)].value) == \
        ("Y", "999")
    assert got[(green.sheet, green.row)].accept == "N"      # said Wrong stays Wrong
    untouched = held[2]
    assert got[(untouched.sheet, untouched.row)].accept == ""
    assert lv.carry_answers(lv.load(review), lv.load(new)) == 0   # already carried


def test_the_filled_template_window_re_runs_with_the_answers(qapp, sample_review):
    from council_qt.tabs.ledger_review_dialog import FilledTemplateDialog
    review, csv, src = sample_review
    assert FilledTemplateDialog(lv.load(review)).rerun_btn.isHidden()
    asked = []
    dlg = FilledTemplateDialog(lv.load(review), on_rerun=lambda r, p: asked.append((r, p)))
    assert not dlg.rerun_btn.isHidden()
    held = next(i for i in dlg.review.items if i.needs and not i.clash)
    dlg.select(held)
    dlg.on_right()
    assert dlg.on_rerun_clicked()
    assert asked == [(str(review), dlg.part)]
    busy = FilledTemplateDialog(lv.load(review), on_rerun=lambda r, p: False)
    assert not busy.on_rerun_clicked() and busy.isVisible() is False   # never shown
    assert next(i for i in lv.load(review).items if i.row == held.row).accept == "Y"


def test_a_value_is_put_where_it_belongs_by_right_click(qapp, sample_review):
    """user, 2026-10-26: 'right click the spot where the data actually
    belongs so I can manually correct things'."""
    from council_qt.tabs.ledger_review_dialog import ReviewDialog
    import source_extract as se
    review, csv, src = sample_review
    dlg = ReviewDialog(lv.load(review))
    assert wait_for(qapp, lambda: dlg.template is not None and dlg._exact)
    freq = dlg._exact[("Frequency (kHz)", "")]
    dlg.current = None
    assert not dlg.place_at(*freq)                           # nothing picked
    power = next(i for i in dlg.review.items if i.fields == ["Power (W)"])
    dlg.select(power)
    assert not dlg.place_at(freq[0], (1, 1))                 # not a field's cell
    there = [i for i in dlg.review.items if i.fields == ["Frequency (kHz)"]
             and i.part == power.part and i.item == power.item and i.accept == "Y"]
    assert there                                             # the log's own frequency
    assert dlg.place_at(*freq, teach=True)
    assert power.fields == ["Frequency (kHz)"] and "Frequency" in power.title
    assert all(dlg._accept_of(i) == "N" for i in there)      # the cell takes one value
    dlg.select(power)
    dlg.on_wrong()                                           # an answer keeps the move
    dlg.select(power)
    dlg.on_right()
    assert dlg.answers[(power.sheet, power.row)]["field"] == "Frequency (kHz)"
    assert dlg.on_save()
    again = next(i for i in lv.load(review).items if i.row == power.row)
    assert (again.fields, again.accept, again.value) == (["Frequency (kHz)"], "Y", "285")
    assert all(next(i for i in lv.load(review).items if i.row == o.row).accept == "N"
               for o in there)
    mapping, _ = se.load_aliases(csv)
    assert mapping[se.normalize_label("Laser Power (W)")] == "Frequency (kHz)"
    dlg.close()

    li.run_apply(review, use_images=False)
    row = ParentCSV(csv).get("101", "")
    assert row["Frequency (kHz)"] == "285" and row["Power (W)"] != "285"
    s = li.run_extract(csv, [src], use_images=False)          # learned for next time
    moved = [i for i in lv.load(s.review).items if i.where == power.where]
    assert [i.fields for i in moved] == [["Frequency (kHz)"]]


def test_the_filled_template_window_offers_put_it_here(qapp, sample_review):
    from council_qt.tabs.ledger_review_dialog import FilledTemplateDialog
    review, csv, src = sample_review
    dlg = FilledTemplateDialog(lv.load(review))
    assert wait_for(qapp, lambda: dlg.template is not None and dlg._exact)
    sheet, rc = dlg._exact[("Frequency (kHz)", "")]
    dlg.current = None
    assert [d for _t, d in dlg.menu_actions(sheet, rc)] == [None]   # pick a value first
    power = next(i for i in dlg.review.items
                 if i.fields == ["Power (W)"] and i.part == dlg.part)
    dlg.select(power)
    assert [d for _t, d in dlg.menu_actions(sheet, (1, 1))] == [None]
    acts = dlg.menu_actions(sheet, rc)
    assert [t for t, _d in acts] == ["Put this value here (Frequency (kHz))",
                                     "Put it here and always map 'Laser Power (W)' here"]
    acts[0][1]()
    assert power.fields == ["Frequency (kHz)"]
    assert dlg.answers[(power.sheet, power.row)] == {"accept": "Y", "field": "Frequency (kHz)",
                                                     "item": ""}
    assert dlg.on_save()
    from source_extract import load_aliases, normalize_label
    assert normalize_label("Laser Power (W)") not in load_aliases(csv)[0]   # not taught
    dlg.close()
