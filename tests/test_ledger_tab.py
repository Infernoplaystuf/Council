"""
The 📑 Ledger tab, driven like a user would: template -> parent CSV, a folder
of sources -> review workbook, the review checked -> apply with filled
templates. Offscreen (no window), no Excel, no file dialogs, a stand-in image
reader.

Run:  python -m pytest tests/test_ledger_tab.py -q
"""
from __future__ import annotations

import os
import sys
import threading
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")
from PySide6.QtWidgets import QApplication  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
from ledger_fixtures import (CONCEALMENT_IMAGE_ROWS, ITEMS, fake_reader,  # noqa: E402
                             make_laser_log, make_nde_book, make_summary, make_template)
from test_ledger_ingest import accept  # noqa: E402

from council_qt.tabs.ledger import LedgerActions, LedgerTab, build_ledger  # noqa: E402
from sample_ledger import ParentCSV  # noqa: E402


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def books(tmp_path):
    src = tmp_path / "sources"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    make_summary(src / "summary.xlsx")
    return make_template(tmp_path / "LMM Template.xlsx"), src, tmp_path


def make_tab(qapp, **kw):
    opened = []
    actions = LedgerActions(image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS),
                            opener=opened.append)
    for name, fn in kw.items():
        setattr(actions, name, fn)
    tab = LedgerTab(None, actions)
    return tab, opened


def test_the_whole_flow_from_the_tab(qapp, books):
    template, src, tmp = books
    tab, opened = make_tab(qapp)
    assert not tab.create_btn.isEnabled() and not tab.extract_btn.isEnabled()

    tab.template.setText(str(template))
    tab.create_btn.click()
    assert tab.wait()
    csv = Path(tab.csv.text())
    assert csv == tmp / "LMM_Template_parent.csv" and csv.exists()
    assert "created" in tab.status.text()

    tab.add_sources([src])
    tab.add_sources([src])                               # once only
    assert tab.source_paths() == [str(src)] and tab.extract_btn.isEnabled()
    tab.extract_btn.click()
    assert tab.wait()
    review = Path(tab.review.text())
    assert review.exists() and review.parent == tmp
    values, pre, todo = (int(tab.stat[k].text()) for k in ("values", "pre_accepted",
                                                            "to_review"))
    assert values == pre + todo and pre > 0 and todo > 0
    assert tab.stat["images"].text() == "1 / 1"
    assert tab.stat["problems"].text() == "0"
    assert "values found" in tab.log.toPlainText()

    tab.open_btn.click()
    assert opened == [str(review)]

    accept(review, [("101", "", "Power (W)", "285"),
                    ("101", None, "Flaw Attributes: Laser Team FOD (um)", None),
                    ("101", "", "Concealment 1: Media", None)])
    tab.fill_dir.setText(str(tmp / "filled"))
    tab.apply_btn.click()
    assert tab.wait()
    assert "rows added" in tab.status.text()
    ledger = ParentCSV(csv)
    assert ledger.sample_ids() == ["101", "102"]
    assert [r["Flaw #"] for r in ledger.rows_for("101")] == ITEMS
    r = ledger.get("101", "2B")
    assert r["Power (W)"] == "285" and r["Concealment 1: Media"] == "ceramic"
    assert r["Flaw Attributes: Laser Team FOD (um)"] == "404"
    assert sorted(p.name for p in (tmp / "filled").iterdir()) == [
        "LMM_Template_101.xlsx", "LMM_Template_102.xlsx"]


def test_work_runs_off_the_gui_thread_with_buttons_locked(qapp, books):
    template, src, tmp = books
    gate, ran_on = threading.Event(), []

    def slow_extract(csv, sources, part, use_images, progress,
                     pattern_mask=None, vision_model=None, cancel=None):
        ran_on.append(threading.current_thread())
        progress("reading workbooks: 1/2 workbooks (50%)", 0.5)
        gate.wait(10)
        raise RuntimeError("stopped by the test")

    tab, _ = make_tab(qapp, extract=slow_extract)
    tab.csv.setText(str(tmp / "x.csv"))
    tab.add_sources([src])
    tab.extract_btn.click()
    assert tab.busy and tab.progress.isVisibleTo(tab)
    assert not any(b.isEnabled() for b in tab._busy_buttons)
    tab.extract_btn.click()                              # ignored while busy
    for _ in range(50):
        QApplication.processEvents()
    assert "1/2 workbooks (50%)" in tab.status.text()
    assert (tab.progress.maximum(), tab.progress.value()) == (100, 50)   # a real percentage
    gate.set()
    assert tab.wait()
    assert ran_on and ran_on[0] is not threading.main_thread() and len(ran_on) == 1
    assert tab.status.text().startswith("failed: RuntimeError")
    assert "Traceback" in tab.log.toPlainText()
    assert tab.extract_btn.isEnabled() and not tab.progress.isVisibleTo(tab)


def test_pipeline_errors_are_shown_in_plain_words(qapp, books):
    _t, src, tmp = books
    tab, _ = make_tab(qapp)
    tab.csv.setText(str(tmp / "missing.csv"))
    tab.add_sources([src])
    tab.extract_btn.click()
    assert tab.wait()
    assert tab.status.text().startswith("error: parent CSV not found")
    assert "Traceback" not in tab.log.toPlainText()


#: What the tab says for a Cancel after every workbook was read.
_CANCELLED_READ_ALL = ("Cancelled after every workbook was read: every value the rules found "
                       "is in the review. Open it, check the amber and red rows, map the "
                       "leftover labels, save it, then Apply.")


def test_cancel_while_councils_model_is_asked_says_cancelled(qapp, books, monkeypatch):
    """Cancel pressed while Council's model was asked about the labels left
    over: the run stopped, but the tab said 'Review written…' and its log
    that the model had not answered - '(none answered)', 'did not answer:
    stopped by the caller' (review, 2026-10-20). Then it said 'Cancelled.
    The partial review…' of a review that had every value the rules found,
    and one pressed a moment earlier, before the model was asked, 'Review
    written…' (second review, 2026-10-20): both say the same, not
    'partial'."""
    import time
    import openpyxl
    import source_extract as se
    from council_core import ledger_llm
    template, src, _tmp = books
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["A2"], wb.active["B2"] = "P/N: 3307", "Zorp rate", 7
    wb.save(src / "extra.xlsx")                           # a label no rule places
    tab, _ = make_tab(qapp)
    tab.template.setText(str(template))
    tab.create_btn.click()
    assert tab.wait()
    tab.add_sources([src])
    tab.ask_model.setChecked(True)
    tab._pick_model = lambda parent=None: "ollama:granite3.3:8b"
    asked = threading.Event()

    def chat_for(model_id, timeout=120, should_stop=None):
        def chat(prompt, schema):                 # council_engine's Ollama call
            if "Which part" in prompt:
                return '{"part": "", "evidence": ""}', "granite3.3:8b"
            asked.set()
            end = time.monotonic() + 30
            while not should_stop() and time.monotonic() < end:     # generating
                time.sleep(0.02)
            raise se.LedgerError("granite3.3:8b (Ollama) did not answer: stopped by the "
                                 "caller")
        return chat
    monkeypatch.setattr(ledger_llm, "chat_for", chat_for)
    tab.extract_btn.click()
    end = time.monotonic() + 60
    while not asked.is_set() and time.monotonic() < end:
        QApplication.processEvents()
        time.sleep(0.01)
    assert asked.is_set() and tab.cancel_btn.isEnabled()
    tab.cancel_btn.click()
    assert tab.wait()
    assert tab.status.text() == _CANCELLED_READ_ALL
    log = tab.log.toPlainText()
    assert "did not answer" not in log and "none answered" not in log
    assert "Cancel stopped it while it was asked about the labels left over" in log
    assert "CANCELLED after every workbook was read" in log
    assert ("Cancelled by user while Council's model was asked about the labels left over; "
            "every value the rules found is in this review.") in log


def test_cancel_before_councils_model_is_asked_says_the_same(qapp, books, monkeypatch):
    """Pressed after the last workbook was read and before Council's model
    was asked: the same words as pressed while it is asked - it said
    'Review written…' (second review, 2026-10-20)."""
    import openpyxl
    from council_core import ledger_ingest as li
    from council_core import ledger_llm
    template, src, _tmp = books
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["A2"], wb.active["B2"] = "P/N: 3307", "Zorp rate", 7
    wb.save(src / "extra.xlsx")                           # a label no rule places
    tab, _ = make_tab(qapp)
    tab.template.setText(str(template))
    tab.create_btn.click()
    assert tab.wait()
    tab.add_sources([src])
    tab.ask_model.setChecked(True)
    tab._pick_model = lambda parent=None: "ollama:granite3.3:8b"
    asked = []

    def chat_for(model_id, timeout=120, should_stop=None):
        def chat(prompt, schema):
            asked.append(prompt)
            return '{"part": "", "evidence": ""}', "granite3.3:8b"
        return chat
    monkeypatch.setattr(ledger_llm, "chat_for", chat_for)
    real = li.extract

    def read_then_cancel(*a, **k):                         # pressed as the last one is read
        run = real(*a, **k)
        k["cancel"].set()
        return run
    monkeypatch.setattr(li, "extract", read_then_cancel)
    tab.extract_btn.click()
    assert tab.wait()
    assert not any("Which part" not in p for p in asked)   # not asked about the labels
    assert tab.status.text() == _CANCELLED_READ_ALL
    log = tab.log.toPlainText()
    assert "CANCELLED after every workbook was read" in log and "partial" not in log
    assert ("Cancelled before Council's model was asked about the labels left over; "
            "every value the rules found is in this review.") in log


def test_sources_can_be_removed_and_cleared(qapp, books):
    _t, src, _tmp = books
    tab, _ = make_tab(qapp)
    tab.csv.setText("p.csv")
    tab.add_sources([src / "summary.xlsx", src / "laser_log.xlsx"])
    tab.sources.item(0).setSelected(True)
    tab.remove_selected()
    assert tab.source_paths() == [str(src / "laser_log.xlsx")]
    tab.clear_sources()
    assert tab.source_paths() == [] and not tab.extract_btn.isEnabled()


def test_registry_factory(qapp):
    assert isinstance(build_ledger(None), LedgerTab)


def test_diagnose_a_file_runs_it_apart_and_opens_the_report(qapp, books, monkeypatch, tmp_path):
    """"Diagnose a file…" (council_core.ledger_diagnose): the file picked is
    read in a separate process, each line shows in the log as it comes, and
    the report is opened to send back."""
    monkeypatch.setenv("COUNCIL_APP_DIR", str(tmp_path / "app"))
    template, src, _ = books
    tab, opened = make_tab(qapp)
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    picked = src / "laser_log.xlsx"
    monkeypatch.setattr(tab, "_browse_diagnose_file", lambda start: str(picked))
    assert tab.diagnose_btn.isEnabled()
    tab.on_diagnose()
    assert tab.wait(300)
    log = tab.log.toPlainText()
    assert "== step: table reader" in log
    assert "Diagnosis of laser_log.xlsx: finished normally" in log
    assert opened and Path(opened[-1]).name == "report.txt"
    assert "whole Extract into" in Path(opened[-1]).read_text(encoding="utf-8")


def test_review_by_clicking_opens_over_the_review(qapp, books):
    """After Extract, "Review by clicking…" reads the review workbook and
    opens the click-through window over it (ledger_review_dialog)."""
    template, src, _ = books
    tab, _ = make_tab(qapp)
    assert not tab.review_btn.isEnabled()
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.on_extract()
    assert tab.wait(120)
    assert tab.review_btn.isEnabled()
    tab.on_review_window()
    assert tab.wait(60)
    dlg = tab.review_dialog
    assert dlg.review.path == tab.review.text()
    assert "need a decision" in tab.log.toPlainText()
    dlg.close()


def test_the_operator_to_prefer_goes_into_the_extract(qapp, books, monkeypatch):
    """user, 2026-10-23: one workbook's several readings - the operator
    typed is used first (source_extract.decide), and the review keeps it."""
    template, src, _ = books
    tab, _ = make_tab(qapp)
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.prefer_operator.setText("jb, AK")
    from council_core import ledger_ingest
    seen = {}
    real = ledger_ingest.run_extract

    def spy(*a, **kw):
        seen["prefer"] = kw.get("prefer")
        return real(*a, **kw)
    monkeypatch.setattr(ledger_ingest, "run_extract", spy)
    tab.on_extract()
    assert tab.wait(120)
    assert seen["prefer"] == ["jb", "AK"]
    import json
    import openpyxl
    meta = json.loads(openpyxl.load_workbook(tab.review.text())["_run"]["A1"].value)
    assert meta["prefer"] == ["jb", "AK"]


def test_a_filled_example_is_picked_and_saved_with_the_ledger(qapp, books, monkeypatch,
                                                              tmp_path):
    template, _src, _ = books
    tab, _ = make_tab(qapp)
    tab.on_pick_filled_example()                           # no parent CSV yet
    assert "parent CSV first" in tab.status.text()
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    import sample_ledger as sl
    schema = sl.ParentCSV(tab.csv.text()).schema
    example = tmp_path / "filled.xlsx"
    sl.fill_template(schema, [{schema.id_field: "101", schema.item_field: it,
                               "Power (W)": "285"} for it in schema.items], example)
    monkeypatch.setattr(tab, "_browse_filled_example", lambda: str(example))
    tab.on_pick_filled_example()
    assert tab.filled_example.text() == str(example)
    assert "9 row(s) for 101" in tab.log.toPlainText()
    from council_core import ledger_examples
    assert ledger_examples.load_filled_example(tab.csv.text()) == str(example)
    tab.forget_example_btn.click()
    assert tab.filled_example.text() == "" and not tab.forget_example_btn.isEnabled()


def test_the_filled_template_opens_after_extract(qapp, books):
    """user, 2026-10-23: 'a pop-out of the output filled in template excel
    so I can click on cells and see their source'."""
    template, src, _ = books
    tab, _ = make_tab(qapp)
    assert not tab.filled_btn.isEnabled()
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.on_extract()
    assert tab.wait(120)
    assert "The template filled in for" in tab.log.toPlainText()
    assert tab.filled_btn.isEnabled()
    tab.on_filled_window()
    assert tab.wait(60)
    dlg = tab.filled_dialog
    assert dlg.part and dlg.review.path == tab.review.text()
    dlg.close()


def test_re_run_with_my_answers(qapp, books):
    """user, 2026-10-23: 'look at a few values and then give an answer and
    have it rerun and fill things in more confidently'."""
    from council_core import ledger_review as lv
    template, src, _ = books
    tab, _ = make_tab(qapp)
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.on_extract()
    assert tab.wait(120)
    first = tab.review.text()
    held = [i for i in lv.load(first).items if i.needs]
    tab.on_filled_window()
    assert tab.wait(60)
    dlg = tab.filled_dialog
    fod = [i for i in held if i.part == "101" and i.fields[0].startswith("Flaw Attributes")]
    date = next(i for i in held if i.part == "101" and i.fields == ["Date"])
    assert len(fod) == 9
    for it in fod[:3] + [date]:                       # a few values said Right
        dlg.select(it)
        dlg.on_right()
    assert dlg.on_rerun_clicked()
    assert tab.wait(180)
    again = tab.review.text()
    assert again != first and Path(first).exists()     # a new review; the old one kept
    assert ParentCSV(tab.csv.text()).get("101", "")["Date"] == "2026-10-01"
    now = {(i.part, i.item, tuple(i.fields), i.where): i for i in lv.load(again).items}
    # learned from the three: every reading of that label is pre-accepted now
    assert all(now[(i.part, i.item, tuple(i.fields), i.where)].accept == "Y" for i in fod)
    assert now[("101", "", ("Date",), date.where)].accept == "Y"     # not asked again
    assert int(tab.stat["to_review"].text()) < len(held)
    assert "Re-run with your answers" in tab.log.toPlainText()
    assert tab.wait(60)
    assert tab.filled_dialog is not dlg and tab.filled_dialog.review.path == again
    assert tab.filled_dialog.part == dlg.part
    tab.filled_dialog.close()


def test_re_run_with_my_answers_from_the_tab(qapp, books):
    """user, 2026-10-26: 'Why don't we put that button in the extract and
    review section in the tab as well'."""
    from council_core import ledger_review as lv
    template, src, _ = books
    tab, _ = make_tab(qapp)
    assert not tab.rerun_btn.isEnabled()                    # no review yet
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.on_extract()
    assert tab.wait(120)
    first = tab.review.text()
    assert tab.rerun_btn.isEnabled()
    held = [i for i in lv.load(first).items if i.needs]
    date = next(i for i in held if i.part == "101" and i.fields == ["Date"])
    lv.save(first, {(date.sheet, date.row): {"accept": "Y"}})   # an answer saved
    assert tab.on_rerun_from_tab()
    assert tab.wait(180)
    again = tab.review.text()
    assert again != first and Path(first).exists()
    assert ParentCSV(tab.csv.text()).get("101", "")["Date"] == "2026-10-01"
    now = {(i.part, i.item, tuple(i.fields), i.where): i for i in lv.load(again).items}
    assert now[("101", "", ("Date",), date.where)].accept == "Y"     # not asked again
    assert "Re-run with your answers" in tab.log.toPlainText()
    assert getattr(tab, "filled_dialog", None) is None      # no window opened


def test_a_saved_ledger_is_picked_again_and_updated_with_new_files(qapp, tmp_path):
    """user, 2026-10-23: 'the folders it was run on get more data in them I
    want to be able to rerun from that stage'; runs 'saved, catergorized,
    and named by the council', grouped 'by template and or project'."""
    from council_core import ledger_saved
    from council_qt.tabs.ledger_runs_dialog import LedgerRunsDialog
    src = tmp_path / "sources"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    template = make_template(tmp_path / "LMM Template.xlsx")
    tab, _ = make_tab(qapp)
    assert not tab.update_btn.isEnabled() and not tab.runs_btn.isEnabled()
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.prefer_operator.setText("JB")
    tab.on_extract()
    assert tab.wait(120)
    csv = tab.csv.text()
    assert "Saved as \"LMM Template · sources ·" in tab.log.toPlainText()
    assert tab.update_btn.isEnabled() and tab.runs_btn.isEnabled()
    tab.ledger_project.setText("Apollo")
    tab._save_details()
    tab.group_by.setCurrentIndex(list(ledger_saved.GROUPS).index("project"))
    texts = [tab.saved_choice.itemText(i) for i in range(tab.saved_choice.count())]
    assert texts[0] == "Apollo" and texts[1].strip().startswith("LMM Template")
    assert tab.saved_choice.currentIndex() == 1
    accept(tab.review.text(), [(None, None, None, None)])
    tab.apply_btn.click()
    assert tab.wait(60)

    tab.on_update()                                          # nothing new yet
    assert tab.wait(60)
    assert tab.status.text().startswith("Nothing new")

    other, _ = make_tab(qapp)                                # another day: pick it
    assert other.load_saved(csv)
    assert other.source_paths() == [str(src)] and other.prefer_operator.text() == "JB"
    assert other.template.text() == str(template) and other.ledger_project.text() == "Apollo"
    make_summary(src / "summary.xlsx")
    other.on_update()
    assert other.wait(120)
    assert other.status.text().startswith("Updated with 1 new and 0 changed file(s)")
    assert other.stat["values"].text() == "4"                # the new file's values only
    assert other.wait(60)
    dlg = other.filled_dialog
    assert dlg.review.path == other.review.text() and "already in the Ledger" in \
        dlg.template_note.text()
    dlg.close()
    led = ledger_saved.load(csv)
    assert [(r.kind, r.files, r.new_files) for r in led.runs] == [("extract", 2, 0),
                                                                 ("update", 1, 1)]
    assert led.runs[-1].name.endswith("· 1 new file")

    runs = LedgerRunsDialog(csv, open_filled=lambda r: None, opener=lambda r: None)
    assert runs.table.rowCount() == 2 and runs.table.item(0, 2).text() == \
        "Update with new files"
    runs.rename(0, "October summary")
    assert ledger_saved.load(csv).runs[-1].name == "October summary"
    runs.close()


def test_groups_are_made_in_review_and_apply_writes_their_workbooks(qapp, books):
    """user, 2026-10-26: sub-groups made in review, each with smaller output
    workbooks 'in addition to the full workbook', a filled-template sheet per
    part in numerical order."""
    import openpyxl
    from council_core import ledger_groups as lg
    template, src, tmp = books
    tab, _ = make_tab(qapp)
    tab.template.setText(str(template))
    tab.on_create()
    assert tab.wait(60)
    tab.add_sources([src])
    tab.on_extract()
    assert tab.wait(120)
    assert tab.groups_btn.isEnabled()
    tab.on_filled_window()
    assert tab.wait(60)
    filled = tab.filled_dialog
    assert not filled.groups_btn.isHidden()
    filled.groups_btn.click()                          # from the review step
    dlg = tab.groups_dialog
    assert sorted(dlg.facts) == ["101", "102"]         # the review's parts
    g = dlg.new_group("Old 101")
    dlg.set_rule("part", pattern="101")
    assert dlg.members() == ["101"]
    dlg.tick("102", True)
    assert dlg.members() == ["101", "102"] and g.include == ["102"]
    dlg.tick("102", False)
    dlg.set_field("Concealment 1: Media", lg.YES_NO, "ceramic")
    dlg.set_field("Date", lg.OFF)
    assert dlg.on_save()
    dlg.close()
    filled.close()
    assert [x.name for x in lg.load_groups(tab.csv.text())] == ["Old 101"]

    accept(tab.review.text(), [(None, None, None, None)])
    tab.fill_dir.setText(str(tmp / "filled"))
    tab.apply_btn.click()
    assert tab.wait(60)
    assert "Workbooks (2), a filled-template sheet per part" in tab.log.toPlainText()
    out = Path(tab.csv.text()).with_name("LMM_Template_parent_workbooks")
    full = openpyxl.load_workbook(out / "LMM Template - all parts.xlsx")
    group = openpyxl.load_workbook(out / "LMM Template - Old 101.xlsx")
    assert full.sheetnames == ["101", "102"] and group.sheetnames == ["101"]
    from sample_ledger import TemplateSchema, schema_path_for
    schema = TemplateSchema.load(schema_path_for(tab.csv.text()))
    media = next(f for f in schema.fields if f.name == "Concealment 1: Media").cells[""]
    assert full["101"][media].value == "ceramic" and group["101"][media].value == "Yes"
