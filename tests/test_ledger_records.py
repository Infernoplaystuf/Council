"""The table Ledger (council_core.ledger_records): a template whose rows are
records, filled from the tables of many workbooks. Synthetic workbooks only;
the BTB set is rebuilt here in the shape of the user's artificial example
(2026-10-20): a highlighted sample, one with no highlights and fewer columns,
a lookup list, and a template whose blue column comes from the file name."""
from __future__ import annotations

import json
import time
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.styles.colors import Color

from council_core import ledger_ingest as li
from council_core import ledger_records as lr

BLUE, YELLOW, GREEN, CYAN = "0000FF", "FFFF00", "00FF00", "00FFFF"
STATS = ["Health", "Stam", "Oxygen", "Food", "Weight", "Melee", "Gender"]


def paint(cell, rgb):
    cell.fill = PatternFill("solid", fgColor="FF" + rgb)


def creature_rows(prefix: str, n: int = 31, fourteenth: str = ""):
    rows = []
    for k in range(1, n + 1):
        name = fourteenth if (k == 14 and fourteenth) else f"{prefix} #{k}"
        rows.append({"Name": name, "Health": [6300.0, 4480.0, 7700.0, 7980.0][k % 4],
                     "Stam": 1820.0 if k in (22, 31) else 1295.0,
                     "Oxygen": {25: 650.0, 29: 780.0, 31: 2860.0}.get(k, 3575.0),
                     "Food": 7020.0 if k in (1, 22) else 10400.0,
                     "Weight": 532.0 if k == 31 else 574.0,
                     "Melee": [337.5, 125.8, 349.3][k % 3], "Gender": "MF"[k % 2]})
    return rows


def make_template(path: Path, rows: int = 12) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = [("Tame", BLUE), ("Name", YELLOW)] + [(s, GREEN) for s in STATS] + \
        [("Saddle BP", CYAN)]
    for c, (label, rgb) in enumerate(cols, start=1):
        ws.cell(1, c, label)
        for r in range(1, rows + 2):
            paint(ws.cell(r, c), rgb)
    wb.save(path)
    return path


def make_sample(path: Path, rows, columns, highlight: bool, top: int = 2, left: int = 1,
                title: str = "") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    if title:
        ws.cell(1, left, title)
    for j, h in enumerate(columns):
        ws.cell(top, left + j, h)
    for i, row in enumerate(rows, start=1):
        for j, h in enumerate(columns):
            ws.cell(top + i, left + j, row.get(h))
    if highlight:
        for i in range(len(rows) + 1):
            for j, h in enumerate(columns):
                paint(ws.cell(top + i, left + j), YELLOW if h == "Name" else GREEN)
    wb.save(path)
    return path


def make_saddles(path: Path, highlight: bool = True) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    data = [("Tame Saddle", "wood", "metal", "fiber"), ("Greg", 35, 12, 1125),
            ("Anky", 38, 450, 24234), ("Spino", 74, 269, 2456), ("Woop Woop", 37, 420, 876),
            ("Rex", 29, 280, 345), ("Trevor", 80085, 169, 420),
            ("Harrison", "N/A", "N/A", "N/A")]
    for r, row in enumerate(data, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
        if highlight:
            paint(ws.cell(r, 1), CYAN)
    wb.save(path)
    return path


@pytest.fixture
def btb(tmp_path):
    src = tmp_path / "samples"
    src.mkdir()
    woop = creature_rows("Woop Woop", fourteenth="Woop Woop Fourt")
    greg = creature_rows("Greg")
    make_sample(src / "BTB_Woopentide.xlsx", woop, ["Name"] + STATS, highlight=True)
    make_sample(src / "BTB_Spider.xlsx", greg,
                ["Name", "Health", "Stam", "Weight", "Melee", "Gender"], highlight=False)
    make_saddles(src / "BTB_Saddle_BP.xlsx")
    tpl = make_template(tmp_path / "BTB_Template.xlsx")
    return tpl, src, tmp_path, {"Spider": greg, "Woopentide": woop}


def decide_all(review: Path, answer: str = "Y", only=None) -> None:
    wb = openpyxl.load_workbook(review)
    for row in wb["Decisions"].iter_rows(min_row=2):
        if only is None or str(row[0].value).startswith(only):
            row[1].value = answer
    wb.save(review)


def set_accept(review: Path, source: str, row_no: int, answer: str) -> None:
    wb = openpyxl.load_workbook(review)
    ws = wb["Records"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[head.index("Source")].value == source and row[head.index("Row")].value == row_no:
            row[0].value = answer
    wb.save(review)


def records_sheet(review: Path):
    wb = openpyxl.load_workbook(review)
    ws = wb["Records"]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    fills = {}
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if c.fill and c.fill.fill_type == "solid":
                fills[(c.row, head[c.column - 1])] = c.fill.fgColor.rgb[-6:]
    return head, rows, fills


# ── the BTB example ─────────────────────────────────────────────────────
def test_btb_reproduces_the_filled_template_cell_for_cell(btb):
    tpl, src, tmp, data = btb
    s = li.create_parent(tpl)                        # no sample-ID column: a table Ledger
    assert s.records_mode and "one row per record" in s.describe()
    csv = Path(s.csv)
    ex = li.run_extract(csv, [src], use_images=False)
    assert ex.records == 62 and ex.files == 3
    assert any("highlights in BTB_Woopentide.xlsx" in x for x in ex.learned)
    assert any("Saddle BP is looked up in BTB_Saddle_BP.xlsx" in x for x in ex.learned)
    assert any(x.startswith("Tame: nothing in the workbooks") for x in ex.learned)
    decide_all(Path(ex.review))                       # the user: yes, Tame is the file name
    ap = li.run_apply(ex.review, tmp / "filled", use_images=False)
    assert ap.rows == 62
    out = openpyxl.load_workbook(tmp / "filled" / "BTB_Template_filled.xlsx").active
    head = [c.value for c in out[1]]
    expect = []
    for tame in ("Spider", "Woopentide"):          # files in name order, rows in sheet order
        for row in data[tame]:
            saddle = "Greg" if row["Name"].startswith("Greg") else "Woop Woop"
            expect.append({"Tame": tame, "Name": row["Name"],
                           **{k: (row[k] if not (tame == "Spider" and k in ("Oxygen", "Food"))
                                  else None) for k in STATS},
                           "Saddle BP": saddle})
    got = [dict(zip(head, [c.value for c in r])) for r in out.iter_rows(min_row=2)
           if any(c.value is not None for c in r)]
    assert len(got) == 62
    for g, e in zip(got, expect):
        assert g == e
    colours = {c.column: c.fill.fgColor.rgb for c in out[2]}
    for r in range(2, 64):                          # rows past 13 take the template's style
        assert {c.column: c.fill.fgColor.rgb for c in out[r]} == colours
    assert isinstance(out.cell(2, 3).value, int)      # 4480.0 in the source, 4480 here
    assert isinstance(out.cell(2, 8).value, float)    # 125.8 stays 125.8
    assert "wood" not in {str(c.value) for row in out.iter_rows() for c in row}
    assert out.parent.properties.keywords == lr.FILLED_MARK


def test_btb_waits_for_the_user_on_what_was_only_inferred(btb):
    tpl, src, tmp, _ = btb
    csv = Path(li.create_parent(tpl).csv)
    ex = li.run_extract(csv, [src], use_images=False)
    # every value but Tame is sure; Tame (inferred from the file name) is amber
    assert ex.pre_accepted == 0 and ex.decisions == 1
    head, rows, fills = records_sheet(Path(ex.review))
    assert {r["Accept"] for r in rows} == {None}
    assert all(fills.get((i + 2, "Tame")) == lr.AMBER for i in range(len(rows)))
    assert not any(fills.get((i + 2, "Health")) for i in range(len(rows)))
    # nothing accepted: nothing applied
    assert li.run_apply(ex.review, use_images=False).rows == 0
    # one record refused, the rule accepted: the rest go in
    set_accept(Path(ex.review), "BTB_Spider.xlsx", 3, "N")
    decide_all(Path(ex.review))
    ap = li.run_apply(ex.review, use_images=False)
    assert ap.rows == 61 and "Tame rule" in ap.learned
    # learned: the next Extract pre-accepts every record
    ex2 = li.run_extract(csv, [src], use_images=False)
    assert ex2.decisions == 0 and ex2.pre_accepted == ex2.values


def test_btb_noise_and_lookup_rows_are_not_records(btb):
    tpl, src, tmp, _ = btb
    csv = Path(li.create_parent(tpl).csv)
    run = lr.extract(csv, sorted(src.glob("*.xlsx")))
    names = [r.values["Name"].value for r in run.records]
    assert not {"Greg", "Anky", "Rex", "Trevor", "Harrison"} & set(names)
    fourt = next(r for r in run.records if r.values["Name"].value == "Woop Woop Fourt")
    assert fourt.values["Saddle BP"].value == "Woop Woop"
    assert fourt.values["Saddle BP"].how == "starts with"
    spider = [r for r in run.records if r.file == "BTB_Spider.xlsx"]
    assert len(spider) == 31 and not any("Oxygen" in r.values for r in spider)
    # the same values under two names are two records, not a duplicate
    assert not any(r.flags for r in run.records)
    assert run.key_fields == ["Tame", "Name"]


def test_the_review_is_never_read_back_as_a_source(btb):
    tpl, src, tmp, _ = btb
    csv = Path(li.create_parent(tpl).csv)
    ex = li.run_extract(csv, [src], review_path=src / "review.xlsx", use_images=False)
    decide_all(Path(ex.review))
    li.run_apply(ex.review, src, use_images=False)    # filled template in the same folder
    kept, ignored = li.collect_sources([src], template_path=str(tpl))
    assert {p.name for p in ignored} == {"review.xlsx", "BTB_Template_filled.xlsx"}
    ex2 = li.run_extract(csv, [src], use_images=False)
    assert ex2.records == 62


def test_the_parent_csv_reads_as_a_ledger(btb):
    from council_core import kg_excel
    from sample_ledger import ParentCSV
    tpl, src, tmp, _ = btb
    csv = Path(li.create_parent(tpl).csv)
    assert kg_excel.is_ledger_csv(csv)
    ex = li.run_extract(csv, [src], use_images=False)
    decide_all(Path(ex.review))
    li.run_apply(ex.review, use_images=False)
    ledger = ParentCSV(csv)
    assert len(ledger.rows) == 62 and ledger.rows[0]["Name"] == "Greg #1"
    assert ledger.rows[0]["source_file"] == "BTB_Spider.xlsx"
    assert ledger.rows[0]["source_row"] == "3"


def test_applying_again_updates_records_in_place(btb):
    tpl, src, tmp, _ = btb
    csv = Path(li.create_parent(tpl).csv)
    ex = li.run_extract(csv, [src], use_images=False)
    decide_all(Path(ex.review))
    li.run_apply(ex.review, use_images=False)
    wb = openpyxl.load_workbook(src / "BTB_Spider.xlsx")
    wb.active["B3"] = 9999
    wb.save(src / "BTB_Spider.xlsx")
    ex2 = li.run_extract(csv, [src], use_images=False)
    ap = li.run_apply(ex2.review, use_images=False)
    assert ap.rows == 62 and "62 updated" in ap.csv_result
    rows = lr._read_rows(csv)[1]
    assert rows[0]["Health"] == "9999"


# ── layouts the BTB example does not have ───────────────────────────────
def simple_ledger(tmp_path, fields=("Name", "Health", "Weight", "Gender"), rows=6):
    wb = openpyxl.Workbook()
    ws = wb.active
    thin = Side(style="thin")
    for c, f in enumerate(fields, start=1):
        ws.cell(1, c, f)
        for r in range(2, 2 + rows):
            ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    csv = tmp_path / "T_parent.csv"
    lr.create_ledger(tpl, csv)
    return tpl, csv


def book(path: Path, sheets):
    """sheets: {name: [(row, col, value), ...]}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for r, c, v in cells:
            ws.cell(r, c, v)
    wb.save(path)
    return path


def table(top, left, header, rows):
    out = [(top, left + j, h) for j, h in enumerate(header)]
    for i, row in enumerate(rows, start=1):
        out += [(top + i, left + j, v) for j, v in enumerate(row) if v is not None]
    return out


def test_columns_moved_renamed_and_extra_are_matched_by_their_words(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    p = book(tmp_path / "a.xlsx", {"Log": [(1, 2, "Creature log 2026")] + table(
        5, 2, ["Gender", "Notes", "WEIGHT:", "Name", "health"],
        [["M", "ok", 574, "Rex 1", 4480], ["F", "", 532, "Rex 2", 7700]])})
    run = lr.extract(csv, [p])
    assert [(r.values["Name"].value, r.values["Health"].value, r.values["Weight"].value,
             r.values["Gender"].value) for r in run.records] == [
        ("Rex 1", 4480, 574, "M"), ("Rex 2", 7700, 532, "F")]
    assert all(r.accept == "Y" for r in run.records)


def test_a_similar_header_is_a_guess_until_the_user_says_yes(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    p = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "Healths", "Weight"],
                                              [["A", 1, 2], ["B", 3, 4]])})
    out = tmp_path / "r.xlsx"
    run = lr.extract(csv, [p])
    assert run.records[0].values["Health"].check == "C|Health|healths"
    assert all(r.accept == "" for r in run.records)
    lr.write_review(out, run, csv, [p])
    decide_all(out)
    assert lr.apply_review(out).records == 2
    assert lr.load_recipe(csv).aliases["healths"] == "Health"
    run2 = lr.extract(csv, [p])
    assert all(r.accept == "Y" for r in run2.records)


def test_a_no_on_a_guess_is_remembered(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    p = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "Healths", "Weight"],
                                              [["A", 1, 2], ["B", 3, 4]])})
    out = tmp_path / "r.xlsx"
    lr.write_review(out, lr.extract(csv, [p]), csv, [p])
    decide_all(out, "N")
    lr.apply_review(out)
    run = lr.extract(csv, [p])
    assert "Health" not in run.records[0].values


class FakeChat:
    def __init__(self, answers):
        self.answers, self.prompts = answers, []

    def __call__(self, prompt, schema):
        self.prompts.append(prompt)
        heads = schema["properties"]["mappings"]["items"]["properties"]["header"]["enum"]
        return json.dumps({"mappings": [{"header": h, "field": self.answers.get(h)}
                                        for h in heads]}), "fake-granite"


def test_the_model_places_unknown_headers_once_per_layout(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    files = [book(tmp_path / f"{n}.xlsx", {"S": table(1, 1, ["Name", "HP", "Mass", "Sex"],
                                                      [[f"{n}1", 1, 2, "M"],
                                                       [f"{n}2", 3, 4, "F"]])})
             for n in ("a", "b", "c")]
    chat = FakeChat({"hp": "Health", "mass": "Weight", "sex": "Gender"})
    run = lr.extract(csv, files, chat=chat)
    assert len(chat.prompts) == 1                    # three files, one header layout
    assert "Template fields" in chat.prompts[0] and '"hp"' in chat.prompts[0]
    rec = run.records[0]
    assert rec.values["Health"].value == 1 and rec.values["Health"].how == "model"
    assert rec.values["Health"].check and rec.accept == ""
    # remembered: asked no more
    chat2 = FakeChat({})
    run2 = lr.extract(csv, files, chat=chat2)
    assert chat2.prompts == [] and run2.records[0].values["Weight"].how == "model"


def test_a_failing_model_is_one_problem_and_the_rules_still_read(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    p = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "HP", "Weight"],
                                              [["A", 1, 2], ["B", 3, 4]])})

    def broken(prompt, schema):
        raise RuntimeError("no model here")
    ex = li._extract_records(csv, [p], tmp_path / "r.xlsx", lr._Quiet(), None,
                             "ollama:none", chat=broken)
    assert ex.records == 2 and any("no model here" in x for x in ex.problems)


def test_totals_repeated_headers_and_a_second_table(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    cells = table(1, 1, ["Name", "Health", "Weight"],
                  [["A", 1, 2], ["B", 3, 4], ["Name", "Health", "Weight"], ["C", 5, 6],
                   ["Total", 9, 12]])
    cells += table(9, 1, ["Name", "Gender"], [["D", "M"], ["E", "F"]])
    p = book(tmp_path / "a.xlsx", {"S": cells})
    run = lr.extract(csv, [p])
    names = [r.values["Name"].value for r in run.records]
    assert names == ["A", "B", "C", "D", "E"]
    assert run.records[3].values["Gender"].value == "M"


def test_a_table_written_sideways(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    cells = [(1, 1, "Name"), (2, 1, "Health"), (3, 1, "Weight"),
             (1, 2, "Ann"), (2, 2, 10), (3, 2, 20), (1, 3, "Bob"), (2, 3, 11), (3, 3, 21),
             (1, 4, "Cid"), (2, 4, 12), (3, 4, 22)]
    p = book(tmp_path / "a.xlsx", {"S": cells})
    run = lr.extract(csv, [p])
    assert [(r.values["Name"].value, r.values["Health"].value) for r in run.records] == [
        ("Ann", 10), ("Bob", 11), ("Cid", 12)]


def test_a_label_beside_the_table_fills_every_record(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["D1"] = "Inspector:", None
    for c, f in enumerate(["Name", "Health"], start=1):
        ws.cell(3, c, f)
        ws.cell(4, c).border = Border(bottom=Side(style="thin"))
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    csv = tmp_path / "T.csv"
    made = lr.create_ledger(tpl, csv)
    assert made.forms == ["Inspector"]
    p = book(tmp_path / "a.xlsx", {"S": [(1, 1, "Inspector: MK")] + table(
        3, 1, ["Name", "Health"], [["A", 1], ["B", 2]])})
    run = lr.extract(csv, [p])
    assert [r.values["Inspector"].value for r in run.records] == ["MK", "MK"]
    assert all(r.accept == "Y" for r in run.records)


def test_the_same_record_in_two_workbooks(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    a = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "Health"],
                                              [["Rex 1", 1], ["Rex 2", 2]])})
    b = book(tmp_path / "b.xlsx", {"S": table(1, 1, ["Name", "Health"],
                                              [["Rex 1", 1], ["Rex 2", 5]])})
    run = lr.extract(csv, [a, b])
    same = [r for r in run.records if r.values["Name"].value == "Rex 1"]
    assert same[0].accept == "Y" and same[1].accept == "N" and "duplicate" in same[1].flags
    clash = [r for r in run.records if r.values["Name"].value == "Rex 2"]
    assert all("clash" in r.flags and r.accept == "" for r in clash)
    out = tmp_path / "r.xlsx"
    lr.write_review(out, run, csv, [a, b])
    _head, _rows, fills = records_sheet(out)
    assert lr.RED in fills.values()
    decide_all(out)                                     # clashes are not decided there
    assert lr.apply_review(out).records == 1


def test_lookups_longest_key_ambiguous_and_mistyped(tmp_path):
    tpl, csv = simple_ledger(tmp_path, fields=("Name", "Saddle"))
    lk = book(tmp_path / "list.xlsx", {"L": table(1, 1, ["Kind", "Cost"], [
        ["Woop", 1], ["Woop Woop", 2], ["Greg", 3], ["Rex", 4], ["Rex", 5]])})
    lr.set_source(csv, "Saddle", "lookup", lookup_file=str(lk), key_header="Kind",
                  value_header="Cost", join="Name")
    p = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "Notes"], [
        ["Woop Woop #3", "x"], ["Woop 2", "x"], ["Gregg #1", "x"], ["Rex", "x"],
        ["Zed", "x"]])})
    run = lr.extract(csv, [p])
    got = {r.values["Name"].value: r.values.get("Saddle") for r in run.records}
    assert got["Woop Woop #3"].value == 2 and got["Woop Woop #3"].how == "starts with"
    assert got["Woop 2"].value == 1
    assert got["Gregg #1"].value == 3 and got["Gregg #1"].check.startswith("L|")
    assert got["Rex"].check and "2 times" in got["Rex"].note
    assert got["Zed"] is None
    # the lookup workbook is outside the sources: read from the recipe's path,
    # and its rows are never records
    assert not any(r.file == "list.xlsx" for r in run.records)


def test_a_field_from_the_file_name_by_example(tmp_path):
    tpl, csv = simple_ledger(tmp_path, fields=("Lot", "Name", "Health"))
    lr.set_source(csv, "Lot", "filename", example=("Run_L-101_final.xlsx", "L-101"))
    a = book(tmp_path / "Run_L-207_final.xlsx", {"S": table(1, 1, ["Name", "Health"],
                                                            [["A", 1]])})
    run = lr.extract(csv, [a])
    assert run.records[0].values["Lot"].value == "L-207"
    assert run.records[0].accept == "Y"


def test_odd_values_under_a_right_header_are_a_guess(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    good = [book(tmp_path / f"g{i}.xlsx", {"S": table(1, 1, ["Name", "Gender", "Weight"], [
        [f"N{i}{k}", "MF"[k % 2], 500 + k] for k in range(12)])}) for i in range(3)]
    bad = book(tmp_path / "bad.xlsx", {"S": table(1, 1, ["Name", "Gender", "Weight"], [
        [f"B{k}", 500 + k, "M"] for k in range(12)])})
    run = lr.extract(csv, good + [bad])
    rec = next(r for r in run.records if r.file == "bad.xlsx")
    assert rec.values["Gender"].check.startswith("V|Gender")
    assert rec.accept == "" and all(r.accept == "Y" for r in run.records
                                    if r.file != "bad.xlsx")


def test_a_chart_is_not_read_as_records(tmp_path):
    from openpyxl.chart import BarChart, Reference
    tpl, csv = simple_ledger(tmp_path)
    p = tmp_path / "c.xlsx"
    book(p, {"S": table(1, 1, ["Name", "Health"], [["A", 1], ["B", 2]])})
    wb = openpyxl.load_workbook(p)
    ch = BarChart()
    ch.title = "Health by name"
    ch.add_data(Reference(wb["S"], min_col=2, min_row=1, max_row=3), titles_from_data=True)
    wb["S"].add_chart(ch, "E2")
    wb.save(p)
    run = lr.extract(csv, [p])
    assert len(run.records) == 2


def test_a_formatted_cell_far_down_is_quick(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    p = tmp_path / "far.xlsx"
    book(p, {"S": table(1, 1, ["Name", "Health"], [["A", 1], ["B", 2]])})
    wb = openpyxl.load_workbook(p)
    paint(wb["S"].cell(1_000_000, 1), YELLOW)
    wb.save(p)
    t0 = time.monotonic()
    run = lr.extract(csv, [p], use_cache=False)
    assert len(run.records) == 2 and time.monotonic() - t0 < 20


def test_unchanged_workbooks_come_from_the_cache(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    files = [book(tmp_path / f"{n}.xlsx", {"S": table(1, 1, ["Name", "Health"], [[n, 1]])})
             for n in "abc"]
    assert lr.extract(csv, files).from_cache == 0
    assert lr.extract(csv, files).from_cache == 3
    time.sleep(0.01)
    book(files[0], {"S": table(1, 1, ["Name", "Health"], [["a", 7]])})
    run = lr.extract(csv, files)
    assert run.from_cache == 2 and run.records[0].values["Health"].value == 7


def test_theme_colours_link_like_plain_ones(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, f in enumerate(["Name", "Power"], start=1):
        ws.cell(1, c, f)
    ws["A1"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.4))
    ws["B1"].fill = PatternFill("solid", fgColor=Color(theme=5, tint=0.4))
    for c in (1, 2):
        ws.cell(2, c).border = Border(bottom=Side(style="thin"))
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    csv = tmp_path / "T.csv"
    lr.create_ledger(tpl, csv)
    s = tmp_path / "s.xlsx"
    book(s, {"S": table(1, 1, ["Label", "Watts"], [["x", 1], ["y", 2]])})
    wb = openpyxl.load_workbook(s)
    wb["S"]["A1"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.4))
    wb["S"]["B1"].fill = PatternFill("solid", fgColor=Color(theme=5, tint=0.4))
    wb.save(s)
    run = lr.extract(csv, [s])
    assert [(r.values["Name"].value, r.values["Power"].value) for r in run.records] == [
        ("x", 1), ("y", 2)]
    assert all(r.accept == "Y" for r in run.records)


# ── filling the template ────────────────────────────────────────────────
def test_fill_grows_the_table_keeps_bands_formulas_validation_and_what_is_below(tmp_path):
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, f in enumerate(["Name", "Health", "Double"], start=1):
        ws.cell(1, c, f)
    for r in range(2, 6):
        for c in (1, 2, 3):
            paint(ws.cell(r, c), "DDEBF7" if r % 2 else "FFFFFF" if False else "FCE4D6")
            if r % 2:
                paint(ws.cell(r, c), "DDEBF7")
        ws.cell(r, 3, f"=B{r}*2")
    dv = DataValidation(type="whole")
    dv.add("B2:B5")
    ws.add_data_validation(dv)
    ws["A7"] = "Signed off by:"
    ws.merge_cells("A8:C8")
    ws["A8"] = "footer"
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    tt = lr.analyze_template(tpl)
    assert [f.name for f in tt.fields] == ["Name", "Health", "Double"]
    assert tt.fields[2].formula == "=B2*2" and tt.content_below and tt.styled_rows == 4
    rows = [{"Name": f"N{i}", "Health": str(i)} for i in range(10)]
    out = tmp_path / "out" / "filled.xlsx"
    lr.fill_records(tt, rows, out)
    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(r, 1).value for r in range(2, 12)] == [f"N{i}" for i in range(10)]
    assert ws["B11"].value == 9 and ws["C11"].value == "=B11*2"
    band = [ws.cell(r, 1).fill.fgColor.rgb[-6:] for r in range(2, 12)]
    assert band == ["FCE4D6", "DDEBF7"] * 5
    assert ws["A13"].value == "Signed off by:" and ws["A14"].value == "footer"
    assert "A14:C14" in {str(m) for m in ws.merged_cells.ranges}
    dv_ranges = " ".join(str(d.sqref) for d in ws.data_validations.dataValidation)
    assert "B11" in dv_ranges or "B2:B11" in dv_ranges
    with pytest.raises(lr.LedgerError):
        lr.fill_records(tt, rows, tpl)


def test_a_template_with_a_sample_id_stays_a_sample_ledger(tmp_path):
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    from ledger_fixtures import make_template as lmm_template
    tpl = lmm_template(tmp_path / "LMM Template.xlsx")
    s = li.create_parent(tpl)
    assert not s.records_mode and not lr.is_record_ledger(s.csv)


def test_a_template_with_no_table_says_so(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "just a title"
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    with pytest.raises(lr.LedgerError, match="no table"):
        lr.create_ledger(tpl, tmp_path / "T.csv")


def test_records_mode_can_be_asked_for(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, f in enumerate(["Sample ID", "Hardness", "Width"], start=1):
        ws.cell(1, c, f)
        ws.cell(2, c).border = Border(bottom=Side(style="thin"))
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    # a sample-ID table: one sample per workbook unless asked
    assert not li.create_parent(tpl, tmp_path / "a.csv").records_mode
    s = li.create_parent(tpl, tmp_path / "b.csv", mode="records")
    assert s.records_mode and lr.is_record_ledger(tmp_path / "b.csv")
    # and it stays one when created again
    assert li.create_parent(tpl, tmp_path / "b.csv").records_mode


def test_teach_says_what_extract_will_read(btb):
    tpl, src, tmp, _ = btb
    csv = Path(li.create_parent(tpl).csv)
    text = lr.teach(csv, [src / "BTB_Woopentide.xlsx", src / "BTB_Saddle_BP.xlsx"])
    assert "Learned from your highlights in BTB_Woopentide.xlsx" in text
    assert "31 records" in text and "a lookup table" in text
    recipe = lr.load_recipe(csv)
    assert recipe.aliases["health"] == "Health"
    assert recipe.sources["Saddle BP"]["kind"] == "lookup"


def test_a_record_named_max_and_a_remark_row_stay_in_the_table(tmp_path):
    tpl, csv = simple_ledger(tmp_path)
    p = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "Health", "Weight", "Gender"], [
        ["Rex", 1, 2, "M"], ["Max", 3, 4, "F"], ["Dead one", "not measured", None, None],
        ["Bo", 5, 6, "M"], ["Average", 3, 4, None]])})
    names = [r.values["Name"].value for r in lr.extract(csv, [p]).records]
    assert names == ["Rex", "Max", "Dead one", "Bo"]


def test_an_unrelated_table_with_two_generic_headers_waits(tmp_path):
    fields = ("Name", "Date", "Health", "Weight", "Length", "Width", "Depth", "Result")
    tpl, csv = simple_ledger(tmp_path, fields=fields)
    good = book(tmp_path / "good.xlsx", {"S": table(1, 1, list(fields), [
        ["A", "2026-01-01", 1, 2, 3, 4, 5, "ok"]])})
    other = book(tmp_path / "agenda.xlsx", {"S": table(1, 1, ["Name", "Date", "Topic"], [
        ["Kickoff", "2026-01-02", "plans"]])})
    run = lr.extract(csv, [good, other])
    by = {r.file: r for r in run.records}
    assert by["good.xlsx"].accept == "Y"
    assert by["agenda.xlsx"].accept == "" and by["agenda.xlsx"].checks[0].startswith("T|")


def test_a_fuzzy_lookup_answer_is_remembered(tmp_path):
    tpl, csv = simple_ledger(tmp_path, fields=("Name", "Saddle"))
    lk = book(tmp_path / "list.xlsx", {"L": table(1, 1, ["Kind", "Cost"], [["Greg", 3]])})
    lr.set_source(csv, "Saddle", "lookup", lookup_file=str(lk), key_header="Kind",
                  value_header="Cost", join="Name")
    p = book(tmp_path / "a.xlsx", {"S": table(1, 1, ["Name", "Notes"], [["Gregg #1", "x"]])})
    out = tmp_path / "r.xlsx"
    lr.write_review(out, lr.extract(csv, [p]), csv, [p])
    decide_all(out)
    lr.apply_review(out)
    rec = lr.extract(csv, [p]).records[0]
    assert rec.values["Saddle"].value == 3 and rec.values["Saddle"].how == "your answer"
    assert rec.accept == "Y"


def test_a_lookup_painted_on_two_columns(tmp_path):
    tpl = make_template(tmp_path / "T.xlsx")
    csv = tmp_path / "T.csv"
    lr.create_ledger(tpl, csv)
    src = tmp_path / "s"
    src.mkdir()
    make_sample(src / "BTB_A.xlsx", creature_rows("Greg", 3), ["Name"] + STATS, True)
    make_sample(src / "BTB_B.xlsx", creature_rows("Rex", 3), ["Name"] + STATS, True)
    sad = make_saddles(src / "BTB_Saddle_BP.xlsx")
    wb = openpyxl.load_workbook(sad)
    for r in range(1, 9):
        paint(wb.active.cell(r, 2), CYAN)                 # the wood column too
    wb.save(sad)
    run = lr.extract(csv, sorted(src.glob("*.xlsx")))
    rule = lr.load_recipe(csv).sources["Saddle BP"]
    assert rule["key_header"] == "tame saddle" and rule["value_header"] == "wood"
    greg = next(r for r in run.records if r.values["Name"].value == "Greg #1")
    assert greg.values["Saddle BP"].value == 35


# ── a grid of tables on one sheet ───────────────────────────────────────
GRID_FIELDS = ("Tank", "Inspector", "Position", "Reading", "Result")


def grid_block(ws, top, left, tank, insp, n=4, title=None):
    if title is None:
        ws.cell(top, left, f"Tank: {tank}")
    else:
        ws.cell(top, left, title)
    ws.cell(top + 1, left, "Inspector:")
    ws.cell(top + 1, left + 1, insp)
    for j, h in enumerate(["Position", "Reading", "Result"]):
        ws.cell(top + 2, left + j, h)
    for i in range(n):
        ws.cell(top + 3 + i, left, f"P{i + 1}")
        ws.cell(top + 3 + i, left + 1, 10 * tank + i)
        ws.cell(top + 3 + i, left + 2, "OK")


def grid_book(path, gap_cols=1, gap_rows=2, stagger=0, titles=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    right = 4 + gap_cols
    below = 8 + gap_rows
    places = [(1, 1), (1 + stagger, right), (below, 1), (below + stagger, right)]
    for n, ((top, left), insp) in enumerate(zip(places, ["JB", "MK", "AL", "TS"]), start=1):
        grid_block(ws, top, left, n, insp, title=titles[n - 1] if titles else None)
    wb.save(path)
    return path


GRID_EXPECT = sorted((str(t), i, f"P{p + 1}", 10 * t + p) for t, i in
                     zip(range(1, 5), ["JB", "MK", "AL", "TS"]) for p in range(4))


@pytest.mark.parametrize("layout", [
    dict(gap_cols=1, gap_rows=2), dict(gap_cols=1, gap_rows=1), dict(gap_cols=0, gap_rows=2),
    dict(gap_cols=1, gap_rows=2, stagger=2), dict(gap_cols=3, gap_rows=4, stagger=1)],
    ids=["spaced", "one-blank-row", "touching-columns", "staggered", "wide-gaps"])
def test_a_grid_of_tables_is_one_block_each(tmp_path, layout):
    tpl, csv = simple_ledger(tmp_path, fields=GRID_FIELDS)
    p = grid_book(tmp_path / "grid.xlsx", **layout)
    run = lr.extract(csv, [p])
    got = sorted((cell_or_none(r, "Tank"), cell_or_none(r, "Inspector"),
                  r.values["Position"].value, r.values["Reading"].value) for r in run.records)
    assert got == GRID_EXPECT
    assert all(r.accept == "Y" and not r.flags for r in run.records)
    assert run.key_fields[:2] == ["Tank", "Inspector"]


def cell_or_none(rec, name):
    v = rec.values.get(name)
    return lr.cell_text(v.value) if v else None


def test_a_grid_applies_without_one_block_overwriting_another(tmp_path):
    tpl, csv = simple_ledger(tmp_path, fields=("Position", "Reading", "Result"))
    p = grid_book(tmp_path / "grid.xlsx")
    out = tmp_path / "r.xlsx"
    lr.write_review(out, lr.extract(csv, [p]), csv, [p])
    assert lr.apply_review(out).rows == 16              # P1..P4 four times, all kept
    again = lr.apply_review(out)
    assert again.rows == 16 and "16 updated" in again.csv_result
    rows = lr._read_rows(csv)[1]
    assert {r["source_table"] for r in rows} == {"A3", "E3", "A12", "E12"}


def test_one_label_over_the_whole_sheet_goes_to_every_block(tmp_path):
    tpl, csv = simple_ledger(tmp_path, fields=("Lot", "Position", "Reading"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Lot: L-77"
    for n, (top, left) in enumerate([(3, 1), (3, 6), (12, 1)], start=1):
        grid_block(ws, top, left, n, "JB")
    p = tmp_path / "g.xlsx"
    wb.save(p)
    run = lr.extract(csv, [p])
    assert len(run.records) == 12 and {r.values["Lot"].value for r in run.records} == {"L-77"}


def test_a_painted_block_title_is_learned_by_its_place(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, (f, rgb) in enumerate([("Tank", BLUE), ("Position", GREEN), ("Reading", GREEN)],
                                 start=1):
        ws.cell(1, c, f)
        paint(ws.cell(1, c), rgb)
        paint(ws.cell(2, c), rgb)
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    csv = tmp_path / "T.csv"
    lr.create_ledger(tpl, csv)
    titles = ["North tank", "South tank", "East tank", "West tank"]
    sample = grid_book(tmp_path / "a.xlsx", titles=titles)
    wb = openpyxl.load_workbook(sample)
    paint(wb.active["A1"], BLUE)                        # the user paints ONE block's title
    for c in (1, 2):
        paint(wb.active.cell(3, c), GREEN)
    wb.save(sample)
    other = grid_book(tmp_path / "b.xlsx", stagger=2, titles=["Tank A", "Tank B", "Tank C",
                                                              "Tank D"])
    run = lr.extract(csv, [sample, other])
    assert lr.load_recipe(csv).sources["Tank"]["kind"] == "near"
    by_file = {}
    for r in run.records:
        by_file.setdefault(r.file, set()).add(r.values["Tank"].value)
    assert by_file == {"a.xlsx": set(titles), "b.xlsx": {"Tank A", "Tank B", "Tank C",
                                                         "Tank D"}}
    assert all(r.accept == "Y" for r in run.records)
