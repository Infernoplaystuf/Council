"""PowerPoint decks as Ledger sources (pptx_open, user 2026-10-21: "a
powerpoint with slides in it with an image, textbox, and multiple tables.
The Textbox gives a hint for what the tables store"). Decks are built here
with python-pptx; Council itself reads them with the standard library."""
from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
import pptx_open  # noqa: E402
import source_extract as se  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_records as lr  # noqa: E402
from ledger_fixtures import (CONCEALMENT_IMAGE_ROWS, ITEMS, _png, fake_reader,  # noqa: E402
                             make_template)


def textbox(slide, x, y, w, h, *lines):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    return tb


def table(slide, x, y, rows, w=4.0):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(0.3 * len(rows)))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            shape.table.cell(i, j).text = str(v)
    return shape.table


def visual_rows(miss=None):
    return [["Defect", "Result"]] + [[it, "not seen" if it == miss else "seen"]
                                     for it in ITEMS]


def lmm_deck(path: Path, operators=("JB", "MK")) -> Path:
    """A title slide with the part; a Visual Inspection slide with one table
    per operator, each under its own text box, and a picture of a table; an
    NDE slide whose text box says what its table holds and who read it."""
    prs = Presentation()
    layout = prs.slide_layouts[5]                          # title only
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Part 101 inspection summary"
    textbox(s, 1, 2, 4, 1, "P/N: 101", "Scan Speed (mm/s): 1200")
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Visual Inspection"
    for k, who in enumerate(operators):
        textbox(s, 0.3 + 3.3 * k, 1.4, 3, 0.4, f"Inspector: {who}")
        table(s, 0.3 + 3.3 * k, 1.9, visual_rows(miss="3C" if k == 1 else None), w=3)
    s.shapes.add_picture(io.BytesIO(_png(CONCEALMENT_IMAGE_ROWS)), Inches(3), Inches(5.6),
                         Inches(3))
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "NDE"
    textbox(s, 0.3, 1.3, 6, 0.4, "Eddy Current Response (%FSH) — Operator: TS")
    table(s, 0.3, 1.8, [["Flaw"] + ITEMS, ["Pre-Conceal"] + [60 + i for i in range(9)],
                        ["Post-Conceal"] + [20 + i for i in range(9)]], w=9)
    prs.save(path)
    return path


# ── the deck as sheets ──────────────────────────────────────────────────
def lines_of(sheet):
    return [v for (r, c), v in sorted(sheet.cells.items()) if c == 1]


def test_each_table_gets_its_own_text_above_it(tmp_path):
    deck = pptx_open.read_pptx(lmm_deck(tmp_path / "d.pptx"))
    names = [s.name for s in deck.sheets]
    assert names == ["Slide 1", "Slide 2 table 1", "Slide 2 table 2", "Slide 3"]
    one, two = deck.sheets[1], deck.sheets[2]
    assert lines_of(one)[:2] == ["Visual Inspection", "Inspector: JB"]
    assert lines_of(two)[:2] == ["Visual Inspection", "Inspector: MK"]
    assert "Inspector: MK" not in lines_of(one)
    # a blank row between the heading and the table's header
    assert (3, 1) not in one.cells and one.cells[(4, 1)] == "Defect"
    # 'heading — Operator: TS' is two lines: the heading and who read it
    assert lines_of(deck.sheets[3])[:3] == ["NDE", "Eddy Current Response (%FSH)",
                                            "Operator: TS"]
    assert deck.sheets[0].cells[(2, 1)] == "P/N: 101"


def test_pictures_carry_their_slide_title(tmp_path):
    deck = pptx_open.read_pptx(lmm_deck(tmp_path / "d.pptx"))
    assert [(p.sheet, p.anchor) for p in deck.pictures] == [
        ("Slide 2 · Visual Inspection", "picture 1")]
    assert deck.pictures[0].data[:4] == b"\x89PNG"


def test_text_over_two_tables_belongs_to_both(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])         # blank
    textbox(s, 0.3, 0.5, 9, 0.4, "Lot: L-77")
    textbox(s, 0.3, 1.2, 4, 0.4, "Inspector: JB")
    table(s, 0.3, 1.7, [["Name", "Health"], ["A", 1]])
    textbox(s, 5.0, 1.2, 4, 0.4, "Inspector: MK")
    table(s, 5.0, 1.7, [["Name", "Health"], ["B", 2]])
    prs.save(tmp_path / "d.pptx")
    a, b = pptx_open.read_pptx(tmp_path / "d.pptx").sheets
    assert lines_of(a)[:2] == ["Lot: L-77", "Inspector: JB"]
    assert lines_of(b)[:2] == ["Lot: L-77", "Inspector: MK"]


def test_groups_charts_fills_and_hidden_slides(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()
    textbox(grp, 0.5, 0.5, 3, 0.4, "Inspector: JB")
    t = table(s, 0.5, 1.2, [["Name", "Health"], ["A", 1]])
    t.cell(0, 1).fill.solid()
    t.cell(0, 1).fill.fore_color.rgb = RGBColor(0x00, 0xFF, 0x00)
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(3),
                       Inches(2), data)
    hidden = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(hidden, 1, 1, 3, 1, "draft: do not read")
    hidden._element.set("show", "0")
    prs.save(tmp_path / "d.pptx")
    deck = pptx_open.read_pptx(tmp_path / "d.pptx")
    assert [s.name for s in deck.sheets] == ["Slide 1"]      # the hidden slide is not read
    sheet = deck.sheets[0]
    assert lines_of(sheet)[0] == "Inspector: JB"            # inside a group
    header = min(r for (r, c), v in sheet.cells.items() if v == "Name")
    assert sheet.colours[(header, 2)] == "00FF00"
    assert any("chart is not read" in n for n in deck.notes)


def test_not_a_deck_says_so(tmp_path):
    p = tmp_path / "old.pptx"
    p.write_bytes(b"\xd0\xcf\x11\xe0 not a zip")
    with pytest.raises(pptx_open.PptxError, match="save it as .pptx"):
        pptx_open.read_pptx(p)


def test_a_folder_scan_finds_decks(tmp_path):
    lmm_deck(tmp_path / "d.pptx")
    (tmp_path / "~$d.pptx").write_bytes(b"lock")
    assert [p.name for p in se.workbook_files(tmp_path)] == ["d.pptx"]


# ── the sample Ledger (one part per file, flaw grids) ───────────────────
def lmm_extract(tmp_path, operators=("JB", "MK")):
    src = tmp_path / "src"
    src.mkdir()
    lmm_deck(src / "part101.pptx", operators)
    tpl = make_template(tmp_path / "LMM Template.xlsx")
    csv = Path(li.create_parent(tpl, tmp_path / "lmm.csv").csv)
    ex = li.run_extract(csv, [src], image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS))
    wb = openpyxl.load_workbook(ex.review)
    rows = [r for r in wb["Values"].iter_rows(min_row=2, values_only=True)]
    return csv, ex, rows


def test_each_operators_table_takes_its_own_visual_slot(tmp_path):
    csv, ex, rows = lmm_extract(tmp_path)
    vis = {(r[2], r[3]): r for r in rows if r[3] and r[3].startswith("Visual Inspection")}
    assert vis[("1A", "Visual Inspection: Inspector 1")][4] == "seen"
    assert vis[("1A", "Visual Inspection: Inspector 1 — Inspector Initials")][4] == "JB"
    assert vis[("3C", "Visual Inspection: Inspector 2")][4] == "not seen"
    assert vis[("3C", "Visual Inspection: Inspector 2 — Inspector Initials")][4] == "MK"
    # inferred from the order the tables were read: amber, said so
    assert all(r[0] is None for r in vis.values())
    assert all("person 2 to sign Visual Inspection" in (r[10] or "")
               for k, r in vis.items() if "Inspector 2" in k[1])
    # what else the deck held
    nde = [r for r in rows if r[3] == "NDE Inspection: Eddy Current Response Pre-Conceal"]
    assert len(nde) == 9 and all(r[0] == "Y" and r[1] == "101" for r in nde)
    assert any(r[3] == "Concealment 1: Media" and r[4] == "ceramic" for r in rows)
    assert ex.unmatched == 0


def test_the_filled_template_has_both_operators_and_learns_no_alias(tmp_path):
    csv, ex, _rows = lmm_extract(tmp_path)
    wb = openpyxl.load_workbook(ex.review)
    for row in wb["Values"].iter_rows(min_row=2):
        row[0].value = "Y"
    wb.save(ex.review)
    ap = li.run_apply(ex.review, tmp_path / "filled", use_images=False)
    assert not ap.learned                                  # 'Result' is not 'Inspector 1'
    ws = openpyxl.load_workbook(tmp_path / "filled" / "LMM_Template_101.xlsx").active
    blocks = [[ws.cell(r, c).value for c in range(3, 7)]
              for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value in
              ("Inspector 1", "Inspector 2")]
    # blocks A, B, C: Inspector 1 then Inspector 2 in each; MK did not see 3C
    assert blocks == [["seen", "seen", "seen", "JB"], ["seen", "seen", "seen", "MK"]] * 2 + [
        ["seen", "seen", "seen", "JB"], ["seen", "seen", "not seen", "MK"]]


def test_more_operators_than_slots_are_not_placed(tmp_path):
    _csv, _ex, rows = lmm_extract(tmp_path, operators=("JB", "MK", "AL"))
    signed = {r[4] for r in rows if r[3] and r[3].endswith("Inspector Initials")
              and r[3].startswith("Visual")}
    assert signed == {"JB", "MK"}


# ── the table Ledger (one row per record) ───────────────────────────────
def test_a_deck_of_record_tables_with_a_label_per_slide(tmp_path):
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, f in enumerate(["Tame", "Name", "Health", "Weight"], start=1):
        ws.cell(1, c, f)
        ws.cell(2, c).fill = PatternFill("solid", fgColor="FFDDDDDD")
    ws["C1"].fill = ws["C2"].fill = PatternFill("solid", fgColor="FF00FF00")
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    csv = tmp_path / "T.csv"
    lr.create_ledger(tpl, csv)
    prs = Presentation()
    for tame, names in (("Spider", ["Greg #1", "Greg #2"]), ("Woopentide", ["Woop #1"])):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"{tame} stats"
        textbox(s, 0.5, 1.4, 4, 0.4, f"Tame: {tame}")
        t = table(s, 0.5, 1.9, [["Name", "HP", "Weight"]]
                  + [[n, 4000 + i, 574] for i, n in enumerate(names)])
        t.cell(0, 1).fill.solid()                          # 'HP' painted Health's green
        t.cell(0, 1).fill.fore_color.rgb = RGBColor(0x00, 0xFF, 0x00)
    deck = tmp_path / "tames.pptx"
    prs.save(deck)
    run = lr.extract(csv, [deck])
    got = [(r.values["Tame"].value, r.values["Name"].value, r.values["Health"].value)
           for r in run.records]
    assert got == [("Spider", "Greg #1", 4000), ("Spider", "Greg #2", 4001),
                   ("Woopentide", "Woop #1", 4000)]
    assert all(r.accept == "Y" for r in run.records)
    assert lr.load_recipe(csv).aliases["hp"] == "Health"
