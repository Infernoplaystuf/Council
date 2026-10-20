"""
source_extract — source workbooks of other layouts -> review -> parent CSV and
filled templates.

The sources (tests/ledger_fixtures.py) are laid out the ways real ones differ
from the template: a part number written inline, flaws down the side in mm, a
sheet named for the part with codes written 1-A, results transposed, a table
that exists only as a picture, several parts in rows, and a value two sources
disagree on. No OCR engine is needed: images go through a stand-in reader.

Run:  python -m pytest tests/test_source_extract.py -q
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
from ledger_fixtures import (CONCEALMENT_IMAGE_ROWS, ITEMS, fake_reader,  # noqa: E402
                             make_laser_log, make_nde_book, make_summary, make_template)

from sample_ledger import ParentCSV, read_template  # noqa: E402
from source_extract import (ItemCodes, Matcher, apply_review, boxes_to_rows,  # noqa: E402
                            decide, extract, load_aliases, parse_vision_text,
                            split_unit, tokens, write_review)


@pytest.fixture
def setup(tmp_path):
    schema = read_template(make_template(tmp_path / "tmpl.xlsx"))
    csv = tmp_path / "parent.csv"
    ParentCSV.create(csv, schema)
    sources = [make_laser_log(tmp_path / "laser_log.xlsx"),
               make_nde_book(tmp_path / "NDE results.xlsx"),
               make_summary(tmp_path / "summary.xlsx")]
    return schema, csv, sources, tmp_path


def _run(schema, sources, tmp_path, **kw):
    return extract(schema, sources, image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS),
                   image_dir=tmp_path / "imgs", **kw)


def _by(decisions):
    return {(d.obs.part, d.obs.item, d.obs.field, d.obs.value): d for d in decisions}


# ── pieces ──────────────────────────────────────────────────────────────
def test_labels_normalise():
    assert tokens("Part #") == tokens("P/N") == tokens("Part Number") == ["part", "number"]
    assert tokens("EC Pre") == ["eddy", "current", "pre"]
    assert tokens("Laser Team FOD A (um)") == ["laser", "team", "fod"]
    assert split_unit("Flaw Length (mm)") == ("Flaw Length", "mm")
    assert split_unit("FOD um") == ("FOD", "um")


def test_item_codes():
    codes = ItemCodes(ITEMS)
    for raw in ("1A", "1a", "Flaw 1A", "1-A", "A1", "#1A", "flaw no. 1A"):
        assert codes.canon(raw) == "1A", raw
    assert codes.canon("4A") is None and codes.canon("Power") is None


def test_matcher_keeps_numbered_rows_apart(setup):
    schema = setup[0]
    m = Matcher(schema)
    assert m.match("Inspector 2", per_item=True).target.name == "Visual Inspection: Inspector 2"
    assert m.match("EC Post", ["NDE"], per_item=True).target.name == \
        "NDE Inspection: Eddy Current Response Post-Conceal"
    date = m.match("Date", per_item=False)
    assert date.target.name == "Date" and not date.sure        # generic: review it
    assert m.match("Operator").target is None


def test_ocr_boxes_become_aligned_rows():
    boxes = [(10, 10, 60, 20, "Media"), (110, 11, 180, 21, "Amplitude"),
             (10, 40, 60, 50, "ceramic"), (110, 41, 130, 51, "40"),
             (110, 71, 130, 81, "35")]                     # a blank first cell
    assert boxes_to_rows(boxes) == [["Media", "Amplitude"], ["ceramic", "40"], ["", "35"]]


def test_vision_text_parsing():
    md = "| Media | Amplitude |\n|---|---|\n| ceramic | 40 |"
    assert parse_vision_text(md) == [["Media", "Amplitude"], ["ceramic", "40"]]
    assert parse_vision_text("Power: 285 W\nOperator: JB") == [["Power", "285 W"],
                                                                ["Operator", "JB"]]
    assert parse_vision_text("NONE") is None


# ── a whole extraction ──────────────────────────────────────────────────
def test_every_layout_is_read(setup):
    schema, _, sources, tmp = setup
    dec = _by(decide(_run(schema, sources, tmp), schema))
    # inline "P/N: 101" names the part; a value two cells right of its label
    assert dec[("101", "", "Frequency (kHz)", "50")].accept
    assert ("101", "", "Power (W)", "285") in dec
    # flaws down the side, in mm: converted, with a note
    fod = dec[("101", "2B", "Flaw Attributes: Laser Team FOD (um)", "404")]
    assert fod.obs.note == "converted mm -> um"
    assert dec[("101", "3C", "Flaw Attributes: Flaw Length (in)", "0.0787402")].accept
    # grid with codes written 1-A, on a sheet named "Part 101"
    pre = dec[("101", "2B", "NDE Inspection: Eddy Current Response Pre-Conceal", "64")]
    assert pre.accept and pre.obs.where.endswith("Part 101!G4")
    # transposed visual results
    assert dec[("101", "3C", "Visual Inspection: Inspector 2", "not seen")].accept
    # a table of parts: each row its own part
    assert dec[("102", "", "Power (W)", "300")].accept


def test_disagreeing_sources_are_never_pre_accepted(setup):
    schema, _, sources, tmp = setup
    dec = _by(decide(_run(schema, sources, tmp), schema))
    a, b = dec[("101", "", "Power (W)", "285")], dec[("101", "", "Power (W)", "290")]
    assert not a.accept and not b.accept
    assert a.conflict == "290" and b.conflict == "285"


def test_image_values_are_found_but_left_for_review(setup):
    schema, _, sources, tmp = setup
    run = _run(schema, sources, tmp)
    dec = _by(decide(run, schema))
    media = dec[("101", "", "Concealment 2: Media", "steel")]
    assert media.obs.via == "image" and not media.accept
    assert len(run.images) == 1 and run.images[0].values == 4
    assert Path(run.images[0].saved_as).exists()


def test_no_reader_means_images_are_saved_not_read(setup):
    schema, _, sources, tmp = setup
    run = extract(schema, sources, image_dir=tmp / "imgs", use_images=False)
    assert run.images[0].reader == "none installed"
    assert not any(o.from_image for o in run.obs)


def test_no_pictures_folder_is_made_for_a_workbook_without_pictures(setup):
    schema, _, sources, tmp = setup
    run = extract(schema, [sources[0]], image_dir=tmp / "imgs", use_images=False)
    assert run.obs and not run.images
    assert not (tmp / "imgs").exists()


def test_a_pictures_folder_path_too_long_is_one_problem_not_a_failed_extract(
        setup, monkeypatch):
    """A review folder whose '<review>_images' path was 248+ characters
    failed the whole Extract with a raw FileNotFoundError [WinError 206],
    pictures ticked or not: the folder was made before anything else
    (review, 2026-10-19). Now the pictures are not saved, once said, and
    the values are all read."""
    import errno
    schema, _, sources, tmp = setup
    nde2 = make_nde_book(tmp / "NDE again.xlsx")            # a second picture
    real = Path.mkdir

    def too_long(self, *a, **kw):
        if self.name == "imgs":
            raise OSError(errno.ENAMETOOLONG, "The filename or extension is too long")
        return real(self, *a, **kw)
    monkeypatch.setattr(Path, "mkdir", too_long)
    for use_images in (False, True):
        reader = fake_reader(CONCEALMENT_IMAGE_ROWS) if use_images else None
        run = extract(schema, sources + [nde2], image_dir=tmp / "imgs", image_reader=reader,
                      use_images=use_images)
        said = [p for p in run.problems if "path too long" in p]
        n = len(str((tmp / "imgs").absolute())) + 1
        assert len(said) == 1, run.problems
        assert said[0].startswith("pictures not saved: path too long (")
        assert said[0].endswith(" characters) — enable Windows long paths or move the folder")
        assert int(said[0].split("(")[1].split()[0]) > n
        assert len(run.images) == 2 and not any(i.saved_as for i in run.images)
        assert any(o.field == "Frequency (kHz)" for o in run.obs)
        assert any(o.from_image for o in run.obs) == use_images     # still read


@pytest.mark.skipif(sys.platform != "win32", reason="Windows' MAX_PATH")
def test_a_real_pictures_folder_path_too_long_does_not_fail_the_extract(tmp_path):
    """The tester's repro: <review>_images of 248+ characters, Windows long
    paths off."""
    import os
    from council_core import ledger_ingest as li
    probe = tmp_path / ("p" * 120) / ("q" * 120)
    try:
        os.makedirs(probe)
        pytest.skip("Windows long paths are on here")
    except OSError:
        pass
    suffix = len("\\review_20261019_123456_images")
    folder = tmp_path / ("d" * (250 - suffix - len(str(tmp_path)) - 1))
    folder.mkdir()
    review = folder / "review_20261019_123456.xlsx"
    assert len(str(review.with_name(review.stem + "_images"))) == 250
    tpl = make_template(folder / "t.xlsx")
    csv = folder / "p.csv"
    li.create_parent(str(tpl), str(csv))
    src = make_nde_book(folder / "nde.xlsx")                  # has a picture
    for use_images in (False, True):
        s = li.run_extract(str(csv), [str(src)], str(review), use_images=use_images,
                           image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS))
        assert Path(s.review).exists() and s.values
        assert [p for p in s.problems if p.startswith("pictures not saved: path too long")]


def test_unknown_labels_and_problems(setup):
    schema, _, sources, tmp = setup
    run = _run(schema, sources, tmp)
    # "Operator: JB" says who logged the sheet: not a value, not a label to map
    assert run.unmatched == [] and run.problems == []
    assert not any(o.value == "JB" for o in run.obs)


def test_missing_part_number_is_reported(setup):
    schema, _, _, tmp = setup
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "Power (W)", 280
    wb.save(tmp / "nameless.xlsx")
    run = extract(schema, [tmp / "nameless.xlsx"], use_images=False)
    assert run.obs[0].part == "" and "no Part Number" in run.problems[0]
    run = extract(schema, [tmp / "nameless.xlsx"], use_images=False, part="105")
    assert run.obs[0].part == "105"


# ── review -> apply ─────────────────────────────────────────────────────
def _edit_review(path, accept=(), map_labels=None, ignore=()):
    wb = openpyxl.load_workbook(path)
    ws = wb["Values"]
    for row in ws.iter_rows(min_row=2):
        key = (str(row[1].value or ""), str(row[2].value or ""), row[3].value,
               str(row[4].value))
        for want in accept:
            if all(w is None or w == k for w, k in zip(want, key)):
                row[0].value = "Y"
    lab = wb["Labels to map"]
    for row in lab.iter_rows(min_row=2):
        if map_labels and row[2].value in map_labels:
            row[0].value = map_labels[row[2].value]
        if row[2].value in ignore:
            row[1].value = "Y"
    wb.save(path)


def test_review_then_apply_fills_csv_and_template(setup):
    schema, csv, sources, tmp = setup
    run = _run(schema, sources, tmp)
    review = tmp / "review.xlsx"
    write_review(review, run, decide(run, schema), schema, str(csv))
    ws = openpyxl.load_workbook(review)["Values"]
    assert [c.value for c in ws[1]][:5] == ["Accept", "Part", "Item", "Field", "Value"]

    _edit_review(review,
                 accept=[("101", "", "Power (W)", "285"),            # settle the clash
                         ("101", None, "Flaw Attributes: Laser Team FOD (um)", None),
                         ("101", "", "Concealment 1: Media", None),  # checked the image
                         ("101", "", "Date", None), ("102", "", "Date", None)],
                 ignore=["Operator"])
    res = apply_review(review, fill_dir=tmp / "filled",
                       image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS))

    ledger = ParentCSV(csv)
    assert ledger.sample_ids() == ["101", "102"]
    assert [r["Flaw #"] for r in ledger.rows_for("101")] == ITEMS
    r = ledger.get("101", "2B")
    assert r["Power (W)"] == "285" and r["Frequency (kHz)"] == "50"
    assert r["Flaw Attributes: Laser Team FOD (um)"] == "404"
    assert r["NDE Inspection: Eddy Current Response Pre-Conceal"] == "64"
    assert r["Concealment 1: Media"] == "ceramic"
    assert r["Concealment 2: Media"] == ""                       # not accepted
    assert ledger.get("102", "1A")["Power (W)"] == "300"

    # accepted unsure matches are remembered; the ignored label too
    mapping, ignore = load_aliases(csv)
    assert mapping["fod"] == "Flaw Attributes: Laser Team FOD (um)"
    assert "orientation" not in mapping          # it was a sure match: nothing to learn
    assert "date" not in mapping and ignore == []       # nothing left to ignore

    # the filled template has the values in the template's own cells
    assert sorted(Path(p).name for p in res.filled) == ["tmpl_101.xlsx", "tmpl_102.xlsx"]
    out = openpyxl.load_workbook(tmp / "filled" / "tmpl_101.xlsx").active
    fod = schema.field("Flaw Attributes: Laser Team FOD (um)")
    assert out[fod.cells["2B"]].value == 404
    assert out[schema.field("Power (W)").cells[""]].value == 285
    assert str(out["L7"].value).startswith("=")                  # formulas kept


def test_second_file_uses_what_was_learned(setup):
    schema, csv, sources, tmp = setup
    run = _run(schema, sources[:1], tmp)
    review = tmp / "r1.xlsx"
    write_review(review, run, decide(run, schema), schema, str(csv))
    _edit_review(review, accept=[("101", None, "Flaw Attributes: Laser Team FOD (um)", None)])
    apply_review(review, use_images=False)
    # a new part, same layout: FOD is now pre-accepted
    wb = openpyxl.load_workbook(sources[0])
    wb["Build"]["A3"] = "P/N: 103"
    wb.save(tmp / "laser_log_103.xlsx")
    mapping, ignore = load_aliases(csv)
    dec = decide(extract(schema, [tmp / "laser_log_103.xlsx"], mapping, ignore,
                         use_images=False), schema)
    fod = [d for d in dec if d.obs.field == "Flaw Attributes: Laser Team FOD (um)"]
    assert len(fod) == 9 and all(d.accept and d.obs.via == "alias" for d in fod)


def test_reapplying_a_review_changes_nothing(setup):
    schema, csv, sources, tmp = setup
    run = _run(schema, sources, tmp)
    review = tmp / "review.xlsx"
    write_review(review, run, decide(run, schema), schema, str(csv))
    apply_review(review, image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS))
    before = csv.read_text(encoding="utf-8-sig").splitlines()
    res = apply_review(review, image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS))
    assert "0 rows added, 0 updated" in res.csv_result
    after = csv.read_text(encoding="utf-8-sig").splitlines()
    strip = lambda rows: [",".join(r.split(",")[:-1]) for r in rows]   # updated_at
    assert strip(before) == strip(after)


def test_cli_extract_and_apply(setup, capsys):
    from source_extract import main
    schema, csv, sources, tmp = setup
    out = tmp / "cli_review.xlsx"
    assert main(["extract", str(csv), str(sources[0]), "-o", str(out), "--no-images"]) == 0
    assert out.exists() and "pre-accepted" in capsys.readouterr().out
    assert main(["apply", str(out), "--no-images"]) == 0
    assert "rows" in capsys.readouterr().out
    meta = json.loads(openpyxl.load_workbook(out)["_run"]["A1"].value)
    assert meta["csv"] == str(csv)


REAL = Path.home() / "Downloads" / "LMM Template.xlsx"


@pytest.mark.skipif(not REAL.exists(), reason="the LMM template is not on this machine")
def test_against_the_real_template(tmp_path):
    """The synthetic laser log and NDE book, read against the real 30-flaw
    template: only the 9 flaws they hold are filled, in the right fields."""
    schema = read_template(REAL)
    run = extract(schema, [make_laser_log(tmp_path / "l.xlsx"),
                           make_nde_book(tmp_path / "n.xlsx")],
                  image_reader=fake_reader(CONCEALMENT_IMAGE_ROWS))
    dec = _by(decide(run, schema))
    assert dec[("101", "2B", "NDE Inspection: Eddy Current Response Pre-Conceal", "64")].accept
    assert dec[("101", "3C", "Visual Inspection: Inspector 2", "not seen")].accept
    assert ("101", "1A", "Flaw Attributes: Laser Team FOD (um)", "400") in dec


# ── OCR line repair (what an OCR engine really returned on a "Label: value" image) ─
def test_ocr_split_line_is_joined_and_overlap_deduped():
    boxes = [(16, 20, 167, 38, "PartNumber:2041"),
             (16, 54, 129, 73, "LaserPower("), (120, 54, 191, 74, "(W):285"),
             (16, 88, 185, 107, "Spot Size (mm): 0.08")]
    assert boxes_to_rows(boxes) == [["PartNumber:2041"], ["LaserPower(W):285"],
                                    ["Spot Size (mm): 0.08"]]
    assert tokens("PartNumber") == ["part", "number"]


def test_inline_label_without_space():
    from source_extract import _split_inline
    assert _split_inline("PartNumber:2041") == ("PartNumber", "2041")
    assert _split_inline("10:30") == ("10:30", "")
    assert _split_inline(r"C:\data\run.xlsx") == (r"C:\data\run.xlsx", "")


# ── the real OCR engine, when installed ─────────────────────────────────
def _table_png(path, rows, widths, lines=True):
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype("arial.ttf", 18)
    except OSError:
        pytest.skip("no Arial font to draw test images with")
    im = Image.new("RGB", (sum(widths) + 20, 34 * len(rows) + 20), "white")
    d = ImageDraw.Draw(im)
    for r, row in enumerate(rows):
        x = 10
        for c, v in enumerate(row):
            if lines:
                d.rectangle([x, 10 + 34 * r, x + widths[c], 44 + 34 * r], outline=(160,) * 3)
            d.text((x + 6, 17 + 34 * r), v, font=font, fill="black")
            x += widths[c]
    im.save(path)
    return path


def test_real_ocr_reads_pictures_only_workbook(setup):
    """Part number, settings, a run table and flaw values exist only as
    pictures. Read by Tesseract, every value lands in the right field, a blank
    cell stays blank, "0.030" keeps its zero, and nothing is pre-accepted."""
    from openpyxl.drawing.image import Image as XLImage
    from source_extract import tesseract_reader
    reader = tesseract_reader()
    if reader is None:
        pytest.skip("Tesseract is not installed here")
    schema, _, _, tmp = setup
    pics = {
        "A3": _table_png(tmp / "form.png", [["Part Number: 2041"], ["Laser Power (W): 285"]],
                         [300], lines=False),
        "A10": _table_png(tmp / "conceal.png", [["Media", "Amplitude", "Notes"],
                                                ["ceramic", "40", ""], ["steel", "35", "re-run"]],
                          [110, 110, 100]),
        "H3": _table_png(tmp / "flaws.png", [["Flaw", "FOD (um)", "Length (in)"],
                                             ["1A", "412", "0.031"], ["3A", "", "0.030"]],
                         [80, 110, 120]),
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    for anchor, p in pics.items():
        ws.add_image(XLImage(str(p)), anchor)
    wb.save(tmp / "pictures.xlsx")

    run = extract(schema, [tmp / "pictures.xlsx"], image_reader=reader,
                  image_dir=tmp / "imgs")
    got = {(d.obs.part, d.obs.item, d.obs.field): d.obs.value for d in decide(run, schema)}
    assert got == {
        ("2041", "", "Power (W)"): "285",
        ("2041", "", "Concealment 1: Media"): "ceramic",
        ("2041", "", "Concealment 2: Media"): "steel",
        ("2041", "", "Concealment 1: Amplitude"): "40",
        ("2041", "", "Concealment 2: Amplitude"): "35",
        ("2041", "", "Concealment 2: Notes"): "re-run",
        ("2041", "1A", "Flaw Attributes: Laser Team FOD (um)"): "412",
        ("2041", "1A", "Flaw Attributes: Flaw Length (in)"): "0.031",
        ("2041", "3A", "Flaw Attributes: Flaw Length (in)"): "0.030",
    }
    assert not any(d.accept for d in decide(run, schema))
    assert [n.values for n in run.images] == [1, 5, 3]


# ── which image reader is used ──────────────────────────────────────────
TSV = "\n".join([
    "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext",
    "5\t1\t1\t1\t1\t1\t10\t10\t50\t20\t96\tLaser",
    "5\t1\t1\t1\t1\t2\t66\t10\t60\t20\t95\tPower",
    "5\t1\t1\t1\t1\t3\t132\t10\t30\t20\t95\t(W)",
    "5\t1\t2\t1\t1\t1\t300\t10\t40\t20\t91\t285",
    "4\t1\t2\t1\t1\t0\t300\t10\t40\t20\t-1\t",
    "5\t1\t3\t1\t1\t1\t10\t50\t60\t20\t93\tMedia",
    "5\t1\t4\t1\t1\t1\t300\t50\t70\t20\t90\tceramic",
])


def test_tesseract_words_become_phrases_and_columns():
    from source_extract import tsv_to_boxes
    boxes = tsv_to_boxes(TSV)
    assert sorted(b[4] for b in boxes) == ["285", "Laser Power (W)", "Media", "ceramic"]
    assert boxes_to_rows(boxes) == [["Laser Power (W)", "285"], ["Media", "ceramic"]]


def test_no_tesseract_means_no_reader_not_a_vision_model(monkeypatch, tmp_path):
    import source_extract as se
    monkeypatch.setenv(se.TESSERACT_ENV, str(tmp_path / "missing.exe"))
    asked = []
    monkeypatch.setattr(se, "_ollama", lambda *a, **k: asked.append(a) or {})
    assert se.find_tesseract() is None and se.default_image_reader() is None
    assert asked == []                     # Ollama was not even asked


def test_a_named_tesseract_is_used(monkeypatch, tmp_path):
    import source_extract as se
    exe = tmp_path / "tesseract.exe"
    exe.write_bytes(b"")
    monkeypatch.setenv(se.TESSERACT_ENV, str(exe))
    assert se.find_tesseract() == exe
    assert se.default_image_reader().name == "Tesseract (local)"


def test_only_us_origin_vision_models_are_offered(monkeypatch):
    import source_extract as se
    shown = []

    def fake(path, payload=None, timeout=10):
        if path == "/api/tags":
            return {"models": [{"name": "qwen2.5vl:7b"}, {"name": "minicpm-v:8b"},
                               {"name": "llama3.2-vision:11b"}]}
        shown.append(payload["model"])
        return {"capabilities": ["completion", "vision"]}
    monkeypatch.setattr(se, "_ollama", fake)
    assert se.ollama_vision_model() == "llama3.2-vision:11b"
    assert shown == ["llama3.2-vision:11b"]


def test_table_borders_are_removed_but_a_filled_bar_is_kept():
    pytest.importorskip("numpy")
    from PIL import Image, ImageDraw
    from source_extract import _strip_table_lines
    im = Image.new("L", (400, 200), 255)
    d = ImageDraw.Draw(im)
    d.line([(10, 20), (390, 20)], fill=0, width=2)            # a cell border
    d.line([(200, 5), (200, 195)], fill=0, width=2)           # a column border
    d.rectangle([10, 100, 390, 140], fill=0)                  # a dark header bar
    d.text((20, 60), "412", fill=0)                           # text
    out = _strip_table_lines(im)
    assert out.getpixel((100, 20)) == 255 and out.getpixel((200, 60)) == 255
    assert out.getpixel((100, 120)) == 0                      # the bar is kept
    assert min(out.crop((18, 58, 60, 75)).getdata()) < 128    # the text is kept


# ── several inspectors' readings of one flaw ────────────────────────────
def _two_inspectors(path):
    """JB and MK each logged EC Pre-Conceal for part 101 on one sheet, in the
    template's grid shape; they agree on 1A and disagree on 2A."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EC"
    ws["A1"] = "Part Number: 101"
    r = 3
    for who, pre in (("JB", [60, 61, 62]), ("MK", [60, 66, 62])):
        ws.cell(r, 1, f"Inspector: {who}")
        ws.cell(r + 1, 2, "Flaw #")
        for j, it in enumerate(["1A", "2A", "3A"]):
            ws.cell(r + 1, 3 + j, it)
            ws.cell(r + 2, 3 + j, pre[j])
        ws.cell(r + 2, 2, "Eddy Current Response Pre-Conceal")
        r += 5
    wb.save(path)
    return path


def test_each_inspectors_reading_is_kept_under_their_initials(setup):
    schema, csv, _, tmp = setup
    pre = "NDE Inspection: Eddy Current Response Pre-Conceal"
    assert schema.takes_inspectors(pre)
    assert not schema.takes_inspectors("Visual Inspection: Inspector 1")
    run = extract(schema, [_two_inspectors(tmp / "ec.xlsx")], use_images=False)
    dec = decide(run, schema)
    got = {(d.obs.item, d.obs.field, d.obs.value): d for d in dec}
    # each person's own column: their value, pre-accepted
    for who, vals in (("JB", ["60", "61", "62"]), ("MK", ["60", "66", "62"])):
        for it, v in zip(["1A", "2A", "3A"], vals):
            assert got[(it, f"{pre} [{who}]", v)].accept
    # the template's one cell: agreed values in, a disagreement left to the reviewer
    assert got[("1A", pre, "60")].accept and got[("1A", pre, "60")].obs.inspector == "JB/MK"
    jb, mk = got[("2A", pre, "61")], got[("2A", pre, "66")]
    assert not jb.accept and not mk.accept
    assert jb.conflict == "66 (MK)" and mk.conflict == "61 (JB)"
    assert not any(d.obs.field.startswith("%pec") or d.obs.value in ("JB", "MK") for d in dec)

    review = tmp / "review.xlsx"
    write_review(review, run, dec, schema, str(csv))
    _edit_review(review, accept=[("101", "2A", pre, "66")])     # MK's reading goes in
    res = apply_review(review, fill_dir=tmp / "filled", use_images=False)
    ledger = ParentCSV(csv)
    cols = ledger.columns
    at = cols.index(pre)
    assert cols[at + 1:at + 3] == [f"{pre} [JB]", f"{pre} [MK]"]
    r2 = ledger.get("101", "2A")
    assert (r2[pre], r2[f"{pre} [JB]"], r2[f"{pre} [MK]"]) == ("66", "61", "66")
    assert r2[pre + " — Inspector Initials"] == "MK"
    assert ledger.get("101", "1A")[pre + " — Inspector Initials"] == "JB/MK"
    # the filled template gets the one value; the per-person columns stay in the CSV
    out = openpyxl.load_workbook(tmp / "filled" / "tmpl_101.xlsx").active
    assert out[schema.field(pre).cells["2A"]].value == 66
    assert "rows added" in res.csv_result
    again = apply_review(review, use_images=False)
    assert "0 rows added, 0 updated" in again.csv_result
    # a loosely spelt template column still means the template column
    res = ParentCSV(csv).add_samples([{"Part Number": "101", "Flaw #": "3A",
                                       "NDE inspection - eddy current response pre conceal": "63"}])
    assert res.conflicts == [("101 / 3A", pre, "62", "63")] and not res.unknown_columns


def test_initials_are_never_a_template_word(setup):
    from source_extract import _initials
    assert _initials("j.b.") == "JB" and _initials("Jane Brown") == "Jane Brown"
    assert _initials("12") == "" and _initials("a long sentence about the run") == ""


# ── layouts the sample bench found missing (2026-10-16) ─────────────────
def _decided(schema, wb, tmp, name):
    wb.save(tmp / name)
    run = extract(schema, [tmp / name], use_images=False)
    return run, {(d.obs.part, d.obs.item, d.obs.field): d for d in decide(run, schema)}


def test_runs_for_several_parts_are_numbered_by_run(setup):
    schema, _, _, tmp = setup
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Part", "Run #", "Media", "Amplitude"])
    for row in ([101, 1, "steel", 40], [102, 1, "ceramic", 30], [101, 2, "plastic", 41]):
        ws.append(row)
    _, got = _decided(schema, wb, tmp, "runs.xlsx")
    assert got[("101", "", "Concealment 2: Media")].obs.value == "plastic"
    assert got[("102", "", "Concealment 1: Amplitude")].obs.value == "30"
    assert not any(k[0] in ("1", "2") for k in got)       # "Run #" is not the part


def test_numbers_typed_with_units_are_converted(setup):
    schema, _, _, tmp = setup
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Part Number: 101"
    ws.append([])
    ws.append(["Flaw", "FOD", "Orientation", "Length"])
    ws.append(["1A", "0.245 mm", "45°", '0.028"'])
    ws.append([])                                          # a gap left while typing
    ws.append(["2A", "0.3 mm", "90°", '0.031"'])
    _, got = _decided(schema, wb, tmp, "typed.xlsx")
    fod = got[("101", "1A", "Flaw Attributes: Laser Team FOD (um)")]
    assert fod.obs.value == "245" and "converted mm -> um" in fod.obs.note
    assert got[("101", "1A", "Flaw Attributes: Flaw Orientation (deg)")].obs.value == "45"
    assert got[("101", "1A", "Flaw Attributes: Flaw Length (in)")].obs.value == "0.028"
    assert got[("101", "2A", "Flaw Attributes: Laser Team FOD (um)")].obs.value == "300"


def test_settings_beside_a_table_are_read(setup):
    schema, _, _, tmp = setup
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "P/N: 101"
    ws["B3"], ws["C3"] = "Freq (kHz)", 50
    ws["B4"], ws["C4"] = "Laser Power [W]", 285
    for j, h in enumerate(["Flaw ID", "FOD (um)", "Passes"]):
        ws.cell(3, 5 + j, h)
    for i, it in enumerate(["1A", "2A"], start=4):
        ws.cell(i, 5, it), ws.cell(i, 6, 400 + i), ws.cell(i, 7, 300)
    _, got = _decided(schema, wb, tmp, "beside.xlsx")
    assert got[("101", "", "Frequency (kHz)")].obs.value == "50"
    assert got[("101", "", "Power (W)")].obs.value == "285"
    assert got[("101", "2A", "Flaw Attributes: Laser Team FOD (um)")].obs.value == "405"


def test_a_long_reading_log(setup):
    schema, _, _, tmp = setup
    pre = "NDE Inspection: Eddy Current Response Pre-Conceal"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Inspector", "Part No.", "Flaw", "Stage", "Response"])
    ws.append(["2026-09-25", "JB", 101, "1A", "Pre-Conceal", 70])
    ws.append(["2026-09-25", "MK", 101, "1A", "Pre-Conceal", 72])
    ws.append(["2026-09-25", "JB", 101, "1A", "Post-Conceal", 22])
    run, got = _decided(schema, wb, tmp, "log.xlsx")
    assert got[("101", "1A", f"{pre} [JB]")].obs.value == "70"
    assert got[("101", "1A", f"{pre} [MK]")].obs.value == "72"
    assert got[("101", "1A", "NDE Inspection: Eddy Current Response Post-Conceal")].obs.value == "22"
    assert not any(o.field.endswith("Date") or o.value in ("JB", "MK") for o in run.obs)


def test_an_inspection_date_is_not_the_build_date(setup):
    schema, _, _, tmp = setup
    wb = _two_inspectors(tmp / "ec.xlsx")
    wb = openpyxl.load_workbook(wb)
    wb.active["D3"] = "Date: 2026-09-12"
    run, got = _decided(schema, wb, tmp, "ec_dated.xlsx")
    assert not any(k[2] == "Date" for k in got)
    assert got[("101", "1A", "NDE Inspection: Eddy Current Response Pre-Conceal [JB]")].accept
