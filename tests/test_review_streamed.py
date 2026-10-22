"""The sample Ledger's review workbook is written streaming (openpyxl
write-only): built in memory it was the peak of a long Extract (user,
2026-10-22: "High RAM usage" - 322 MB at the end of 80 messy workbooks, 79 MB
streamed). It must still look as it always has."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
from council_core import ledger_ingest as li  # noqa: E402
from ledger_fixtures import make_laser_log, make_template  # noqa: E402
from test_ledger_examples import messy  # noqa: E402


@pytest.fixture
def review(tmp_path, monkeypatch):
    modes = []
    real = openpyxl.Workbook

    def workbook(*a, **k):
        modes.append(k.get("write_only", False))
        return real(*a, **k)
    monkeypatch.setattr(openpyxl, "Workbook", workbook)
    src = tmp_path / "src"
    src.mkdir()
    csv = Path(li.create_parent(make_template(tmp_path / "T.xlsx"), tmp_path / "l.csv").csv)
    make_laser_log(src / "log.xlsx")
    messy(src / "m.xlsx", 101, [("Zap level", 285), ("Operator note", "fine")])
    ex = li.run_extract(csv, [src], use_images=False)
    monkeypatch.setattr(openpyxl, "Workbook", real)
    return ex, openpyxl.load_workbook(ex.review), modes


def test_it_is_written_streaming(review):
    _ex, _wb, modes = review
    assert True in modes                     # the review: write_only


def test_it_looks_as_it_always_has(review):
    ex, wb, _modes = review
    assert wb.sheetnames == ["Read me", "Values", "Labels to map", "People", "Missing",
                             "Images", "Fields", "_run"]
    assert wb.active.title == "Values"
    assert [wb[n].sheet_state for n in ("Fields", "_run")] == ["hidden", "hidden"]
    readme = wb["Read me"]
    assert readme["A1"].font.b and readme["A1"].font.sz == 14
    assert readme.column_dimensions["A"].width == 120
    values = wb["Values"]
    assert values.freeze_panes == "A2" and all(c.font.b for c in values[1])
    assert values.auto_filter.ref == f"A1:M{values.max_row}"
    assert values.column_dimensions["D"].width == 48
    lists = {str(dv.sqref): dv.formula1 for dv in values.data_validations.dataValidation}
    n = values.max_row
    assert lists == {f"A2:A{n + 200}": '"Y,N"',
                     f"D2:D{n + 200}": f"=Fields!$A$1:$A${wb['Fields'].max_row}"}
    # an amber row is one to decide: every cell of it
    amber = [r for r in values.iter_rows(min_row=2) if r[0].value != "Y"]
    assert amber and all(c.fill.fgColor.rgb.endswith("FFF2CC") for c in amber[0])
    labels = wb["Labels to map"]
    assert "Operator note" in [c.value for c in labels["C"]]
    assert {str(dv.sqref) for dv in labels.data_validations.dataValidation} == {
        f"A2:A{labels.max_row + 100}", f"B2:B{labels.max_row + 100}"}
    assert li.run_apply(ex.review, use_images=False)          # Apply still reads it
