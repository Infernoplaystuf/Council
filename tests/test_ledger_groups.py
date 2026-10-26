"""
A Ledger's sub-groups and their workbooks (council_core.ledger_groups) -
user, 2026-10-26: "create sub groups ... earlier data has some datapoints
that aren't tracked anymore or got condensed down to just yes or no ...
slice the data into smaller output workbooks in addition to the full
workbook"; "separate sheets of the filled out template for each part number
... ordered in numerical order. Data is most commonly separated by part name
but can be date or source folder".

Run:  python -m pytest tests/test_ledger_groups.py -q
"""
from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).parent))

from council_core import ledger_groups as lg  # noqa: E402
from council_core import ledger_ingest as li  # noqa: E402
from council_core import ledger_review as lv  # noqa: E402
from council_core import ledger_saved as ls  # noqa: E402
from ledger_fixtures import make_laser_log, make_nde_book, make_summary, make_template  # noqa: E402,E501
from sample_ledger import FILLED_MARK, LedgerError, ParentCSV, TemplateSchema, \
    schema_path_for  # noqa: E402


@pytest.fixture
def app_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("COUNCIL_APP_DIR", str(tmp_path / "app"))
    return tmp_path / "app"


@pytest.fixture
def parts_csv(tmp_path, app_dir):
    """A Ledger of parts S-10, S-2, S-1, LMX-7: power, date, from two folders."""
    csv = li.create_parent(make_template(tmp_path / "T.xlsx"), tmp_path / "l.csv").csv
    ledger = ParentCSV(csv)
    rows = [("S-10", "300", "2023-05-01", "old.xlsx"), ("S-2", "250", "1/15/2025", "new.xlsx"),
            ("S-1", "", "2022-12-31", "old.xlsx"), ("LMX-7", "410", "2025-03-03", "new.xlsx")]
    ledger.add_samples([{"Part Number": p, "Flaw #": "1A", "Power (W)": w, "Date": d,
                         "source_file": src} for p, w, d, src in rows])
    for name, folder in (("old.xlsx", "2023 data"), ("new.xlsx", "2025 data")):
        (tmp_path / folder).mkdir()
        (tmp_path / folder / name).write_bytes(b"x")
    led = ls.SavedLedger(csv=str(csv), files={
        ls.file_key(tmp_path / f / n): [1, 1] for n, f in (("old.xlsx", "2023 data"),
                                                          ("new.xlsx", "2025 data"))})
    ls.save(led)
    return Path(csv)


def cell_of(csv, fname):
    schema = TemplateSchema.load(schema_path_for(csv))
    return next(f for f in schema.fields if f.name == fname).cells[""]


def test_parts_by_name_date_folder_and_by_hand(parts_csv, tmp_path):
    facts = lg.facts_from_ledger(ParentCSV(parts_csv))
    assert lg.members(lg.Group("s", pattern="S-*"), facts) == ["S-1", "S-2", "S-10"]
    assert lg.members(lg.Group("m", pattern="@@@-#"), facts) == ["LMX-7"]
    old = lg.Group("old", by="date", date_field="Date", date_to="2023-12-31")
    assert lg.members(old, facts) == ["S-1", "S-10"]
    new = lg.Group("new", by="date", date_field="Date", date_from="2025-01-01")
    assert lg.members(new, facts) == ["LMX-7", "S-2"]
    by_folder = lg.Group("f", by="folder", folder=str(tmp_path / "2023 data"))
    assert lg.members(by_folder, facts, parts_csv) == ["S-1", "S-10"]
    hand = lg.Group("h", pattern="S-*", include=["LMX-7"], exclude=["S-2"])
    assert lg.members(hand, facts) == ["LMX-7", "S-1", "S-10"]


def test_dates_and_yes_no_as_people_write_them():
    assert lg.parse_date("2026-10-01") == date(2026, 10, 1)
    assert lg.parse_date("10/1/2026") == date(2026, 10, 1)            # month first
    assert lg.parse_date("1 Oct 2026") == date(2026, 10, 1)
    assert lg.parse_date("2026-10-01 00:00:00") == date(2026, 10, 1)
    assert lg.parse_date("46296") == date(2026, 10, 1)                # an Excel day
    assert lg.parse_date("n/a") is None and lg.parse_date("") is None
    assert lg.as_yes_no("", "") == "No" and lg.as_yes_no("ceramic", "") == "Yes"
    assert lg.as_yes_no("PASS", "pass, ok") == "Yes" and lg.as_yes_no("fail", "pass") == "No"


def test_a_workbook_has_a_sheet_per_part_in_numerical_order(parts_csv):
    out = parts_csv.with_name("all.xlsx")
    lg.write_workbook(parts_csv, ["S-10", "LMX-7", "S-2", "S-1"], out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["LMX-7", "S-1", "S-2", "S-10"]
    power = cell_of(parts_csv, "Power (W)")
    assert wb["S-10"][power].value == 300 and wb["S-1"][power].value is None
    assert wb.properties.keywords == FILLED_MARK
    lg.write_workbook(parts_csv, ["S-1"], out)                 # its own: written again
    assert openpyxl.load_workbook(out).sheetnames == ["S-1"]
    mine = parts_csv.with_name("mine.xlsx")
    openpyxl.Workbook().save(mine)
    with pytest.raises(LedgerError, match="did not write it"):
        lg.write_workbook(parts_csv, ["S-1"], mine)           # a file of the user's


def test_a_template_of_several_sheets_gives_each_part_each_sheet(parts_csv):
    schema = TemplateSchema.load(schema_path_for(parts_csv))
    wb = openpyxl.load_workbook(schema.template_path)
    first = wb.worksheets[0].title
    wb.create_sheet("Notes")["A1"] = "notes"
    wb.save(schema.template_path)
    out = lg.write_workbook(parts_csv, ["S-2", "S-1"], parts_csv.with_name("two.xlsx"))
    got = openpyxl.load_workbook(out)
    assert got.sheetnames == [f"S-1 {first}", "S-1 Notes", f"S-2 {first}", "S-2 Notes"]
    assert got["S-2 Notes"]["A1"].value == "notes"


def test_a_group_leaves_out_its_untracked_fields_and_says_yes_or_no(parts_csv):
    old = lg.Group("Old data", by="date", date_field="Date", date_to="2023-12-31",
                   off=["Date"], yes_no={"Power (W)": ""})
    lg.save_groups(parts_csv, [old])
    assert [g.name for g in lg.load_groups(parts_csv)] == ["Old data"]
    books = lg.write_all(parts_csv)
    assert [Path(b).name for b in books] == ["T - all parts.xlsx", "T - Old data.xlsx"]
    assert Path(books[0]).parent == parts_csv.with_name("l_workbooks")
    full, group = (openpyxl.load_workbook(b) for b in books)
    assert full.sheetnames == ["LMX-7", "S-1", "S-2", "S-10"]
    assert group.sheetnames == ["S-1", "S-10"]
    power, when = cell_of(parts_csv, "Power (W)"), cell_of(parts_csv, "Date")
    assert (group["S-10"][power].value, group["S-1"][power].value) == ("Yes", "No")
    assert group["S-10"][when].value is None and full["S-10"][when].value == "2023-05-01"
    with pytest.raises(LedgerError):
        lg.save_groups(parts_csv, [lg.Group("All parts")])
    with pytest.raises(LedgerError):
        lg.save_groups(parts_csv, [lg.Group("a"), lg.Group("A")])


def test_review_leaves_out_a_field_no_group_of_the_part_tracks(tmp_path, app_dir):
    src = tmp_path / "src"
    src.mkdir()
    make_laser_log(src / "laser_log.xlsx")
    make_nde_book(src / "NDE results.xlsx")
    make_summary(src / "summary.xlsx")
    csv = li.create_parent(make_template(tmp_path / "T.xlsx"), tmp_path / "l.csv").csv
    lg.save_groups(csv, [lg.Group("101 only", pattern="101", off=["Power (W)"])])
    s = li.run_extract(csv, [src], use_images=False)
    items = [i for i in lv.load(s.review).items if i.fields == ["Power (W)"]]
    assert {i.accept for i in items if i.part == "101"} == {"N"}
    assert all(i.accept != "N" for i in items if i.part == "102")
    assert any("do not track" in line for line in s.learned)
    assert "other readings set aside" not in s.review_reasons
