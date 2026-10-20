"""
excel_open — workbooks openpyxl cannot open (rights-protected, .xls, .xlsb),
opened through the installed Excel; and text boxes / charts.

Most of this runs anywhere: a stand-in for Excel.Application checks how Excel
is driven (macros off, read-only, closed unsaved) and how its values come
back. The tests marked `real Excel` drive the installed Excel (hidden, no
window) and are skipped where there is none.

Run:  python -m pytest tests/test_excel_open.py -q
"""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent))
import excel_open  # noqa: E402
from excel_open import (OLE_MAGIC, ExcelOpenError, _chart_from_xml, chart_cells,  # noqa: E402
                        kind_of, read_with_excel, text_cells)


# ── telling files apart ────────────────────────────────────────────────
def _ole(path, *names):
    path.write_bytes(OLE_MAGIC + b"\0" * 504 + b"".join(n.encode("utf-16-le") + b"\0\0"
                                                        for n in names))
    return path


def test_kinds_of_file(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "a.xlsx")
    assert kind_of(tmp_path / "a.xlsx") == "zip"
    assert kind_of(_ole(tmp_path / "old.xls", "Workbook")) == "legacy"
    assert kind_of(_ole(tmp_path / "irm.xlsx", "EncryptedPackage",
                        "DRMEncryptedDataSpace")) == "rights"
    assert kind_of(_ole(tmp_path / "pw.xlsx", "EncryptedPackage", "EncryptionInfo",
                        "StrongEncryptionDataSpace")) == "password"
    (tmp_path / "x.xlsx").write_bytes(b"not a workbook")
    assert kind_of(tmp_path / "x.xlsx") == "unknown"


def test_a_password_file_is_refused_before_excel_is_asked(tmp_path):
    pw = _ole(tmp_path / "pw.xlsx", "EncryptedPackage", "StrongEncryptionDataSpace")
    asked = []
    with pytest.raises(ExcelOpenError, match="password to open"):
        read_with_excel(pw, app_factory=lambda: asked.append(1))
    assert asked == []


# ── a password-to-open .xls (review, 2026-10-19) ───────────────────────
def _cfb(streams: dict, pad: int = 0, ss: int = 512) -> bytes:
    """A compound file (MS-CFB; 512-byte sectors, or 4096 for version 4)
    with `streams` in its root - those under 4096 bytes in the mini stream,
    as Excel writes them. `pad` bytes go before the directory (it then lies
    past the first 8 MB that kind_of scans for stream names)."""
    import struct
    end, free, fatsect, none = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    per = ss // 4
    small = {n: d for n, d in streams.items() if len(d) < 4096}
    mini, minifat, mini_at = b"", [], {}
    for n, d in small.items():
        k = max(1, -(-len(d) // 64))
        mini_at[n] = len(minifat)
        minifat += [len(minifat) + i + 1 for i in range(k - 1)] + [end]
        mini += d.ljust(k * 64, b"\0")
    blobs = []
    if minifat:
        blobs.append(("minifat", struct.pack(f"<{len(minifat)}I", *minifat)))
        blobs.append(("mini", mini))
    blobs += [(n, d) for n, d in streams.items() if n not in small]
    if pad:
        blobs.append(("pad", b"\0" * pad))
    blobs.append(("dir", b"\0" * (128 * (1 + len(streams)))))
    count = {k: max(1, -(-len(b) // ss)) for k, b in blobs}
    n_fat = 1
    while n_fat * per < n_fat + sum(count.values()):
        n_fat += 1
    fat, at, s = [fatsect] * n_fat, {}, n_fat
    for k, _b in blobs:
        at[k] = s
        fat += [s + i + 1 for i in range(count[k] - 1)] + [end]
        s += count[k]
    fat += [free] * (n_fat * per - len(fat))

    def entry(name, typ, right, child, first, size):
        nm = name.encode("utf-16-le") + b"\0\0" if name else b""
        return (nm.ljust(64, b"\0") + struct.pack("<HBBIII", len(nm), typ, 1, none, right, child)
                + b"\0" * 36 + struct.pack("<IQ", first, size))
    names = list(streams)
    d = entry("Root Entry", 5, none, 1 if names else none, at.get("mini", end), len(mini))
    for i, n in enumerate(names):
        d += entry(n, 2, i + 2 if i + 1 < len(names) else none, none,
                   mini_at[n] if n in small else at[n], len(streams[n]))
    d = d.ljust(count["dir"] * ss, b"\0")
    head = (OLE_MAGIC + b"\0" * 16 + struct.pack("<HHHHH", 0x3E, 3 if ss == 512 else 4, 0xFFFE,
                                                  9 if ss == 512 else 12, 6)
            + b"\0" * 6 + struct.pack("<9I", 0, n_fat, at["dir"], 0, 4096,
                                      at.get("minifat", end), count.get("minifat", 0), end, 0)
            + struct.pack("<109I", *(list(range(n_fat)) + [free] * (109 - n_fat))))
    body = head.ljust(ss, b"\0") + struct.pack(f"<{len(fat)}I", *fat)
    for k, b in blobs:
        body += (d if k == "dir" else b).ljust(count[k] * ss, b"\0")
    return body


def _biff(*records) -> bytes:
    import struct
    return b"".join(struct.pack("<HH", t, len(b)) + b for t, b in records)


BOF = (0x0809, bytes.fromhex("000605005a4fcd07c100020006080000"))
EOF = (0x000A, b"")
#: FILEPASS records Excel 16 wrote. RC4 CryptoAPI for SaveAs(Password="x");
#: RC4 under Excel's default password (VelvetSweatshop) for a workbook whose
#: structure is protected with a password - Excel opens that one unasked.
FILEPASS_PASSWORD = bytes.fromhex(
    "0100040002000c0000007e0000000c000000000000000168000004800000800000000100"
    "000000000000000000004d006900630072006f0073006f0066007400200045006e006800"
    "61006e006300650064002000430072007900700074006f00670072006100700068006900"
    "63002000500072006f00760069006400650072002000760031002e003000000010000000"
    "3b138bb43654c22b71844d5a9f9b5e63fd6cfd611e4428ad045b0d86bfabff8014000000"
    "a38113351641add8b8eaead5b9ab6c850c3c0b4a")
FILEPASS_DEFAULT = bytes.fromhex(
    "010001000100a0935bbdfab947a4cbc6970230f0975d343e73bb1ff12d80139c3327f625"
    "cbce8de7320c771af5402412f57e77a189ae")


def _xls(path, *records, size=6000, stream="Workbook"):
    """An .xls whose Workbook stream holds `records` (padded to `size`)."""
    path.write_bytes(_cfb({stream: _biff(*records).ljust(size, b"\0"),
                           "\x05SummaryInformation": b"\xfe\xff" + b"\0" * 200}))
    return path


def test_a_password_xls_is_told_by_its_filepass_record(tmp_path):
    """Excel showed a VISIBLE 'Password' dialog for a password .xls (the
    Password= argument does not stop it for .xls) and the read sat on it
    for 2 minutes, reported as a sign-in timeout."""
    interface = (0x00E1, b"\xb0\x04")
    for size in (6000, 300):                         # a regular and a mini stream
        pw = _xls(tmp_path / f"pw{size}.xls", BOF, (0x002F, FILEPASS_PASSWORD), interface,
                  size=size)
        assert kind_of(pw) == "password"
        asked = []
        with pytest.raises(ExcelOpenError, match="password to open") as e:
            read_with_excel(pw, app_factory=lambda: asked.append(1))
        assert e.value.kind == "password" and asked == []
        assert kind_of(_xls(tmp_path / f"plain{size}.xls", BOF, interface, EOF,
                            size=size)) == "legacy"
    # Excel's own default password: Excel opens it without asking, so it is read
    assert kind_of(_xls(tmp_path / "wbprot.xls", BOF, (0x002F, FILEPASS_DEFAULT),
                        interface)) == "legacy"
    # Excel 5 (a 'Book' stream): XOR obfuscation, a 4-byte FILEPASS
    xor_default = (0x002F, b"\x34\x12" + (0x9A0A).to_bytes(2, "little"))
    xor_other = (0x002F, b"\x34\x12" + (0xCEBA).to_bytes(2, "little"))
    assert kind_of(_xls(tmp_path / "b5.xls", BOF, xor_default, stream="Book")) == "legacy"
    assert kind_of(_xls(tmp_path / "b5pw.xls", BOF, xor_other, stream="Book")) == "password"
    # what is not understood is left to Excel, as before
    assert kind_of(_xls(tmp_path / "odd.xls", (0x1234, b"zz"), (0x002F, b"x"))) == "legacy"
    assert kind_of(_ole(tmp_path / "torn.xls", "Workbook")) == "legacy"


def test_the_xor_verifier_is_excels_password_hash():
    from openpyxl.utils.protection import hash_password
    for pw in ("VelvetSweatshop", "x", "Password1", "a longer pass phrase"):
        assert f"{excel_open._xor_verifier(pw):X}" == hash_password(pw)


def test_a_password_xls_is_counted_as_a_password_in_the_run(tmp_path, monkeypatch):
    import source_extract as se
    from ledger_fixtures import make_template
    from sample_ledger import read_template
    # never the real Excel, even if the file were not told apart
    monkeypatch.setattr(excel_open, "excel_available", lambda: False)
    pw = _xls(tmp_path / "locked.xls", BOF, (0x002F, FILEPASS_PASSWORD))
    run = se.extract(read_template(make_template(tmp_path / "t.xlsx")), [pw], use_images=False)
    assert run.open_failures == [("locked.xls", "password")]
    assert any("password to open" in p for p in run.problems)


def test_an_encrypted_package_is_found_past_the_first_8_mb(tmp_path):
    """kind_of looked for the stream names only in the first 8 MB; a big
    encrypted .xlsx keeps its directory at the end."""
    big = _cfb({"EncryptionInfo": b"\x04\x00\x04\x00" + b"\0" * 5000,
                "EncryptedPackage": b"\0" * 5000}, pad=9 << 20, ss=4096)
    (tmp_path / "big.xlsx").write_bytes(big)
    assert kind_of(tmp_path / "big.xlsx") == "password"


def _cfb_with_storage(streams: dict, storage: str, inner: dict) -> bytes:
    """A compound file with `streams` in its root and a storage `storage`
    holding `inner` - an embedded object's, as an .xls keeps one ('MBD0001').
    Each stream a whole number of 512-byte sectors, 4096 bytes or more (none
    in the mini stream)."""
    import struct
    end, free, fatsect, none = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    blobs = list(streams.values()) + list(inner.values())
    assert all(len(b) >= 4096 and len(b) % 512 == 0 for b in blobs)
    fat, first, s = [fatsect], [], 1                 # sector 0: the FAT
    for b in blobs:
        first.append(s)
        fat += [s + i + 1 for i in range(len(b) // 512 - 1)] + [end]
        s += len(b) // 512

    def entry(name, typ, right, child, at, size):
        nm = name.encode("utf-16-le") + b"\0\0"
        return (nm.ljust(64, b"\0") + struct.pack("<HBBIII", len(nm), typ, 1, none, right, child)
                + b"\0" * 36 + struct.pack("<IQ", at, size))
    n, m = len(streams), len(inner)
    d = entry("Root Entry", 5, none, 1, end, 0)
    for i, (name, b) in enumerate(streams.items()):    # root's children, the storage last
        d += entry(name, 2, i + 2, none, first[i], len(b))
    d += entry(storage, 1, none, n + 2 if m else none, 0, 0)
    for j, (name, b) in enumerate(inner.items()):
        d += entry(name, 2, n + 3 + j if j + 1 < m else none, none, first[n + j], len(b))
    d = d.ljust(-(-len(d) // 512) * 512, b"\0")
    dir_at = s
    fat += [dir_at + i + 1 for i in range(len(d) // 512 - 1)] + [end]
    assert len(fat) <= 128
    fat += [free] * (128 - len(fat))
    head = (OLE_MAGIC + b"\0" * 16 + struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6)
            + b"\0" * 6 + struct.pack("<9I", 0, 1, dir_at, 0, 4096, end, 0, end, 0)
            + struct.pack("<109I", 0, *([free] * 108)))
    return head.ljust(512, b"\0") + struct.pack("<128I", *fat) + b"".join(blobs) + d


def test_an_encrypted_embedded_object_does_not_lock_the_workbook(tmp_path):
    """The scan of the first 8 MB for 'EncryptedPackage' ran before the
    directory was read, so an .xls holding a password-protected document as
    an embedded object was refused as 'password' - and the knowledge graph
    marked it unreadable for good - though Excel opens it without asking
    (review, 2026-10-19). Only a stream in the root counts."""
    book = _biff(BOF, (0x00E1, b"\xb0\x04"), EOF).ljust(4096, b"\0")
    sealed = {"EncryptionInfo": b"\x04\x00\x04\x00".ljust(4096, b"\0"),
              "EncryptedPackage": b"\0" * 4096}
    embedded = tmp_path / "embedded.xls"
    embedded.write_bytes(_cfb_with_storage({"Workbook": book}, "MBD0001", sealed))
    assert kind_of(embedded) == "legacy"
    # the same streams in the root: the file itself is encrypted
    locked = tmp_path / "locked.xlsx"
    locked.write_bytes(_cfb_with_storage(sealed, "MBD0001", {"Workbook": book}))
    assert kind_of(locked) == "password"
    # an embedded password .xls (FILEPASS in ITS Workbook stream) is that object's too
    inner = _biff(BOF, (0x002F, FILEPASS_PASSWORD), EOF).ljust(4096, b"\0")
    host = tmp_path / "host.xls"
    host.write_bytes(_cfb_with_storage({"Workbook": book}, "MBD0002", {"Workbook": inner}))
    assert kind_of(host) == "legacy"


# ── text boxes and charts as tables ─────────────────────────────────────
def test_text_box_lines_become_label_value_cells():
    assert text_cells("Scan Speed (mm/s): 1800\r\nSpot Size (mm): 0.05\n") == {
        (1, 1): "Scan Speed (mm/s): 1800", (2, 1): "Spot Size (mm): 0.05"}


def test_a_chart_of_inspectors_becomes_one_row_per_reading():
    cells = chart_cells("EC Mid-Conceal", "", ["1A", "2A"], [("JB", [30, 31]), ("MK", [32, None])])
    rows = sorted((r, c, v) for (r, c), v in cells.items())
    assert cells[(3, 1)] == "Flaw" and cells[(3, 2)] == "Inspector" and cells[(3, 3)] == "EC Mid-Conceal"
    assert [(cells[(r, 1)], cells[(r, 2)], cells[(r, 3)]) for r in (4, 5, 6)] == [
        ("1A", "JB", 30), ("2A", "JB", 31), ("1A", "MK", 32)]
    assert (7, 1) not in cells and rows                  # a blank reading is not a row


def test_a_chart_of_measurements_is_a_wide_table():
    cells = chart_cells("", "Flaw", ["1A", "2A"], [("EC Pre", [70, 71]), ("EC Post", [20, 21])])
    assert [cells[(1, c)] for c in (1, 2, 3)] == ["Flaw", "EC Pre", "EC Post"]
    assert [cells[(3, c)] for c in (1, 2, 3)] == ["2A", 71, 21]


CHART_XML = b"""<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
              xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>EC Pre-Conceal</a:t></a:r></a:p></c:rich></c:tx></c:title>
  <c:plotArea><c:barChart>
   <c:ser><c:tx><c:v>JB</c:v></c:tx>
    <c:cat><c:strLit><c:ptCount val="2"/><c:pt idx="0"><c:v>1A</c:v></c:pt><c:pt idx="1"><c:v>2A</c:v></c:pt></c:strLit></c:cat>
    <c:val><c:numLit><c:ptCount val="2"/><c:pt idx="0"><c:v>70</c:v></c:pt><c:pt idx="1"><c:v>71.5</c:v></c:pt></c:numLit></c:val></c:ser>
   <c:ser><c:tx><c:strRef><c:f>S!B1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>MK</c:v></c:pt></c:strCache></c:strRef></c:tx>
    <c:val><c:numRef><c:f>S!B2:B3</c:f><c:numCache><c:ptCount val="2"/><c:pt idx="1"><c:v>73</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser>
  </c:barChart></c:plotArea></c:chart></c:chartSpace>"""


def test_chart_values_come_from_the_saved_cache_or_literals():
    cells = _chart_from_xml(CHART_XML)
    got = {(cells[(r, 1)], cells[(r, 2)]): cells[(r, 3)] for r in range(4, 8) if (r, 1) in cells}
    assert got == {("1A", "JB"): 70, ("2A", "JB"): 71.5, ("2A", "MK"): 73}


# ── how Excel is driven: a stand-in for Excel.Application ───────────────
class _Range:
    def __init__(self, row=1, col=1, value=None, merged=None):
        self.Row, self.Column, self.Value = row, col, value
        self._merged = merged
        self.MergeCells = merged is not None

    @property
    def MergeArea(self):
        r, c, nr, nc = self._merged
        area = _Range(r, c)
        area.Rows = type("n", (), {"Count": nr})()
        area.Columns = type("n", (), {"Count": nc})()
        return area


class _Shape:
    def __init__(self, typ, text="", chart=None):
        self.Type, self.HasChart = typ, chart is not None
        self.Chart = chart
        # as real Excel: Address is a STRING property; its arguments go to
        # GetAddress (Address(False, False) raises - the user's "?" anchors)
        self.TopLeftCell = type("c", (), {"Address": "$E$2",
                                          "GetAddress": lambda self, a, b: "E2"})()
        tr = type("tr", (), {"Text": text})()
        self.TextFrame2 = type("tf", (), {"HasText": bool(text), "TextRange": tr})()


class _Series:
    def __init__(self, name, x, y):
        self.Name, self.XValues, self.Values = name, x, y


class _Coll:
    def __init__(self, items):
        self._items = items
        self.Count = len(items)

    def Item(self, i):
        return self._items[i - 1]


class _Chart:
    HasTitle = True
    ChartTitle = type("t", (), {"Text": "EC Pre-Conceal"})()

    def SeriesCollection(self):
        return _Coll([_Series("JB", ("1A", "2A"), (70.0, 71.0))])

    def Axes(self, i):
        return type("ax", (), {"HasTitle": False})()


class _Sheet:
    Name = "EC"

    def __init__(self):
        when = dt.datetime(2026, 9, 12, 0, 0, tzinfo=dt.timezone.utc)
        self.UsedRange = _Range(1, 1, (("Flaw report", None, None),
                                       ("Part Number", 2041.0, when),
                                       ("Power (W)", -2146826281, 285.5)))
        self.UsedRange.MergeCells = None        # as Excel says: some of it (A1:C1) is merged
        self.Shapes = _Coll([_Shape(17, "Spot Size (mm): 0.05"), _Shape(3, chart=_Chart())])

    def Cells(self, r, c):
        return _Range(r, c, merged=(1, 1, 1, 3) if (r, c) == (1, 1) else None)


class _App:
    def __init__(self, log, fail=None):
        self.log, self.fail = log, fail
        app = self

        class Books:
            def Open(self, path, **kw):
                app.log.append(("open", Path(path).name, kw))
                if app.fail:
                    raise Exception(app.fail)
                return app._book
        self.Workbooks = Books()

        class Book:
            Worksheets = [_Sheet()]

            def Close(self, SaveChanges=True):
                app.log.append(("close", SaveChanges))
        self._book = Book()

    def __setattr__(self, k, v):
        if k[0].isupper() and k != "Workbooks":
            self.log.append(("set", k, v))
        object.__setattr__(self, k, v)

    def Quit(self):
        self.log.append(("quit",))


def test_excel_is_driven_read_only_with_macros_off_and_closed_unsaved(tmp_path):
    f = tmp_path / "book.xlsb"
    f.write_bytes(b"PK\x03\x04")
    log = []
    book = read_with_excel(f, want_pictures=False, app_factory=lambda: _App(log))
    sets = {e[1]: e[2] for e in log if e[0] == "set"}
    assert sets["AutomationSecurity"] == 3 and sets["Visible"] is False
    assert sets["EnableEvents"] is False and sets["DisplayAlerts"] is False
    opened = next(e for e in log if e[0] == "open")[2]
    assert opened["ReadOnly"] is True and opened["UpdateLinks"] == 0 and opened["AddToMru"] is False
    assert ("close", False) in log and log[-1] == ("quit",)
    (name, cells), = book.sheets
    assert cells[(2, 2)] == 2041 and cells[(3, 3)] == 285.5          # floats as Excel shows them
    assert cells[(2, 3)] == dt.datetime(2026, 9, 12)                 # COM dates, naive
    assert (3, 2) not in cells                                       # #DIV/0! is no value
    assert cells[(1, 2)] == ("__merged__", "Flaw report")            # a merged title
    assert [e.name for e in book.extras] == ["EC › text box 1", "EC › chart 2"]
    assert book.extras[0].cells == {(1, 1): "Spot Size (mm): 0.05"}


class _GridSheet:
    """A sheet of text cells, some of them merged, that counts what Excel is
    asked: MergeCells is True / False / None (all / none / some of the range
    merged), as Excel answers. `asked` counts the MergeCells questions,
    `calls` every call into Excel (each costs ~1-1.6 ms in a real one).
    `data` ({(r, c): value}) replaces the text grid."""
    Name = "Grid"

    def __init__(self, rows, cols, merged=(), data=None):
        self.merged = list(merged)                       # (r1, c1, r2, c2)
        self.asked = self.calls = 0
        self.at = {(r, c): a for a in self.merged
                   for r in range(a[0], a[2] + 1) for c in range(a[1], a[3] + 1)}
        sheet = self
        if data is None:
            data = {(r, c): f"t{r}.{c}" for r in range(1, rows + 1) for c in range(1, cols + 1)
                    if (r, c) not in self.at or self.at[(r, c)][:2] == (r, c)}
        grid = tuple(tuple(data.get((r, c)) for c in range(1, cols + 1))
                     for r in range(1, rows + 1))

        def counted(f):
            def get(self):
                sheet.calls += 1
                return f(self)
            return property(get)

        class Rng:
            def __init__(self, r1, c1, r2, c2, value=None):
                self.box, self.Value = (r1, c1, r2, c2), value

            Row = counted(lambda self: self.box[0])
            Column = counted(lambda self: self.box[1])
            Rows = counted(lambda self: _GridSheet._n(sheet, self.box[2] - self.box[0] + 1))
            Columns = counted(lambda self: _GridSheet._n(sheet, self.box[3] - self.box[1] + 1))

            @counted
            def Address(self):
                from openpyxl.utils.cell import get_column_letter as col
                r1, c1, r2, c2 = self.box
                a = f"${col(c1)}${r1}"
                return a if (r1, c1) == (r2, c2) else f"{a}:${col(c2)}${r2}"

            @counted
            def MergeCells(self):
                sheet.asked += 1
                r1, c1, r2, c2 = self.box
                inside = [(r, c) in sheet.at for r in range(r1, r2 + 1)
                          for c in range(c1, c2 + 1)]
                return True if all(inside) else (False if not any(inside) else None)

            @counted
            def MergeArea(self):
                a = sheet.at.get(self.box[:2])
                return Rng(*a) if a else self
        self._Rng = Rng
        self.UsedRange = Rng(1, 1, rows, cols, grid)
        self.Shapes = _Coll([])

    @staticmethod
    def _n(sheet, n):
        class N:
            @property
            def Count(self):
                sheet.calls += 1
                return n
        return N()

    def Cells(self, r, c):
        self.calls += 1
        return self._Rng(r, c, r, c)

    def Range(self, address):
        from openpyxl.utils.cell import range_boundaries
        self.calls += 1
        c1, r1, c2, r2 = range_boundaries(address)
        return self._Rng(r1, c1, r2, c2)


def _read_sheet(tmp_path, sheet):
    app = _App([])
    app._book.Worksheets = [sheet]
    f = tmp_path / "grid.xls"
    f.write_bytes(OLE_MAGIC)
    (_name, cells), = read_with_excel(f, want_pictures=False, app_factory=lambda: app).sheets
    return cells


def test_merged_labels_are_found_without_asking_cell_by_cell(tmp_path):
    """Two COM calls per text cell took 24-39 s for a 12,000-cell .xls and
    set off the stall watchdog in ordinary reads (review, 2026-10-19)."""
    plain = _GridSheet(60, 40)
    cells = _read_sheet(tmp_path, plain)
    assert len(cells) == 2400 and plain.asked == 1               # one question: none merged
    areas = [(2, 2, 2, 5), (10, 1, 14, 1), (30, 30, 32, 33), (60, 39, 60, 40)]
    sheet = _GridSheet(60, 40, areas)
    cells = _read_sheet(tmp_path, sheet)
    fill = {k: v for k, v in cells.items() if isinstance(v, tuple)}
    want = {(r, c): ("__merged__", f"t{a[0]}.{a[1]}") for a in areas
            for r in range(a[0], a[2] + 1) for c in range(a[1], a[3] + 1) if (r, c) != a[:2]}
    assert fill == want                                          # the same labels as before
    assert sheet.asked < 150, sheet.asked                        # not ~2,400


def test_a_merged_cell_in_every_row_takes_a_handful_of_questions(tmp_path):
    """A Remarks column merged G:H in every row, and a form (A:B label, C
    value, D:E label, F value): halving the rows found every half mixed down
    to single rows - about one question per text cell, sometimes more - and
    each merged cell's extent took 7-8 more calls. Such a table still took
    tens of seconds (review, 2026-10-19)."""
    table = [(r, 7, r, 8) for r in range(1, 301)]
    form = [a for r in range(1, 301) for a in ((r, 1, r, 2), (r, 4, r, 5))]
    for cols, areas in ((8, table), (6, form)):
        sheet = _GridSheet(300, cols, areas)
        cells = _read_sheet(tmp_path, sheet)
        texts = sum(isinstance(v, str) for v in cells.values())
        fill = {k: v for k, v in cells.items() if isinstance(v, tuple)}
        assert fill == {(a[0], a[3]): ("__merged__", f"t{a[0]}.{a[1]}") for a in areas}
        assert texts >= 1200 and sheet.asked <= 20, (texts, sheet.asked)
        # each range's extent from its address: Cells, MergeArea, Address
        assert sheet.calls <= 3 * len(areas) + 60, sheet.calls


def test_merged_cells_never_cost_much_more_than_asking_cell_by_cell(tmp_path):
    """Halving asked up to twice as many questions as there are text cells
    on a sheet merged every which way (review, 2026-10-19). Past one per text
    cell and _MERGE_SLACK more, the rest is asked cell by cell."""
    # two-cell ranges, staggered from row to row: halves are mixed both ways
    areas = [(r, c, r, c + 1) for r in range(1, 41) for c in range(1 + 2 * (r % 2), 15, 4)]
    sheet = _GridSheet(40, 16, areas)
    cells = _read_sheet(tmp_path, sheet)
    texts = sum(isinstance(v, str) for v in cells.values())
    assert sheet.asked <= texts + excel_open._MERGE_SLACK, (sheet.asked, texts)
    assert {k for k, v in cells.items() if isinstance(v, tuple)} == {(a[0], a[3]) for a in areas}


def test_merged_labels_match_asking_every_text_cell():
    """The labels the per-cell loop gave, on random sheets: some merged
    ranges keep a value in more than their first cell (written by something
    other than Excel), numbers are never repeated, and a range is asked
    about once however many text cells it holds."""
    import random
    rnd = random.Random(19)
    for _ in range(200):
        rows, cols = rnd.randint(1, 30), rnd.randint(1, 12)
        taken, areas = set(), []
        for _ in range(rnd.randint(0, 20)):
            r1, c1 = rnd.randint(1, rows), rnd.randint(1, cols)
            r2, c2 = min(rows, r1 + rnd.randint(0, 3)), min(cols, c1 + rnd.randint(0, 3))
            box = {(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}
            if len(box) > 1 and not box & taken:
                taken |= box
                areas.append((r1, c1, r2, c2))
        data = {}
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                x = rnd.random()
                if x < 0.5:
                    data[(r, c)] = f"t{r}.{c}"
                elif x < 0.7:
                    data[(r, c)] = float(r * c)
        sheet = _GridSheet(rows, cols, areas, data)
        got = dict(data)
        excel_open._fill_merged(sheet, sheet.UsedRange, got)
        want = dict(data)                        # the loop it replaced
        for (r, c), v in data.items():
            if isinstance(v, str) and (r, c) in sheet.at:
                r1, c1, r2, c2 = sheet.at[(r, c)]
                for rr in range(r1, r2 + 1):
                    for cc in range(c1, c2 + 1):
                        want.setdefault((rr, cc), ("__merged__", v))
        assert got == want, areas


def test_a_picture_that_cannot_be_saved_is_reported_and_the_run_goes_on(tmp_path, monkeypatch):
    """The user's run stopped on 'OSError 22 ... Sheet1_?_1.png'. Now the name
    is cleaned, and a save that still fails is a problem, not a crash."""
    import source_extract as se
    from ledger_fixtures import make_template, make_nde_book
    from sample_ledger import read_template
    schema = read_template(make_template(tmp_path / "t.xlsx"))
    src = make_nde_book(tmp_path / "NDE results.xlsx")           # has one picture
    real = se.load_workbook_grids

    def odd_anchor(path, notes=None):
        grids, images = real(path, notes)
        for im in images:
            im.anchor, im.sheet = "?", "Sheet:1"                 # what Excel gave back
        return grids, images
    monkeypatch.setattr(se, "load_workbook_grids", odd_anchor)
    run = se.extract(schema, [src], image_dir=tmp_path / "imgs", use_images=False)
    assert [Path(n.saved_as).name for n in run.images] == ["NDE_results_Sheet_1___1.png"]

    def refuse(self, data):
        raise OSError(22, "Invalid argument")
    monkeypatch.setattr(Path, "write_bytes", refuse)
    run = se.extract(schema, [src], image_dir=tmp_path / "imgs2", use_images=False)
    assert any("could not be saved" in p for p in run.problems)
    assert run.obs                                              # the rest was still read


def _busy_clipboard(monkeypatch) -> list:
    """_copy_picture as when another program holds the clipboard open."""
    tried = []

    def busy(shape):
        tried.append(shape)
        why = getattr(excel_open, "_COPY_WHY", None)
        if why is not None:
            why.why = "busy"
        return None
    monkeypatch.setattr(excel_open, "_copy_picture", busy)
    monkeypatch.setattr(excel_open, "_clip_text", lambda: None)
    monkeypatch.setattr(excel_open, "_clip_restore", lambda text: None)   # the real one stays
    return tried


def _picture_app(n=1):
    sheet = _Sheet()
    sheet.Shapes = _Coll([_Shape(13) for _ in range(n)])     # pictures at E2
    app = _App([])
    app._book.Worksheets = [sheet]
    return app


def test_a_picture_the_busy_clipboard_kept_out_is_still_listed(tmp_path, monkeypatch):
    """It left the picture count and the Images sheet, and the note blamed the
    workbook ('may forbid copying') - review, 2026-10-19."""
    tried = _busy_clipboard(monkeypatch)
    f = tmp_path / "pics.xls"
    f.write_bytes(OLE_MAGIC)
    app = _picture_app(2)
    book = read_with_excel(f, app_factory=lambda: app)
    assert [(p.sheet, p.anchor, p.data) for p in book.pictures] == [("EC", "E2", b"")] * 2
    assert len(tried) == 1                                   # the second waits for nothing
    assert [n for n in book.notes if "clipboard was busy" in n] == [
        "pics.xls › EC!E2: a picture could not be copied out: the clipboard was busy "
        "(another program held it open); read the workbook again later"]
    assert not any("may forbid" in n for n in book.notes)


def test_a_picture_excel_could_not_copy_is_counted_as_not_read(tmp_path, monkeypatch):
    import source_extract as se
    from council_core.ledger_ingest import _images_read
    from ledger_fixtures import make_template
    from sample_ledger import read_template
    _busy_clipboard(monkeypatch)
    app = _picture_app()
    real = excel_open.read_with_excel
    monkeypatch.setattr(excel_open, "read_with_excel", lambda path, want_pictures=True, **kw:
                        real(path, want_pictures=want_pictures, app_factory=lambda: app))
    f = tmp_path / "pics.xls"
    f.write_bytes(OLE_MAGIC)
    read = []

    def reader(data):
        read.append(data)
        return []
    reader.name = "OCR"
    run = se.extract(read_template(make_template(tmp_path / "t.xlsx")), [f],
                     image_dir=tmp_path / "imgs", image_reader=reader)
    assert [(n.sheet, n.anchor, n.saved_as, n.reader) for n in run.images] == [("EC", "E2", "", "")]
    assert run.images[0].text.startswith("not read") and _images_read(run.images) == 0
    assert read == [] and not (tmp_path / "imgs").exists()      # no empty file saved or read
    assert any("clipboard was busy" in p for p in run.problems)


def test_copy_picture_says_when_the_clipboard_was_busy(monkeypatch):
    class Refused:
        def CopyPicture(self, **kw):
            raise Exception("CopyPicture method of Shape class failed")
    monkeypatch.setattr(excel_open, "_clipboard_busy", lambda timeout=0.5: True)
    assert excel_open._copy_picture(Refused()) is None and excel_open._COPY_WHY.why == "busy"
    monkeypatch.setattr(excel_open, "_clipboard_busy", lambda timeout=0.5: False)
    assert excel_open._copy_picture(Refused()) is None and excel_open._COPY_WHY.why == ""


def test_a_busy_clipboard_is_not_tried_again_through_imagegrab(monkeypatch):
    """_clipboard_dib waited 5 s for the clipboard and gave up ('busy');
    ImageGrab then tried again, holding the GIL - one more freeze of the
    whole app per workbook, for a copy that could not succeed (review,
    2026-10-19)."""
    from PIL import ImageGrab
    grabbed = []
    monkeypatch.setattr(ImageGrab, "grabclipboard", lambda *a, **kw: grabbed.append(1))

    def never_opens(timeout=5.0):
        excel_open._COPY_WHY.why = "busy"
        return None

    class Copied:
        def CopyPicture(self, **kw):
            pass
    monkeypatch.setattr(excel_open, "_clipboard_dib", never_opens)
    assert excel_open._copy_picture(Copied()) is None
    assert excel_open._COPY_WHY.why == "busy" and grabbed == []
    # it opened but held no bitmap: ImageGrab is still asked, as before
    monkeypatch.setattr(excel_open, "_clipboard_dib", lambda timeout=5.0: None)
    assert excel_open._copy_picture(Copied()) is None and grabbed == [1]


def test_a_clipboard_that_never_opens_is_busy(monkeypatch):
    win32clipboard = pytest.importorskip("win32clipboard")

    def held(*a):
        raise OSError("OpenClipboard: access is denied")
    monkeypatch.setattr(win32clipboard, "OpenClipboard", held)
    excel_open._COPY_WHY.why = ""
    assert excel_open._clipboard_dib(timeout=0.1) is None
    assert excel_open._COPY_WHY.why == "busy"


def test_an_excel_failure_says_what_happened(tmp_path):
    f = tmp_path / "x.xls"
    f.write_bytes(OLE_MAGIC)
    log = []
    with pytest.raises(ExcelOpenError, match="Excel could not open x.xls: corrupt"):
        read_with_excel(f, app_factory=lambda: _App(log, fail="corrupt file"))
    assert log[-1] == ("quit",)                                      # Excel still closed


def test_folders_bring_every_workbook_kind(tmp_path):
    from council_core.ledger_ingest import collect_sources
    for n in ("a.xlsx", "b.xlsm", "c.xls", "d.xlsb", "~$a.xlsx", "notes.txt"):
        (tmp_path / n).write_bytes(b"")
    kept, _ignored = collect_sources([tmp_path])
    assert sorted(p.name for p in kept) == [
        "a.xlsx", "b.xlsm", "c.xls", "d.xlsb"]


# ── the installed Excel itself ──────────────────────────────────────────
real_excel = pytest.mark.skipif(not excel_open.excel_available(),
                                reason="no Excel automation here (Windows + Excel + pywin32)")


def _record_excel_pids(monkeypatch) -> list:
    """The process id of every hidden Excel started from now on in this
    process — win32com.client.DispatchEx wrapped (excel_open and
    ledger_bench both import it when they start one), each instance's
    process read from its window, as excel_open._watchdog does.

    BY PROCESS ID, NOT BY COUNTING EXCEL.EXE. The test counted every
    EXCEL.EXE on the PC against the count before it, and a whole-suite run
    8 files at a time failed it with "assert 3 <= 0" (2026-10-18):
    tests/test_ledger_bench.py was building its workbooks in Excel at the
    same time, in another process. Whether THIS test's Excels were closed
    is the claim, so they are the ones looked for."""
    import win32com.client
    import win32process
    real = win32com.client.DispatchEx
    pids: list = []

    def recording(*args, **kwargs):
        app = real(*args, **kwargs)
        try:
            pids.append(int(win32process.GetWindowThreadProcessId(int(app.Hwnd))[1]))
        except Exception as exc:                 # an Excel this test cannot follow
            pids.append(repr(exc))
        return app
    monkeypatch.setattr(win32com.client, "DispatchEx", recording)
    return pids


def _excel_running(pid: int) -> bool:
    import subprocess
    out = subprocess.run(["tasklist", "/NH", "/FI", f"PID eq {pid}",
                          "/FI", "IMAGENAME eq EXCEL.EXE"], capture_output=True,
                         text=True).stdout
    return "EXCEL.EXE" in out.upper()


@real_excel
def test_real_excel_xls_xlsb_text_box_chart_and_password(tmp_path, monkeypatch):
    from ledger_bench import _excel_save
    from source_extract import LedgerError, load_workbook_grids
    started = _record_excel_pids(monkeypatch)
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "Power (W)", 285
    wb.save(tmp_path / "s.xlsx")

    from PIL import Image
    Image.new("RGB", (120, 40), "white").save(tmp_path / "pic.png")

    def box_and_chart(sheet):
        sheet.Shapes.AddPicture(str(tmp_path / "pic.png"), False, True, 150, 60, 120, 40)
        sheet.Shapes.AddTextbox(1, 20, 20, 200, 40).TextFrame2.TextRange.Text = "Spot Size (mm): 0.05"
        ch = sheet.ChartObjects().Add(240, 20, 300, 200).Chart
        ser = ch.SeriesCollection().NewSeries()
        ser.Name, ser.Values, ser.XValues = "EC Pre", (70, 71), ("1A", "2A")
    _excel_save(tmp_path / "s.xlsx", tmp_path / "old.xls", 56, edit=box_and_chart)
    grids, images = load_workbook_grids(tmp_path / "old.xls")
    assert [(i.sheet, i.anchor) for i in images] == [("Sheet", "D5")] and images[0].data
    # the whole run: the picture is saved under a name Windows accepts
    from sample_ledger import read_template
    from ledger_fixtures import make_template
    from source_extract import extract
    run = extract(read_template(make_template(tmp_path / "t.xlsx")), [tmp_path / "old.xls"],
                  image_dir=tmp_path / "imgs", use_images=False)
    assert [Path(n.saved_as).name for n in run.images] == ["old_Sheet_D5_1.png"]
    assert (tmp_path / "imgs" / "old_Sheet_D5_1.png").exists()
    names = {g.sheet: g.cells for g in grids}
    assert names["Sheet"][(1, 2)] == 285
    assert any("text box" in n and c == {(1, 1): "Spot Size (mm): 0.05"} for n, c in names.items())
    chart = next(c for n, c in names.items() if "chart" in n)
    # Excel titles a one-series chart after its series; the table sits under it
    assert chart[(1, 1)] == "EC Pre" and chart[(3, 1)] == "Flaw" and chart[(3, 2)] == "EC Pre"
    assert (chart[(5, 1)], chart[(5, 2)]) == ("2A", 71)

    wb.save(tmp_path / "s.xlsx")
    _excel_save(tmp_path / "s.xlsx", tmp_path / "pw.xlsx", 51, password="x")
    with pytest.raises(LedgerError, match="password to open"):
        load_workbook_grids(tmp_path / "pw.xlsx")
    # every hidden Excel this test started was closed
    assert len(started) >= 4 and all(isinstance(p, int) for p in started), started
    import time                                         # Excel takes a moment to exit
    for _ in range(40):
        if not any(_excel_running(p) for p in started):
            break
        time.sleep(0.5)
    left = [p for p in started if _excel_running(p)]
    assert not left, f"{len(left)} of the {len(started)} hidden Excels left running: {left}"


#: RPC_E_CALL_REJECTED (0x80010001) as pywintypes.com_error reports it
_RPC_E_CALL_REJECTED = -2147418111


def test_an_excel_that_will_not_quit_is_ended_after_building_a_workbook(tmp_path, monkeypatch):
    """ledger_bench._excel_save, which the real-Excel tests build their
    workbooks with, gave up when Quit was refused ('Call was rejected by
    callee': Excel's license agreement waiting) and left that hidden Excel
    running, showing the dialog on the user's desktop (review, 2026-10-19).
    It is now tied to this process and ended by its process id."""
    pytest.importorskip("pythoncom")
    win32process = pytest.importorskip("win32process")
    import subprocess
    import win32com.client
    import ledger_bench
    log = []

    class Book:
        def SaveAs(self, path, **kw):
            Path(path).write_bytes(b"saved")

        def Close(self, SaveChanges=True):
            log.append("close")

    class App:
        Hwnd = 77

        class Workbooks:
            @staticmethod
            def Open(path):
                return Book()

        def Quit(self):
            log.append("quit")
            raise Exception(_RPC_E_CALL_REJECTED, "Call was rejected by callee.", None, None)
    monkeypatch.setattr(win32com.client, "DispatchEx", lambda name: App())
    monkeypatch.setattr(win32process, "GetWindowThreadProcessId",
                        lambda hwnd: (1, 4242) if hwnd == 77 else (0, 0))
    monkeypatch.setattr(excel_open, "tie_to_this_process", lambda pid: log.append(("tied", pid)))
    monkeypatch.setattr(subprocess, "run", lambda cmd, **kw: log.append(cmd))
    (tmp_path / "s.xlsx").write_bytes(b"PK")
    ledger_bench._excel_save(tmp_path / "s.xlsx", tmp_path / "out.xls", 56)
    assert log == [("tied", 4242), "close", "quit", ["taskkill", "/F", "/T", "/PID", "4242"]]
    assert (tmp_path / "out.xls").read_bytes() == b"saved"


@real_excel
def test_real_excel_password_xls_is_refused_at_once_and_excel_is_not_started(
        tmp_path, monkeypatch):
    """Excel 16 itself. A password .xls showed a VISIBLE 'Password' dialog on
    the user's desktop and was read as a 2-minute sign-in timeout (review,
    2026-10-19): it, and a password .xlsb, are now refused before any Excel
    starts. A workbook Excel encrypts with its own default password (its
    structure protected) still opens unasked, so it is still read - merged
    label and all."""
    import time
    from ledger_bench import _excel_save
    from source_extract import load_workbook_grids

    def made(name, fmt, password="", edit=None):
        wb = openpyxl.Workbook()
        wb.active["A1"], wb.active["B1"] = "Power (W)", 285
        wb.active["A3"] = "Flaw report"
        wb.save(tmp_path / f"{name}.xlsx")
        _excel_save(tmp_path / f"{name}.xlsx", tmp_path / name, fmt, password=password,
                    edit=edit)
        return tmp_path / name

    def merge_and_protect(sheet):
        sheet.Range("A3:C3").Merge()
        sheet.Parent.Protect("x", True, False)           # the workbook's structure
    try:
        locked = [made("pw.xls", 56, password="x"), made("pw.xlsb", 50, password="x")]
        protected = made("prot.xls", 56, edit=merge_and_protect)
    except Exception as e:
        # _excel_save ended that Excel; a dialog of its own (its license
        # agreement, measured) makes it refuse every call until answered
        if excel_open._com_error_hresult(e) == _RPC_E_CALL_REJECTED:
            pytest.skip("Excel refuses automation calls here ('Call was rejected by callee'): "
                        "a dialog of its own is waiting - open Excel once and answer it")
        raise
    started = _record_excel_pids(monkeypatch)
    for f in locked:
        assert kind_of(f) == "password"
        t = time.monotonic()
        with pytest.raises(ExcelOpenError) as e:
            read_with_excel(f, timeout=10)           # a regression waits 10 s, not 120
        assert e.value.kind == "password" and time.monotonic() - t < 5
    assert started == []                                 # no Excel for them
    assert kind_of(protected) == "legacy"
    grids, _images = load_workbook_grids(protected)
    cells = grids[0].cells
    assert cells[(1, 2)] == 285 and cells[(3, 1)] == "Flaw report"
    assert cells[(3, 2)] == cells[(3, 3)] == "Flaw report"          # the merged label
    assert len(started) == 1 and isinstance(started[0], int), started
    for _ in range(40):                                  # Excel takes a moment to exit
        if not _excel_running(started[0]):
            break
        time.sleep(0.5)
    assert not _excel_running(started[0])


# ── an Excel that cannot start, or refuses every call (review, 2026-10-20) ──
_DOWN = ("Excel could not be started: {} — open Excel once yourself (it may be asking "
         "you to accept its license), then {} again")


def _server_execution_failed():
    """DispatchEx's error when Excel never comes up (~93 s on this PC, its
    license agreement waiting), shaped like pywintypes.com_error."""
    return Exception(-2146959355, "Server execution failed", None, None)


def _call_rejected():
    return Exception(_RPC_E_CALL_REJECTED, "Call was rejected by callee.", None, None)


class _RefusingApp:
    """An Excel.Application whose own dialog (its license agreement) is up:
    every call is refused - app.Hwnd for its first `hwnd_refusals` reads
    (then it answers 77), every property set, Workbooks and Quit."""

    def __init__(self, log, hwnd_refusals=10 ** 6):
        object.__setattr__(self, "log", log)
        object.__setattr__(self, "left", hwnd_refusals)
        object.__setattr__(self, "registered", [])

    def __getattr__(self, name):
        self.log.append(("get", name))
        if name == "Hwnd" and self.left <= 0:
            return 77
        if name == "Hwnd":
            object.__setattr__(self, "left", self.left - 1)
        if name == "Workbooks" and self.left <= 0:
            class Books:
                @staticmethod
                def Open(path, **kw):
                    return type("Book", (), {"Worksheets": [], "Close": lambda s, **k: None})()
            return Books()
        raise _call_rejected()

    def __setattr__(self, name, value):
        # what Cancel could end while the read runs
        self.registered.append(set(excel_open._ACTIVE_EXCEL_PIDS))
        if self.left > 0:
            raise _call_rejected()

    def Quit(self):
        self.log.append("quit")
        raise _call_rejected()


def test_an_excel_that_cannot_start_is_tried_once_per_extract(tmp_path):
    """Each .xls / .xlsb cost ~93 s of 'Server execution failed', at Extract
    and again at Apply: now the extract's first says why and the rest fail
    at once. The next extract tries again; reads outside one (the knowledge
    graph's) are not stopped by it."""
    files = []
    for n in ("a.xls", "b.xlsb"):
        (tmp_path / n).write_bytes(OLE_MAGIC)
        files.append(tmp_path / n)
    tries = []

    def factory():
        tries.append(1)
        raise _server_execution_failed()
    with excel_open.cancel_scope():
        for f in files:
            with pytest.raises(ExcelOpenError) as e:
                read_with_excel(f, want_pictures=False, app_factory=factory)
            assert e.value.kind == "unavailable" and e.value.down == "Server execution failed"
            assert str(e.value) == _DOWN.format("Server execution failed", "extract")
    assert len(tries) == 1
    with excel_open.cancel_scope():                      # Apply's extract: once more
        with pytest.raises(ExcelOpenError):
            read_with_excel(files[0], want_pictures=False, app_factory=factory)
    assert len(tries) == 2
    for f in files:                                      # no extract: each its own
        with pytest.raises(ExcelOpenError) as e:
            read_with_excel(f, want_pictures=False, app_factory=factory)
        assert str(e.value) == _DOWN.format("Server execution failed", "try")
    assert len(tries) == 4


def test_an_excel_refusing_every_call_opens_nothing_and_is_not_asked_again(
        tmp_path, monkeypatch):
    """'Call was rejected by callee' (a dialog of its own up): its macro
    setting refused too, so nothing is opened in it; the extract's other
    workbooks do not start another Excel."""
    monkeypatch.setattr(excel_open, "_HWND_WAIT_S", 0.3, raising=False)
    f = tmp_path / "a.xls"
    f.write_bytes(OLE_MAGIC)
    log, apps = [], []

    def factory():
        apps.append(_RefusingApp(log))
        return apps[-1]
    with excel_open.cancel_scope():
        for _ in range(2):
            with pytest.raises(ExcelOpenError) as e:
                read_with_excel(f, want_pictures=False, app_factory=factory)
            assert e.value.kind == "unavailable"
            assert str(e.value) == _DOWN.format("Call was rejected by callee", "extract")
    assert len(apps) == 1
    assert ("get", "Workbooks") not in log and log[-1] == "quit"
    assert log.count(("get", "Hwnd")) > 1               # asked again, not once


def _dispatch(monkeypatch, app, snapshots):
    """read_with_excel's own DispatchEx path, Excel being `app`: the EXCEL.EXE
    processes it sees are `snapshots` (before, after DispatchEx), and what
    it would do to the process - tie, hold, end - is logged, never done."""
    pytest.importorskip("pythoncom")
    win32com_client = pytest.importorskip("win32com.client")
    win32process = pytest.importorskip("win32process")
    log = app.log
    monkeypatch.setattr(excel_open, "excel_available", lambda: True)
    monkeypatch.setattr(win32com_client, "DispatchEx", lambda name: app)
    monkeypatch.setattr(win32process, "GetWindowThreadProcessId",
                        lambda hwnd: (1, 4242) if hwnd == 77 else (0, 0))
    monkeypatch.setattr(excel_open, "_HWND_WAIT_S", 0.5, raising=False)
    shots = iter(snapshots)
    patch = dict(raising=False)          # so the same test shows what came before
    monkeypatch.setattr(excel_open, "_excel_processes", lambda: next(shots), **patch)
    monkeypatch.setattr(excel_open, "tie_to_this_process", lambda pid: log.append(("tied", pid)))
    monkeypatch.setattr(excel_open, "_hold_excel",
                        lambda pid, created=None: log.append(("held", pid, created)) or "H",
                        **patch)
    monkeypatch.setattr(excel_open, "_let_go",
                        lambda held, end: log.append(("let go", held, end)), **patch)
    monkeypatch.setattr(excel_open, "_ACTIVE_EXCEL_PIDS", set())


def test_an_excel_refusing_every_call_is_found_tied_registered_and_ended(tmp_path, monkeypatch):
    """Its window (app.Hwnd) was read once: refused, so the Excel was neither
    tied to this process nor registered for Cancel, and its refused Quit
    left it running - a hidden Excel sat on its license agreement for half
    an hour (review, 2026-10-20). It is now the one EXCEL.EXE that started
    during DispatchEx, and is ended through a handle taken then."""
    app = _RefusingApp([])
    _dispatch(monkeypatch, app, [{(100, 1)}, {(100, 1), (4242, 9)}])
    f = tmp_path / "a.xls"
    f.write_bytes(OLE_MAGIC)
    with pytest.raises(ExcelOpenError) as e:
        read_with_excel(f, want_pictures=False, timeout=30)
    assert e.value.kind == "unavailable"
    steps = [x for x in app.log if x[0] != "get"]
    assert steps == [("tied", 4242), ("held", 4242, 9), "quit", ("let go", "H", True)]
    assert app.registered and app.registered[0] == {4242}       # Cancel could end it
    assert excel_open._ACTIVE_EXCEL_PIDS == set()


def test_an_excel_whose_window_answers_late_is_tied_and_ended_when_quit_is_refused(
        tmp_path, monkeypatch):
    """Refusing calls for a moment, then answering: its window is asked again,
    so it is tied and registered by it; Quit refused, it is ended."""
    app = _RefusingApp([], hwnd_refusals=2)
    _dispatch(monkeypatch, app, [{(100, 1)}, {(100, 1)}])
    f = tmp_path / "a.xls"
    f.write_bytes(OLE_MAGIC)
    book = read_with_excel(f, want_pictures=False, timeout=30)
    assert book.sheets == []
    steps = [x for x in app.log if x[0] != "get"]
    assert steps == [("tied", 4242), ("held", 4242, None), "quit", ("let go", "H", True)]
    assert app.registered and app.registered[0] == {4242}


def test_two_excels_starting_at_once_are_not_guessed_at(tmp_path, monkeypatch):
    """Another EXCEL.EXE started during DispatchEx too: which is this read's
    cannot be told, so neither is tied, held or ended."""
    app = _RefusingApp([])
    _dispatch(monkeypatch, app, [{(100, 1)}, {(100, 1), (4242, 9), (5150, 9)}])
    f = tmp_path / "a.xls"
    f.write_bytes(OLE_MAGIC)
    with pytest.raises(ExcelOpenError):
        read_with_excel(f, want_pictures=False, timeout=30)
    assert [x for x in app.log if x[0] != "get"] == ["quit"]


def test_the_process_this_read_started_is_held_and_ended_by_handle(tmp_path):
    """The process helpers on a real process (a Python child standing in for
    EXCEL.EXE): found as the one new process of its image; a handle is not
    taken when the start time differs (the PID given to another process
    since); ending it through the handle ends it."""
    import subprocess
    if sys.platform != "win32":
        pytest.skip("Windows processes")
    image = Path(sys.executable).name
    before = excel_open._processes_named(image)
    child = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(60)"])
    try:
        new = set()
        for _ in range(50):
            new = {p for p in excel_open._processes_named(image) - before if p[0] == child.pid}
            if new:
                break
            import time
            time.sleep(0.1)
        (pid, created), = new
        assert excel_open._hold_excel(pid, created + 1) is None
        held = excel_open._hold_excel(pid, created)
        assert held
        excel_open._let_go(held, end=True)
        assert child.wait(10) == 1
    finally:
        child.kill()


def test_an_extract_whose_excel_cannot_start_says_so_on_one_line(tmp_path, monkeypatch):
    """One 93-s failure and one line per .xls / .xlsb ('could not be read
    (com_error(...))'): now one try, and one line naming every workbook it
    left unread. Apply's extract tries once more."""
    import source_extract as se
    from ledger_fixtures import make_laser_log, make_template
    from sample_ledger import read_template
    schema = read_template(make_template(tmp_path / "t.xlsx"))
    a = make_laser_log(tmp_path / "a.xlsx")
    olds = [tmp_path / n for n in ("b.xls", "c.xls", "d.xlsb")]
    for f in olds:
        f.write_bytes(OLE_MAGIC)
    tries = []

    def factory():
        tries.append(1)
        raise _server_execution_failed()
    real = excel_open.read_with_excel
    monkeypatch.setattr(excel_open, "read_with_excel", lambda path, **kw: real(
        path, want_pictures=False, app_factory=factory))
    run = se.extract(schema, [a, *olds], use_images=False)
    assert run.obs and len(tries) == 1                   # a.xlsx read; Excel tried once
    said = [p for p in run.problems if any(f.name in p for f in olds)]
    assert said == [_DOWN.format("Server execution failed", "extract")
                    + " (not read: b.xls, c.xls, d.xlsb)"]
    assert [k for _n, k in run.open_failures] == ["unavailable"] * 3
    se.extract(schema, [a, *olds], use_images=False)     # Apply reads them again
    assert len(tries) == 2
