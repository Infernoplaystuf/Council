"""The Ledger tab with a table template (council_core.ledger_records): the
same buttons as for a sample Ledger, offscreen."""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")
from PySide6.QtWidgets import QApplication  # noqa: E402

from council_qt.tabs.ledger import LedgerActions, LedgerTab  # noqa: E402
from council_core import ledger_records as lr  # noqa: E402

sys.path.insert(0, str(Path(__file__).parent))
from test_ledger_records import btb, decide_all  # noqa: E402,F401


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


def test_a_table_template_from_the_tab(qapp, btb):  # noqa: F811
    tpl, src, tmp, _ = btb
    tab = LedgerTab(None, LedgerActions(opener=lambda _p: None))
    tab.template.setText(str(tpl))
    tab.create_btn.click()
    assert tab.wait()
    csv = Path(tab.csv.text())
    assert lr.is_record_ledger(csv)
    assert "one row per record" in tab.log.toPlainText()

    # Teach: the selected sources' highlights
    tab.add_sources([src / "BTB_Woopentide.xlsx", src / "BTB_Saddle_BP.xlsx"])
    for i in range(tab.sources.count()):
        tab.sources.item(i).setSelected(True)
    tab.teach_btn.click()
    assert tab.wait()
    log = tab.log.toPlainText()
    assert "Learned from your highlights in BTB_Woopentide.xlsx" in log
    assert "a lookup table" in log

    tab.clear_sources()
    tab.add_sources([src])
    tab.extract_btn.click()
    assert tab.wait()
    assert tab.stat["values"].text() == "558"
    assert "62 records" in tab.log.toPlainText()
    review = Path(tab.review.text())
    decide_all(review)
    tab.fill_dir.setText(str(tmp / "filled"))
    tab.apply_btn.click()
    assert tab.wait()
    assert "62 records added" in tab.status.text()
    assert (tmp / "filled" / "BTB_Template_filled.xlsx").is_file()


def test_the_records_box_makes_a_table_ledger(qapp, tmp_path):
    import openpyxl
    from openpyxl.styles import Border, Side
    wb = openpyxl.Workbook()
    for c, f in enumerate(["Sample ID", "Hardness", "Width"], start=1):
        wb.active.cell(1, c, f)
        wb.active.cell(2, c).border = Border(bottom=Side(style="thin"))
    tpl = tmp_path / "T.xlsx"
    wb.save(tpl)
    tab = LedgerTab(None, LedgerActions(opener=lambda _p: None))
    tab.template.setText(str(tpl))
    tab.records_mode.setChecked(True)
    tab.create_btn.click()
    assert tab.wait()
    assert lr.is_record_ledger(tab.csv.text())
